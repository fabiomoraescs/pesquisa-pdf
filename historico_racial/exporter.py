"""Exportação documental em quatro abas; não executa análise ou recodificação."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .schemas import CODIFICACAO, COOCORRENCIAS, DOCUMENTOS, OCORRENCIAS


DOCUMENT_HEADERS = (
    "ID projeto", "ID documento", "Arquivo PDF", "ID arquivo", "Página PDF inicial",
    "Página PDF final", "OCR utilizado", "Idioma OCR utilizado", "Publicação",
    "Organização", "Ano", "Autor", "Título", "Página impressa", "Método de análise",
    "Limiar semântico", "Modelo semântico", "Versão do vocabulário", "Hash do vocabulário",
    "Bibliotecas de origem", "Data do processamento",
)
OCCURRENCE_HEADERS = (
    "ID ocorrência", "ID documento", "Arquivo PDF", "Página PDF", "Página PDF final",
    "Categoria de busca", "Tipo de entidade", "Identificador da entidade", "Entidade canônica",
    "Termo encontrado", "Forma original no texto", "Grupo(s)", "Bibliotecas de origem",
    "Método de análise", "Tipo de correspondência", "Similaridade semântica",
    "Trecho anterior", "Trecho da ocorrência", "Trecho posterior", "Contexto completo",
    "Tradição intelectual", "País/região matriz", "Validação", "Forma de referência",
    "Operação de repertório",
)
CODING_HEADERS = ("ID ocorrência", "Validação manual", "Forma de referência", "Operação de repertório", "Nota analítica do pesquisador")
COOCCURRENCE_HEADERS = ("ID ocorrência A", "ID ocorrência B", "ID documento", "Observações do pesquisador")


def _valor(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, list):
        return "; ".join(str(item) for item in value)
    return value


def _append(sheet, values) -> None:
    sheet.append([_valor(value) for value in values])
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"  # texto de PDF nunca deve ser interpretado como fórmula


def gerar_xlsx(resultado: dict[str, Any]) -> BytesIO:
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = {
        name: workbook.create_sheet(name)
        for name in (DOCUMENTOS, OCORRENCIAS, CODIFICACAO, COOCORRENCIAS)
    }
    for name, headers in (
        (DOCUMENTOS, DOCUMENT_HEADERS), (OCORRENCIAS, OCCURRENCE_HEADERS),
        (CODIFICACAO, CODING_HEADERS), (COOCORRENCIAS, COOCCURRENCE_HEADERS),
    ):
        _append(sheets[name], headers)
    metodo = "Híbrido" if resultado.get("metodo_analise") == "hibrido" else "Lexical"
    for document in resultado["documentos"]:
        _append(sheets[DOCUMENTOS], (
            document.get("id_project"), document["id_documento"], document["arquivo_pdf"],
            document["id_arquivo"], document["pagina_pdf_inicio"], document["pagina_pdf_fim"],
            "Sim" if document["ocr_utilizado"] else "Não", document.get("ocr_idioma_utilizado"),
            document.get("publicacao"), document.get("organizacao"), document.get("ano"),
            document.get("autor"), document.get("titulo"), document.get("pagina_impressa"),
            metodo, resultado.get("limiar_semantico"), resultado.get("modelo_semantico"),
            resultado.get("vocabulario_version"), resultado.get("vocabulario_hash"),
            resultado.get("library_names", []), resultado.get("data_processamento"),
        ))
    for item in resultado["ocorrencias"]:
        _append(sheets[OCORRENCIAS], (
            item["id_ocorrencia"], item["id_documento"], item["arquivo_pdf"], item["pagina_pdf"],
            item.get("pagina_pdf_final"), item["categoria_busca"], item["tipo_entidade"],
            item["id_entidade"], item["entidade_canonica"], item["termo_encontrado"],
            item["forma_original_no_texto"], item["grupo"], item.get("bibliotecas_origem", []),
            metodo, item.get("tipo_correspondencia", "Lexical"), item.get("similaridade_semantica"),
            item["trecho_anterior"], item["trecho_ocorrencia"], item["trecho_posterior"],
            item["contexto_completo"], item.get("tradicao_intelectual"), item.get("pais_regiao_matriz"),
            item.get("validar"), item.get("forma_referencia"), item.get("operacao_repertorio"),
        ))
        _append(sheets[CODIFICACAO], (item["id_ocorrencia"], "", "", "", ""))
    for sheet in sheets.values():
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 28
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=False)
        for column in range(1, sheet.max_column + 1):
            heading = sheet.cell(1, column).value or ""
            sheet.column_dimensions[get_column_letter(column)].width = (
                55 if heading in {"Trecho anterior", "Trecho da ocorrência", "Trecho posterior", "Contexto completo"}
                else 26 if "ID " in heading or "Hash" in heading else 22
            )
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
