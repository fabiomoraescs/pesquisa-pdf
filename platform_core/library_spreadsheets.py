"""Leitura estrita de modelos XLSX para bibliotecas oficiais em rascunho."""

from __future__ import annotations

import io
import re
import secrets
import time
import unicodedata
from dataclasses import dataclass
from threading import RLock
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from historico_racial.dictionaries import carregar_categorias
from historico_racial.vocabulary import VocabularioError, _contagens, validar_vocabulario


MAX_XLSX_BYTES = 2 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_DATA_ROWS = 20_000
REQUIRED_SHEETS = ("BIBLIOTECA", "GRUPOS", "ENTIDADES", "VARIANTES")
HEADERS = {
    "BIBLIOTECA": ("nome", "descricao"),
    "GRUPOS": ("id_grupo", "nome", "descricao", "ativo"),
    "ENTIDADES": ("identificador_entidade", "forma_canonica", "tipo_entidade", "grupos", "ativo",
                  "tradicao_intelectual", "pais_regiao", "observacoes"),
    "VARIANTES": ("identificador_entidade", "variante", "ativo"),
}
ID_PATTERN = re.compile(r"[a-zA-Z][a-zA-Z0-9_]{0,79}\Z")
TEMPLATE_REFERENCE_NAME = "Relações raciais — cópia de referência"


class SpreadsheetImportError(ValueError):
    """Erro público em português, com localização na aba e linha."""


@dataclass(frozen=True)
class ImportedLibrary:
    name: str
    description: str
    snapshot: dict
    counts: dict[str, int]


def _text(value: object, location: str, *, required: bool = True, limit: int = 200) -> str:
    if value is None and not required:
        return ""
    if not isinstance(value, str):
        raise SpreadsheetImportError(f"{location}: informe um texto válido.")
    text = value.strip()
    if required and not text:
        raise SpreadsheetImportError(f"{location}: campo obrigatório vazio.")
    if len(text) > limit:
        raise SpreadsheetImportError(f"{location}: excede o limite de {limit} caracteres.")
    return text


def _active(value: object, location: str) -> bool:
    if type(value) is bool:
        return value
    if isinstance(value, str):
        normalized = value.strip().casefold()
        if normalized in {"sim", "ativo", "true"}:
            return True
        if normalized in {"não", "nao", "inativo", "false"}:
            return False
    raise SpreadsheetImportError(f"{location}: use Sim ou Não para ativo.")


def _variant_key(text: str) -> str:
    folded = unicodedata.normalize("NFKD", text.casefold())
    return "".join(char for char in folded if not unicodedata.combining(char))


def _rows(workbook, sheet_name: str) -> list[tuple[int, dict[str, object]]]:
    sheet = workbook[sheet_name]
    expected = HEADERS[sheet_name]
    if (sheet.max_row is not None and sheet.max_row > MAX_DATA_ROWS + 1) or (sheet.max_column is not None and sheet.max_column > len(expected) + 8):
        raise SpreadsheetImportError(f"{sheet_name}: a planilha excede o limite de linhas ou colunas.")
    iterator = sheet.iter_rows(values_only=False)
    first = next(iterator, None)
    if not first:
        raise SpreadsheetImportError(f"{sheet_name}, linha 1: cabeçalhos ausentes.")
    names = [cell.value for cell in first]
    normalized = ["identificador_entidade" if name == "entity_key" else name for name in names]
    if any(not isinstance(name, str) for name in names) or len(set(normalized)) != len(normalized) or set(normalized) != set(expected):
        raise SpreadsheetImportError(f"{sheet_name}, linha 1: use exatamente os cabeçalhos do modelo oficial.")
    result = []
    for line, cells in enumerate(iterator, start=2):
        if line > MAX_DATA_ROWS + 1:
            raise SpreadsheetImportError(f"{sheet_name}: a planilha excede o limite de linhas.")
        if all(cell.value is None or cell.value == "" for cell in cells):
            continue
        if any(cell.data_type == "f" for cell in cells):
            raise SpreadsheetImportError(f"{sheet_name}, linha {line}: fórmulas não são permitidas.")
        result.append((line, dict(zip(normalized, (cell.value for cell in cells)))))
    return result


def parse_library_xlsx(data: bytes) -> ImportedLibrary:
    """Valida o arquivo integralmente antes de retornar qualquer snapshot persistível."""
    if not data or len(data) > MAX_XLSX_BYTES:
        raise SpreadsheetImportError("Envie uma planilha XLSX de até 2 MB.")
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            items = archive.infolist()
            if len(items) > 1000 or sum(item.file_size for item in items) > MAX_UNCOMPRESSED_BYTES:
                raise SpreadsheetImportError("A planilha descompactada excede o limite permitido.")
            if any("vbaproject.bin" in item.filename.casefold() or "externallinks/" in item.filename.casefold()
                   or "oleobjects/" in item.filename.casefold() or "embeddings/" in item.filename.casefold()
                   for item in items):
                raise SpreadsheetImportError("A planilha contém conteúdo externo ou incorporado não permitido.")
            for item in items:
                if item.filename.casefold().endswith(".rels"):
                    relationships = ElementTree.fromstring(archive.read(item))
                    if any(relation.attrib.get("TargetMode", "").casefold() == "external"
                           for relation in relationships):
                        raise SpreadsheetImportError("A planilha contém links externos não permitidos.")
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except SpreadsheetImportError:
        raise
    except (BadZipFile, ValueError, OSError, KeyError, ElementTree.ParseError) as error:
        raise SpreadsheetImportError("O arquivo enviado não é um XLSX válido.") from error
    try:
        missing = set(REQUIRED_SHEETS) - set(workbook.sheetnames)
        unexpected = set(workbook.sheetnames) - set(REQUIRED_SHEETS) - {"INSTRUÇÕES"}
        if missing or unexpected:
            raise SpreadsheetImportError(
                "Abas inválidas. São necessárias BIBLIOTECA, GRUPOS, ENTIDADES e VARIANTES; INSTRUÇÕES é opcional."
            )
        tables = {name: _rows(workbook, name) for name in REQUIRED_SHEETS}
    except SpreadsheetImportError:
        raise
    except (ValueError, OSError, KeyError, IndexError, TypeError, AttributeError, ElementTree.ParseError) as error:
        raise SpreadsheetImportError("O arquivo enviado não é um XLSX válido.") from error
    finally:
        workbook.close()

    if len(tables["BIBLIOTECA"]) != 1:
        raise SpreadsheetImportError("BIBLIOTECA: preencha exatamente uma linha de dados.")
    library_line, library_data = tables["BIBLIOTECA"][0]
    name = _text(library_data["nome"], f"BIBLIOTECA, linha {library_line}, nome", limit=160)
    if name.casefold() == TEMPLATE_REFERENCE_NAME.casefold():
        raise SpreadsheetImportError(
            f"BIBLIOTECA, linha {library_line}: altere o nome da cópia de referência "
            "antes de importar uma nova biblioteca."
        )
    description = _text(library_data["descricao"], f"BIBLIOTECA, linha {library_line}, descrição",
                        required=False, limit=4000)
    groups: dict[str, dict] = {}
    group_names: set[str] = set()
    for line, item in tables["GRUPOS"]:
        location = f"GRUPOS, linha {line}"
        group_id = _text(item["id_grupo"], f"{location}, id_grupo", limit=80)
        if not ID_PATTERN.fullmatch(group_id):
            raise SpreadsheetImportError(f"{location}: identificador de grupo inválido.")
        if group_id.casefold() in {key.casefold() for key in groups}:
            raise SpreadsheetImportError(f"{location}: identificador de grupo duplicado: '{group_id}'.")
        group_name = _text(item["nome"], f"{location}, nome", limit=160)
        if group_name.casefold() in group_names:
            raise SpreadsheetImportError(f"{location}: nome de grupo duplicado: '{group_name}'.")
        group_names.add(group_name.casefold())
        groups[group_id] = {"id_grupo": group_id, "nome": group_name,
                            "descricao": _text(item["descricao"], f"{location}, descrição", required=False, limit=1000),
                            "ativo": _active(item["ativo"], f"{location}, ativo")}
    if not groups:
        raise SpreadsheetImportError("GRUPOS: cadastre ao menos um grupo.")

    categories = carregar_categorias()
    entities: list[dict] = []
    by_key: dict[str, dict] = {}
    canonical_names: set[str] = set()
    for line, item in tables["ENTIDADES"]:
        location = f"ENTIDADES, linha {line}"
        key = _text(item["identificador_entidade"], f"{location}, identificador_entidade", limit=80)
        if not ID_PATTERN.fullmatch(key):
            raise SpreadsheetImportError(f"{location}: identificador de entidade inválido.")
        if key.casefold() in {saved.casefold() for saved in by_key}:
            raise SpreadsheetImportError(f"{location}: identificador de entidade duplicado: '{key}'.")
        canonical = _text(item["forma_canonica"], f"{location}, forma_canonica")
        if canonical.casefold() in canonical_names:
            raise SpreadsheetImportError(f"{location}: forma canônica duplicada: '{canonical}'.")
        canonical_names.add(canonical.casefold())
        entity_type = _text(item["tipo_entidade"], f"{location}, tipo_entidade", required=False, limit=80) or "outro"
        if entity_type not in categories["tipos_entidade"]:
            raise SpreadsheetImportError(f"{location}: tipo de entidade inválido: '{entity_type}'.")
        raw_groups = _text(item["grupos"], f"{location}, grupos", limit=1000)
        memberships = [value.strip() for value in raw_groups.split(";")]
        if not all(memberships) or len(set(memberships)) != len(memberships):
            raise SpreadsheetImportError(f"{location}: informe grupos distintos separados por ponto e vírgula.")
        for group_id in memberships:
            if group_id not in groups:
                raise SpreadsheetImportError(f"{location}: o grupo '{group_id}' não existe.")
        tradition = _text(item["tradicao_intelectual"], f"{location}, tradição", required=False, limit=100)
        if tradition and tradition not in categories["tradicoes_intelectuais"]:
            raise SpreadsheetImportError(f"{location}: tradição intelectual inválida: '{tradition}'.")
        entity = {
            "id_entidade": key, "entity_key": key, "forma_canonica": canonical,
            "tipo_entidade": entity_type, "grupo": memberships,
            "ativo": _active(item["ativo"], f"{location}, ativo"),
            "tradicao_intelectual": tradition,
            "pais_regiao": _text(item["pais_regiao"], f"{location}, país/região", required=False, limit=160),
            "observacoes": _text(item["observacoes"], f"{location}, observações", required=False, limit=4000),
            "variantes": [],
        }
        entities.append(entity)
        by_key[key] = entity
    if not entities:
        raise SpreadsheetImportError("ENTIDADES: cadastre ao menos uma entidade.")

    variants_seen: dict[str, str] = {}
    for line, item in tables["VARIANTES"]:
        location = f"VARIANTES, linha {line}"
        key = _text(item["identificador_entidade"], f"{location}, identificador_entidade", limit=80)
        if key not in by_key:
            raise SpreadsheetImportError(f"{location}: a entidade '{key}' não existe.")
        variant = _text(item["variante"], f"{location}, variante")
        folded = _variant_key(variant)
        if folded in variants_seen:
            previous = variants_seen[folded]
            if previous == key:
                raise SpreadsheetImportError(f"{location}: variante duplicada na entidade '{key}': '{variant}'.")
            raise SpreadsheetImportError(f"{location}: a variante '{variant}' já pertence a outra entidade: '{previous}'.")
        variants_seen[folded] = key
        by_key[key]["variantes"].append({"texto": variant, "ativo": _active(item["ativo"], f"{location}, ativo")})
    for entity in entities:
        if not entity["variantes"]:
            raise SpreadsheetImportError(f"VARIANTES: a entidade '{entity['id_entidade']}' não possui variante.")
        if entity["forma_canonica"].casefold() not in {variant["texto"].casefold() for variant in entity["variantes"]}:
            raise SpreadsheetImportError(f"VARIANTES: inclua a forma canônica da entidade '{entity['id_entidade']}'.")
    snapshot = {"grupos": groups, "entidades": entities}
    try:
        validated = validar_vocabulario(snapshot)
    except VocabularioError as error:
        raise SpreadsheetImportError(f"Vocabulário inválido: {error}") from error
    return ImportedLibrary(name, description, validated, _contagens(validated))


class ImportPreviews:
    """Prévia efêmera por admin; nenhum registro parcial é inserido no banco."""

    def __init__(self, ttl_seconds: int = 900):
        self._ttl = ttl_seconds
        self._items: dict[str, tuple[float, str, ImportedLibrary]] = {}
        self._lock = RLock()

    def put(self, user_id: str, imported: ImportedLibrary) -> str:
        token = secrets.token_urlsafe(24)
        with self._lock:
            now = time.monotonic()
            self._items = {key: item for key, item in self._items.items() if now - item[0] <= self._ttl}
            self._items[token] = (now, user_id, imported)
        return token

    def get(self, token: str, user_id: str) -> ImportedLibrary | None:
        with self._lock:
            item = self._items.get(token)
            if item is None or item[1] != user_id:
                return None
            if time.monotonic() - item[0] > self._ttl:
                self._items.pop(token, None)
                return None
            return item[2]

    def discard(self, token: str) -> None:
        with self._lock:
            self._items.pop(token, None)


previews = ImportPreviews()
