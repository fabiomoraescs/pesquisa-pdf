"""Fluxo mínimo de projetos e importação inicial do Instagram."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
import secrets
import time
from uuid import UUID

from flask import Blueprint, abort, current_app, flash, jsonify, redirect, render_template, request, send_file, session, url_for
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import joinedload

from .extensions import db
from .instagram_import import (
    MAX_IMPORT_BYTES,
    MAX_IMPORT_ROWS,
    IMPORT_FIELDS,
    InstagramImportError,
    build_preview,
    parse_instagram_import,
)
from .instagram_comment_import import (
    COMMENT_IMPORT_FIELDS,
    build_comment_preview,
    parse_instagram_comment_import,
)
from .instagram_comment_previews import (
    InstagramCommentPreviewExpired,
    comment_previews,
)
from .instagram_comment_storage import (
    InstagramCommentImportPersistenceError,
    comment_import_conflict,
    persist_instagram_comment_import,
)
from .instagram_comments import (
    ASSESSMENT_SOURCE_LABELS,
    INTENSITY_LABELS,
    INTENT_LABELS,
    RISK_LABELS,
    SENTIMENT_LABELS,
    CommentAssessmentValidationError,
    build_instagram_comments_dashboard,
    classify_comment,
    current_comment_observation,
    current_text_version,
    effective_assessment,
)
from .instagram_comment_classification import (
    COMMENT_CLASSIFIER_LABELS,
    CommentClassificationError,
    CommentClassificationPersistenceError,
    LOCAL_COMMENT_CLASSIFIER_LABEL,
    classify_project_comments,
    get_comment_classifier,
    plan_comment_classification,
)
from .instagram_lifecycle import (
    InstagramProjectDeletionError,
    delete_instagram_project_permanently,
)
from .instagram_meta import (
    META_CONNECTION_PROVIDERS,
    InstagramMetaError,
    InstagramMetaSettings,
    choose_candidate,
    connection_provider,
    disconnect_connection,
    log_meta_operation,
    pending_meta_reconciliations,
    pending_meta_selections,
    persist_validated_connection,
    provider_scopes,
    reconciliation_required,
    run_capability_test,
    safe_meta_error_message,
)
from .instagram_meta_sync_jobs import (
    DEFAULT_META_SYNC_STALE_SECONDS,
    META_SYNC_COMPLETED,
    META_SYNC_FAILED,
    active_meta_sync_run,
    enqueue_meta_sync_run,
    recover_stale_meta_sync_runs,
)
from .instagram_meta_sync import (
    SYNC_PERIOD_OPTIONS,
    SYNC_PERIOD_PRESETS,
    InstagramMetaSyncError,
    resolve_publication_window,
)
from .instagram_periods import (
    TEMPORAL_PERIOD_OPTIONS,
    effective_run_period,
    form_period as instagram_period_form,
    period_label as instagram_period_label,
    resolve_calendar_dates,
)
from .pagination import PAGE_SIZE, PlatformPage, paginate_sequence, requested_page
from .instagram_insights import (
    InstagramInsightsError,
    InstagramInsightsPersistenceError,
    evidence_views,
    finding_for_project,
    generate_insights,
    grouped_findings,
    latest_insight_run,
)
from .instagram_recommendations import (
    InstagramRecommendationsError,
    InstagramRecommendationsPersistenceError,
    generate_recommendations,
    linked_findings,
    ordered_recommendations,
    recommendation_for_project,
    recommendation_run_for_insight,
)
from .instagram_actions import (
    ACTION_KINDS,
    ACTION_KIND_LABELS,
    ACTION_OUTCOME_LABELS,
    ACTION_STATUSES,
    ACTION_STATUS_LABELS,
    InstagramActionError,
    InstagramActionPersistenceError,
    action_counts,
    action_prefill,
    actions_for_project,
    create_recommendation_action,
    delete_recommendation_action,
    parse_planned_datetime,
    recommendation_action_for_project,
    recommendation_action_has_evaluation,
    transition_recommendation_action,
)
from .instagram_action_evaluations import (
    ASSESSMENT_LABELS,
    InstagramActionEvaluationDuplicateError,
    InstagramActionEvaluationError,
    InstagramActionEvaluationPersistenceError,
    create_action_evaluation,
    eligible_followup_runs,
    evaluation_view,
    evaluations_for_action,
    metric_label,
    preview_action_evaluation,
    update_action_evaluation,
)
from .instagram_learnings import (
    LEARNING_ASSESSMENT_LABELS,
    InstagramLearningError,
    InstagramLearningPersistenceError,
    comparable_evaluations,
    create_learning,
    delete_learning,
    learning_cohort_label,
    learning_for_project,
    learnings_for_project,
    update_learning,
)
from .instagram_reports import (
    InstagramReportError,
    InstagramReportPersistenceError,
    artifact_path,
    delete_report,
    generate_report,
    report_chart_data,
    report_download_name,
    report_summary_cards,
    report_for_project,
    reports_page_for_project,
    reports_for_project,
)
from .instagram_previews import InstagramPreviewExpired, previews
from .presentation import format_local_datetime, to_local_datetime
from .instagram_dashboard import PERIOD_DAYS, TABLE_SORTS, build_instagram_dashboard
from .instagram_taxonomy import (
    CLASSIFICATION_STATUSES,
    DEFAULT_TAXONOMY_VERSION,
    TaxonomyValidationError,
    classification_states,
    classify_media,
    create_category,
    create_dimension,
    create_initial_taxonomy,
    current_classifications,
    set_category_active,
    set_dimension_active,
    taxonomy_dimensions,
)
from .instagram_storage import (
    InstagramImportPersistenceError,
    account_conflict,
    persist_instagram_import,
)
from .models import (
    AnalyticsProject,
    AnalyticsRun,
    InstagramAccount,
    InstagramComment,
    InstagramCommentTextVersion,
    InstagramInsightRun,
    InstagramMedia,
    InstagramRecommendationActionEvaluation,
    InstagramReport,
    InstagramTaxonomyCategory,
    InstagramTaxonomyDimension,
)
from .services import ANALYTICS_INSTAGRAM_TOOL, can_use_tool


instagram_bp = Blueprint("instagram", __name__, url_prefix="/analytics/instagram")

_META_OAUTH_STATE_KEY = "instagram_meta_oauth_state"
_META_OAUTH_STATE_TTL_SECONDS = 600


def _require_tool() -> None:
    if not can_use_tool(current_user, ANALYTICS_INSTAGRAM_TOOL):
        abort(403)


def _project_for_current_user(project_id: UUID, *, active: bool = False) -> AnalyticsProject:
    _require_tool()
    query = select(AnalyticsProject).where(
        AnalyticsProject.id == str(project_id),
        AnalyticsProject.owner_user_id == current_user.id,
        AnalyticsProject.module_key == "instagram",
    )
    if active:
        query = query.where(AnalyticsProject.status == "active", AnalyticsProject.archived_at.is_(None))
    project = db.session.scalar(query)
    if project is None:
        abort(404)
    return project


def _meta_settings() -> InstagramMetaSettings:
    return InstagramMetaSettings.from_mapping(current_app.config)


def _meta_client(settings: InstagramMetaSettings, provider: str):
    factory = current_app.extensions.get("instagram_meta_client_factory")
    return factory(settings, provider) if factory is not None else connection_provider(settings, provider)


def _sync_period_form(values: dict | None = None, *, active_run: AnalyticsRun | None = None) -> dict[str, str]:
    """Mantém o formulário de janela previsível sem aceitar presets arbitrários."""
    values = values or {}
    if "sync_period" not in values and active_run is not None:
        parameters = active_run.parameters_json if isinstance(active_run.parameters_json, dict) else {}
        preset = parameters.get("period_preset")
        if isinstance(preset, str):
            values = {"sync_period": preset}
            # Um reload enquanto uma sincronização personalizada está ativa
            # reconstrói o formulário da janela já congelada, sem depender do
            # estado do navegador nem alterar os dados do run.
            resolved = effective_run_period(active_run)
            if preset == "custom" and resolved.is_bounded:
                timezone_name = resolved.timezone_name or current_app.config["PRESENTATION_TIMEZONE"]
                start = to_local_datetime(resolved.start, timezone_name)
                end = to_local_datetime(resolved.end, timezone_name)
                values.update({
                    "published_from": start.date().isoformat() if start is not None else "",
                    "published_to": end.date().isoformat() if end is not None else "",
                })
    preset = values.get("sync_period", "last_90_days")
    return {
        "key": preset if preset in SYNC_PERIOD_PRESETS else "last_90_days",
        "start": str(values.get("published_from", "") or ""),
        "end": str(values.get("published_to", "") or ""),
    }


def _meta_sync_completion_url(project: AnalyticsProject, run: AnalyticsRun) -> str:
    """Retorno de polling que preserva a janela efetivamente congelada."""
    period_form = _sync_period_form({}, active_run=run)
    values = {
        "project_id": project.id,
        "meta_sync_completed": run.id,
        "sync_period": period_form["key"],
    }
    if period_form["key"] == "custom":
        values.update({"published_from": period_form["start"], "published_to": period_form["end"]})
    return url_for("instagram.new_import", **values)


def _runs_history_context(project: AnalyticsProject, page_value: object) -> dict:
    """Uma listagem paginada compartilhada para importações e sincronizações."""
    total = int(db.session.scalar(select(func.count()).select_from(AnalyticsRun).where(
        AnalyticsRun.analytics_project_id == project.id
    )) or 0)
    requested = requested_page(page_value)
    current = min(requested, max(1, (total + PAGE_SIZE - 1) // PAGE_SIZE))
    runs = db.session.scalars(
        select(AnalyticsRun).where(AnalyticsRun.analytics_project_id == project.id)
        .order_by(AnalyticsRun.completed_at.desc(), AnalyticsRun.started_at.desc())
        .limit(PAGE_SIZE).offset((current - 1) * PAGE_SIZE)
    ).all()
    return {
        "runs": runs,
        "runs_pagination": PlatformPage(tuple(runs), total, current),
        "run_count": total,
        "run_period_label": _run_period_label,
        "run_history_endpoint": "instagram.project_detail",
        "run_history_page_param": "page",
    }


def _import_template_context(project: AnalyticsProject, *, sync_period_form: dict | None = None) -> dict:
    account = project.instagram_account
    active_sync = active_meta_sync_run(project.id)
    form_values = sync_period_form if sync_period_form is not None else request.args
    history = {
        **_runs_history_context(project, request.args.get("history_page", "1")),
        "run_history_endpoint": "instagram.new_import",
        "run_history_page_param": "history_page",
    }
    return {
        "project": project,
        "import_fields": IMPORT_FIELDS,
        "meta_settings": _meta_settings(),
        "meta_connection": account.api_connection if account is not None else None,
        "meta_account": account,
        "active_meta_sync": active_sync,
        "meta_sync_completion_url": _meta_sync_completion_url(project, active_sync) if active_sync else "",
        "sync_period_options": SYNC_PERIOD_OPTIONS,
        "sync_period_form": _sync_period_form(form_values, active_run=active_sync),
        **history,
        "can_mutate": project.status == "active" and project.archived_at is None,
    }


def _meta_sync_completion_flash(project: AnalyticsProject) -> None:
    """Transforma o retorno assíncrono em feedback humano após reload seguro."""
    raw_run_id = request.args.get("meta_sync_completed", "")
    try:
        run_id = str(UUID(raw_run_id))
    except (TypeError, ValueError):
        return
    run = db.session.scalar(select(AnalyticsRun).where(
        AnalyticsRun.id == run_id,
        AnalyticsRun.analytics_project_id == project.id,
        AnalyticsRun.source_kind == "meta_api",
        AnalyticsRun.status == META_SYNC_COMPLETED,
    ))
    if run is None:
        return
    count = run.record_count
    content_text = "1 conteúdo foi processado" if count == 1 else f"{count} conteúdos foram processados"
    flash(f"Sincronização concluída. {content_text}.", "success")


def _run_period_label(run: AnalyticsRun) -> str:
    """Mostra a janela Meta congelada sem reinterpretar runs legados."""
    resolved = effective_run_period(run)
    if resolved.kind == "unknown" or run.source_kind != "meta_api":
        return "Não informado"
    parameters = run.parameters_json if isinstance(run.parameters_json, dict) else {}
    preset = parameters.get("period_preset")
    labels = dict(SYNC_PERIOD_OPTIONS)
    if not isinstance(preset, str) or preset not in labels:
        return "Não informado"
    label = labels[preset]
    if resolved.kind == "all_history":
        return label
    if not resolved.is_bounded or not resolved.timezone_name:
        return label
    try:
        start_text = to_local_datetime(resolved.start, resolved.timezone_name).strftime("%d/%m/%Y")
        end_text = to_local_datetime(resolved.end, resolved.timezone_name).strftime("%d/%m/%Y")
    except (TypeError, ValueError):
        return label
    return f"{label} · {start_text} a {end_text}"


def _persist_meta_candidate_or_stage_reconciliation(
    project: AnalyticsProject,
    *,
    candidate,
    provider: str,
    granted_scopes: tuple[str, ...] | list[str],
    settings: InstagramMetaSettings,
):
    """Persiste uma conexão validada ou guarda sua credencial no servidor até confirmação."""
    if reconciliation_required(project, candidate.instagram_profile):
        reconciliation_token = pending_meta_reconciliations.put(
            current_user.id,
            project,
            provider,
            candidate,
            granted_scopes,
        )
        return redirect(url_for(
            "instagram.reconcile_meta_account",
            project_id=project.id,
            reconciliation_token=reconciliation_token,
        ))
    persist_validated_connection(
        project,
        access_token=candidate.connection_access_token,
        profile=candidate.instagram_profile,
        granted_scopes=granted_scopes,
        settings=settings,
        provider=provider,
    )
    return None


def _restore_reconciliation_preconfirmation_state(project: AnalyticsProject, reconciliation) -> None:
    """Restaura apenas o estado capturado antes da confirmação pendente.

    A etapa de seleção não deve persistir uma conexão. A restauração é uma
    defesa adicional para que um cancelamento jamais deixe o status da conta
    divergente do estado que existia antes da tentativa. Se outra operação
    alterou a conta ou a conexão enquanto o nonce estava pendente, não
    sobrescrevemos esse estado concorrente.
    """
    account = project.instagram_account
    connection = account.api_connection if account is not None else None
    if (
        account is None
        or str(account.id) != reconciliation.account_id
        or account.external_id != reconciliation.account_external_id
        or (str(connection.id) if connection is not None else None) != reconciliation.connection_id
    ):
        raise InstagramMetaError(
            "reconciliation_expired",
            "O estado da conexão foi alterado antes da confirmação.",
        )
    account.connection_status = reconciliation.account_connection_status


def _store_meta_oauth_state(project: AnalyticsProject, provider: str) -> str:
    if provider not in META_CONNECTION_PROVIDERS:
        abort(400)
    state = secrets.token_urlsafe(32)
    session[_META_OAUTH_STATE_KEY] = {
        "state": state,
        "project_id": project.id,
        "user_id": current_user.id,
        "provider": provider,
        "issued_at": int(time.time()),
    }
    session.modified = True
    return state


def _consume_meta_oauth_state(state: str) -> dict:
    saved = session.pop(_META_OAUTH_STATE_KEY, None)
    session.modified = True
    if not isinstance(saved, dict) or not isinstance(state, str):
        abort(400)
    expected = saved.get("state")
    issued_at = saved.get("issued_at")
    if (
        not isinstance(expected, str)
        or not secrets.compare_digest(expected, state)
        or saved.get("user_id") != current_user.id
        or not isinstance(issued_at, int)
        or int(time.time()) - issued_at > _META_OAUTH_STATE_TTL_SECONDS
    ):
        abort(400)
    return saved


def _report_for_project(project: AnalyticsProject, report_id: UUID) -> InstagramReport:
    report = report_for_project(project, str(report_id))
    if report is None:
        abort(404)
    return report


def _report_date(value: object) -> str:
    """Apresenta datas ISO guardadas no snapshot sem recalcular o relatório."""
    if isinstance(value, str):
        try:
            if len(value) == 10:
                return datetime.strptime(value, "%Y-%m-%d").strftime("%d/%m/%Y")
        except ValueError:
            return "Não informado"
        try:
            value = datetime.fromisoformat(value)
        except ValueError:
            return "Não informado"
    return format_local_datetime(value, current_app.config["PRESENTATION_TIMEZONE"])


@instagram_bp.get("")
@login_required
def home():
    """Lista somente os projetos Instagram pertencentes ao usuário atual."""
    _require_tool()
    query = select(AnalyticsProject).where(
        AnalyticsProject.owner_user_id == current_user.id,
        AnalyticsProject.module_key == "instagram",
        AnalyticsProject.status == "active",
        AnalyticsProject.archived_at.is_(None),
    )
    total = int(db.session.scalar(select(func.count()).select_from(AnalyticsProject).where(
        AnalyticsProject.owner_user_id == current_user.id,
        AnalyticsProject.module_key == "instagram",
        AnalyticsProject.status == "active",
        AnalyticsProject.archived_at.is_(None),
    )) or 0)
    requested = requested_page(request.args.get("page", "1"))
    current = min(requested, max(1, (total + PAGE_SIZE - 1) // PAGE_SIZE))
    projects = db.session.scalars(
        query.order_by(AnalyticsProject.updated_at.desc()).limit(PAGE_SIZE).offset((current - 1) * PAGE_SIZE)
    ).all()
    return render_template(
        "platform/instagram/index.html", projects=projects,
        projects_pagination=PlatformPage(tuple(projects), total, current),
    )


@instagram_bp.get("/projects/archived")
@login_required
def archived_projects():
    _require_tool()
    query = select(AnalyticsProject).where(
        AnalyticsProject.owner_user_id == current_user.id,
        AnalyticsProject.module_key == "instagram",
        AnalyticsProject.status == "archived",
        AnalyticsProject.archived_at.is_not(None),
    )
    total = int(db.session.scalar(select(func.count()).select_from(AnalyticsProject).where(
        AnalyticsProject.owner_user_id == current_user.id,
        AnalyticsProject.module_key == "instagram",
        AnalyticsProject.status == "archived",
        AnalyticsProject.archived_at.is_not(None),
    )) or 0)
    requested = requested_page(request.args.get("page", "1"))
    current = min(requested, max(1, (total + PAGE_SIZE - 1) // PAGE_SIZE))
    projects = db.session.scalars(
        query.order_by(AnalyticsProject.archived_at.desc()).limit(PAGE_SIZE).offset((current - 1) * PAGE_SIZE)
    ).all()
    return render_template(
        "platform/instagram/archived_projects.html", projects=projects,
        projects_pagination=PlatformPage(tuple(projects), total, current),
    )


@instagram_bp.post("/projects/<uuid:project_id>/archive")
@login_required
def archive_project(project_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    if request.form.get("confirm") != "yes":
        abort(400)
    project.status = "archived"
    project.archived_at = datetime.now(timezone.utc)
    db.session.commit()
    flash("Projeto arquivado. O histórico foi preservado.", "success")
    return redirect(url_for("instagram.home", _anchor="instagram-project-list"))


@instagram_bp.post("/projects/<uuid:project_id>/restore")
@login_required
def restore_project(project_id: UUID):
    project = _project_for_current_user(project_id)
    if project.status != "archived" or project.archived_at is None:
        abort(404)
    if request.form.get("confirm") != "yes":
        abort(400)
    project.status = "active"
    project.archived_at = None
    db.session.commit()
    flash("Projeto restaurado.", "success")
    return redirect(url_for("instagram.archived_projects", _anchor="instagram-archived-project-list"))


@instagram_bp.route("/projects/<uuid:project_id>/delete", methods=("GET", "POST"))
@login_required
def delete_project(project_id: UUID):
    project = _project_for_current_user(project_id)
    if project.status != "archived" or project.archived_at is None:
        abort(400)
    if request.method == "POST":
        try:
            delete_instagram_project_permanently(
                project, current_user, request.form.get("confirmation", ""),
            )
        except InstagramProjectDeletionError as error:
            flash(str(error), "danger")
            return render_template(
                "platform/instagram/project_delete.html", project=project,
            ), 400
        except Exception:
            current_app.logger.exception("Falha ao excluir projeto Instagram permanentemente")
            db.session.rollback()
            flash("Não foi possível excluir o projeto permanentemente. Nenhum dado foi removido.", "danger")
            return render_template(
                "platform/instagram/project_delete.html", project=project,
            ), 500
        flash("Projeto excluído permanentemente.", "success")
        return redirect(url_for("instagram.archived_projects"))
    return render_template("platform/instagram/project_delete.html", project=project)


@instagram_bp.get("/projects/new")
@login_required
def new_project():
    _require_tool()
    return render_template("platform/instagram/project_new.html")


@instagram_bp.post("/projects")
@login_required
def create_project():
    _require_tool()
    name = request.form.get("name", "").strip()
    description = request.form.get("description", "").strip()
    if not name or len(name) > 200:
        flash("Informe o nome do projeto (até 200 caracteres).", "danger")
        return render_template("platform/instagram/project_new.html"), 400
    if len(description) > 4_000:
        flash("A descrição excede o limite de 4.000 caracteres.", "danger")
        return render_template("platform/instagram/project_new.html"), 400
    project = AnalyticsProject(
        owner_user_id=current_user.id,
        module_key="instagram",
        name=name,
        description=description,
    )
    db.session.add(project)
    db.session.commit()
    flash("Projeto Instagram criado.", "success")
    return redirect(url_for("instagram.project_detail", project_id=project.id))


@instagram_bp.get("/projects/<uuid:project_id>")
@login_required
def project_detail(project_id: UUID):
    project = _project_for_current_user(project_id)
    account = project.instagram_account
    history = _runs_history_context(project, request.args.get("page", "1"))
    return render_template(
        "platform/instagram/project_detail.html",
        project=project,
        account=account,
        **history,
        last_run=db.session.scalar(
            select(AnalyticsRun).where(AnalyticsRun.analytics_project_id == project.id)
            .order_by(AnalyticsRun.completed_at.desc(), AnalyticsRun.started_at.desc()).limit(1)
        ),
    )


@instagram_bp.get("/projects/<uuid:project_id>/dashboard")
@login_required
def dashboard(project_id: UUID):
    project = _project_for_current_user(project_id)
    if request.args.get("period", "30d") not in PERIOD_DAYS:
        abort(400)
    period_form = instagram_period_form(request.args, default="30d")
    period = period_form["key"]
    sort = request.args.get("sort", "recent")
    if period not in PERIOD_DAYS or sort not in TABLE_SORTS:
        abort(400)
    try:
        summary = build_instagram_dashboard(
            project, period, current_app.config["PRESENTATION_TIMEZONE"], sort=sort,
            taxonomy_dimension_id=request.args.get("taxonomy_dimension") or None,
            period_start=period_form["start"], period_end=period_form["end"],
            table_page=request.args.get("page", "1"),
        )
    except ValueError:
        abort(400)
    return render_template(
        "platform/instagram/dashboard.html", project=project, account=project.instagram_account,
        dashboard=summary,
        period_options=TEMPORAL_PERIOD_OPTIONS,
        period_form=period_form,
        sort_options=(("recent", "Mais recentes"), ("reach", "Maior alcance"),
                      ("interactions", "Mais interações"),
                      ("interaction_rate", "Maior taxa de interação por alcance")),
    )


@instagram_bp.get("/projects/<uuid:project_id>/insights")
@login_required
def insights(project_id: UUID):
    """Consulta a última execução explícita de achados para um recorte."""
    project = _project_for_current_user(project_id)
    if request.args.get("period", "30d") not in PERIOD_DAYS:
        abort(400)
    period_form = instagram_period_form(request.args, default="30d")
    period = period_form["key"]
    if period not in PERIOD_DAYS:
        abort(400)
    action_kind = request.args.get("action_kind", "all")
    action_status = request.args.get("action_status", "all")
    try:
        filtered_actions = actions_for_project(project, action_kind=action_kind, status=action_status)
    except InstagramActionError:
        abort(400)
    try:
        run = latest_insight_run(
            project, period,
            timezone_name=current_app.config["PRESENTATION_TIMEZONE"],
            period_start=period_form["start"], period_end=period_form["end"],
        )
    except InstagramInsightsError:
        abort(400)
    recommendation_run = recommendation_run_for_insight(project, run.id) if run else None
    raw_findings_by_group = grouped_findings(run.findings) if run else {}
    findings_by_group: dict[str, tuple] = {}
    finding_paginations: dict[str, PlatformPage] = {}
    for index, (group, items) in enumerate(raw_findings_by_group.items(), start=1):
        page = paginate_sequence(items, request.args.get(f"findings_page_{index}", "1"))
        findings_by_group[group] = page.items
        finding_paginations[group] = page
    recommendations_page = paginate_sequence(
        ordered_recommendations(recommendation_run), request.args.get("recommendations_page", "1"),
    )
    actions_page = paginate_sequence(filtered_actions, request.args.get("actions_page", "1"))
    learnings_page = paginate_sequence(learnings_for_project(project), request.args.get("learnings_page", "1"))
    eligible_learning_evaluations = comparable_evaluations(project)
    period_options = TEMPORAL_PERIOD_OPTIONS
    return render_template(
        "platform/instagram/insights.html",
        project=project,
        insight_run=run,
        findings_by_group=findings_by_group,
        finding_paginations=finding_paginations,
        recommendation_run=recommendation_run,
        recommendations=recommendations_page.items,
        recommendations_pagination=recommendations_page,
        period=period, period_options=period_options, period_form=period_form,
        period_label=instagram_period_label(period),
        actions=actions_page.items,
        actions_pagination=actions_page,
        action_counts=action_counts(actions_for_project(project)),
        action_kind=action_kind,
        action_status=action_status,
        action_kind_labels=ACTION_KIND_LABELS,
        action_status_labels=ACTION_STATUS_LABELS,
        learnings=learnings_page.items,
        learnings_pagination=learnings_page,
        eligible_learning_evaluations=eligible_learning_evaluations,
        learning_assessment_labels=LEARNING_ASSESSMENT_LABELS,
        learning_cohort_label=learning_cohort_label,
        metric_label=metric_label,
        can_generate=project.status == "active" and project.archived_at is None,
        can_mutate=project.status == "active" and project.archived_at is None,
    )


@instagram_bp.post("/projects/<uuid:project_id>/insights/generate")
@login_required
def generate_project_insights(project_id: UUID):
    """Cria um run analítico somente mediante ação protegida por CSRF."""
    project = _project_for_current_user(project_id, active=True)
    period = request.form.get("period", "30d")
    if period not in PERIOD_DAYS:
        abort(400)
    try:
        run = generate_insights(
            project,
            period,
            current_app.config["PRESENTATION_TIMEZONE"],
            period_start=request.form.get("period_start"),
            period_end=request.form.get("period_end"),
        )
    except InstagramInsightsError as error:
        flash(str(error), "danger")
    except InstagramInsightsPersistenceError:
        current_app.logger.exception("Falha ao gerar insights do Instagram")
        flash("Não foi possível gerar os insights. Nenhuma execução parcial foi gravada.", "danger")
    else:
        flash(f"Insights gerados: {run.finding_count} achado(s) descritivo(s).", "success")
    return_values = {"project_id": project.id, "period": period, "_anchor": "instagram-insights-content"}
    if period == "custom":
        return_values["period_start"] = request.form.get("period_start", "")
        return_values["period_end"] = request.form.get("period_end", "")
    return redirect(url_for("instagram.insights", **return_values))


@instagram_bp.post("/projects/<uuid:project_id>/insights/recommendations/generate")
@login_required
def generate_project_recommendations(project_id: UUID):
    """Gera recomendações somente para um InsightRun existente e pertencente ao projeto."""
    project = _project_for_current_user(project_id, active=True)
    insight_run_id = request.form.get("insight_run_id", "")
    insight_run = db.session.scalar(
        select(InstagramInsightRun).where(
            InstagramInsightRun.id == insight_run_id,
            InstagramInsightRun.analytics_project_id == project.id,
        )
    )
    if insight_run is None:
        abort(404)
    existing = recommendation_run_for_insight(project, insight_run.id)
    try:
        recommendation_run = generate_recommendations(project, insight_run)
    except InstagramRecommendationsError as error:
        flash(str(error), "danger")
    except InstagramRecommendationsPersistenceError:
        current_app.logger.exception("Falha ao gerar recomendações do Instagram")
        flash("Não foi possível gerar as recomendações. Nenhuma execução parcial foi gravada.", "danger")
    else:
        if existing is not None:
            flash("As recomendações desta versão já haviam sido geradas para este run de Insights.", "info")
        elif recommendation_run.number_of_recommendations_generated:
            flash(
                "Recomendações geradas: "
                f"{recommendation_run.number_of_recommendations_generated} a partir de "
                f"{recommendation_run.number_of_findings_considered} achado(s) considerado(s); "
                f"{recommendation_run.parameters_json.get('findings_ignored_by_rules', 0)} sem recomendação nesta versão.",
                "success",
            )
        else:
            flash(
                "Os achados deste período ainda não sustentam recomendações segundo os critérios desta versão.",
                "info",
            )
    return redirect(url_for(
        "instagram.insights", project_id=project.id, period=insight_run.period_key,
        _anchor="instagram-recommendations",
    ))


@instagram_bp.get("/projects/<uuid:project_id>/insights/<uuid:finding_id>")
@login_required
def insight_detail(project_id: UUID, finding_id: UUID):
    project = _project_for_current_user(project_id)
    finding = finding_for_project(project, str(finding_id))
    if finding is None:
        abort(404)
    return render_template(
        "platform/instagram/insight_detail.html",
        project=project,
        finding=finding,
        insight_run=finding.insight_run,
        evidence=evidence_views(project, finding),
    )


@instagram_bp.get("/projects/<uuid:project_id>/insights/recommendations/<uuid:recommendation_id>")
@login_required
def recommendation_detail(project_id: UUID, recommendation_id: UUID):
    project = _project_for_current_user(project_id)
    recommendation = recommendation_for_project(project, str(recommendation_id))
    if recommendation is None:
        abort(404)
    return render_template(
        "platform/instagram/recommendation_detail.html",
        project=project,
        recommendation=recommendation,
        recommendation_run=recommendation.recommendation_run,
        insight_run=recommendation.recommendation_run.insight_run,
        findings=linked_findings(recommendation),
        action_kind_labels=ACTION_KIND_LABELS,
        action_status_labels=ACTION_STATUS_LABELS,
        can_mutate=project.status == "active" and project.archived_at is None,
    )


def _action_form_values(prefill, submitted: bool = False) -> dict[str, object]:
    """Mantém valores submetidos na validação e usa o prefill determinístico no GET."""
    if submitted:
        return {
            "title": request.form.get("title", ""),
            "description": request.form.get("description", ""),
            "objective": request.form.get("objective", ""),
            "hypothesis": request.form.get("hypothesis", ""),
            "target_metric": request.form.get("target_metric", ""),
            "baseline_value": request.form.get("baseline_value", ""),
            "target_value": request.form.get("target_value", ""),
            "success_criterion": request.form.get("success_criterion", ""),
            "planned_start_at": request.form.get("planned_start_at", ""),
            "planned_end_at": request.form.get("planned_end_at", ""),
        }
    return {
        "title": prefill.title,
        "description": prefill.description,
        "objective": prefill.objective,
        "hypothesis": prefill.hypothesis,
        "target_metric": prefill.target_metric or "",
        "baseline_value": "" if prefill.baseline_value is None else prefill.baseline_value,
        "target_value": "",
        "success_criterion": "",
        "planned_start_at": "",
        "planned_end_at": "",
    }


@instagram_bp.route(
    "/projects/<uuid:project_id>/insights/recommendations/<uuid:recommendation_id>/actions/<action_kind>/new",
    methods=("GET", "POST"),
)
@login_required
def new_recommendation_action(project_id: UUID, recommendation_id: UUID, action_kind: str):
    project = _project_for_current_user(project_id, active=True)
    recommendation = recommendation_for_project(project, str(recommendation_id))
    if recommendation is None or action_kind not in ACTION_KIND_LABELS:
        abort(404)
    prefill = action_prefill(recommendation, action_kind)
    if request.method == "POST":
        values = _action_form_values(prefill, submitted=True)
        try:
            values["planned_start_at"] = parse_planned_datetime(
                str(values["planned_start_at"]), current_app.config["PRESENTATION_TIMEZONE"],
            )
            values["planned_end_at"] = parse_planned_datetime(
                str(values["planned_end_at"]), current_app.config["PRESENTATION_TIMEZONE"],
            )
            action = create_recommendation_action(
                project, recommendation, action_kind, values, actor_id=current_user.id,
            )
        except InstagramActionError as error:
            flash(str(error), "danger")
            return render_template(
                "platform/instagram/recommendation_action_form.html",
                project=project, recommendation=recommendation, prefill=prefill,
                form_values=_action_form_values(prefill, submitted=True),
                action_kind=action_kind, action_kind_label=ACTION_KIND_LABELS[action_kind],
            ), 400
        except InstagramActionPersistenceError:
            current_app.logger.exception("Falha ao criar ação/teste Instagram")
            flash("Não foi possível criar a ação ou teste. Nenhum registro parcial foi gravado.", "danger")
            return render_template(
                "platform/instagram/recommendation_action_form.html",
                project=project, recommendation=recommendation, prefill=prefill,
                form_values=_action_form_values(prefill, submitted=True),
                action_kind=action_kind, action_kind_label=ACTION_KIND_LABELS[action_kind],
            ), 500
        flash(f"{ACTION_KIND_LABELS[action_kind]} criada e vinculada à recomendação.", "success")
        return redirect(url_for(
            "instagram.recommendation_action_detail", project_id=project.id, action_id=action.id,
        ))
    return render_template(
        "platform/instagram/recommendation_action_form.html",
        project=project, recommendation=recommendation, prefill=prefill,
        form_values=_action_form_values(prefill), action_kind=action_kind,
        action_kind_label=ACTION_KIND_LABELS[action_kind],
    )


@instagram_bp.get("/projects/<uuid:project_id>/insights/actions/<uuid:action_id>")
@login_required
def recommendation_action_detail(project_id: UUID, action_id: UUID):
    project = _project_for_current_user(project_id)
    action = recommendation_action_for_project(project, str(action_id))
    if action is None:
        abort(404)
    evaluations = evaluations_for_action(project, action)
    evaluations_pagination = paginate_sequence(
        evaluations,
        page=requested_page(request.args.get("evaluations_page")),
    )
    return render_template(
        "platform/instagram/recommendation_action_detail.html",
        project=project, action=action, recommendation=action.recommendation,
        recommendation_run=action.recommendation.recommendation_run,
        insight_run=action.recommendation.recommendation_run.insight_run,
        findings=linked_findings(action.recommendation),
        action_kind_labels=ACTION_KIND_LABELS, action_status_labels=ACTION_STATUS_LABELS,
        action_outcome_labels=ACTION_OUTCOME_LABELS,
        evaluations=evaluations_pagination.items,
        evaluations_pagination=evaluations_pagination,
        eligible_followup_runs=eligible_followup_runs(project, action),
        assessment_labels=ASSESSMENT_LABELS,
        can_mutate=project.status == "active" and project.archived_at is None,
    )


@instagram_bp.get("/projects/<uuid:project_id>/insights/actions/<uuid:action_id>/evaluations/new")
@login_required
def new_recommendation_action_evaluation(project_id: UUID, action_id: UUID):
    """Escolha explícita do período e preview sem persistência."""
    project = _project_for_current_user(project_id, active=True)
    action = recommendation_action_for_project(project, str(action_id))
    if action is None:
        abort(404)
    if action.status != "completed":
        abort(400, description="Acompanhamento formal só está disponível após a conclusão.")
    runs = eligible_followup_runs(project, action)
    selected_run_id = request.args.get("analytics_run_id", "")
    preview = None
    if selected_run_id:
        try:
            preview = preview_action_evaluation(project, action, selected_run_id)
        except InstagramActionEvaluationError as error:
            abort(400, description=str(error))
    return render_template(
        "platform/instagram/recommendation_action_evaluation_form.html",
        project=project, action=action, runs=runs, selected_run_id=selected_run_id,
        preview=preview, assessment_labels=ASSESSMENT_LABELS,
    )


@instagram_bp.post("/projects/<uuid:project_id>/insights/actions/<uuid:action_id>/evaluations")
@login_required
def create_recommendation_action_evaluation(project_id: UUID, action_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    action = recommendation_action_for_project(project, str(action_id))
    if action is None:
        abort(404)
    run_id = request.form.get("analytics_run_id", "")
    try:
        evaluation = create_action_evaluation(
            project, action, run_id, request.form, actor_id=current_user.id,
        )
    except InstagramActionEvaluationDuplicateError as error:
        flash("Este período já possui acompanhamento para esta ação/teste.", "info")
        return redirect(url_for(
            "instagram.recommendation_action_evaluation_detail", project_id=project.id,
            action_id=action.id, evaluation_id=error.evaluation_id,
        ))
    except InstagramActionEvaluationError as error:
        flash(str(error), "danger")
        return redirect(url_for(
            "instagram.new_recommendation_action_evaluation", project_id=project.id,
            action_id=action.id, analytics_run_id=run_id,
        ))
    except InstagramActionEvaluationPersistenceError:
        current_app.logger.exception("Falha ao criar acompanhamento de ação/teste Instagram")
        flash("Não foi possível salvar o acompanhamento. Nenhum registro parcial foi gravado.", "danger")
        return redirect(url_for(
            "instagram.new_recommendation_action_evaluation", project_id=project.id,
            action_id=action.id, analytics_run_id=run_id,
        ))
    flash("Acompanhamento posterior registrado.", "success")
    return redirect(url_for(
        "instagram.recommendation_action_evaluation_detail", project_id=project.id,
        action_id=action.id, evaluation_id=evaluation.id,
    ))


def _evaluation_for_project_action(
    project: AnalyticsProject, action, evaluation_id: UUID,
) -> InstagramRecommendationActionEvaluation | None:
    return db.session.scalar(
        select(InstagramRecommendationActionEvaluation)
        .join(AnalyticsRun)
        .where(
            InstagramRecommendationActionEvaluation.id == str(evaluation_id),
            InstagramRecommendationActionEvaluation.instagram_recommendation_action_id == action.id,
            AnalyticsRun.analytics_project_id == project.id,
        )
        .options(
            joinedload(InstagramRecommendationActionEvaluation.analytics_run),
            joinedload(InstagramRecommendationActionEvaluation.action),
        )
    )


@instagram_bp.get("/projects/<uuid:project_id>/insights/actions/<uuid:action_id>/evaluations/<uuid:evaluation_id>")
@login_required
def recommendation_action_evaluation_detail(project_id: UUID, action_id: UUID, evaluation_id: UUID):
    project = _project_for_current_user(project_id)
    action = recommendation_action_for_project(project, str(action_id))
    if action is None:
        abort(404)
    evaluation = _evaluation_for_project_action(project, action, evaluation_id)
    if evaluation is None:
        abort(404)
    return render_template(
        "platform/instagram/recommendation_action_evaluation_detail.html",
        project=project, action=action, evaluation=evaluation,
        view=evaluation_view(evaluation), assessment_labels=ASSESSMENT_LABELS,
        can_mutate=project.status == "active" and project.archived_at is None,
    )


@instagram_bp.post("/projects/<uuid:project_id>/insights/actions/<uuid:action_id>/evaluations/<uuid:evaluation_id>")
@login_required
def update_recommendation_action_evaluation(project_id: UUID, action_id: UUID, evaluation_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    action = recommendation_action_for_project(project, str(action_id))
    if action is None:
        abort(404)
    evaluation = _evaluation_for_project_action(project, action, evaluation_id)
    if evaluation is None:
        abort(404)
    try:
        update_action_evaluation(project, evaluation, request.form)
    except InstagramActionEvaluationError as error:
        flash(str(error), "danger")
    except InstagramActionEvaluationPersistenceError:
        current_app.logger.exception("Falha ao atualizar aprendizado Instagram")
        flash("Não foi possível atualizar o aprendizado registrado.", "danger")
    else:
        flash("Aprendizado atualizado sem alterar o snapshot quantitativo.", "success")
    return redirect(url_for(
        "instagram.recommendation_action_evaluation_detail", project_id=project.id,
        action_id=action.id, evaluation_id=evaluation.id,
        _anchor="instagram-evaluation-learning",
    ))


@instagram_bp.route("/projects/<uuid:project_id>/insights/learnings/new", methods=("GET", "POST"))
@login_required
def new_learning(project_id: UUID):
    """Cria uma síntese humana a partir de snapshots comparáveis selecionados."""
    project = _project_for_current_user(project_id, active=True)
    period = request.values.get("period", "30d")
    if period not in PERIOD_DAYS:
        abort(400)
    evaluations = comparable_evaluations(project)
    selected_ids = request.form.getlist("evaluation_ids") if request.method == "POST" else request.args.getlist("evaluation_id")
    form_values = {
        "title": request.form.get("title", "") if request.method == "POST" else "",
        "learning_summary": request.form.get("learning_summary", "") if request.method == "POST" else "",
        "assessment": request.form.get("assessment", "inconclusive") if request.method == "POST" else "inconclusive",
    }
    if request.method == "POST":
        try:
            learning = create_learning(project, request.form, selected_ids, actor_id=current_user.id)
        except InstagramLearningError as error:
            flash(str(error), "danger")
            return render_template(
                "platform/instagram/learning_form.html", project=project, evaluations=evaluations,
                selected_ids=set(selected_ids), form_values=form_values,
                learning_assessment_labels=LEARNING_ASSESSMENT_LABELS, metric_label=metric_label, period=period,
            ), 400
        except InstagramLearningPersistenceError:
            current_app.logger.exception("Falha ao criar aprendizado consolidado Instagram")
            flash("Não foi possível salvar o aprendizado consolidado. Nenhum registro parcial foi gravado.", "danger")
            return render_template(
                "platform/instagram/learning_form.html", project=project, evaluations=evaluations,
                selected_ids=set(selected_ids), form_values=form_values,
                learning_assessment_labels=LEARNING_ASSESSMENT_LABELS, metric_label=metric_label, period=period,
            ), 500
        flash("Aprendizado consolidado registrado.", "success")
        return redirect(url_for(
            "instagram.learning_detail", project_id=project.id, learning_id=learning.id,
        ))
    return render_template(
        "platform/instagram/learning_form.html", project=project, evaluations=evaluations,
        selected_ids=set(selected_ids), form_values=form_values,
        learning_assessment_labels=LEARNING_ASSESSMENT_LABELS, metric_label=metric_label, period=period,
    )


@instagram_bp.get("/projects/<uuid:project_id>/insights/learnings/<uuid:learning_id>")
@login_required
def learning_detail(project_id: UUID, learning_id: UUID):
    project = _project_for_current_user(project_id)
    learning = learning_for_project(project, str(learning_id))
    if learning is None:
        abort(404)
    return render_template(
        "platform/instagram/learning_detail.html", project=project, learning=learning,
        learning_assessment_labels=LEARNING_ASSESSMENT_LABELS,
        learning_cohort_label=learning_cohort_label,
        metric_label=metric_label,
        can_mutate=project.status == "active" and project.archived_at is None,
    )


@instagram_bp.post("/projects/<uuid:project_id>/insights/learnings/<uuid:learning_id>")
@login_required
def update_learning_route(project_id: UUID, learning_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    learning = learning_for_project(project, str(learning_id))
    if learning is None:
        abort(404)
    try:
        update_learning(project, learning, request.form)
    except InstagramLearningError as error:
        flash(str(error), "danger")
    except InstagramLearningPersistenceError:
        current_app.logger.exception("Falha ao atualizar aprendizado consolidado Instagram")
        flash("Não foi possível atualizar o aprendizado consolidado.", "danger")
    else:
        flash("Aprendizado consolidado atualizado.", "success")
    return redirect(url_for(
        "instagram.learning_detail", project_id=project.id, learning_id=learning.id,
        _anchor="instagram-learning-summary",
    ))


@instagram_bp.route(
    "/projects/<uuid:project_id>/insights/learnings/<uuid:learning_id>/delete",
    methods=("GET", "POST"),
)
@login_required
def delete_learning_route(project_id: UUID, learning_id: UUID):
    """Reutiliza a confirmação permanente compartilhada para um Learning."""
    project = _project_for_current_user(project_id, active=True)
    learning = learning_for_project(project, str(learning_id))
    if learning is None:
        abort(404)
    period = request.values.get("period", "30d")
    if period not in PERIOD_DAYS:
        abort(400)
    if request.method == "GET":
        return render_template(
            "platform/instagram/learning_delete.html", project=project, learning=learning, period=period,
        )
    try:
        delete_learning(project, learning, request.form.get("confirmation", ""))
    except InstagramLearningError as error:
        flash(str(error), "danger")
        return render_template(
            "platform/instagram/learning_delete.html", project=project, learning=learning, period=period,
        ), 400
    except InstagramLearningPersistenceError:
        current_app.logger.exception("Falha ao excluir aprendizado consolidado Instagram")
        flash("Não foi possível excluir o aprendizado consolidado.", "danger")
        return render_template(
            "platform/instagram/learning_delete.html", project=project, learning=learning, period=period,
        ), 500
    flash("Aprendizado consolidado excluído com sucesso.", "success")
    return redirect(url_for(
        "instagram.insights", project_id=project.id, period=period, _anchor="instagram-learnings",
    ))


@instagram_bp.get("/projects/<uuid:project_id>/reports")
@login_required
def reports(project_id: UUID):
    """Histórico de relatórios imutáveis, inclusive em projetos arquivados."""
    project = _project_for_current_user(project_id)
    reports_page = reports_page_for_project(project, request.args.get("page", "1"))
    return render_template(
        "platform/instagram/reports.html",
        project=project,
        reports_page=reports_page,
        reports=reports_page.items,
        report_date=_report_date,
        can_mutate=project.status == "active" and project.archived_at is None,
    )


@instagram_bp.route("/projects/<uuid:project_id>/reports/new", methods=("GET", "POST"))
@login_required
def new_report(project_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    defaults = instagram_period_form({}, default="90d")
    if request.method == "POST":
        period_form = instagram_period_form(request.form, default="90d")
        # Compatibilidade com o formulário anterior, que enviava somente as
        # duas datas e não possuía o select de preset.
        if "period" not in request.form and (
            request.form.get("period_start") or request.form.get("period_end")
        ):
            period_form = {**period_form, "key": "custom"}
        try:
            requested_dates = resolve_calendar_dates(
                period_form["key"], current_app.config["PRESENTATION_TIMEZONE"],
                period_start=period_form["start"], period_end=period_form["end"],
            )
            values = {
                "period_preset": period_form["key"],
                **({
                    "period_start": requested_dates[0].isoformat(),
                    "period_end": requested_dates[1].isoformat(),
                } if requested_dates is not None else {}),
            }
            report = generate_report(project, values, actor_id=current_user.id)
        except (InstagramReportError, ValueError) as error:
            flash(str(error), "danger")
            return render_template(
                "platform/instagram/report_form.html", project=project,
                period_form=period_form, period_options=TEMPORAL_PERIOD_OPTIONS,
            ), 400
        except InstagramReportPersistenceError:
            current_app.logger.exception("Falha ao gerar relatório Instagram")
            flash("Não foi possível gerar o relatório. Nenhum relatório concluído foi criado.", "danger")
            return render_template(
                "platform/instagram/report_form.html", project=project,
                period_form=period_form, period_options=TEMPORAL_PERIOD_OPTIONS,
            ), 500
        flash("Relatório gerado com snapshot reprodutível.", "success")
        return redirect(url_for("instagram.report_detail", project_id=project.id, report_id=report.id))
    return render_template(
        "platform/instagram/report_form.html", project=project,
        period_form=defaults, period_options=TEMPORAL_PERIOD_OPTIONS,
    )


@instagram_bp.get("/projects/<uuid:project_id>/reports/<uuid:report_id>")
@login_required
def report_detail(project_id: UUID, report_id: UUID):
    project = _project_for_current_user(project_id)
    report = _report_for_project(project, report_id)
    return render_template(
        "platform/instagram/report_detail.html", project=project, report=report,
        snapshot=report.snapshot_json or {},
        report_charts=report_chart_data(report.snapshot_json or {}) if report.snapshot_json else {},
        report_summary_cards=report_summary_cards(report.snapshot_json or {}) if report.snapshot_json else [],
        report_date=_report_date,
        can_mutate=project.status == "active" and project.archived_at is None,
    )


@instagram_bp.get("/projects/<uuid:project_id>/reports/<uuid:report_id>/download/<string:format_key>")
@login_required
def download_report(project_id: UUID, report_id: UUID, format_key: str):
    project = _project_for_current_user(project_id)
    report = _report_for_project(project, report_id)
    if report.status != "completed" or format_key not in set(report.formats_json or []):
        abort(404)
    try:
        path = artifact_path(project, report, format_key)
    except InstagramReportError:
        abort(404)
    if not path.is_file():
        abort(404)
    mimetype = "application/pdf" if format_key == "pdf" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return send_file(
        path, as_attachment=True, download_name=report_download_name(report, format_key), mimetype=mimetype,
    )


@instagram_bp.route("/projects/<uuid:project_id>/reports/<uuid:report_id>/delete", methods=("GET", "POST"))
@login_required
def delete_report_route(project_id: UUID, report_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    report = _report_for_project(project, report_id)
    if request.method == "GET":
        return render_template("platform/instagram/report_delete.html", project=project, report=report)
    try:
        delete_report(project, report, request.form.get("confirmation", ""))
    except InstagramReportError as error:
        flash(str(error), "danger")
        return render_template("platform/instagram/report_delete.html", project=project, report=report), 400
    except InstagramReportPersistenceError:
        current_app.logger.exception("Falha ao excluir relatório Instagram")
        flash("Não foi possível excluir o relatório.", "danger")
        return render_template("platform/instagram/report_delete.html", project=project, report=report), 500
    flash("Relatório excluído permanentemente.", "success")
    return redirect(url_for("instagram.reports", project_id=project.id))


def _transition_action_route(project_id: UUID, action_id: UUID, status: str):
    project = _project_for_current_user(project_id, active=True)
    action = recommendation_action_for_project(project, str(action_id))
    if action is None:
        abort(404)
    try:
        transition_recommendation_action(action, status, request.form)
    except InstagramActionError as error:
        abort(400, description=str(error))
    except InstagramActionPersistenceError:
        current_app.logger.exception("Falha ao atualizar ação/teste Instagram")
        abort(500)
    flash(f"Status atualizado para {ACTION_STATUS_LABELS[status].casefold()}.", "success")
    return redirect(url_for(
        "instagram.recommendation_action_detail", project_id=project.id, action_id=action.id,
        _anchor="instagram-action-status",
    ))


@instagram_bp.post("/projects/<uuid:project_id>/insights/actions/<uuid:action_id>/start")
@login_required
def start_recommendation_action(project_id: UUID, action_id: UUID):
    return _transition_action_route(project_id, action_id, "in_progress")


@instagram_bp.post("/projects/<uuid:project_id>/insights/actions/<uuid:action_id>/complete")
@login_required
def complete_recommendation_action(project_id: UUID, action_id: UUID):
    return _transition_action_route(project_id, action_id, "completed")


@instagram_bp.post("/projects/<uuid:project_id>/insights/actions/<uuid:action_id>/cancel")
@login_required
def cancel_recommendation_action(project_id: UUID, action_id: UUID):
    return _transition_action_route(project_id, action_id, "cancelled")


@instagram_bp.route(
    "/projects/<uuid:project_id>/insights/actions/<uuid:action_id>/delete",
    methods=("GET", "POST"),
)
@login_required
def delete_recommendation_action_route(project_id: UUID, action_id: UUID):
    """Confirma e executa a exclusão definitiva de uma ação/teste elegível."""
    project = _project_for_current_user(project_id, active=True)
    action = recommendation_action_for_project(project, str(action_id))
    if action is None:
        abort(404)
    period = request.values.get("period", "30d")
    action_kind = request.values.get("action_kind", "all")
    action_status = request.values.get("action_status", "all")
    if (
        period not in PERIOD_DAYS
        or action_kind not in {*ACTION_KINDS, "all"}
        or action_status not in {*ACTION_STATUSES, "all"}
    ):
        abort(400)

    def insights_return_url() -> str:
        return url_for(
            "instagram.insights",
            project_id=project.id,
            period=period,
            action_kind=action_kind,
            action_status=action_status,
            _anchor="instagram-actions",
        )

    if recommendation_action_has_evaluation(action):
        flash("Esta ação/teste possui acompanhamento registrado e não pode ser excluída.", "danger")
        return redirect(insights_return_url())
    if request.method == "GET":
        return render_template(
            "platform/instagram/recommendation_action_delete.html",
            project=project,
            action=action,
            period=period,
            action_kind=action_kind,
            action_status=action_status,
        )
    try:
        delete_recommendation_action(project, action, request.form.get("confirmation", ""))
    except InstagramActionError as error:
        flash(str(error), "danger")
        return render_template(
            "platform/instagram/recommendation_action_delete.html",
            project=project,
            action=action,
            period=period,
            action_kind=action_kind,
            action_status=action_status,
        ), 400
    except InstagramActionPersistenceError:
        current_app.logger.exception("Falha ao excluir ação/teste Instagram")
        flash("Não foi possível excluir a ação ou teste. Nenhum registro relacionado foi removido.", "danger")
        return render_template(
            "platform/instagram/recommendation_action_delete.html",
            project=project,
            action=action,
            period=period,
            action_kind=action_kind,
            action_status=action_status,
        ), 500
    else:
        flash(f"{ACTION_KIND_LABELS[action.action_kind]} excluída com sucesso.", "success")
    return redirect(insights_return_url())


def _taxonomy_dimension_anchor(dimension_id: UUID | str) -> str:
    return f"instagram-taxonomy-dimension-{dimension_id}"


def _taxonomy_category_anchor(project: AnalyticsProject, category_id: UUID) -> str:
    dimension_id = db.session.scalar(
        select(InstagramTaxonomyCategory.instagram_taxonomy_dimension_id)
        .join(InstagramTaxonomyDimension)
        .where(
            InstagramTaxonomyCategory.id == str(category_id),
            InstagramTaxonomyDimension.analytics_project_id == project.id,
        )
    )
    return _taxonomy_dimension_anchor(dimension_id) if dimension_id else "instagram-taxonomy-dimensions"


def _taxonomy_error_redirect(
    project: AnalyticsProject, error: TaxonomyValidationError, *, anchor: str = "instagram-taxonomy-dimensions",
):
    flash(str(error), "danger")
    return redirect(url_for("instagram.taxonomies", project_id=project.id, _anchor=anchor))


def _classification_media(project: AnalyticsProject) -> tuple[InstagramMedia, ...]:
    return tuple(db.session.scalars(
        select(InstagramMedia)
        .join(InstagramAccount, InstagramAccount.id == InstagramMedia.instagram_account_id)
        .where(InstagramAccount.analytics_project_id == project.id)
        .order_by(InstagramMedia.published_at.desc(), InstagramMedia.id.desc())
    ).all())


def _classification_media_for_project(project: AnalyticsProject, media_id: UUID) -> InstagramMedia:
    media = db.session.scalar(
        select(InstagramMedia)
        .join(InstagramAccount, InstagramAccount.id == InstagramMedia.instagram_account_id)
        .where(
            InstagramMedia.id == str(media_id),
            InstagramAccount.analytics_project_id == project.id,
        )
    )
    if media is None:
        abort(404)
    return media


@instagram_bp.get("/projects/<uuid:project_id>/taxonomies")
@login_required
def taxonomies(project_id: UUID):
    project = _project_for_current_user(project_id)
    return render_template(
        "platform/instagram/taxonomies.html",
        project=project,
        dimensions=taxonomy_dimensions(project.id),
        can_mutate=project.status == "active" and project.archived_at is None,
        taxonomy_version=DEFAULT_TAXONOMY_VERSION,
    )


@instagram_bp.post("/projects/<uuid:project_id>/taxonomies/initial")
@login_required
def initialize_taxonomy(project_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    create_initial_taxonomy(project)
    db.session.commit()
    flash("Estrutura inicial de taxonomias criada.", "success")
    return redirect(url_for("instagram.taxonomies", project_id=project.id, _anchor="instagram-taxonomy-dimensions"))


@instagram_bp.post("/projects/<uuid:project_id>/taxonomies/dimensions")
@login_required
def add_taxonomy_dimension(project_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    try:
        create_dimension(
            project,
            name=request.form.get("name", ""),
            key=request.form.get("key") or None,
            description=request.form.get("description") or None,
            selection_mode=request.form.get("selection_mode", ""),
            position=request.form.get("position"),
        )
        db.session.commit()
    except TaxonomyValidationError as error:
        db.session.rollback()
        return _taxonomy_error_redirect(project, error, anchor="instagram-new-dimension")
    flash("Dimensão criada.", "success")
    return redirect(url_for("instagram.taxonomies", project_id=project.id, _anchor="instagram-taxonomy-dimensions"))


@instagram_bp.post("/projects/<uuid:project_id>/taxonomies/dimensions/<uuid:dimension_id>/categories")
@login_required
def add_taxonomy_category(project_id: UUID, dimension_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    try:
        create_category(
            project, str(dimension_id),
            name=request.form.get("name", ""),
            key=request.form.get("key") or None,
            description=request.form.get("description") or None,
            position=request.form.get("position"),
        )
        db.session.commit()
    except TaxonomyValidationError as error:
        db.session.rollback()
        return _taxonomy_error_redirect(project, error, anchor=_taxonomy_dimension_anchor(dimension_id))
    flash("Categoria criada.", "success")
    return redirect(url_for(
        "instagram.taxonomies", project_id=project.id, _anchor=_taxonomy_dimension_anchor(dimension_id),
    ))


def _active_form_value() -> bool:
    value = request.form.get("active")
    if value == "true":
        return True
    if value == "false":
        return False
    abort(400)


@instagram_bp.post("/projects/<uuid:project_id>/taxonomies/dimensions/<uuid:dimension_id>/state")
@login_required
def set_taxonomy_dimension_state(project_id: UUID, dimension_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    try:
        set_dimension_active(project, str(dimension_id), active=_active_form_value())
        db.session.commit()
    except TaxonomyValidationError as error:
        db.session.rollback()
        return _taxonomy_error_redirect(project, error, anchor=_taxonomy_dimension_anchor(dimension_id))
    flash("Estado da dimensão atualizado.", "success")
    return redirect(url_for(
        "instagram.taxonomies", project_id=project.id, _anchor=_taxonomy_dimension_anchor(dimension_id),
    ))


@instagram_bp.post("/projects/<uuid:project_id>/taxonomies/categories/<uuid:category_id>/state")
@login_required
def set_taxonomy_category_state(project_id: UUID, category_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    anchor = _taxonomy_category_anchor(project, category_id)
    try:
        set_category_active(project, str(category_id), active=_active_form_value())
        db.session.commit()
    except TaxonomyValidationError as error:
        db.session.rollback()
        return _taxonomy_error_redirect(project, error, anchor=anchor)
    flash("Estado da categoria atualizado.", "success")
    return redirect(url_for("instagram.taxonomies", project_id=project.id, _anchor=anchor))


@instagram_bp.get("/projects/<uuid:project_id>/classifications")
@login_required
def classifications(project_id: UUID):
    project = _project_for_current_user(project_id)
    status_filter = request.args.get("status", "all")
    if status_filter not in CLASSIFICATION_STATUSES:
        abort(400)
    media = _classification_media(project)
    states = classification_states(project, tuple(item.id for item in media))
    if status_filter != "all":
        media = tuple(item for item in media if states[item.id].status == status_filter)
    media_pagination = paginate_sequence(media, request.args.get("page", "1"))
    return render_template(
        "platform/instagram/classifications.html",
        project=project,
        media=media_pagination.items,
        media_pagination=media_pagination,
        states=states,
        status_filter=status_filter,
        can_mutate=project.status == "active" and project.archived_at is None,
    )


@instagram_bp.route("/projects/<uuid:project_id>/classifications/<uuid:media_id>", methods=("GET", "POST"))
@login_required
def edit_classification(project_id: UUID, media_id: UUID):
    project = _project_for_current_user(project_id, active=request.method == "POST")
    media = _classification_media_for_project(project, media_id)
    dimensions = taxonomy_dimensions(project.id, active_only=True)
    if request.method == "POST":
        selections = {
            dimension.id: request.form.getlist(f"dimension_{dimension.id}")
            for dimension in dimensions
        }
        try:
            classify_media(project, media.id, selections, current_user)
            db.session.commit()
        except TaxonomyValidationError as error:
            db.session.rollback()
            flash(str(error), "danger")
        else:
            flash("Classificação atualizada. O histórico anterior foi preservado.", "success")
            return redirect(url_for(
                "instagram.classifications",
                project_id=project.id,
                _anchor=f"instagram-classification-{media.id}",
            ))
    current = current_classifications(project.id, media_ids=(media.id,))
    current_category_ids = {item.instagram_taxonomy_category_id for item in current}
    return render_template(
        "platform/instagram/classification_edit.html",
        project=project,
        media=media,
        dimensions=dimensions,
        current=current,
        current_category_ids=current_category_ids,
        can_mutate=project.status == "active" and project.archived_at is None,
    )


def _comment_for_project(project: AnalyticsProject, comment_id: UUID) -> InstagramComment:
    comment = db.session.scalar(
        select(InstagramComment)
        .join(InstagramMedia, InstagramMedia.id == InstagramComment.instagram_media_id)
        .join(InstagramAccount, InstagramAccount.id == InstagramMedia.instagram_account_id)
        .where(
            InstagramComment.id == str(comment_id),
            InstagramAccount.analytics_project_id == project.id,
        )
    )
    if comment is None:
        abort(404)
    # As relações são consumidas por funções centrais de estado e pelo
    # histórico no template. Carregá-las explicitamente evita queries por linha.
    _ = comment.instagram_media, comment.text_versions, comment.observations, comment.assessments, comment.replies
    return comment


@instagram_bp.get("/projects/<uuid:project_id>/comments")
@login_required
def comments(project_id: UUID):
    project = _project_for_current_user(project_id)
    if request.args.get("period", "30d") not in PERIOD_DAYS:
        abort(400)
    period_form = instagram_period_form(request.args, default="30d")
    filters = {
        "period": period_form["key"],
        "sort": request.args.get("sort", "recent"),
        "status": request.args.get("status", "all"),
        "sentiment": request.args.get("sentiment", "all"),
        "needs_response": request.args.get("needs_response", "all"),
        "risk": request.args.get("risk", "all"),
    }
    try:
        summary = build_instagram_comments_dashboard(
            project,
            **filters,
            timezone_name=current_app.config["PRESENTATION_TIMEZONE"],
            period_start=period_form["start"],
            period_end=period_form["end"],
            page=request.args.get("page", "1"),
        )
    except ValueError:
        abort(400)
    return render_template(
        "platform/instagram/comments.html",
        project=project,
        comments_dashboard=summary,
        filters=filters,
        period_options=TEMPORAL_PERIOD_OPTIONS,
        period_form=period_form,
        sentiment_labels=SENTIMENT_LABELS,
        intent_labels=INTENT_LABELS,
        risk_labels=RISK_LABELS,
        can_mutate=project.status == "active" and project.archived_at is None,
    )


@instagram_bp.get("/projects/<uuid:project_id>/comments/classify")
@login_required
def preview_comment_classification(project_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    classifier = get_comment_classifier()
    try:
        summary = plan_comment_classification(project, classifier)
    except CommentClassificationError as error:
        flash(str(error), "danger")
        return redirect(url_for("instagram.comments", project_id=project.id, _anchor="instagram-comments-content"))
    return render_template(
        "platform/instagram/comment_classification_preview.html",
        project=project,
        classifier=classifier,
        classifier_label=LOCAL_COMMENT_CLASSIFIER_LABEL,
        summary=summary,
    )


@instagram_bp.post("/projects/<uuid:project_id>/comments/classify")
@login_required
def confirm_comment_classification(project_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    reprocess = request.form.get("reprocess") == "yes"
    try:
        summary = classify_project_comments(project, reprocess=reprocess)
    except CommentClassificationError as error:
        flash(str(error), "danger")
    except CommentClassificationPersistenceError:
        current_app.logger.exception("Falha ao persistir classificação automática de comentários")
        flash("Não foi possível concluir a classificação automática. Nenhuma avaliação foi gravada.", "danger")
    else:
        if summary.classified:
            flash(
                "Classificação automática concluída: "
                f"{summary.eligible} elegível(is), {summary.classified} classificado(s), "
                f"{summary.skipped} ignorado(s), {summary.already_classified} já classificado(s), "
                f"{summary.human_reviewed} com revisão humana e {summary.errors} erro(s).",
                "success",
            )
        elif summary.already_classified:
            flash("Nenhum comentário novo para classificar; os elegíveis já possuem esta versão do classificador.", "info")
        elif summary.skipped or summary.errors:
            flash(
                "Nenhuma avaliação foi criada: "
                f"{summary.skipped} comentário(s) ignorado(s) e {summary.errors} erro(s).",
                "info",
            )
        else:
            flash("Nenhum comentário disponível para classificação automática.", "info")
        for message in summary.error_messages[:3]:
            flash(message, "warning")
    return redirect(url_for("instagram.comments", project_id=project.id, _anchor="instagram-comments-content"))


@instagram_bp.route("/projects/<uuid:project_id>/comments/<uuid:comment_id>", methods=("GET", "POST"))
@login_required
def comment_detail(project_id: UUID, comment_id: UUID):
    project = _project_for_current_user(project_id, active=request.method == "POST")
    comment = _comment_for_project(project, comment_id)
    if request.method == "POST":
        try:
            classify_comment(comment, current_user.id, {
                "sentiment": request.form.get("sentiment"),
                "intensity": request.form.get("intensity"),
                "intent": request.form.get("intent"),
                "topic": request.form.get("topic"),
                "topic_normalized": request.form.get("topic_normalized"),
                "needs_response": request.form.get("needs_response"),
                "risk": request.form.get("risk"),
                "rationale": request.form.get("rationale"),
            })
            db.session.commit()
        except CommentAssessmentValidationError as error:
            db.session.rollback()
            flash(str(error), "danger")
        else:
            flash("Classificação do comentário atualizada. O histórico anterior foi preservado.", "success")
            return redirect(url_for(
                "instagram.comment_detail", project_id=project.id, comment_id=comment.id,
                _anchor="instagram-comment-classification",
            ))
    assessment = effective_assessment(comment.assessments)
    return render_template(
        "platform/instagram/comment_detail.html",
        project=project,
        comment=comment,
        media=comment.instagram_media,
        text_version=current_text_version(comment.text_versions),
        observation=current_comment_observation(comment.observations),
        assessment=assessment,
        history=sorted(comment.assessments, key=lambda item: (item.created_at, item.id), reverse=True),
        text_versions=sorted(comment.text_versions, key=lambda item: (item.first_observed_at, item.id), reverse=True),
        sentiment_labels=SENTIMENT_LABELS,
        intensity_labels=INTENSITY_LABELS,
        intent_labels=INTENT_LABELS,
        risk_labels=RISK_LABELS,
        assessment_source_labels=ASSESSMENT_SOURCE_LABELS,
        classifier_labels=COMMENT_CLASSIFIER_LABELS,
        can_mutate=project.status == "active" and project.archived_at is None,
    )


@instagram_bp.get("/projects/<uuid:project_id>/comments/imports/new")
@login_required
def new_comment_import(project_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    return render_template(
        "platform/instagram/comment_import_new.html", project=project, import_fields=COMMENT_IMPORT_FIELDS,
    )


@instagram_bp.get("/projects/<uuid:project_id>/comments/imports/modelo")
@login_required
def comment_import_template(project_id: UUID):
    _project_for_current_user(project_id, active=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Comentários Instagram"
    for index, field in enumerate(COMMENT_IMPORT_FIELDS, start=1):
        cell = sheet.cell(1, index, field.name)
        cell.fill = PatternFill("solid", fgColor="184775" if field.required else "52677A")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.comment = Comment(field.note, "Análysis")
        sheet.column_dimensions[get_column_letter(index)].width = min(34, max(19, len(field.name) + 3))
        if field.kind == "datetime":
            for row in range(2, MAX_IMPORT_ROWS + 2):
                sheet.cell(row, index).number_format = "@"
    sheet.row_dimensions[1].height = 31
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COMMENT_IMPORT_FIELDS))}1"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return send_file(
        output, as_attachment=True, download_name="modelo_importacao_comentarios_instagram_analysis.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _comment_known_hashes(project: AnalyticsProject) -> dict[tuple[str, str], set[str]]:
    account = project.instagram_account
    if account is None:
        return {}
    comments = db.session.execute(
        select(InstagramComment, InstagramMedia.external_id)
        .join(InstagramMedia, InstagramMedia.id == InstagramComment.instagram_media_id)
        .where(InstagramMedia.instagram_account_id == account.id)
    ).all()
    hashes = {(media_external_id, comment.external_id): set() for comment, media_external_id in comments}
    key_by_comment_id = {
        comment.id: (media_external_id, comment.external_id)
        for comment, media_external_id in comments
    }
    if comments:
        ids = list(key_by_comment_id)
        for version in db.session.scalars(select(InstagramCommentTextVersion).where(
            InstagramCommentTextVersion.instagram_comment_id.in_(ids)
        )):
            hashes[key_by_comment_id[version.instagram_comment_id]].add(version.content_hash)
    return hashes


@instagram_bp.post("/projects/<uuid:project_id>/comments/imports/preview")
@login_required
def preview_comment_import(project_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    uploaded = request.files.get("dataset")
    if uploaded is None or not uploaded.filename:
        flash("Selecione um arquivo CSV ou XLSX.", "danger")
        return render_template(
            "platform/instagram/comment_import_new.html", project=project, import_fields=COMMENT_IMPORT_FIELDS,
        ), 400
    data = uploaded.read(MAX_IMPORT_BYTES + 1)
    try:
        dataset = parse_instagram_comment_import(uploaded.filename, data)
        conflict = comment_import_conflict(project, dataset)
        if conflict:
            raise InstagramImportError(conflict)
        repeated_source_hash = db.session.scalar(
            select(AnalyticsRun.id).where(
                AnalyticsRun.analytics_project_id == project.id,
                AnalyticsRun.source_hash == dataset.source_hash,
            ).limit(1)
        ) is not None
        preview = build_comment_preview(
            dataset, known_comment_hashes=_comment_known_hashes(project),
            repeated_source_hash=repeated_source_hash,
        )
    except InstagramImportError as error:
        flash(str(error), "danger")
        return render_template(
            "platform/instagram/comment_import_new.html", project=project, import_fields=COMMENT_IMPORT_FIELDS,
        ), 400
    try:
        token = comment_previews.put(current_user.id, project.id, preview)
    except OSError:
        current_app.logger.exception("Não foi possível guardar a prévia de comentários Instagram")
        flash("Não foi possível preparar a prévia. Tente enviar o arquivo novamente.", "danger")
        return render_template(
            "platform/instagram/comment_import_new.html", project=project, import_fields=COMMENT_IMPORT_FIELDS,
        ), 500
    return render_template("platform/instagram/comment_import_preview.html", project=project, preview=preview, token=token)


@instagram_bp.post("/projects/<uuid:project_id>/comments/imports/confirm")
@login_required
def confirm_comment_import(project_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    token = request.form.get("preview_token", "")
    try:
        claim = comment_previews.claim(token, current_user.id, project.id)
    except InstagramCommentPreviewExpired:
        flash("A prévia expirou. Envie o arquivo novamente.", "warning")
        return redirect(url_for("instagram.new_comment_import", project_id=project.id))
    if claim is None:
        flash("Prévia indisponível. Envie o arquivo novamente.", "warning")
        return redirect(url_for("instagram.new_comment_import", project_id=project.id))
    preview = claim.preview
    try:
        result = persist_instagram_comment_import(project, preview.dataset)
    except InstagramCommentImportPersistenceError as error:
        claim.release()
        flash(str(error), "danger")
        return render_template(
            "platform/instagram/comment_import_preview.html", project=project, preview=preview, token=token,
        ), 409
    except Exception:
        claim.release()
        current_app.logger.exception("Falha ao confirmar importação de comentários Instagram")
        flash("Não foi possível concluir a importação. Nenhum dado foi salvo.", "danger")
        return render_template(
            "platform/instagram/comment_import_preview.html", project=project, preview=preview, token=token,
        ), 500
    try:
        claim.complete()
    except OSError:
        current_app.logger.exception("Importação de comentários concluída; falha ao remover prévia consumida")
    return render_template("platform/instagram/comment_import_complete.html", project=project, result=result, preview=preview)


@instagram_bp.post("/projects/<uuid:project_id>/comments/imports/cancel")
@login_required
def cancel_comment_import(project_id: UUID):
    project = _project_for_current_user(project_id)
    comment_previews.discard(request.form.get("preview_token", ""), current_user.id, project.id)
    return redirect(url_for("instagram.comments", project_id=project.id))


@instagram_bp.get("/projects/<uuid:project_id>/imports/new")
@login_required
def new_import(project_id: UUID):
    # Arquivados são somente leitura: exibem conexão e diagnóstico, sem POSTs.
    project = _project_for_current_user(project_id)
    _meta_sync_completion_flash(project)
    return render_template("platform/instagram/import_new.html", **_import_template_context(project))


@instagram_bp.post("/projects/<uuid:project_id>/imports/meta/connect", defaults={"provider": "facebook_login"})
@instagram_bp.post("/projects/<uuid:project_id>/imports/meta/connect/<provider>")
@login_required
def connect_meta_account(project_id: UUID, provider: str):
    """Inicia somente o OAuth escolhido; o state associa usuário, projeto e provider."""
    project = _project_for_current_user(project_id, active=True)
    if provider not in META_CONNECTION_PROVIDERS:
        abort(404)
    settings = _meta_settings()
    if not settings.provider_available(provider):
        label = "Instagram direto" if provider == "instagram_login" else "Facebook"
        flash(f"A integração {label} não está configurada neste ambiente.", "warning")
        return redirect(url_for("instagram.new_import", project_id=project.id))
    state = _store_meta_oauth_state(project, provider)
    try:
        authorization_url = _meta_client(settings, provider).authorization_url(state)
    except InstagramMetaError as error:
        log_meta_operation(project.id, "connect_start", error.category, http_status=error.http_status)
        flash(safe_meta_error_message(error), "danger")
        return redirect(url_for("instagram.new_import", project_id=project.id))
    log_meta_operation(project.id, "connect_start", "redirected")
    return redirect(authorization_url)


@instagram_bp.get("/imports/meta/callback")
@login_required
def meta_callback():
    """Consome state e deixa o provider resolver a credencial de conexão."""
    saved = _consume_meta_oauth_state(request.args.get("state", ""))
    try:
        project = _project_for_current_user(UUID(str(saved["project_id"])), active=True)
    except (KeyError, ValueError):
        abort(400)
    provider = saved.get("provider")
    if provider not in META_CONNECTION_PROVIDERS:
        abort(400)
    if request.args.get("error"):
        log_meta_operation(project.id, "callback", "cancelled")
        flash("A conexão com o Instagram foi cancelada ou não foi autorizada.", "warning")
        return redirect(url_for("instagram.new_import", project_id=project.id))
    code = request.args.get("code", "")
    if not code:
        abort(400)
    settings = _meta_settings()
    if not settings.provider_available(provider):
        flash("O método de conexão escolhido não está configurado neste ambiente.", "warning")
        return redirect(url_for("instagram.new_import", project_id=project.id))
    started_at = time.monotonic()
    try:
        client = _meta_client(settings, provider)
        token_payload = client.exchange_code(code)
        candidates = client.resolve_connection_candidates(token_payload["access_token"])
        candidate = choose_candidate(project, candidates)
        if candidate is None:
            selection_token = pending_meta_selections.put(
                current_user.id, project.id, provider, candidates, list(provider_scopes(provider))
            )
            return redirect(url_for(
                "instagram.select_meta_account", project_id=project.id, selection_token=selection_token
            ))
        reconciliation_redirect = _persist_meta_candidate_or_stage_reconciliation(
            project,
            candidate=candidate,
            provider=provider,
            granted_scopes=provider_scopes(provider),
            settings=settings,
        )
        if reconciliation_redirect is not None:
            log_meta_operation(project.id, "reconciliation", "pending")
            return reconciliation_redirect
        db.session.commit()
    except InstagramMetaError as error:
        db.session.rollback()
        log_meta_operation(
            project.id, "callback", error.category, http_status=error.http_status,
            duration_ms=round((time.monotonic() - started_at) * 1000),
        )
        flash(safe_meta_error_message(error), "danger")
        return redirect(url_for("instagram.new_import", project_id=project.id))
    except Exception:
        db.session.rollback()
        current_app.logger.error("instagram_meta project=%s operation=callback status=unexpected_error", project.id)
        flash("Não foi possível concluir a conexão com o Instagram. Nenhuma credencial foi substituída.", "danger")
        return redirect(url_for("instagram.new_import", project_id=project.id))
    log_meta_operation(
        project.id, "callback", "connected",
        duration_ms=round((time.monotonic() - started_at) * 1000),
    )
    flash("Conta profissional do Instagram conectada. Nenhum dado foi importado automaticamente.", "success")
    return redirect(url_for("instagram.new_import", project_id=project.id))


@instagram_bp.route("/projects/<uuid:project_id>/imports/meta/select", methods=("GET", "POST"))
@login_required
def select_meta_account(project_id: UUID):
    """Seleciona explicitamente uma conta quando a descoberta é ambígua."""
    project = _project_for_current_user(project_id, active=True)
    selection_token = request.values.get("selection_token", "")
    try:
        selection = pending_meta_selections.get(selection_token, current_user.id, project.id)
        if request.method == "POST":
            index = int(request.form.get("candidate_index", ""))
            candidate, granted_scopes = pending_meta_selections.claim(
                selection_token, current_user.id, project.id, index
            )
            reconciliation_redirect = _persist_meta_candidate_or_stage_reconciliation(
                project,
                candidate=candidate,
                provider=selection.provider,
                granted_scopes=granted_scopes,
                settings=_meta_settings(),
            )
            if reconciliation_redirect is not None:
                log_meta_operation(project.id, "reconciliation", "pending")
                return reconciliation_redirect
            db.session.commit()
            log_meta_operation(project.id, "selection", "connected")
            flash("Conta profissional do Instagram conectada. Nenhum dado foi importado automaticamente.", "success")
            return redirect(url_for("instagram.new_import", project_id=project.id))
    except InstagramMetaError as error:
        db.session.rollback()
        if error.category == "unknown":
            flash("A seleção da conta expirou ou não é válida. Inicie a conexão novamente.", "warning")
        else:
            flash(safe_meta_error_message(error), "danger")
        return redirect(url_for("instagram.new_import", project_id=project.id))
    except ValueError:
        db.session.rollback()
        flash("A seleção da conta expirou ou não é válida. Inicie a conexão novamente.", "warning")
        return redirect(url_for("instagram.new_import", project_id=project.id))
    return render_template(
        "platform/instagram/meta_select_account.html",
        project=project, selection_token=selection_token, candidates=selection.candidates,
    )


@instagram_bp.route("/projects/<uuid:project_id>/imports/meta/reconcile", methods=("GET", "POST"))
@login_required
def reconcile_meta_account(project_id: UUID):
    """Exige confirmação explícita antes de associar uma identidade Meta divergente."""
    project = _project_for_current_user(project_id, active=True)
    reconciliation_token = request.values.get("reconciliation_token", "")
    try:
        reconciliation = pending_meta_reconciliations.get(
            reconciliation_token, current_user.id, project.id
        )
        if request.method == "POST":
            action = request.form.get("reconciliation_action", "")
            if action not in {"confirm", "cancel"}:
                raise InstagramMetaError("reconciliation_expired", "A confirmação do vínculo não é válida.")
            reconciliation = pending_meta_reconciliations.claim(
                reconciliation_token, current_user.id, project.id
            )
            if action == "cancel":
                _restore_reconciliation_preconfirmation_state(project, reconciliation)
                db.session.commit()
                log_meta_operation(project.id, "reconciliation", "cancelled")
                flash("A vinculação da conta foi cancelada. Nenhuma alteração foi feita.", "warning")
                return redirect(url_for("instagram.new_import", project_id=project.id))
            persist_validated_connection(
                project,
                access_token=reconciliation.candidate.connection_access_token,
                profile=reconciliation.candidate.instagram_profile,
                granted_scopes=reconciliation.granted_scopes,
                settings=_meta_settings(),
                provider=reconciliation.provider,
                reconciliation_confirmed=True,
            )
            db.session.commit()
            log_meta_operation(project.id, "reconciliation", "connected")
            flash("Conta profissional do Instagram vinculada ao projeto. Os dados existentes foram preservados.", "success")
            return redirect(url_for("instagram.new_import", project_id=project.id))
    except InstagramMetaError as error:
        db.session.rollback()
        log_meta_operation(project.id, "reconciliation", error.category)
        flash(safe_meta_error_message(error), "danger")
        return redirect(url_for("instagram.new_import", project_id=project.id))
    return render_template(
        "platform/instagram/meta_reconcile_account.html",
        project=project,
        reconciliation_token=reconciliation_token,
        provider=reconciliation.provider,
        provider_profile=reconciliation.candidate.instagram_profile,
    )


@instagram_bp.post("/projects/<uuid:project_id>/imports/meta/capabilities")
@login_required
def test_meta_capabilities(project_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    account = project.instagram_account
    connection = account.api_connection if account is not None else None
    if connection is None or connection.status != "connected":
        abort(404)
    settings = _meta_settings()
    started_at = time.monotonic()
    try:
        run_capability_test(connection, settings, _meta_client(settings, connection.provider))
        db.session.commit()
    except InstagramMetaError as error:
        db.session.commit()
        log_meta_operation(
            project.id, "capability_test", error.category, http_status=error.http_status,
            duration_ms=round((time.monotonic() - started_at) * 1000),
        )
        flash(safe_meta_error_message(error), "danger")
    except Exception:
        db.session.rollback()
        current_app.logger.error("instagram_meta project=%s operation=capability_test status=unexpected_error", project.id)
        flash("Não foi possível concluir a verificação dos dados disponíveis. Nenhum dado analítico foi importado.", "danger")
    else:
        log_meta_operation(
            project.id, "capability_test", "completed",
            duration_ms=round((time.monotonic() - started_at) * 1000),
        )
        flash("Verificação dos dados disponíveis concluída. Nenhum dado analítico foi importado.", "success")
    return redirect(url_for("instagram.new_import", project_id=project.id))


@instagram_bp.post("/projects/<uuid:project_id>/imports/meta/sync")
@login_required
def sync_meta_data(project_id: UUID):
    """Agenda sincronização durável; a coleta não ocupa a requisição web."""
    project = _project_for_current_user(project_id, active=True)
    account = project.instagram_account
    connection = account.api_connection if account is not None else None
    if connection is None or connection.status != "connected":
        abort(404)
    stale_seconds = current_app.config.get("PESQUISAPDF_META_SYNC_STALE_SECONDS", DEFAULT_META_SYNC_STALE_SECONDS)
    try:
        recover_stale_meta_sync_runs(stale_seconds=int(stale_seconds))
        # Uma execução já ativa conserva a janela congelada no próprio run. A
        # seleção enviada no reload não pode alterar o trabalho em andamento.
        active_run = active_meta_sync_run(project.id)
        if active_run is not None:
            run, created = active_run, False
        else:
            publication_window = resolve_publication_window(
                request.form.get("sync_period", "last_90_days"),
                request.form.get("published_from"),
                request.form.get("published_to"),
                current_app.config["PRESENTATION_TIMEZONE"],
            )
            run, created = enqueue_meta_sync_run(
                project, connection.id, connection.provider, publication_window,
            )
    except InstagramMetaSyncError as error:
        if request.accept_mimetypes.best == "application/json":
            return jsonify({"message": str(error)}), 400
        flash(str(error), "danger")
        return render_template(
            "platform/instagram/import_new.html",
            **_import_template_context(project, sync_period_form=request.form),
        ), 400
    except (TypeError, ValueError):
        current_app.logger.error("instagram_meta project=%s operation=sync status=invalid_stale_config", project.id)
        flash("Não foi possível iniciar a sincronização.", "danger")
        return redirect(url_for("instagram.new_import", project_id=project.id))
    except Exception:
        db.session.rollback()
        current_app.logger.error("instagram_meta project=%s operation=sync status=enqueue_error", project.id)
        flash("Não foi possível iniciar a sincronização.", "danger")
        return redirect(url_for("instagram.new_import", project_id=project.id))

    status_url = url_for("instagram.meta_sync_status", project_id=project.id, run_id=run.id)
    # Inclusive quando o POST reutiliza um job já pendente/em execução, o
    # reload precisa representar a janela congelada daquele job, não uma nova
    # escolha que não foi aplicada a ele.
    completion_url = _meta_sync_completion_url(project, run)
    if request.accept_mimetypes.best == "application/json":
        return jsonify({
            "status": run.status,
            "status_url": status_url,
            "completion_url": completion_url,
            "reused_active_run": not created,
        }), 202
    flash("Sincronização iniciada." if created else "Uma sincronização já está em andamento.", "success")
    return redirect(completion_url)


@instagram_bp.get("/projects/<uuid:project_id>/imports/meta/sync/<uuid:run_id>/status")
@login_required
def meta_sync_status(project_id: UUID, run_id: UUID):
    """Estado seguro de um job Meta pertencente ao projeto do usuário atual."""
    project = _project_for_current_user(project_id)
    run = db.session.scalar(select(AnalyticsRun).where(
        AnalyticsRun.id == str(run_id),
        AnalyticsRun.analytics_project_id == project.id,
        AnalyticsRun.source_kind == "meta_api",
    ))
    if run is None:
        abort(404)
    payload = {
        "status": run.status,
        "record_count": run.record_count,
        "started_at": run.started_at.isoformat() if run.started_at else None,
        "completed_at": run.completed_at.isoformat() if run.completed_at else None,
    }
    if run.status == META_SYNC_FAILED:
        payload["message"] = run.error_message or "Não foi possível concluir a sincronização. Nenhum dado foi salvo."
    elif run.status == META_SYNC_COMPLETED:
        count = run.record_count
        payload["message"] = "Sincronização concluída. " + (
            "1 conteúdo foi processado." if count == 1 else f"{count} conteúdos foram processados."
        )
    return jsonify(payload)


@instagram_bp.route("/projects/<uuid:project_id>/imports/meta/disconnect", methods=("GET", "POST"))
@login_required
def disconnect_meta_account(project_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    account = project.instagram_account
    connection = account.api_connection if account is not None else None
    if connection is None or connection.status == "revoked":
        abort(404)
    if request.method == "POST":
        if request.form.get("confirm") != "yes":
            abort(400)
        disconnect_connection(connection)
        db.session.commit()
        log_meta_operation(project.id, "disconnect", "revoked")
        flash("A conexão com a API do Instagram foi desconectada. O histórico do projeto foi preservado.", "success")
        return redirect(url_for("instagram.new_import", project_id=project.id))
    return render_template(
        "platform/instagram/meta_disconnect.html", project=project, connection=connection,
    )


@instagram_bp.get("/projects/<uuid:project_id>/imports/modelo")
@login_required
def import_template(project_id: UUID):
    _project_for_current_user(project_id, active=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dados Instagram"
    for index, field in enumerate(IMPORT_FIELDS, start=1):
        cell = sheet.cell(1, index, field.name)
        cell.fill = PatternFill("solid", fgColor="184775" if field.required else "52677A")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.comment = Comment(field.note, "Análysis")
        sheet.column_dimensions[get_column_letter(index)].width = min(34, max(19, len(field.name) + 3))
        if field.kind == "datetime":
            for row in range(2, MAX_IMPORT_ROWS + 2):
                sheet.cell(row, index).number_format = "@"
    sheet.row_dimensions[1].height = 31
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(IMPORT_FIELDS))}1"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="modelo_importacao_instagram_analysis.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@instagram_bp.post("/projects/<uuid:project_id>/imports/preview")
@login_required
def preview_import(project_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    uploaded = request.files.get("dataset")
    if uploaded is None or not uploaded.filename:
        flash("Selecione um arquivo CSV ou XLSX.", "danger")
        return render_template("platform/instagram/import_new.html", **_import_template_context(project)), 400
    data = uploaded.read(MAX_IMPORT_BYTES + 1)
    try:
        dataset = parse_instagram_import(uploaded.filename, data)
        conflict = account_conflict(project, dataset)
        if conflict:
            raise InstagramImportError(conflict)
        known_media_external_ids = set()
        if project.instagram_account is not None:
            known_media_external_ids = set(db.session.scalars(
                select(InstagramMedia.external_id).where(
                    InstagramMedia.instagram_account_id == project.instagram_account.id,
                    InstagramMedia.external_id.is_not(None),
                )
            ))
        repeated_source_hash = db.session.scalar(
            select(AnalyticsRun.id).where(
                AnalyticsRun.analytics_project_id == project.id,
                AnalyticsRun.source_hash == dataset.source_hash,
            ).limit(1)
        ) is not None
        preview = build_preview(
            dataset,
            known_media_external_ids=known_media_external_ids,
            repeated_source_hash=repeated_source_hash,
        )
    except InstagramImportError as error:
        flash(str(error), "danger")
        return render_template("platform/instagram/import_new.html", **_import_template_context(project)), 400
    try:
        token = previews.put(current_user.id, project.id, preview)
    except OSError:
        current_app.logger.exception("Não foi possível guardar a prévia Instagram")
        flash("Não foi possível preparar a prévia. Tente enviar o arquivo novamente.", "danger")
        return render_template("platform/instagram/import_new.html", **_import_template_context(project)), 500
    return render_template("platform/instagram/import_preview.html", project=project, preview=preview, token=token)


@instagram_bp.post("/projects/<uuid:project_id>/imports/confirm")
@login_required
def confirm_import(project_id: UUID):
    project = _project_for_current_user(project_id, active=True)
    token = request.form.get("preview_token", "")
    try:
        claim = previews.claim(token, current_user.id, project.id)
    except InstagramPreviewExpired:
        flash("A prévia expirou. Envie o arquivo novamente.", "warning")
        return redirect(url_for("instagram.new_import", project_id=project.id))
    if claim is None:
        flash("Prévia indisponível. Envie o arquivo novamente.", "warning")
        return redirect(url_for("instagram.new_import", project_id=project.id))
    preview = claim.preview
    try:
        result = persist_instagram_import(project, preview.dataset)
    except InstagramImportPersistenceError as error:
        claim.release()
        flash(str(error), "danger")
        return render_template("platform/instagram/import_preview.html", project=project, preview=preview, token=token), 409
    except Exception:
        claim.release()
        flash("Não foi possível concluir a importação. Nenhum dado foi salvo.", "danger")
        return render_template("platform/instagram/import_preview.html", project=project, preview=preview, token=token), 500
    try:
        claim.complete()
    except OSError:
        # O .claim permanece inacessível: uma falha de limpeza não permite novo run.
        current_app.logger.exception("Importação concluída; falha ao remover prévia consumida")
    return render_template("platform/instagram/import_complete.html", project=project, result=result, preview=preview)


@instagram_bp.post("/projects/<uuid:project_id>/imports/cancel")
@login_required
def cancel_import(project_id: UUID):
    project = _project_for_current_user(project_id)
    token = request.form.get("preview_token", "")
    previews.discard(token, current_user.id, project.id)
    return redirect(url_for("instagram.project_detail", project_id=project.id))
