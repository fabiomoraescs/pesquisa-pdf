"""Ambiente qualitativo: corpus, leitura e registros manuais do projeto."""

from __future__ import annotations

import logging
import os
import shutil
import time
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import RLock
from uuid import UUID, uuid4

from flask import Blueprint, abort, current_app, flash, jsonify, redirect, render_template, request, send_file, url_for
from flask_login import current_user, login_required
from sqlalchemy import delete, select, update
from sqlalchemy.exc import IntegrityError
from werkzeug.utils import secure_filename
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from historico_racial.pdf import PDFInvalidoError, contar_paginas

from .analyses import analysis_dir, create_analysis, preserve_documents, save_error
from .extensions import db
from .models import (
    Analysis, AnalysisDocument, Project, QualitativeCode, QualitativeCoding,
    QualitativeExcerpt, QualitativeMemo, QualitativeRejection, normalized_qualitative_code_name, utcnow,
)
from .qualitative import (
    QUALITATIVE_STRATEGIES, code_excerpt_counts, get_qualitative_analysis, get_qualitative_code,
    get_qualitative_document, get_qualitative_excerpt, get_qualitative_memo,
)
from .qualitative_corpus import (
    CorpusUnavailableError, QualitativePageNotFoundError, append_qualitative_corpus,
    delete_qualitative_document, load_qualitative_manifest,
    prepare_qualitative_corpus, qualitative_corpus_dir, read_qualitative_page,
    validate_qualitative_corpus,
)
from .qualitative_colors import QUALITATIVE_CODE_PALETTE, code_color_style, code_palette_options
from .qualitative_layout import read_qualitative_layout
from .qualitative_annotations import (
    SelectionConflict, SelectionError, apply_codes, find_code_by_name,
    get_or_create_excerpt, lock_annotation_write, page_excerpts, update_excerpt_bounds, validated_selection,
)
from .qualitative_automatic import automatic_coding, remove_coding
from .qualitative_expanded_search import SemanticModelUnavailable
from .qualitative_search import QualitativeSearchError, search_qualitative_document
from .scraping_types import QUALITATIVE, QUALITATIVE_TOOL
from .services import ACCOUNT_LIFECYCLE_LOCK, account_accepts_new_work, can_use_tool, get_project_for_user


qualitative_bp = Blueprint("qualitative", __name__, url_prefix="/analise-qualitativa")
_progress: dict[str, dict] = {}
_progress_lock = RLock()
_progress_retention_seconds = 60 * 60
_semantic_jobs: dict[str, dict] = {}
_semantic_jobs_lock = RLock()
_semantic_stage_labels = {
    "preparing": "Preparando busca semântica",
    "literal": "Localizando correspondências literais",
    "lexical": "Analisando famílias lexicais",
    "model": "Carregando modelo semântico",
    "semantic": "Processando e comparando segmentos",
    "coding": "Registrando codificações",
    "finalizing": "Finalizando análise",
    "complete": "Busca semântica concluída",
    "error": "Busca semântica interrompida",
}


def _semantic_job_update(job_id: str, *, stage: str, percent: int,
                         document_id: str | None = None, page_number: int | None = None,
                         document_names: dict[str, str] | None = None, error: str | None = None) -> None:
    with _semantic_jobs_lock:
        job = _semantic_jobs[job_id]
        job["percent"] = max(job["percent"], min(100, max(0, int(percent))))
        job["stage"] = _semantic_stage_labels[stage]
        if document_id:
            job["document"] = (document_names or {}).get(document_id, job.get("document", ""))
        if page_number:
            job["page_number"] = page_number
        if stage in {"complete", "error"}:
            job["state"] = stage
            job["finished_at"] = time.monotonic()
        if error:
            job["error"] = error


def _semantic_job_begin(job_id: str, analysis: Analysis, user_id: str) -> None:
    with _semantic_jobs_lock:
        for identifier, value in list(_semantic_jobs.items()):
            if value.get("finished_at") and time.monotonic() - value["finished_at"] > _progress_retention_seconds:
                _semantic_jobs.pop(identifier, None)
        if job_id in _semantic_jobs:
            raise ValueError("Esta operação semântica já foi iniciada.")
        _semantic_jobs[job_id] = {"analysis_id": analysis.id, "owner_user_id": user_id,
                                  "state": "running", "stage": _semantic_stage_labels["preparing"],
                                  "percent": 0, "document": "", "page_number": None}


@qualitative_bp.get("/bases/<uuid:analysis_id>/progresso-semantico/<uuid:job_id>")
@login_required
def semantic_job_progress(analysis_id: UUID, job_id: UUID):
    analysis = _base(analysis_id)
    with _semantic_jobs_lock:
        job = _semantic_jobs.get(str(job_id))
        if job is None or job["analysis_id"] != analysis.id or job["owner_user_id"] != current_user.id:
            abort(404)
        return jsonify({key: job.get(key) for key in
                        ("state", "stage", "percent", "document", "page_number", "error")})


def _require_tool() -> None:
    if not can_use_tool(current_user, QUALITATIVE_TOOL):
        abort(403)


def _project(project_id: UUID, *, active: bool = False) -> Project:
    _require_tool()
    project = get_project_for_user(str(project_id), current_user, include_inactive=not active)
    if project.scrape_type != QUALITATIVE or project.deleted_at is not None:
        abort(404)
    return project


def _base(analysis_id: UUID) -> Analysis:
    return get_qualitative_analysis(str(analysis_id), current_user)


def _editable_base(analysis_id: UUID) -> Analysis:
    analysis = _base(analysis_id)
    project = _project(UUID(analysis.project_id), active=True)
    if project.owner_user_id != current_user.id or not account_accepts_new_work(current_user.id):
        abort(403)
    return analysis


def _analytic_records(analysis: Analysis) -> tuple[list[QualitativeCode], list[QualitativeMemo]]:
    codes = db.session.scalars(select(QualitativeCode).where(
        QualitativeCode.analysis_id == analysis.id).order_by(QualitativeCode.created_at, QualitativeCode.id)).all()
    memos = db.session.scalars(select(QualitativeMemo).where(
        QualitativeMemo.analysis_id == analysis.id,
    ).order_by(QualitativeMemo.created_at, QualitativeMemo.id)).all()
    return codes, memos


def _records_payload(analysis: Analysis) -> dict:
    codes, memos = _analytic_records(analysis)
    counts = code_excerpt_counts(analysis)
    return {
        "codes": [{"id": code.id, "name": code.name, "description": code.description,
                   **code_color_style(code.color),
                   "excerpt_count": counts.get(code.id, 0)}
                  for code in codes],
        "memos": [{"id": memo.id, "text": memo.text,
                   "excerpt_id": memo.excerpt_id,
                   "context": ("trecho" if memo.excerpt_id else
                               "documento" if memo.document_id else
                               "código" if memo.code_id else "geral")}
                  for memo in memos],
    }


def _input_text(field: str, limit: int, *, required: bool = True) -> str:
    if request.content_length is None or request.content_length > 131_072:
        raise ValueError("O registro enviado excede o limite permitido.")
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict) or not isinstance(payload.get(field), str):
        raise ValueError("Dados inválidos.")
    value = payload[field].strip()
    if (required and not value) or len(value) > limit:
        label = {"name": "Nome", "text": "Conteúdo"}.get(field, field.capitalize())
        raise ValueError(f"{label} deve ter até {limit} caracteres e não pode estar vazio." if required
                         else f"{label} deve ter até {limit} caracteres.")
    return value


def _code_fields(*, default_description: str = "") -> tuple[str, str]:
    payload = request.get_json(silent=True)
    name = " ".join(_input_text("name", 160).split())
    if not isinstance(payload.get("description", default_description), str):
        raise ValueError("Descrição inválida.")
    description = payload.get("description", default_description).strip()
    if len(description) > 4000:
        raise ValueError("Descrição deve ter até 4000 caracteres.")
    return name, description


def _duplicate_code(analysis: Analysis, name: str, *, except_id: str | None = None) -> bool:
    query = select(QualitativeCode.id).where(
        QualitativeCode.analysis_id == analysis.id,
        QualitativeCode.normalized_name == normalized_qualitative_code_name(name),
    )
    if except_id:
        query = query.where(QualitativeCode.id != except_id)
    existing = db.session.scalar(query)
    return existing is not None


def _project_workspaces(project: Project) -> list[Analysis]:
    return db.session.scalars(select(Analysis).where(
        Analysis.project_id == project.id, Analysis.tool_id == QUALITATIVE_TOOL
    ).order_by(Analysis.created_at, Analysis.id)).all()


def ensure_project_workspace(project: Project) -> Analysis | None:
    """Cria o ambiente técnico vazio somente se não houver histórico ambíguo."""
    with ACCOUNT_LIFECYCLE_LOCK:
        workspaces = _project_workspaces(project)
        if len(workspaces) == 1:
            return workspaces[0]
        if len(workspaces) > 1 or project.status != "active" or project.deleted_at is not None:
            return None
        if not account_accepts_new_work(project.owner_user_id):
            return None
        return create_analysis(user_id=project.owner_user_id, project_id=project.id,
                               tool_id=QUALITATIVE_TOOL, tool_version="manual-v1",
                               name=project.name, parameters={})


def _upload_lock(analysis_id: str) -> Path:
    return analysis_dir(analysis_id) / ".qualitative_upload.lock"


def _current_progress(analysis: Analysis) -> dict:
    with _progress_lock:
        for identifier, value in list(_progress.items()):
            if value.get("finished_at") and time.monotonic() - value["finished_at"] > _progress_retention_seconds:
                _progress.pop(identifier, None)
        current = dict(_progress.get(analysis.id, {}))
    error = current.get("error") or (analysis.error_message if analysis.status == "erro" else None)
    state = "erro" if error else "processando" if _upload_lock(analysis.id).exists() else analysis.status
    percent = 100 if state == "concluida" else min(99, max(0, current.get("percentual", 0)))
    return {"status": state, "percentual": percent, "pages_completed": current.get("pages_completed", 0),
            "total_pages": current.get("total_pages", 0), "document_index": current.get("document_index", 0),
            "document_count": current.get("document_count", analysis.document_count),
            "document_name": current.get("document_name"), "page_number": current.get("page_number", 0),
            "page_count": current.get("page_count", 0), "stage": current.get("stage"),
            "result_url": url_for("qualitative.base", analysis_id=analysis.id) if state == "concluida" else None,
            "error": error}


def _run_corpus_job(app, analysis_id: str) -> None:
    with app.app_context():
        published = False
        def report(event: dict) -> None:
            with _progress_lock:
                _progress.setdefault(analysis_id, {}).update(event)

        try:
            analysis = db.session.get(Analysis, analysis_id)
            if analysis is None:
                return
            prepare_qualitative_corpus(analysis, report)
            published = True
            validate_qualitative_corpus(analysis, allow_processing=True)
            analysis.status = "concluida"  # corpus pronto, não pesquisa interpretativa concluída
            analysis.completed_at = utcnow()
            analysis.error_message = None
            db.session.commit()
        except Exception:
            db.session.rollback()
            logging.getLogger(__name__).exception("Falha ao preparar corpus qualitativo %s", analysis_id)
            if published:
                try:
                    shutil.rmtree(qualitative_corpus_dir(analysis_id))
                except OSError:
                    logging.getLogger(__name__).exception("Falha ao retirar corpus não confirmado %s", analysis_id)
            save_error(analysis_id, "Não foi possível preparar o corpus. Os PDFs originais foram preservados.")
        finally:
            with _progress_lock:
                _progress.setdefault(analysis_id, {})["finished_at"] = time.monotonic()
            db.session.remove()


@qualitative_bp.get("")
@login_required
def index():
    _require_tool()
    return redirect(url_for("projects.list_qualitative_projects"))


@qualitative_bp.get("/projetos/<uuid:project_id>")
@login_required
def project_workspace(project_id: UUID):
    project = _project(project_id)
    workspaces = _project_workspaces(project)
    if not workspaces and project.owner_user_id == current_user.id:
        created = ensure_project_workspace(project)
        if created is not None:
            workspaces = [created]
    if len(workspaces) == 1:
        return redirect(url_for("qualitative.base", analysis_id=workspaces[0].id))
    if not workspaces:
        return render_template("platform/qualitative_empty.html", project=project, analysis=None)
    # Dados históricos não são fundidos nem atribuídos arbitrariamente.
    return render_template("platform/qualitative_legacy_choice.html", project=project,
                           workspaces=workspaces)


@qualitative_bp.get("/projetos/<uuid:project_id>/bases")
@login_required
def project_bases(project_id: UUID):
    project = _project(project_id)
    bases = db.session.scalars(select(Analysis).where(
        Analysis.project_id == project.id, Analysis.tool_id == QUALITATIVE_TOOL
    ).order_by(Analysis.created_at.desc())).all()
    return render_template("platform/qualitative_bases.html", project=project, analyses=bases)


@qualitative_bp.route("/projetos/<uuid:project_id>/bases/nova", methods=["GET", "POST"])
@login_required
def new_base(project_id: UUID):
    project = _project(project_id, active=True)
    if project.owner_user_id != current_user.id:
        abort(403)
    if _project_workspaces(project):
        if request.method == "POST":
            abort(409)
        return redirect(url_for("qualitative.project_workspace", project_id=project.id))
    if request.method == "GET":
        return render_template("platform/qualitative_new_base.html", project=project)

    name = request.form.get("name", "").strip()
    strategy = request.form.get("qualitative_strategy", "")
    uploads = [item for item in request.files.getlist("documents") if item.filename]
    if not name or len(name) > 200 or strategy not in QUALITATIVE_STRATEGIES or not uploads:
        return render_template("platform/qualitative_new_base.html", project=project,
                               error="Informe nome, estratégia e ao menos um PDF."), 400
    with TemporaryDirectory(prefix="qualitative-upload-") as temporary:
        files = []
        try:
            for index, upload in enumerate(uploads, start=1):
                original = upload.filename.replace("\\", "/").rsplit("/", 1)[-1].strip()
                if not original.lower().endswith(".pdf") or len(original) > 255:
                    raise PDFInvalidoError("Envie somente arquivos PDF válidos.")
                safe = secure_filename(original) or "documento.pdf"
                source = Path(temporary) / f"{index:03d}-{safe}"
                upload.save(source)
                contar_paginas(source)  # valida o conteúdo real sem extrair o corpus
                files.append((source, original))
        except (PDFInvalidoError, OSError, ValueError):
            return render_template("platform/qualitative_new_base.html", project=project,
                                   error="Um dos arquivos não é um PDF legível, está vazio ou está protegido."), 400

        with ACCOUNT_LIFECYCLE_LOCK:
            current_project = db.session.get(Project, project.id, populate_existing=True)
            if (not account_accepts_new_work(current_user.id) or current_project is None
                    or current_project.status != "active" or current_project.deleted_at is not None
                    or current_project.owner_user_id != current_user.id):
                abort(403)
            analysis = create_analysis(
                user_id=current_user.id, project_id=project.id, tool_id=QUALITATIVE_TOOL,
                tool_version="manual-v1", name=name,
                parameters={"qualitative_strategy": strategy},
            )
            try:
                preserve_documents(analysis, files)
            except Exception:
                db.session.rollback()
                logging.getLogger(__name__).exception("Falha ao preservar PDFs qualitativos %s", analysis.id)
                partial_documents = analysis_dir(analysis.id) / "documents"
                if partial_documents.exists() and not partial_documents.is_symlink():
                    shutil.rmtree(partial_documents)
                save_error(analysis.id, "Não foi possível salvar os PDFs enviados.")
                return render_template("platform/qualitative_new_base.html", project=project,
                                       error="Não foi possível salvar os PDFs enviados."), 503

    from app import EXECUTOR_ANALISES

    try:
        EXECUTOR_ANALISES.submit(_run_corpus_job, current_app._get_current_object(), analysis.id)
    except RuntimeError:
        save_error(analysis.id, "Não foi possível iniciar a preparação do corpus.")
        return render_template("platform/qualitative_new_base.html", project=project,
                               error="Não foi possível iniciar a preparação do corpus."), 503
    return redirect(url_for("qualitative.base", analysis_id=analysis.id))


def _run_added_documents_job(app, analysis_id: str, incoming_name: str,
                             items: list[tuple[str, str, str]], initial: bool) -> None:
    with app.app_context():
        incoming = analysis_dir(analysis_id) / incoming_name
        try:
            analysis = db.session.get(Analysis, analysis_id)
            if analysis is None:
                return
            if initial:
                preserve_documents(analysis, [(incoming / stored, original)
                                              for _, stored, original in items])
                _run_corpus_job(app, analysis_id)
            else:
                additions = [(AnalysisDocument(id=identifier, analysis_id=analysis_id,
                                               stored_name=stored, original_name=original), incoming / stored)
                             for identifier, stored, original in items]
                def report(event: dict) -> None:
                    with _progress_lock:
                        _progress.setdefault(analysis_id, {}).update(event)
                append_qualitative_corpus(analysis, additions, report)
                with _progress_lock:
                    _progress.setdefault(analysis_id, {})["finished_at"] = time.monotonic()
        except Exception:
            db.session.rollback()
            logging.getLogger(__name__).exception("Falha ao adicionar documentos qualitativos %s", analysis_id)
            if initial:
                save_error(analysis_id, "Não foi possível preparar os documentos enviados.")
            with _progress_lock:
                _progress.setdefault(analysis_id, {}).update(
                    error="Não foi possível preparar os novos documentos; os anteriores foram preservados.",
                    finished_at=time.monotonic())
        finally:
            try:
                if incoming.exists() and not incoming.is_symlink():
                    shutil.rmtree(incoming)
            except OSError:
                logging.getLogger(__name__).exception("Falha ao limpar upload transitório %s", analysis_id)
            finally:
                try:
                    _upload_lock(analysis_id).unlink(missing_ok=True)
                finally:
                    db.session.remove()


@qualitative_bp.post("/bases/<uuid:analysis_id>/documentos")
@login_required
def add_documents(analysis_id: UUID):
    analysis = _base(analysis_id)
    project = _project(UUID(analysis.project_id), active=True)
    if project.owner_user_id != current_user.id:
        abort(403)
    initial = analysis.document_count == 0 and analysis.status == "processando"
    if not initial and analysis.status != "concluida":
        abort(409)
    if not initial:
        try:
            load_qualitative_manifest(analysis)
        except CorpusUnavailableError:
            abort(409)
    uploads = [item for item in request.files.getlist("documents") if item.filename]
    if not uploads:
        flash("Selecione ao menos um PDF.", "danger")
        return redirect(url_for("qualitative.project_workspace", project_id=project.id))
    lock = _upload_lock(analysis.id)
    with ACCOUNT_LIFECYCLE_LOCK:
        # Revalidar depois de adquirir o mesmo lock usado pela exclusão de projetos.
        # A requisição pode ter aguardado enquanto o projeto era arquivado/excluído.
        current_project = db.session.get(Project, project.id, populate_existing=True)
        if (not account_accepts_new_work(current_user.id) or current_project is None
                or current_project.status != "active" or current_project.deleted_at is not None
                or current_project.owner_user_id != current_user.id):
            abort(403)
        try:
            fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
        except FileExistsError:
            abort(409)
        else:
            os.close(fd)
        db.session.refresh(analysis)
        initial = analysis.document_count == 0 and analysis.status == "processando"
        if not initial and analysis.status != "concluida":
            lock.unlink(missing_ok=True)
            abort(409)
        if not initial:
            try:
                load_qualitative_manifest(analysis)
            except CorpusUnavailableError:
                lock.unlink(missing_ok=True)
                abort(409)
    incoming = analysis_dir(analysis.id) / f".qualitative_upload.{uuid4().hex}.tmp"
    try:
        incoming.mkdir()
        items = []
        for upload in uploads:
            original = upload.filename.replace("\\", "/").rsplit("/", 1)[-1].strip()
            if not original.lower().endswith(".pdf") or len(original) > 255:
                raise PDFInvalidoError("Envie somente PDFs válidos.")
            stored = f"{uuid4().hex}-{secure_filename(original) or 'documento.pdf'}"
            path = incoming / stored
            upload.save(path)
            contar_paginas(path)
            items.append((str(uuid4()), stored, original))
        from app import EXECUTOR_ANALISES
        with _progress_lock:
            _progress[analysis.id] = {"document_count": len(items), "document_index": 0}
        EXECUTOR_ANALISES.submit(_run_added_documents_job, current_app._get_current_object(),
                                 analysis.id, incoming.name, items, initial)
    except (PDFInvalidoError, OSError, ValueError, RuntimeError):
        if incoming.exists() and not incoming.is_symlink():
            shutil.rmtree(incoming)
        lock.unlink(missing_ok=True)
        flash("Um dos arquivos não é um PDF legível ou não foi possível iniciar o preparo.", "danger")
        return redirect(url_for("qualitative.project_workspace", project_id=project.id))
    return redirect(url_for("qualitative.upload_pending", analysis_id=analysis.id))


@qualitative_bp.post("/bases/<uuid:analysis_id>/documentos/<uuid:document_id>/excluir")
@login_required
def delete_document(analysis_id: UUID, document_id: UUID):
    analysis = _base(analysis_id)
    project = _project(UUID(analysis.project_id), active=True)
    if project.owner_user_id != current_user.id:
        abort(403)
    document = get_qualitative_document(analysis, str(document_id))
    lock = _upload_lock(analysis.id)
    try:
        fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    except FileExistsError:
        return jsonify({"error": "Aguarde a preparação dos documentos terminar."}), 409
    else:
        os.close(fd)
    try:
        next_id = delete_qualitative_document(analysis, document)
    except CorpusUnavailableError:
        return jsonify({"error": "O corpus está inconsistente. Nenhum documento foi removido."}), 409
    except Exception:
        logging.getLogger(__name__).exception("Falha ao excluir documento %s", document_id)
        return jsonify({"error": "Não foi possível excluir o documento. Tente novamente após revisão."}), 500
    finally:
        lock.unlink(missing_ok=True)
    destination = (url_for("qualitative.page", analysis_id=analysis.id,
                           document_id=next_id, page_number=1) if next_id else
                   url_for("qualitative.project_workspace", project_id=project.id))
    return jsonify({"document_id": document.id, "document_count": analysis.document_count,
                    "next_url": destination, "records": _records_payload(analysis)})


@qualitative_bp.post("/bases/<uuid:analysis_id>/codigos")
@login_required
def create_code(analysis_id: UUID):
    analysis = _editable_base(analysis_id)
    try:
        name, description = _code_fields()
    except ValueError as error:
        return jsonify({"error": str(error)}), 400
    if _duplicate_code(analysis, name):
        return jsonify({"error": "Já existe um código com esse nome neste projeto."}), 409
    db.session.add(QualitativeCode(analysis_id=analysis.id, name=name, description=description,
                                   created_by_user_id=current_user.id))
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return jsonify({"error": "Já existe um código com esse nome neste projeto."}), 409
    return jsonify(_records_payload(analysis)), 201


@qualitative_bp.patch("/bases/<uuid:analysis_id>/codigos/<uuid:code_id>")
@login_required
def edit_code(analysis_id: UUID, code_id: UUID):
    analysis = _editable_base(analysis_id)
    code = get_qualitative_code(analysis, str(code_id))
    if request.content_length is None or request.content_length > 131_072:
        return jsonify({"error": "O registro enviado excede o limite permitido."}), 400
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict) or not ({"name", "color"} & payload.keys()):
        return jsonify({"error": "Informe o nome ou uma cor válida para o código."}), 400
    if "color" in payload and (type(payload["color"]) is not str
                               or payload["color"] not in QUALITATIVE_CODE_PALETTE):
        return jsonify({"error": "Cor inválida. Escolha uma das cores da paleta."}), 400
    try:
        name, description = (_code_fields(default_description=code.description)
                             if "name" in payload else (code.name, code.description))
    except ValueError as error:
        return jsonify({"error": str(error)}), 400
    if _duplicate_code(analysis, name, except_id=code.id):
        return jsonify({"error": "Já existe um código com esse nome neste projeto."}), 409
    code.name, code.description = name, description
    if "color" in payload:
        code.color = payload["color"]
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return jsonify({"error": "Já existe um código com esse nome neste projeto."}), 409
    return jsonify(_records_payload(analysis))


@qualitative_bp.delete("/bases/<uuid:analysis_id>/codigos/<uuid:code_id>")
@login_required
def delete_code(analysis_id: UUID, code_id: UUID):
    analysis = _editable_base(analysis_id)
    code = get_qualitative_code(analysis, str(code_id))
    db.session.execute(delete(QualitativeRejection).where(
        QualitativeRejection.analysis_id == analysis.id, QualitativeRejection.code_id == code.id))
    db.session.execute(delete(QualitativeCoding).where(
        QualitativeCoding.analysis_id == analysis.id, QualitativeCoding.code_id == code.id))
    # Memos contextuais são preservados como memos gerais; trechos mantêm sua identidade.
    db.session.execute(update(QualitativeMemo).where(
        QualitativeMemo.analysis_id == analysis.id, QualitativeMemo.code_id == code.id
    ).values(code_id=None))
    db.session.delete(code)
    db.session.commit()
    return jsonify(_records_payload(analysis))


@qualitative_bp.post("/bases/<uuid:analysis_id>/memos")
@login_required
def create_memo(analysis_id: UUID):
    analysis = _editable_base(analysis_id)
    try:
        content = _input_text("text", 20000)
    except ValueError as error:
        return jsonify({"error": str(error)}), 400
    db.session.add(QualitativeMemo(analysis_id=analysis.id, text=content,
                                   created_by_user_id=current_user.id))
    db.session.commit()
    return jsonify(_records_payload(analysis)), 201


@qualitative_bp.patch("/bases/<uuid:analysis_id>/memos/<uuid:memo_id>")
@login_required
def edit_memo(analysis_id: UUID, memo_id: UUID):
    analysis = _editable_base(analysis_id)
    memo = get_qualitative_memo(analysis, str(memo_id))
    try:
        memo.text = _input_text("text", 20000)
    except ValueError as error:
        return jsonify({"error": str(error)}), 400
    db.session.commit()
    return jsonify(_records_payload(analysis))


@qualitative_bp.delete("/bases/<uuid:analysis_id>/memos/<uuid:memo_id>")
@login_required
def delete_memo(analysis_id: UUID, memo_id: UUID):
    analysis = _editable_base(analysis_id)
    db.session.delete(get_qualitative_memo(analysis, str(memo_id)))
    db.session.commit()
    return jsonify(_records_payload(analysis))


@qualitative_bp.get("/bases/<uuid:analysis_id>/documentos/preparando")
@login_required
def upload_pending(analysis_id: UUID):
    analysis = _base(analysis_id)
    project = _project(UUID(analysis.project_id))
    progress_data = _current_progress(analysis)
    if not _upload_lock(analysis.id).exists():
        if progress_data["error"]:
            return render_template("platform/qualitative_pending.html", analysis=analysis,
                                   project=project, progress=progress_data, upload_error=progress_data["error"])
        return redirect(url_for("qualitative.project_workspace", project_id=project.id))
    return render_template("platform/qualitative_pending.html", analysis=analysis,
                           project=project, progress=progress_data, uploading=True)


@qualitative_bp.get("/bases/<uuid:analysis_id>/progresso")
@login_required
def progress(analysis_id: UUID):
    return jsonify(_current_progress(_base(analysis_id)))


@qualitative_bp.get("/bases/<uuid:analysis_id>")
@login_required
def base(analysis_id: UUID):
    analysis = _base(analysis_id)
    project = _project(UUID(analysis.project_id))
    if analysis.document_count == 0:
        return render_template("platform/qualitative_empty.html", project=project, analysis=analysis,
                               records=_records_payload(analysis))
    if analysis.status != "concluida":
        return render_template("platform/qualitative_pending.html", analysis=analysis,
                               project=project, progress=_current_progress(analysis))
    try:
        manifest = load_qualitative_manifest(analysis)
    except CorpusUnavailableError:
        logging.getLogger(__name__).error("Base %s marcada concluída sem corpus íntegro", analysis.id)
        return render_template("platform/qualitative_pending.html", analysis=analysis,
                               project=project, inconsistent=True), 409
    first = manifest["documents"][0]
    return redirect(url_for("qualitative.page", analysis_id=analysis.id,
                            document_id=first["document_id"], page_number=1))


def _coding_report_rows(analysis: Analysis):
    """Relações Código × Trecho da Base, sempre derivadas do estado atual."""
    codes = db.session.scalars(select(QualitativeCode).where(
        QualitativeCode.analysis_id == analysis.id,
    ).order_by(QualitativeCode.normalized_name, QualitativeCode.id)).all()
    rows = db.session.execute(select(QualitativeCoding, QualitativeExcerpt, AnalysisDocument).join(
        QualitativeExcerpt, QualitativeExcerpt.id == QualitativeCoding.excerpt_id
    ).join(AnalysisDocument, AnalysisDocument.id == QualitativeExcerpt.document_id).where(
        QualitativeCoding.analysis_id == analysis.id,
        QualitativeExcerpt.analysis_id == analysis.id,
        AnalysisDocument.analysis_id == analysis.id,
    ).order_by(AnalysisDocument.original_name, QualitativeExcerpt.page_number,
               QualitativeExcerpt.start_offset, QualitativeExcerpt.id)).all()
    return codes, rows


def _safe_excel_text(value: str) -> str:
    """Exporta texto em uma linha, sem alterar a fonte nem permitir fórmulas."""
    value = " ".join(value.splitlines())
    return f"'{value}" if value.lstrip().startswith(("=", "+", "-", "@")) else value


@qualitative_bp.get("/bases/<uuid:analysis_id>/relatorio-codificacao")
@login_required
def coding_report(analysis_id: UUID):
    """Relatório derivado das codificações atuais, sem cópia do corpus."""
    analysis = _base(analysis_id)
    project = _project(UUID(analysis.project_id))
    codes, rows = _coding_report_rows(analysis)
    grouped = {code.id: [] for code in codes}
    excerpt_ids = {excerpt.id for _, excerpt, _ in rows}
    memo_counts = {identifier: 0 for identifier in excerpt_ids}
    if excerpt_ids:
        for excerpt_id in db.session.scalars(select(QualitativeMemo.excerpt_id).where(
                QualitativeMemo.analysis_id == analysis.id,
                QualitativeMemo.excerpt_id.in_(excerpt_ids))).all():
            memo_counts[excerpt_id] += 1
    seen = set()
    for coding, excerpt, document in rows:
        key = (coding.code_id, excerpt.id)
        if key in seen or coding.code_id not in grouped:
            continue
        seen.add(key)
        grouped[coding.code_id].append({
            "document_name": document.original_name,
            "page_number": excerpt.page_number,
            "text": excerpt.quoted_text,
            "memo_count": memo_counts[excerpt.id],
            "url": url_for("qualitative.page", analysis_id=analysis.id,
                           document_id=document.id, page_number=excerpt.page_number,
                           excerpt=excerpt.id),
        })
    return render_template("platform/qualitative_coding_report.html", analysis=analysis,
                           project=project, codes=codes, grouped=grouped,
                           code_styles={code.id: code_color_style(code.color) for code in codes})


@qualitative_bp.get("/bases/<uuid:analysis_id>/relatorio-codificacao.xlsx")
@login_required
def coding_report_excel(analysis_id: UUID):
    """Exportação dinâmica e autorizada; não persiste arquivo redundante."""
    analysis = _base(analysis_id)
    project = _project(UUID(analysis.project_id))
    codes, rows = _coding_report_rows(analysis)
    if not rows:
        flash("Ainda não há trechos codificados para exportar.", "warning")
        return redirect(url_for("qualitative.coding_report", analysis_id=analysis.id))
    code_names = {code.id: code.name for code in codes}
    code_styles = {code.id: code_color_style(code.color) for code in codes}
    excerpt_ids = {excerpt.id for _, excerpt, _ in rows}
    memos = db.session.scalars(select(QualitativeMemo).where(
        QualitativeMemo.analysis_id == analysis.id,
        QualitativeMemo.excerpt_id.in_(excerpt_ids),
    ).order_by(QualitativeMemo.created_at, QualitativeMemo.id)).all()
    memo_texts: dict[str, list[str]] = {}
    for memo in memos:
        memo_texts.setdefault(memo.excerpt_id, []).append(memo.text)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Codificação"
    sheet.append(["Código", "Documento", "Página", "Trecho codificado", "Memo contextual"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="294A67")
        cell.alignment = Alignment(vertical="top", wrap_text=False)
    seen = set()
    for coding, excerpt, document in rows:
        key = (coding.code_id, excerpt.id)
        if key in seen or coding.code_id not in code_names:
            continue
        seen.add(key)
        sheet.append([_safe_excel_text(code_names[coding.code_id]),
                      _safe_excel_text(document.original_name),
                      excerpt.page_number, _safe_excel_text(excerpt.quoted_text),
                      _safe_excel_text("\n---\n".join(memo_texts.get(excerpt.id, [])))])
        code_cell = sheet.cell(row=sheet.max_row, column=1)
        style = code_styles[coding.code_id]
        code_cell.fill = PatternFill(fill_type="solid", fgColor="FF" + style["color_hex"].lstrip("#"))
        code_cell.font = Font(color="FF" + style["color_text"].lstrip("#"))
    for column, width in {"A": 26, "B": 36, "C": 10, "D": 80, "E": 65}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    project_name = secure_filename(project.name).lower() or "projeto"
    return send_file(output,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=f"relatorio-codificacao-{project_name}.xlsx")


@qualitative_bp.get("/bases/<uuid:analysis_id>/documentos/<uuid:document_id>/paginas/<int:page_number>")
@login_required
def page(analysis_id: UUID, document_id: UUID, page_number: int):
    analysis = _base(analysis_id)
    get_qualitative_document(analysis, str(document_id))
    project = _project(UUID(analysis.project_id))
    try:
        manifest = load_qualitative_manifest(analysis)
        current = read_qualitative_page(analysis, str(document_id), page_number, manifest=manifest)
    except QualitativePageNotFoundError:
        abort(404)
    except CorpusUnavailableError:
        abort(409)
    target_excerpt_id = None
    requested_excerpt = request.args.get("excerpt", "")
    if requested_excerpt:
        try:
            target = db.session.get(QualitativeExcerpt, str(UUID(requested_excerpt)))
        except ValueError:
            target = None
        if target is not None:
            if (target.analysis_id != analysis.id or target.document_id != str(document_id)
                    or target.page_number != page_number):
                abort(404)
            target_excerpt_id = target.id
    return render_template("platform/qualitative_reader.html", analysis=analysis, project=project,
                           manifest=manifest, current=current, records=_records_payload(analysis),
                           target_excerpt_id=target_excerpt_id, code_palette=code_palette_options())


@qualitative_bp.get("/bases/<uuid:analysis_id>/documentos/<uuid:document_id>/ir")
@login_required
def go_to_page(analysis_id: UUID, document_id: UUID):
    analysis = _base(analysis_id)
    get_qualitative_document(analysis, str(document_id))
    try:
        number = int(request.args.get("page_number", ""))
    except ValueError:
        abort(404)
    try:
        read_qualitative_page(analysis, str(document_id), number)
    except QualitativePageNotFoundError:
        abort(404)
    except CorpusUnavailableError:
        abort(409)
    return redirect(url_for("qualitative.page", analysis_id=analysis.id,
                            document_id=document_id, page_number=number))


@qualitative_bp.get("/bases/<uuid:analysis_id>/documentos/<uuid:document_id>/paginas/<int:page_number>/dados")
@login_required
def page_data(analysis_id: UUID, document_id: UUID, page_number: int):
    analysis = _base(analysis_id)
    get_qualitative_document(analysis, str(document_id))
    try:
        return jsonify(read_qualitative_page(analysis, str(document_id), page_number))
    except QualitativePageNotFoundError:
        abort(404)
    except CorpusUnavailableError:
        abort(409)


@qualitative_bp.get("/bases/<uuid:analysis_id>/documentos/<uuid:document_id>/pdf")
@login_required
def original_pdf(analysis_id: UUID, document_id: UUID):
    """PDF original da Base, com autorização e suporte a Range via send_file."""
    analysis = _base(analysis_id)
    document = get_qualitative_document(analysis, str(document_id))
    if analysis.status != "concluida":
        abort(409)
    root = analysis_dir(analysis.id) / "documents"
    source = root / document.stored_name
    if (root.is_symlink() or source.is_symlink() or source.resolve().parent != root.resolve()
            or not source.is_file()):
        abort(409)
    response = send_file(source, mimetype="application/pdf", conditional=True,
                         as_attachment=False, download_name=document.original_name)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Cache-Control"] = "private, no-store"
    return response


@qualitative_bp.get("/bases/<uuid:analysis_id>/documentos/<uuid:document_id>/buscar")
@login_required
def search(analysis_id: UUID, document_id: UUID):
    analysis = _base(analysis_id)
    get_qualitative_document(analysis, str(document_id))
    try:
        result = search_qualitative_document(
            analysis, str(document_id), request.args.get("q", ""),
            grep=request.args.get("grep") == "1", case_sensitive=request.args.get("case") == "1",
        )
        return jsonify(result)
    except QualitativePageNotFoundError:
        abort(404)
    except CorpusUnavailableError:
        abort(409)
    except QualitativeSearchError as error:
        return jsonify({"error": str(error)}), 400


@qualitative_bp.get("/bases/<uuid:analysis_id>/documentos/<uuid:document_id>/paginas/<int:page_number>/layout")
@login_required
def page_layout(analysis_id: UUID, document_id: UUID, page_number: int):
    analysis = _base(analysis_id)
    get_qualitative_document(analysis, str(document_id))
    try:
        return jsonify(read_qualitative_layout(analysis, str(document_id), page_number))
    except QualitativePageNotFoundError:
        abort(404)
    except CorpusUnavailableError:
        abort(409)


@qualitative_bp.post("/bases/<uuid:analysis_id>/documentos/<uuid:document_id>/codificar")
@login_required
def automatically_code(analysis_id: UUID, document_id: UUID):
    """Busca explícita com escrita; GET /buscar permanece estritamente leitura."""
    with ACCOUNT_LIFECYCLE_LOCK:
        analysis = _editable_base(analysis_id)
        get_qualitative_document(analysis, str(document_id))
        if analysis.status != "concluida" or _upload_lock(analysis.id).exists():
            return jsonify({"error": "Aguarde a preparação dos documentos."}), 409
        if request.content_length is None or request.content_length > 131_072:
            return jsonify({"error": "Consulta excede o limite permitido."}), 400
        payload = request.get_json(silent=True)
        if (not isinstance(payload, dict) or not isinstance(payload.get("q"), str)
                or type(payload.get("mode")) is not str
                or payload["mode"] not in {"literal", "lexical", "semantic"}
                or not isinstance(payload.get("scope"), str)
                or payload["scope"] not in {"document", "project"}
                or type(payload.get("grep", False)) is not bool or type(payload.get("case_sensitive", False)) is not bool
                or type(payload.get("contextual_rejection_enabled", False)) is not bool
                or type(payload.get("multiple_terms", False)) is not bool):
            return jsonify({"error": "Configuração de codificação automática inválida."}), 400
        if payload.get("grep") and payload["mode"] != "literal":
            return jsonify({"error": "Regex está disponível somente na autocodificação Literal."}), 400
        if any(field in payload for field in ("code_id", "name", "description")):
            return jsonify({"error": "Na autocodificação, o código é definido pelo termo pesquisado."}), 400
        job_id = None
        if "progress_id" in payload:
            try:
                if payload["mode"] != "semantic" or not isinstance(payload["progress_id"], str):
                    raise ValueError
                job_id = str(UUID(payload["progress_id"]))
            except ValueError:
                return jsonify({"error": "Identificador de progresso semântico inválido."}), 400
        job_started = False
        try:
            if job_id:
                _semantic_job_begin(job_id, analysis, current_user.id)
                job_started = True
                manifest = load_qualitative_manifest(analysis)
                document_names = {item["document_id"]: item["original_name"]
                                  for item in manifest["documents"]}
                def report_progress(event):
                    _semantic_job_update(job_id, stage=event["stage"], percent=event["percent"],
                                         document_id=event.get("document_id"),
                                         page_number=event.get("page_number"),
                                         document_names=document_names)
            else:
                report_progress = None
            lock_annotation_write(analysis)
            result = automatic_coding(analysis, str(document_id) if payload["scope"] == "document" else None,
                payload["q"], current_user.id, grep=payload.get("grep", False),
                case_sensitive=payload.get("case_sensitive", False),
                contextual_rejection_enabled=payload.get("contextual_rejection_enabled", False),
                multiple_terms=payload.get("multiple_terms", False), mode=payload["mode"],
                progress_callback=report_progress)
            records = _records_payload(analysis)
            db.session.commit()
            if job_started:
                _semantic_job_update(job_id, stage="complete", percent=100)
        except QualitativePageNotFoundError:
            db.session.rollback()
            if job_started:
                _semantic_job_update(job_id, stage="error", percent=0, error="Documento não encontrado.")
            abort(404)
        except (CorpusUnavailableError, SelectionConflict) as error:
            db.session.rollback()
            if job_started:
                _semantic_job_update(job_id, stage="error", percent=0, error=str(error))
            return jsonify({"error": str(error)}), 409
        except SemanticModelUnavailable as error:
            db.session.rollback()
            if job_started:
                _semantic_job_update(job_id, stage="error", percent=0, error=str(error))
            return jsonify({"error": str(error)}), 503
        except (ValueError, QualitativeSearchError) as error:
            db.session.rollback()
            if job_started:
                _semantic_job_update(job_id, stage="error", percent=0, error=str(error))
            return jsonify({"error": str(error)}), 400
        except IntegrityError:
            db.session.rollback()
            if job_started:
                _semantic_job_update(job_id, stage="error", percent=0,
                                     error="Outra operação alterou os códigos.")
            return jsonify({"error": "Outra operação alterou os códigos. Tente novamente."}), 409
        except Exception:
            db.session.rollback()
            if job_started:
                _semantic_job_update(job_id, stage="error", percent=0,
                                     error="Não foi possível concluir a busca semântica.")
            raise
        return jsonify({**result, "records": records}), 201


@qualitative_bp.delete("/bases/<uuid:analysis_id>/codificacoes/<uuid:coding_id>")
@login_required
def delete_coding(analysis_id: UUID, coding_id: UUID):
    with ACCOUNT_LIFECYCLE_LOCK:
        analysis = _editable_base(analysis_id)
        if _upload_lock(analysis.id).exists():
            return jsonify({"error": "Aguarde a preparação dos documentos."}), 409
        lock_annotation_write(analysis)
        coding = db.session.get(QualitativeCoding, str(coding_id))
        if coding is None or coding.analysis_id != analysis.id:
            abort(404)
        try:
            excerpt = remove_coding(analysis, coding, current_user.id)
            # Valida a resposta antes do commit: corpus inconsistente não produz sucesso parcial.
            db.session.flush()
            page = page_excerpts(analysis, excerpt.document_id, excerpt.page_number)
            records = _records_payload(analysis)
            db.session.commit()
        except (CorpusUnavailableError, SelectionConflict):
            db.session.rollback()
            abort(409)
        except QualitativePageNotFoundError:
            db.session.rollback()
            abort(404)
        except IntegrityError:
            db.session.rollback()
            return jsonify({"error": "Outra operação alterou esta codificação. Atualize a página."}), 409
        return jsonify({"page": page, "records": records, "page_number": excerpt.page_number,
                        "document_id": excerpt.document_id})


@qualitative_bp.get("/bases/<uuid:analysis_id>/documentos/<uuid:document_id>/paginas/<int:page_number>/trechos")
@login_required
def excerpts_on_page(analysis_id: UUID, document_id: UUID, page_number: int):
    analysis = _base(analysis_id)
    get_qualitative_document(analysis, str(document_id))
    try:
        return jsonify(page_excerpts(analysis, str(document_id), page_number))
    except QualitativePageNotFoundError:
        abort(404)
    except CorpusUnavailableError:
        abort(409)


@qualitative_bp.patch("/bases/<uuid:analysis_id>/trechos/<uuid:excerpt_id>")
@login_required
def edit_excerpt(analysis_id: UUID, excerpt_id: UUID):
    analysis = _editable_base(analysis_id)
    lock_annotation_write(analysis)
    excerpt = get_qualitative_excerpt(analysis, str(excerpt_id))
    get_qualitative_document(analysis, excerpt.document_id)
    if analysis.status != "concluida":
        abort(409)
    if request.content_length is None or request.content_length > 131_072:
        return jsonify({"error": "O ajuste enviado excede o limite permitido."}), 400
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify({"error": "Dados inválidos."}), 400
    try:
        update_excerpt_bounds(analysis, excerpt, payload.get("start"), payload.get("end"),
                              payload.get("page_hash"))
        db.session.commit()
        return jsonify({"excerpt_id": excerpt.id,
                        "page": page_excerpts(analysis, excerpt.document_id, excerpt.page_number)})
    except QualitativePageNotFoundError:
        abort(404)
    except CorpusUnavailableError:
        abort(409)
    except SelectionConflict as error:
        db.session.rollback()
        return jsonify({"error": str(error)}), 409
    except SelectionError as error:
        db.session.rollback()
        return jsonify({"error": str(error)}), 400
    except IntegrityError:
        db.session.rollback()
        return jsonify({"error": "Os limites foram alterados por outra operação. Tente novamente."}), 409


@qualitative_bp.post("/bases/<uuid:analysis_id>/documentos/<uuid:document_id>/paginas/<int:page_number>/contexto")
@login_required
def annotate_selection(analysis_id: UUID, document_id: UUID, page_number: int):
    """Uma transação por gesto: texto canônico → trecho → código(s)/memo."""
    analysis = _editable_base(analysis_id)
    lock_annotation_write(analysis)
    get_qualitative_document(analysis, str(document_id))
    if analysis.status != "concluida":
        abort(409)
    if request.content_length is None or request.content_length > 131_072:
        return jsonify({"error": "A seleção enviada excede o limite permitido."}), 400
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify({"error": "Dados inválidos."}), 400
    if (payload.get("analysis_id", analysis.id) != analysis.id
            or payload.get("document_id", str(document_id)) != str(document_id)
            or payload.get("page_number", page_number) != page_number):
        return jsonify({"error": "A seleção não pertence ao documento solicitado."}), 400
    action = payload.get("action")
    if action not in {"apply_codes", "create_code", "create_and_apply", "in_vivo", "create_memo"}:
        return jsonify({"error": "Ação de codificação inválida."}), 400
    try:
        page, start, end, quote = validated_selection(
            analysis, str(document_id), page_number, payload)
    except QualitativePageNotFoundError:
        abort(404)
    except CorpusUnavailableError:
        abort(409)
    except SelectionConflict as error:
        return jsonify({"error": str(error)}), 409
    except SelectionError as error:
        return jsonify({"error": str(error)}), 400

    codes = []
    new_code = None
    memo_text = None
    if action == "apply_codes":
        code_ids = payload.get("code_ids")
        if (not isinstance(code_ids, list) or not 1 <= len(code_ids) <= 30
                or any(not isinstance(identifier, str) for identifier in code_ids)):
            return jsonify({"error": "Escolha de um a trinta códigos deste projeto."}), 400
        codes = [get_qualitative_code(analysis, identifier) for identifier in dict.fromkeys(code_ids)]
        if any(not code.active for code in codes):
            return jsonify({"error": "Código inativo não pode ser aplicado."}), 400
    elif action in {"create_code", "create_and_apply"}:
        try:
            name, description = _code_fields()
        except ValueError as error:
            return jsonify({"error": str(error)}), 400
        if action == "create_code" and _duplicate_code(analysis, name):
            return jsonify({"error": "Já existe um código com esse nome neste projeto."}), 409
        existing_code = find_code_by_name(analysis, name) if action == "create_and_apply" else None
        if existing_code is not None:
            if not existing_code.active:
                return jsonify({"error": "Esse nome pertence a um código inativo."}), 409
            codes = [existing_code]
        else:
            new_code = QualitativeCode(analysis_id=analysis.id, name=name,
                                       description=description, created_by_user_id=current_user.id)
    elif action == "in_vivo":
        name = " ".join(quote.split())
        if not name or len(name) > 160 or len(normalized_qualitative_code_name(name)) > 160:
            return jsonify({"error": "O trecho é longo para nome de código. Edite uma sugestão no formulário Criar código.",
                            "suggested_name": name}), 422
        code = find_code_by_name(analysis, name)
        if code is not None and not code.active:
            return jsonify({"error": "Esse nome pertence a um código inativo."}), 409
        if code is not None:
            codes = [code]
        else:
            new_code = QualitativeCode(analysis_id=analysis.id, name=name,
                                       description="", created_by_user_id=current_user.id)
    else:
        try:
            memo_text = _input_text("text", 20000)
        except ValueError as error:
            return jsonify({"error": str(error)}), 400

    try:
        if new_code is not None:
            db.session.add(new_code)
            db.session.flush()
            codes = [new_code]
        excerpt = get_or_create_excerpt(analysis, str(document_id), page_number,
                                        start, end, quote, page["sha256"], current_user.id)
        if codes:
            apply_codes(analysis, excerpt, codes, current_user.id)
        if memo_text is not None:
            db.session.add(QualitativeMemo(analysis_id=analysis.id, excerpt_id=excerpt.id,
                                           text=memo_text, created_by_user_id=current_user.id))
        db.session.commit()
    except SelectionConflict as error:
        db.session.rollback()
        return jsonify({"error": str(error)}), 409
    except SelectionError as error:
        db.session.rollback()
        return jsonify({"error": str(error)}), 400
    except IntegrityError:
        db.session.rollback()
        return jsonify({"error": "A seleção ou o código foi alterado por outra operação. Tente novamente."}), 409
    return jsonify({"excerpt_id": excerpt.id, "records": _records_payload(analysis),
                    "page": page_excerpts(analysis, str(document_id), page_number)}), 201
