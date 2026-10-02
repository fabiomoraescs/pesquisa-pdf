"""Parsing estrito e prévias canônicas para importações Análysis Instagram."""

from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime, timezone
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from werkzeug.utils import secure_filename


IMPORT_SCHEMA_VERSION = "instagram-import-v1"
IMPORT_PROCESSOR_VERSION = "instagram-import-v1"
MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 20_000
MAX_CAPTION_CHARS = 8_000
MAX_PERMALINK_CHARS = 512
MAX_UNCOMPRESSED_XLSX_BYTES = 20 * 1024 * 1024
MAX_XLSX_MEMBERS = 1_000
PREVIEW_TTL_SECONDS = 15 * 60
MAX_COUNT_VALUE = 9_223_372_036_854_775_807

@dataclass(frozen=True)
class InstagramImportField:
    name: str
    required: bool
    kind: str
    description: str
    example: str
    group: str = "media"

    @property
    def format_label(self) -> str:
        return {"text": "Texto", "datetime": "ISO 8601 com timezone", "count": "Inteiro não negativo"}[self.kind]

    @property
    def note(self) -> str:
        if self.kind == "datetime":
            rule = ". Digite em ISO 8601 com timezone explícito; não converta para uma data visual comum do Excel."
        elif self.kind == "count":
            rule = ". Deixe vazio quando indisponível; use 0 apenas para zero real."
        else:
            rule = "."
        return (f"{'Obrigatório' if self.required else 'Opcional'}. {self.description} "
                f"Formato: {self.format_label}. Ex.: {self.example}{rule}")


# Única definição de colunas para parser, instruções HTML e modelo XLSX.
IMPORT_FIELDS = (
    InstagramImportField("account_username", True, "text", "Usuário da conta, sem @.", "revistax", "account"),
    InstagramImportField("media_external_id", True, "text", "ID estável da publicação.", "post_123"),
    InstagramImportField("published_at", True, "datetime", "Data e hora da publicação.", "2026-09-20T18:30:00-03:00"),
    InstagramImportField("observed_at", True, "datetime", "Data e hora da observação das métricas.", "2026-10-02T14:30:00-03:00"),
    InstagramImportField("account_display_name", False, "text", "Nome exibido da conta.", "Revista X", "account"),
    InstagramImportField("account_external_id", False, "text", "ID externo da conta.", "17841400000000000", "account"),
    InstagramImportField("media_type", False, "text", "Tipo de publicação.", "IMAGE"),
    InstagramImportField("permalink", False, "text", "Link da publicação.", "https://www.instagram.com/p/exemplo/"),
    InstagramImportField("caption", False, "text", "Legenda da publicação.", "Texto da publicação"),
    InstagramImportField("period_start", False, "datetime", "Início do período da métrica.", "2026-09-01T00:00:00-03:00"),
    InstagramImportField("period_end", False, "datetime", "Fim do período da métrica.", "2026-09-30T23:59:59-03:00"),
    *(InstagramImportField(name, False, "count", description, "1000", "media_metric") for name, description in (
        ("reach", "Contas alcançadas pela publicação."),
        ("impressions", "Impressões da publicação."),
        ("plays", "Reproduções da publicação."),
        ("likes", "Curtidas da publicação."),
        ("comments_count", "Comentários da publicação."),
        ("shares", "Compartilhamentos da publicação."),
        ("saves", "Salvamentos da publicação."),
        ("interactions", "Interações da publicação."),
        ("follows_generated", "Seguidores gerados pela publicação."),
    )),
    *(InstagramImportField(name, False, "count", description, "1000", "account_metric") for name, description in (
        ("followers_count", "Total de seguidores da conta."),
        ("new_followers", "Novos seguidores da conta."),
        ("account_reach", "Contas alcançadas pela conta."),
        ("profile_views", "Visualizações do perfil."),
        ("website_clicks", "Cliques no site da conta."),
    )),
)
MAX_IMPORT_COLUMNS = len(IMPORT_FIELDS)
REQUIRED_HEADERS = tuple(field.name for field in IMPORT_FIELDS if field.required)
OPTIONAL_HEADERS = tuple(field.name for field in IMPORT_FIELDS if not field.required)
CANONICAL_HEADERS = tuple(field.name for field in IMPORT_FIELDS)
MEDIA_METRIC_FIELDS = tuple(field.name for field in IMPORT_FIELDS if field.group == "media_metric")
ACCOUNT_METRIC_FIELDS = tuple(field.name for field in IMPORT_FIELDS if field.group == "account_metric")
_ISO_DATETIME = re.compile(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}(?::\d{2}(?:\.\d{1,6})?)?(?:Z|[+-]\d{2}:\d{2})\Z")


class InstagramImportError(ValueError):
    """Erro público de validação antes de qualquer escrita no banco."""


@dataclass(frozen=True)
class InstagramImportRow:
    line_number: int
    media_external_id: str
    published_at: datetime
    observed_at: datetime
    period_start: datetime | None
    period_end: datetime | None
    media_type: str | None
    permalink: str | None
    caption: str | None
    media_metrics: dict[str, int | None]


@dataclass(frozen=True)
class InstagramAccountSnapshotData:
    observed_at: datetime
    period_start: datetime | None
    period_end: datetime | None
    metrics: dict[str, int | None]


@dataclass(frozen=True)
class InstagramImportDataset:
    schema_version: str
    processor_version: str
    source_kind: str
    original_filename: str
    source_hash: str
    source_row_count: int
    rows: tuple[InstagramImportRow, ...]
    account_username: str
    account_username_normalized: str
    account_display_name: str | None
    account_external_id: str | None
    account_snapshot: InstagramAccountSnapshotData | None
    warnings: tuple[str, ...]

    @property
    def observed_at_min(self) -> datetime:
        return min(row.observed_at for row in self.rows)

    @property
    def observed_at_max(self) -> datetime:
        return max(row.observed_at for row in self.rows)

    @property
    def run_period_start(self) -> datetime | None:
        if any(row.period_start is None for row in self.rows):
            return None
        return min(row.period_start for row in self.rows if row.period_start is not None)

    @property
    def run_period_end(self) -> datetime | None:
        if any(row.period_end is None for row in self.rows):
            return None
        return max(row.period_end for row in self.rows if row.period_end is not None)


@dataclass(frozen=True)
class InstagramImportPreview:
    dataset: InstagramImportDataset
    known_media_count: int
    new_media_count: int
    repeated_source_hash: bool
    warnings: tuple[str, ...]


def normalized_instagram_username(value: str) -> tuple[str, str]:
    """Retorna username de apresentação e chave de identidade sem @ inicial."""
    if not isinstance(value, str):
        raise InstagramImportError("account_username: informe um texto válido.")
    display = value.strip()
    if display.startswith("@"):
        display = display[1:]
    if not display or display.startswith("@") or any(char.isspace() for char in display):
        raise InstagramImportError("account_username: informe um username Instagram válido.")
    normalized = unicodedata.normalize("NFKC", display).casefold()
    if len(display) > 100 or len(normalized) > 100:
        raise InstagramImportError("account_username: excede o limite de 100 caracteres.")
    return display, normalized


def sanitized_import_filename(filename: str) -> str:
    """Remove caminhos e caracteres inseguros antes de registrar a proveniência."""
    original = (filename or "").replace("\\", "/").rsplit("/", 1)[-1].strip()
    safe = secure_filename(original)
    if not safe or len(safe) > 255:
        raise InstagramImportError("Informe um nome de arquivo válido.")
    return safe


def _optional_text(value: object, location: str, *, limit: int) -> str | None:
    if value is None:
        return None
    if not isinstance(value, str):
        raise InstagramImportError(f"{location}: informe um texto válido.")
    text = value.strip()
    if not text:
        return None
    if len(text) > limit:
        raise InstagramImportError(f"{location}: excede o limite de {limit} caracteres.")
    return text


def _required_text(value: object, location: str, *, limit: int) -> str:
    text = _optional_text(value, location, limit=limit)
    if text is None:
        raise InstagramImportError(f"{location}: campo obrigatório vazio.")
    return text


def _iso_datetime(value: object, location: str, *, required: bool = False) -> datetime | None:
    if value is None or value == "":
        if required:
            raise InstagramImportError(f"{location}: campo obrigatório vazio.")
        return None
    if not isinstance(value, str):
        raise InstagramImportError(f"{location}: use data e hora ISO 8601 com timezone.")
    text = value.strip()
    if not text:
        if required:
            raise InstagramImportError(f"{location}: campo obrigatório vazio.")
        return None
    if not _ISO_DATETIME.fullmatch(text):
        raise InstagramImportError(f"{location}: use data e hora ISO 8601 com timezone, por exemplo 2026-10-02T14:30:00Z.")
    try:
        parsed = datetime.fromisoformat(f"{text[:-1]}+00:00" if text.endswith("Z") else text)
    except ValueError as error:
        raise InstagramImportError(f"{location}: data e hora inválida.") from error
    return parsed.astimezone(timezone.utc)


def _optional_count(value: object, location: str) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise InstagramImportError(f"{location}: informe um número inteiro não negativo.")
    if isinstance(value, int):
        number = value
    elif isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        if text.startswith("-") and text[1:].isdecimal():
            raise InstagramImportError(f"{location}: não pode ser negativo.")
        if not text.isdecimal():
            raise InstagramImportError(f"{location}: informe um número inteiro não negativo.")
        number = int(text)
    else:
        raise InstagramImportError(f"{location}: informe um número inteiro não negativo.")
    if number < 0:
        raise InstagramImportError(f"{location}: não pode ser negativo.")
    if number > MAX_COUNT_VALUE:
        raise InstagramImportError(f"{location}: excede o limite de uma métrica inteira.")
    return number


def _validate_headers(headers: list[object], location: str) -> tuple[str, ...]:
    if not headers:
        raise InstagramImportError(f"{location}: cabeçalhos ausentes.")
    if len(headers) > MAX_IMPORT_COLUMNS:
        raise InstagramImportError(f"{location}: o arquivo excede o limite de colunas.")
    if any(not isinstance(header, str) for header in headers):
        raise InstagramImportError(f"{location}: todos os cabeçalhos devem ser texto.")
    normalized = tuple(header.strip() for header in headers)
    if len(set(normalized)) != len(normalized):
        raise InstagramImportError(f"{location}: há cabeçalhos duplicados.")
    missing = [header for header in REQUIRED_HEADERS if header not in normalized]
    if missing:
        raise InstagramImportError(f"Coluna obrigatória ausente: {missing[0]}.")
    unexpected = [header for header in normalized if header not in CANONICAL_HEADERS]
    if unexpected:
        raise InstagramImportError(f"Coluna não reconhecida neste formato: {unexpected[0]}.")
    return normalized


def _csv_rows(data: bytes) -> list[tuple[int, dict[str, object]]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise InstagramImportError("O CSV deve estar codificado em UTF-8.") from error
    try:
        reader = csv.reader(io.StringIO(text, newline=""))
        headers = _validate_headers(next(reader, []), "CSV, linha 1")
        rows: list[tuple[int, dict[str, object]]] = []
        for line_number, values in enumerate(reader, start=2):
            if line_number > MAX_IMPORT_ROWS + 1:
                raise InstagramImportError("O CSV excede o limite de linhas.")
            if not any(value.strip() for value in values):
                continue
            if len(values) != len(headers):
                raise InstagramImportError(f"CSV, linha {line_number}: quantidade de colunas inválida.")
            rows.append((line_number, dict(zip(headers, values))))
    except csv.Error as error:
        raise InstagramImportError("O CSV enviado é inválido.") from error
    return rows


def _safe_xlsx_rows(data: bytes) -> list[tuple[int, dict[str, object]]]:
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if (len(members) > MAX_XLSX_MEMBERS
                    or sum(member.file_size for member in members) > MAX_UNCOMPRESSED_XLSX_BYTES):
                raise InstagramImportError("A planilha descompactada excede o limite permitido.")
            prohibited = ("vbaproject.bin", "externallinks/", "oleobjects/", "embeddings/")
            if any(any(token in member.filename.casefold() for token in prohibited) for member in members):
                raise InstagramImportError("A planilha contém macros, conteúdo externo ou objetos incorporados não permitidos.")
            for member in members:
                if member.filename.casefold().endswith(".rels"):
                    relationships = ElementTree.fromstring(archive.read(member))
                    if any(item.attrib.get("TargetMode", "").casefold() == "external" for item in relationships):
                        raise InstagramImportError("A planilha contém links externos não permitidos.")
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except InstagramImportError:
        raise
    except (BadZipFile, ValueError, OSError, KeyError, ElementTree.ParseError) as error:
        raise InstagramImportError("O arquivo enviado não é um XLSX válido.") from error
    try:
        if len(workbook.sheetnames) != 1:
            raise InstagramImportError("A importação XLSX aceita exatamente uma aba na versão atual.")
        sheet = workbook.active
        if sheet.sheet_state != "visible":
            raise InstagramImportError("A aba da planilha deve estar visível.")
        if ((sheet.max_row is not None and sheet.max_row > MAX_IMPORT_ROWS + 1)
                or (sheet.max_column is not None and sheet.max_column > MAX_IMPORT_COLUMNS)):
            raise InstagramImportError("A planilha excede o limite de linhas ou colunas.")
        iterator = sheet.iter_rows(values_only=False)
        first = next(iterator, None)
        if first is None:
            raise InstagramImportError("XLSX, linha 1: cabeçalhos ausentes.")
        headers = _validate_headers([cell.value for cell in first], "XLSX, linha 1")
        rows = []
        for line_number, cells in enumerate(iterator, start=2):
            if line_number > MAX_IMPORT_ROWS + 1:
                raise InstagramImportError("A planilha excede o limite de linhas.")
            if any(cell.data_type == "f" for cell in cells):
                raise InstagramImportError(f"XLSX, linha {line_number}: fórmulas não são permitidas.")
            if any(getattr(cell, "hyperlink", None) is not None for cell in cells):
                raise InstagramImportError(f"XLSX, linha {line_number}: hyperlinks não são permitidos.")
            values = [cell.value for cell in cells]
            if all(value is None or value == "" for value in values):
                continue
            if len(values) != len(headers):
                raise InstagramImportError(f"XLSX, linha {line_number}: quantidade de colunas inválida.")
            rows.append((line_number, dict(zip(headers, values))))
    except InstagramImportError:
        raise
    except (ValueError, OSError, KeyError, IndexError, TypeError, AttributeError, ElementTree.ParseError) as error:
        raise InstagramImportError("O arquivo enviado não é um XLSX válido.") from error
    finally:
        workbook.close()
    return rows


def _single_optional_value(rows: list[tuple[int, dict[str, object]]], field: str, *, limit: int) -> str | None:
    found: list[tuple[int, str]] = []
    for line_number, raw in rows:
        value = _optional_text(raw.get(field), f"Linha {line_number}: {field}", limit=limit)
        if value is not None:
            found.append((line_number, value))
    unique = {value for _line, value in found}
    if len(unique) > 1:
        raise InstagramImportError(f"O arquivo possui valores divergentes para {field}.")
    return found[0][1] if found else None


def _account_snapshot(rows: list[InstagramImportRow], source_rows: list[tuple[int, dict[str, object]]]) -> InstagramAccountSnapshotData | None:
    snapshots = []
    for line_number, raw in source_rows:
        metrics = {field: _optional_count(raw.get(field), f"Linha {line_number}: {field}")
                   for field in ACCOUNT_METRIC_FIELDS}
        snapshots.append((line_number, metrics))
    if not any(any(value is not None for value in metrics.values()) for _line, metrics in snapshots):
        return None
    first_metrics = snapshots[0][1]
    if any(metrics != first_metrics for _line, metrics in snapshots[1:]):
        raise InstagramImportError("As métricas da conta devem ser idênticas em todas as linhas do arquivo.")
    first = rows[0]
    if any((row.observed_at, row.period_start, row.period_end)
           != (first.observed_at, first.period_start, first.period_end) for row in rows[1:]):
        raise InstagramImportError(
            "Para registrar métricas da conta, observed_at e período devem ser idênticos em todas as linhas."
        )
    return InstagramAccountSnapshotData(
        observed_at=first.observed_at,
        period_start=first.period_start,
        period_end=first.period_end,
        metrics=first_metrics,
    )


def _same_media_observation(left: InstagramImportRow, right: InstagramImportRow) -> bool:
    """A linha física não participa da identidade de uma duplicata de arquivo."""
    return (
        left.media_external_id,
        left.published_at,
        left.observed_at,
        left.period_start,
        left.period_end,
        left.media_type,
        left.permalink,
        left.caption,
        left.media_metrics,
    ) == (
        right.media_external_id,
        right.published_at,
        right.observed_at,
        right.period_start,
        right.period_end,
        right.media_type,
        right.permalink,
        right.caption,
        right.media_metrics,
    )


def _normalized_dataset(
    source_rows: list[tuple[int, dict[str, object]]], *, source_kind: str, original_filename: str, source_hash: str
) -> InstagramImportDataset:
    if not source_rows:
        raise InstagramImportError("O arquivo não possui linhas de dados.")
    accounts: list[tuple[int, str, str]] = []
    normalized_rows: list[InstagramImportRow] = []
    for line_number, raw in source_rows:
        username, normalized_username = normalized_instagram_username(
            _required_text(raw.get("account_username"), f"Linha {line_number}: account_username", limit=100)
        )
        accounts.append((line_number, username, normalized_username))
        published_at = _iso_datetime(raw.get("published_at"), f"Linha {line_number}: published_at", required=True)
        observed_at = _iso_datetime(raw.get("observed_at"), f"Linha {line_number}: observed_at", required=True)
        period_start = _iso_datetime(raw.get("period_start"), f"Linha {line_number}: period_start")
        period_end = _iso_datetime(raw.get("period_end"), f"Linha {line_number}: period_end")
        if (period_start is None) != (period_end is None):
            raise InstagramImportError(f"Linha {line_number}: period_start e period_end devem ser informados juntos.")
        if period_start is not None and period_start > period_end:
            raise InstagramImportError(f"Linha {line_number}: period_start não pode ser posterior a period_end.")
        normalized_rows.append(InstagramImportRow(
            line_number=line_number,
            media_external_id=_required_text(
                raw.get("media_external_id"), f"Linha {line_number}: media_external_id", limit=128
            ),
            published_at=published_at,
            observed_at=observed_at,
            period_start=period_start,
            period_end=period_end,
            media_type=_optional_text(raw.get("media_type"), f"Linha {line_number}: media_type", limit=40),
            permalink=_optional_text(raw.get("permalink"), f"Linha {line_number}: permalink", limit=MAX_PERMALINK_CHARS),
            caption=_optional_text(raw.get("caption"), f"Linha {line_number}: caption", limit=MAX_CAPTION_CHARS),
            media_metrics={field: _optional_count(raw.get(field), f"Linha {line_number}: {field}")
                           for field in MEDIA_METRIC_FIELDS},
        ))
    usernames = {normalized for _line, _username, normalized in accounts}
    if len(usernames) != 1:
        raise InstagramImportError("O arquivo contém mais de uma conta Instagram.")
    unique_rows: list[InstagramImportRow] = []
    seen: dict[str, InstagramImportRow] = {}
    warnings: list[str] = []
    for row in normalized_rows:
        previous = seen.get(row.media_external_id)
        if previous is None:
            seen[row.media_external_id] = row
            unique_rows.append(row)
        elif _same_media_observation(previous, row):
            warnings.append(f"Linha {row.line_number} duplicada da linha {previous.line_number}; será consolidada.")
        else:
            raise InstagramImportError(
                f"Linha {row.line_number}: media_external_id '{row.media_external_id}' conflita com a linha {previous.line_number}."
            )
    # A primeira grafia serve apenas para apresentação; a identidade usa a chave normalizada.
    _line, account_username, account_username_normalized = accounts[0]
    account_display_name = _single_optional_value(source_rows, "account_display_name", limit=200)
    account_external_id = _single_optional_value(source_rows, "account_external_id", limit=128)
    # Também valide linhas de mídia consolidadas: métricas de conta divergentes
    # nelas não podem desaparecer junto com a duplicata.
    snapshot = _account_snapshot(unique_rows, source_rows)
    return InstagramImportDataset(
        schema_version=IMPORT_SCHEMA_VERSION,
        processor_version=IMPORT_PROCESSOR_VERSION,
        source_kind=source_kind,
        original_filename=original_filename,
        source_hash=source_hash,
        source_row_count=len(source_rows),
        rows=tuple(unique_rows),
        account_username=account_username,
        account_username_normalized=account_username_normalized,
        account_display_name=account_display_name,
        account_external_id=account_external_id,
        account_snapshot=snapshot,
        warnings=tuple(warnings),
    )


def parse_instagram_import(filename: str, data: bytes) -> InstagramImportDataset:
    """Lê e valida integralmente CSV/XLSX sem tocar no banco de dados."""
    original_filename = sanitized_import_filename(filename)
    if not data or len(data) > MAX_IMPORT_BYTES:
        raise InstagramImportError("Envie um arquivo CSV ou XLSX de até 2 MB.")
    suffix = original_filename.rsplit(".", 1)[-1].casefold() if "." in original_filename else ""
    if suffix == "csv":
        rows = _csv_rows(data)
        source_kind = "csv"
    elif suffix == "xlsx":
        rows = _safe_xlsx_rows(data)
        source_kind = "xlsx"
    else:
        raise InstagramImportError("Envie somente arquivos com extensão .csv ou .xlsx.")
    return _normalized_dataset(
        rows,
        source_kind=source_kind,
        original_filename=original_filename,
        source_hash=hashlib.sha256(data).hexdigest(),
    )


def build_preview(
    dataset: InstagramImportDataset, *, known_media_external_ids: set[str], repeated_source_hash: bool
) -> InstagramImportPreview:
    known_media_count = sum(row.media_external_id in known_media_external_ids for row in dataset.rows)
    warnings = list(dataset.warnings)
    if repeated_source_hash:
        warnings.append("Este arquivo parece já ter sido importado anteriormente neste projeto.")
    return InstagramImportPreview(
        dataset=dataset,
        known_media_count=known_media_count,
        new_media_count=len(dataset.rows) - known_media_count,
        repeated_source_hash=repeated_source_hash,
        warnings=tuple(warnings),
    )
