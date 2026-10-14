"""Relatórios Instagram reproduzíveis, com artefatos privados e snapshots imutáveis.

Relatórios não recalculam o passado na leitura. O serviço seleciona apenas
``AnalyticsRun`` concluídos, congela valores, limitações e identificadores de
origem e gera os artefatos a partir desse mesmo snapshot.
"""

from __future__ import annotations

from datetime import date, datetime, time, timezone
from io import BytesIO
import os
from pathlib import Path
from shutil import rmtree
from statistics import median
from uuid import UUID, uuid4

from flask import current_app
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import joinedload, selectinload

from .extensions import db
from .instagram_action_evaluations import (
    ASSESSMENT_LABELS,
    RAW_METRIC_LABELS,
    format_metric_value,
    metric_label,
)
from .instagram_actions import ACTION_KIND_LABELS, ACTION_STATUS_LABELS
from .instagram_cohorts import cohort_label
from .instagram_learnings import LEARNING_ASSESSMENT_LABELS, learning_cohort_label
from .instagram_metrics import RATE_METRIC_DEFINITIONS, calculate_derived_metrics
from .instagram_periods import effective_run_period
from .pagination import PlatformPage, requested_page
from .models import (
    AnalyticsProject,
    AnalyticsRun,
    InstagramCommentObservation,
    InstagramInsightFinding,
    InstagramInsightRun,
    InstagramLearning,
    InstagramMediaObservation,
    InstagramRecommendation,
    InstagramRecommendationAction,
    InstagramRecommendationActionEvaluation,
    InstagramRecommendationRun,
    InstagramReport,
)


# Os valores de tipo/modo permanecem apenas porque o schema atual os exige e
# relatórios históricos continuam usando esses campos; o novo fluxo não os
# expõe nem permite escolhê-los manualmente.
INSTAGRAM_REPORT_VERSION = "instagram-report-v2"
REPORT_TYPES = {"executive", "analytical"}
REPORT_TYPE_LABELS = {"executive": "Executivo", "analytical": "Analítico"}
ANALYSIS_MODES = {"single_period", "comparative_evolution"}
ANALYSIS_MODE_LABELS = {
    "single_period": "Período único",
    "comparative_evolution": "Comparativo e evolução",
}
REPORT_FORMATS = {"pdf", "xlsx"}
SECTION_LABELS = {
    "summary": "Resumo executivo",
    "metrics": "Métricas e cobertura",
    "contents": "Conteúdos observados",
    "comments": "Comentários e recepção",
    "findings": "Achados",
    "recommendations": "Recomendações",
    "actions": "Ações e testes",
    "evaluations": "Acompanhamentos",
    "learnings": "Aprendizados",
    "methodology": "Metodologia e limitações",
    "comparison": "Comparação temporal",
    "evolution": "Evolução entre períodos",
    "milestones": "Marcos e decisões",
}
DEFAULT_REPORT_SECTIONS = (
    "summary", "metrics", "contents", "comments", "findings", "recommendations",
    "actions", "evaluations", "learnings", "methodology",
)
DEFAULT_EXECUTIVE_SECTIONS = DEFAULT_REPORT_SECTIONS
DEFAULT_ANALYTICAL_SECTIONS = DEFAULT_REPORT_SECTIONS
RAW_METRICS = tuple(RAW_METRIC_LABELS)


class InstagramReportError(ValueError):
    """Solicitação de relatório inválida ou estado não permitido."""


class InstagramReportPersistenceError(RuntimeError):
    """Falha atômica na criação ou remoção de relatório."""


def _utc(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)


def _iso(value: datetime | None) -> str | None:
    value = _utc(value)
    return value.isoformat() if value is not None else None


def _display_date(value: str | datetime | None) -> str:
    if isinstance(value, str):
        try:
            value = datetime.fromisoformat(value)
        except ValueError:
            return "Não informado"
    value = _utc(value)
    return value.strftime("%d/%m/%Y %H:%M") if value is not None else "Não informado"


def _number(value: float | int | None) -> str:
    return format_metric_value(value, None)


def _percent(value: float | None, *, signed: bool = False) -> str:
    return format_metric_value(value, "%", signed=signed, percentage_fraction=True, fixed_decimals=2)


def _parse_requested_period(values: dict[str, object]) -> tuple[date | None, date | None]:
    if values.get("period_preset") in {"all", "all_history"}:
        return None, None
    try:
        start = date.fromisoformat(str(values.get("period_start") or ""))
        end = date.fromisoformat(str(values.get("period_end") or ""))
    except ValueError as error:
        raise InstagramReportError("Informe datas inicial e final válidas para o relatório.") from error
    if start > end:
        raise InstagramReportError("A data inicial não pode ser posterior à data final.")
    return start, end


def _period_boundaries(start: date, end: date) -> tuple[datetime, datetime]:
    """Limites inclusivos UTC para as datas de calendário enviadas pelo formulário."""
    return (
        datetime.combine(start, time.min, tzinfo=timezone.utc),
        datetime.combine(end, time.max, tzinfo=timezone.utc),
    )


def _run_period_sort_key(run: AnalyticsRun) -> tuple[int, datetime, datetime, str]:
    resolved = effective_run_period(run)
    minimum = datetime.min.replace(tzinfo=timezone.utc)
    maximum = datetime.max.replace(tzinfo=timezone.utc)
    if resolved.is_bounded:
        return 0, resolved.start or minimum, resolved.end or maximum, run.id
    return 1, maximum, maximum, run.id


def _latest_run(winner: AnalyticsRun | None, candidate: AnalyticsRun) -> AnalyticsRun:
    if winner is None:
        return candidate
    minimum = datetime.min.replace(tzinfo=timezone.utc)
    candidate_key = (_utc(candidate.completed_at) or minimum, _utc(candidate.created_at) or minimum, candidate.id)
    winner_key = (_utc(winner.completed_at) or minimum, _utc(winner.created_at) or minimum, winner.id)
    return candidate if candidate_key > winner_key else winner


def eligible_report_runs(project: AnalyticsProject, *, include_all_history: bool = False) -> tuple[AnalyticsRun, ...]:
    """Runs concluídos, temporalmente conhecidos e deduplicados.

    Meta usa a janela congelada no JSON.  ``all_history`` só entra quando o
    usuário pediu esse escopo explicitamente; um legado sem intervalo continua
    fora da seleção para não converter a data da execução em data analítica.
    """
    runs = db.session.scalars(select(AnalyticsRun).where(
        AnalyticsRun.analytics_project_id == project.id,
        AnalyticsRun.status == "completed",
    )).all()
    winners: dict[tuple[object, ...], AnalyticsRun] = {}
    for run in runs:
        resolved = effective_run_period(run)
        if resolved.is_bounded:
            key: tuple[object, ...] = ("bounded", resolved.start, resolved.end)
        elif include_all_history and resolved.kind == "all_history":
            key = ("all_history",)
        else:
            continue
        winners[key] = _latest_run(winners.get(key), run)
    return tuple(sorted(winners.values(), key=_run_period_sort_key))


def resolve_report_runs(project: AnalyticsProject, start: date | None, end: date | None) -> tuple[AnalyticsRun, ...]:
    """Seleciona runs completos para uma janela ou para todo o histórico explícito."""
    if start is None or end is None:
        if start is not None or end is not None:
            raise InstagramReportError("O período do relatório está incompleto.")
        return eligible_report_runs(project, include_all_history=True)
    lower, upper = _period_boundaries(start, end)
    return tuple(
        run for run in eligible_report_runs(project)
        if (resolved := effective_run_period(run)).is_bounded
        and (resolved.start or lower) >= lower and (resolved.end or upper) <= upper
    )


def _metric_value(observation: InstagramMediaObservation, key: str) -> float | int | None:
    if key in RAW_METRIC_LABELS:
        return getattr(observation, key)
    return calculate_derived_metrics({field: getattr(observation, field) for field in RAW_METRICS}).get(key)


def _run_metrics(project: AnalyticsProject, run: AnalyticsRun) -> tuple[dict, list[dict]]:
    observations = db.session.scalars(select(InstagramMediaObservation).options(
        joinedload(InstagramMediaObservation.instagram_media)
    ).where(InstagramMediaObservation.analytics_run_id == run.id)).all()
    observations.sort(key=lambda item: (item.instagram_media.published_at or datetime.min.replace(tzinfo=timezone.utc), item.instagram_media.id))
    metrics: list[dict] = []
    for key in (*RAW_METRICS, *(definition.identifier for definition in RATE_METRIC_DEFINITIONS)):
        values = [_metric_value(item, key) for item in observations]
        known = [item for item in values if item is not None]
        is_rate = key not in RAW_METRIC_LABELS
        metrics.append({
            "key": key,
            "label": metric_label(key) or key,
            "unit": "%" if is_rate else "contagem",
            "aggregation": "mediana por conteúdo observado",
            "value": float(median(known)) if known else None,
            "known_count": len(known),
            "total_count": len(observations),
            "coverage": (len(known) / len(observations)) if observations else None,
        })
    contents = [{
        "media_id": item.instagram_media_id,
        "media_type": item.instagram_media.media_type_raw or "Não informado",
        "published_at": _iso(item.instagram_media.published_at),
        "permalink": item.instagram_media.permalink,
        "caption": (item.instagram_media.caption_original or "")[:500],
        "metrics": {key: _metric_value(item, key) for key in RAW_METRICS},
    } for item in observations]
    return {"run_id": run.id, "metrics": metrics}, contents


def _run_snapshot(run: AnalyticsRun) -> dict:
    resolved = effective_run_period(run)
    return {
        "id": run.id, "source_kind": run.source_kind, "status": run.status,
        "period_start": _iso(resolved.start), "period_end": _iso(resolved.end),
        "period_kind": resolved.kind, "period_source": resolved.source,
        "period_preset": resolved.preset,
        "completed_at": _iso(run.completed_at), "created_at": _iso(run.created_at),
        "record_count": run.record_count, "processor_version": run.processor_version,
    }


def _findings_for_project(project: AnalyticsProject) -> list[dict]:
    findings = db.session.scalars(select(InstagramInsightFinding).join(InstagramInsightRun).where(
        InstagramInsightRun.analytics_project_id == project.id
    ).options(joinedload(InstagramInsightFinding.insight_run)).order_by(
        InstagramInsightFinding.created_at.asc(), InstagramInsightFinding.id.asc()
    )).all()
    return [{
        "id": item.id, "insight_run_id": item.instagram_insight_run_id,
        "created_at": _iso(item.created_at), "title": item.title, "summary": item.summary,
        "metric_key": item.metric_key, "metric_label": metric_label(item.metric_key),
        "support_level": item.support_level, "coverage": item.coverage,
        "group_n": item.group_n, "comparison_n": item.comparison_n,
        "cohort": item.cohort_json, "source_ids": {"insight_run_id": item.instagram_insight_run_id},
    } for item in findings]


def _recommendations_for_project(project: AnalyticsProject) -> list[dict]:
    rows = db.session.scalars(select(InstagramRecommendation).join(InstagramRecommendationRun).where(
        InstagramRecommendationRun.analytics_project_id == project.id
    ).order_by(InstagramRecommendation.created_at.asc(), InstagramRecommendation.id.asc())).all()
    return [{
        "id": item.id, "recommendation_run_id": item.instagram_recommendation_run_id,
        "created_at": _iso(item.created_at), "title": item.title, "summary": item.summary,
        "priority": item.priority, "support_level": item.support_level,
        "limitations": list(item.limitations_json or []),
    } for item in rows]


def _actions_for_project(project: AnalyticsProject) -> tuple[list[dict], list[dict]]:
    actions = db.session.scalars(select(InstagramRecommendationAction).where(
        InstagramRecommendationAction.analytics_project_id == project.id
    ).order_by(InstagramRecommendationAction.created_at.asc(), InstagramRecommendationAction.id.asc())).all()
    action_data = [{
        "id": item.id, "recommendation_id": item.instagram_recommendation_id,
        "created_at": _iso(item.created_at), "title": item.title, "kind": item.action_kind,
        "kind_label": ACTION_KIND_LABELS.get(item.action_kind, item.action_kind), "status": item.status,
        "status_label": ACTION_STATUS_LABELS.get(item.status, item.status),
        "metric_key": item.target_metric, "metric_label": metric_label(item.target_metric),
        "cohort": item.cohort_json, "baseline_value": item.baseline_value,
    } for item in actions]
    action_ids = [item.id for item in actions]
    evaluations = db.session.scalars(select(InstagramRecommendationActionEvaluation).where(
        InstagramRecommendationActionEvaluation.instagram_recommendation_action_id.in_(action_ids or ["-"])
    ).order_by(InstagramRecommendationActionEvaluation.created_at.asc(), InstagramRecommendationActionEvaluation.id.asc())).all()
    evaluation_data = [{
        "id": item.id, "action_id": item.instagram_recommendation_action_id,
        "analytics_run_id": item.analytics_run_id, "created_at": _iso(item.created_at),
        "metric_key": item.metric_key, "metric_label": metric_label(item.metric_key),
        "cohort": item.cohort_json, "baseline_value": item.baseline_value,
        "followup_value": item.followup_value, "absolute_change": item.absolute_change,
        "relative_change": item.relative_change, "assessment": item.assessment,
        "assessment_label": ASSESSMENT_LABELS.get(item.assessment, item.assessment),
        "limitations": list(item.limitations_json or []),
    } for item in evaluations]
    return action_data, evaluation_data


def _learnings_for_project(project: AnalyticsProject) -> list[dict]:
    rows = db.session.scalars(select(InstagramLearning).where(
        InstagramLearning.analytics_project_id == project.id
    ).options(selectinload(InstagramLearning.evaluations)).order_by(
        InstagramLearning.created_at.asc(), InstagramLearning.id.asc()
    )).all()
    return [{
        "id": item.id, "created_at": _iso(item.created_at), "title": item.title,
        "summary": item.learning_summary, "assessment": item.assessment,
        "assessment_label": LEARNING_ASSESSMENT_LABELS.get(item.assessment, item.assessment),
        "metric_key": item.metric_key, "metric_label": metric_label(item.metric_key),
        "cohort": item.cohort_json, "evaluation_ids": [evaluation.id for evaluation in item.evaluations],
    } for item in rows]


def _in_range(value: str | None, start: datetime | None, end: datetime | None) -> bool:
    if not value:
        return False
    try:
        actual = _utc(datetime.fromisoformat(value))
    except ValueError:
        return False
    return actual is not None and start is not None and end is not None and start <= actual <= end


def _milestones(source: dict, baseline: AnalyticsRun, final: AnalyticsRun) -> list[dict]:
    baseline_period, final_period = effective_run_period(baseline), effective_run_period(final)
    start = baseline_period.start if baseline_period.is_bounded else None
    end = final_period.end if final_period.is_bounded else None
    milestone_rows: list[dict] = []
    for kind, values in (
        ("Recomendação", source["recommendations"]), ("Ação/Teste", source["actions"]),
        ("Acompanhamento", source["evaluations"]), ("Aprendizado", source["learnings"]),
    ):
        for item in values:
            # Todo o histórico é explícito e ilimitado; um legado desconhecido
            # nunca chega a este ponto como fonte de relatório.
            if (start is None or end is None) or _in_range(item.get("created_at"), start, end):
                milestone_rows.append({
                    "at": item["created_at"], "kind": kind, "id": item["id"], "title": item.get("title") or item.get("assessment_label") or kind,
                })
    return sorted(milestone_rows, key=lambda item: (item["at"], item["kind"], item["id"]))


def _comparison(baseline_metrics: list[dict], final_metrics: list[dict]) -> list[dict]:
    by_key = {item["key"]: item for item in final_metrics}
    rows = []
    for initial in baseline_metrics:
        final = by_key.get(initial["key"])
        final_value = final.get("value") if final else None
        initial_value = initial.get("value")
        absolute = (final_value - initial_value) if initial_value is not None and final_value is not None else None
        relative = (absolute / initial_value) if absolute is not None and initial_value != 0 else None
        comparable = initial_value is not None and final_value is not None
        rows.append({
            "key": initial["key"], "label": initial["label"], "unit": initial["unit"],
            "baseline_value": initial_value, "final_value": final_value,
            "absolute_change": absolute, "relative_change": relative,
            "comparable": comparable,
            "limitation": "Baseline igual a zero: variação relativa indisponível." if comparable and initial_value == 0 else None,
        })
    return rows


def _chart_number(value: float | int | None, unit: str | None = None) -> str:
    """Apresenta valores dos gráficos sem alterar a representação do snapshot."""
    return _percent(value) if unit == "%" else _number(value)


def report_summary_cards(snapshot: dict, *, limit: int = 3) -> list[dict]:
    """Retorna indicadores executivos exclusivamente a partir do snapshot.

    A ordem explicita as métricas de maior leitura gerencial e evita que a
    ordem incidental do JSON decida quais valores aparecem no resumo.
    """
    final = snapshot.get("last_observation") or snapshot.get("final_run") or {}
    comparison_by_key = {
        item.get("key"): item for item in snapshot.get("comparison", ())
        if item.get("key") and item.get("final_value") is not None
    }
    final_metrics = {
        item.get("key"): item for item in (
            ((snapshot.get("metrics_by_run") or {}).get(final.get("id"), {}) or {}).get("metrics", ()) or ()
        ) if item.get("key") and item.get("value") is not None
    }
    cards = []
    for key in RAW_METRICS:
        item = comparison_by_key.get(key)
        value = item.get("final_value") if item else None
        unit = item.get("unit") if item else None
        label = item.get("label") if item else None
        if value is None:
            metric = final_metrics.get(key)
            if metric is None:
                continue
            value, unit, label = metric.get("value"), metric.get("unit"), metric.get("label")
        if value is None:
            continue
        cards.append({
            "key": key,
            "label": label or RAW_METRIC_LABELS[key],
            "value": value,
            "value_display": _chart_number(value, unit),
            "context": "Última observação selecionada",
        })
        if len(cards) >= limit:
            break
    return cards


def _series_label(run: dict) -> str:
    if run.get("period_kind") == "all_history":
        return "Todo o histórico disponível"
    start, end = _display_date(run.get("period_start")), _display_date(run.get("period_end"))
    return start if start == end else f"{start} — {end}"


def report_chart_data(snapshot: dict) -> dict[str, dict]:
    """Especificações determinísticas de gráficos a partir do snapshot imutável.

    Não há consulta ao banco nesta função. Assim, a tela apresenta somente o
    estado congelado na geração, inclusive depois de novas importações.
    """
    final = snapshot.get("last_observation") or snapshot.get("final_run") or {}
    contents = list((snapshot.get("contents_by_run") or {}).get(final.get("id"), ()) or ())
    charts: dict[str, dict] = {}

    final_metrics = list(((snapshot.get("metrics_by_run") or {}).get(final.get("id"), {}) or {}).get("metrics", ()) or ())
    primary_metrics = [
        item for item in final_metrics
        if item.get("key") in RAW_METRIC_LABELS and item.get("value") is not None
    ][:4]
    if primary_metrics:
        values = [float(item["value"]) for item in primary_metrics]
        charts["period_metrics"] = {
            "kind": "bar", "title": "Indicadores principais",
            "labels": [item["label"] for item in primary_metrics], "values": values,
            "unit": "contagem",
            "description": "Medianas por conteúdo observado na última observação selecionada.",
        }

    distribution: dict[str, int] = {}
    for item in contents:
        label = str(item.get("media_type") or "Não informado")
        distribution[label] = distribution.get(label, 0) + 1
    if 2 <= len(distribution) <= 6:
        labels, values = list(distribution), list(distribution.values())
        total = sum(values)
        charts["media_type_distribution"] = {
            "kind": "pie", "title": "Distribuição por tipo de mídia", "labels": labels, "values": values,
            "total": total, "shares": [value / total * 100 for value in values],
            "description": "Participação dos conteúdos observados na última observação selecionada.",
        }

    for metric_key in ("interactions", "saves", "reach", "impressions", "comments"):
        groups: dict[str, list[float | int]] = {}
        for item in contents:
            value = (item.get("metrics") or {}).get(metric_key)
            if value is not None:
                groups.setdefault(str(item.get("media_type") or "Não informado"), []).append(value)
        if len(groups) >= 2:
            labels = list(groups)
            values = [float(median(groups[label])) for label in labels]
            charts["media_type_performance"] = {
                "kind": "bar", "title": f"Desempenho por tipo de mídia — {RAW_METRIC_LABELS[metric_key]}",
                "labels": labels, "values": values, "unit": "contagem",
                "description": f"Mediana de {RAW_METRIC_LABELS[metric_key].casefold()} entre os conteúdos observados de cada tipo.",
            }
            break

    comparison = [item for item in snapshot.get("comparison", ()) if item.get("relative_change") is not None]
    if comparison:
        comparison = comparison[:6]
        charts["comparison"] = {
            "kind": "diverging_bar", "title": "Comparação no período",
            "labels": [item["label"] for item in comparison],
            "values": [round(float(item["relative_change"]) * 100, 4) for item in comparison], "unit": "%",
            "description": "Variação relativa entre a primeira e a última observação selecionadas.",
        }
        charts["comparison"]["zero_value"] = 0

    evolution = list(snapshot.get("evolution", ()) or ())
    legacy_evolution = snapshot.get("schema_version") == "instagram-report-v1"
    if len(evolution) >= 3 or (legacy_evolution and len(evolution) >= 2):
        preferred = ("interactions", "saves", "reach", "impressions", "comments")
        for metric_key in preferred:
            points = []
            unit = None
            for point in evolution:
                metric = next((item for item in point.get("metrics", ()) if item.get("key") == metric_key), None)
                if metric is not None and metric.get("value") is not None:
                    points.append((point.get("run") or {}, metric))
                    unit = metric.get("unit")
            if len(points) >= 2:
                charts["evolution"] = {
                    "kind": "line", "title": "Evolução no período",
                    "labels": [_series_label(run) for run, _ in points],
                    "values": [float(metric["value"]) for _, metric in points], "unit": unit,
                    "description": "A linha conecta somente pontos observados; lacunas não recebem interpolação.",
                }
                break
    return charts


def _has_material_gap(runs: tuple[AnalyticsRun, ...]) -> bool:
    for left, right in zip(runs, runs[1:]):
        left_period, right_period = effective_run_period(left), effective_run_period(right)
        if left_period.is_bounded and right_period.is_bounded and (
            (right_period.start - left_period.end).total_seconds() > 24 * 60 * 60  # type: ignore[operator]
        ):
            return True
    return False


def build_report_snapshot(
    project: AnalyticsProject,
    runs: tuple[AnalyticsRun, ...],
    *,
    requested_start: date | None,
    requested_end: date | None,
    requested_period_key: str | None = None,
) -> dict:
    """Congela dados e fontes sem inventar observações entre os runs selecionados."""
    if not runs:
        raise InstagramReportError("Não há execuções concluídas no período informado.")
    first, last = runs[0], runs[-1]
    run_data, contents_by_run = {}, {}
    for run in runs:
        metric_data, contents = _run_metrics(project, run)
        run_data[run.id] = metric_data
        contents_by_run[run.id] = contents
    comments_by_run = {
        run.id: len(db.session.scalars(select(InstagramCommentObservation.id).where(
            InstagramCommentObservation.analytics_run_id == run.id
        )).all())
        for run in runs
    }
    source = {
        "findings": _findings_for_project(project),
        "recommendations": _recommendations_for_project(project),
    }
    source["actions"], source["evaluations"] = _actions_for_project(project)
    source["learnings"] = _learnings_for_project(project)
    initial_metrics = run_data[first.id]["metrics"]
    final_metrics = run_data[last.id]["metrics"]
    limitations = [
        "Os resultados são descritivos e não demonstram causalidade.",
        "Métricas ausentes permanecem indisponíveis; zero é valor conhecido.",
    ]
    first_period, last_period = effective_run_period(first), effective_run_period(last)
    if (
        len(runs) >= 2 and first_period.is_bounded and last_period.is_bounded
        and (first_period.end - first_period.start).total_seconds() != (last_period.end - last_period.start).total_seconds()  # type: ignore[operator]
    ):
        limitations.append("As primeiras e últimas observações têm durações diferentes; considere esse contexto na comparação.")
    if _has_material_gap(runs):
        limitations.append("Há lacunas entre observações selecionadas; a evolução mostra somente pontos efetivamente registrados.")
    run_snapshots = [_run_snapshot(run) for run in runs]
    return {
        "schema_version": INSTAGRAM_REPORT_VERSION,
        "generated_at": _iso(datetime.now(timezone.utc)),
        "timezone": current_app.config["PRESENTATION_TIMEZONE"],
        "project": {"id": project.id, "name": project.name, "description": project.description},
        "requested_period": {
            **({"start": requested_start.isoformat(), "end": requested_end.isoformat()} if requested_start is not None and requested_end is not None else {}),
            **({"preset": requested_period_key} if requested_period_key is not None else {}),
        },
        "report_type": "executive", "analysis_mode": "comparative_evolution" if len(runs) >= 2 else "single_period",
        "selected_sections": list(DEFAULT_REPORT_SECTIONS),
        "first_observation": run_snapshots[0], "last_observation": run_snapshots[-1],
        "intermediate_observations": run_snapshots[1:-1], "runs_used": run_snapshots,
        "baseline_run": run_snapshots[0], "final_run": run_snapshots[-1], "source_runs": run_snapshots,
        "metrics_by_run": run_data, "contents_by_run": contents_by_run,
        "comments_by_run": comments_by_run,
        "comparison": _comparison(initial_metrics, final_metrics) if len(runs) >= 2 else [],
        "evolution": ([{"run": _run_snapshot(run), "metrics": run_data[run.id]["metrics"]} for run in runs] if len(runs) >= 3 else []),
        "findings": source["findings"], "recommendations": source["recommendations"],
        "actions": source["actions"], "evaluations": source["evaluations"], "learnings": source["learnings"],
        "milestones": _milestones(source, first, last), "limitations": limitations,
        "methodology": {
            "metric_aggregation": "Mediana calculada somente entre valores conhecidos por conteúdo observado.",
            "coverage": "Cobertura = valores conhecidos / conteúdos observados no run.",
            "causality": "Associações temporais e avaliações humanas não provam causalidade.",
            "comparability": "Comparações usam somente métricas com a mesma agregação e unidades compatíveis.",
            "run_selection": "A seleção automática inclui apenas runs concluídos com período conhecido e inteiramente dentro do recorte; runs Meta usam a janela congelada. Intervalos idênticos usam o run concluído mais recente.",
        },
    }


def _data_root() -> Path:
    return Path(current_app.config["PLATFORM_DATA_DIR"]).resolve()


def artifact_directory(project_id: str, report_id: str, *, create: bool = False) -> Path:
    """Retorna diretório privado com somente UUIDs, nunca título informado."""
    project_key, report_key = str(UUID(project_id)), str(UUID(report_id))
    root = _data_root()
    target = root / "analytics" / project_key / "reports" / report_key
    if create:
        target.mkdir(mode=0o700, parents=True, exist_ok=True)
    resolved_root = root.resolve()
    if target.is_symlink() or (target.exists() and (not target.is_dir() or not target.resolve().is_relative_to(resolved_root))):
        raise InstagramReportPersistenceError("Diretório privado de relatórios inválido.")
    return target


def artifact_path(project: AnalyticsProject, report: InstagramReport, extension: str) -> Path:
    if extension not in REPORT_FORMATS:
        raise InstagramReportError("Formato de download inválido.")
    path = artifact_directory(project.id, report.id) / f"report.{extension}"
    root = _data_root()
    if path.is_symlink() or not path.resolve().is_relative_to(root):
        raise InstagramReportError("Artefato de relatório inválido.")
    return path


def report_download_name(report: InstagramReport, extension: str) -> str:
    """Nome de download estável para históricos e orientado ao período para XLSX novo."""
    snapshot = report.snapshot_json or {}
    period = snapshot.get("requested_period") or {}
    if extension == "xlsx":
        try:
            start = date.fromisoformat(str(period.get("start")))
            end = date.fromisoformat(str(period.get("end")))
            return f"instagram_relatorio_{start.isoformat()}_{end.isoformat()}.xlsx"
        except ValueError:
            pass
    return f"relatorio-instagram-{report.id}.{extension}"


def _safe_cell(value: object) -> object:
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def _workbook(snapshot: dict) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(name: str, headers: list[str], rows: list[list[object]]):
        page = workbook.create_sheet(name)
        page.append(headers)
        for cell in page[1]:
            cell.fill = PatternFill("solid", fgColor="184775")
            cell.font = Font(color="FFFFFF", bold=True)
        for row in rows:
            page.append([_safe_cell(value) for value in row])
        page.freeze_panes = "A2"
        page.auto_filter.ref = page.dimensions
        for column in page.columns:
            letter = get_column_letter(column[0].column)
            page.column_dimensions[letter].width = min(52, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
            for cell in column:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    project = snapshot["project"]
    base = snapshot.get("first_observation") or snapshot["baseline_run"]
    final = snapshot.get("last_observation") or snapshot["final_run"]
    period = snapshot.get("requested_period") or {}
    runs = snapshot.get("runs_used") or snapshot.get("source_runs") or []
    requested_period_text = (
        "Todo o histórico disponível"
        if period.get("preset") in {"all", "all_history"}
        else f"{period.get('start') or _display_date(base['period_start'])} até {period.get('end') or _display_date(final['period_end'])}"
    )
    sheet("Summary", ["Field", "Value"], [
        ["Projeto", project["name"]],
        ["Período solicitado", requested_period_text],
        ["Execuções usadas", len(runs)],
        ["Versão", snapshot["schema_version"]],
    ])
    comparison = list(snapshot.get("comparison") or [])
    if comparison:
        sheet("Metrics", ["Metric", "Baseline", "Final", "Relative change", "Limitation"], [
            [row["label"], row["baseline_value"], row["final_value"], row["relative_change"], row["limitation"] or ""]
            for row in comparison
        ])
    else:
        metrics = ((snapshot.get("metrics_by_run") or {}).get(final["id"], {}) or {}).get("metrics", ()) or ()
        sheet("Metrics", ["Metric", "Median", "Known", "Total", "Coverage"], [
            [item["label"], item["value"], item["known_count"], item["total_count"], item["coverage"]]
            for item in metrics
        ])
    content_rows = []
    for run in runs:
        for item in (snapshot.get("contents_by_run") or {}).get(run["id"], ()):
            content_rows.append([_series_label(run), item["media_type"], _display_date(item["published_at"]), item["caption"], *[item["metrics"].get(key) for key in RAW_METRICS]])
    sheet("Contents", ["Run", "Type", "Published", "Caption", *[RAW_METRIC_LABELS[key] for key in RAW_METRICS]], content_rows)
    sheet("Comments", ["Run", "Observed comments"], [[_series_label(run), (snapshot.get("comments_by_run") or {}).get(run["id"], 0)] for run in runs])
    if "findings" in snapshot.get("selected_sections", DEFAULT_REPORT_SECTIONS):
        sheet("Findings", ["Title", "Summary", "Metric", "Support", "Coverage"], [[
            item["title"], item["summary"], item["metric_label"] or "", item["support_level"], item["coverage"]
        ] for item in snapshot.get("findings", ())])
    if "recommendations" in snapshot.get("selected_sections", DEFAULT_REPORT_SECTIONS):
        sheet("Recommendations", ["Title", "Summary", "Priority", "Support"], [[
            item["title"], item["summary"], item["priority"], item["support_level"]
        ] for item in snapshot["recommendations"]])
    if "actions" in snapshot.get("selected_sections", DEFAULT_REPORT_SECTIONS):
        # Excel proíbe "/" em nomes de abas; o hífen mantém o sentido de Actions/Tests.
        sheet("Actions-Tests", ["Title", "Type", "Status", "Metric", "Cohort"], [[
            item["title"], item["kind_label"], item["status_label"], item["metric_label"] or "", cohort_label(item["cohort"]) or "Não disponível"
        ] for item in snapshot["actions"]])
    if "evaluations" in snapshot.get("selected_sections", DEFAULT_REPORT_SECTIONS):
        sheet("Evaluations", ["Metric", "Baseline", "Follow-up", "Assessment"], [[
            item["metric_label"] or "", item["baseline_value"], item["followup_value"], item["assessment_label"]
        ] for item in snapshot["evaluations"]])
    if "learnings" in snapshot.get("selected_sections", DEFAULT_REPORT_SECTIONS):
        sheet("Learnings", ["Title", "Summary", "Assessment"], [[
            item["title"], item["summary"], item["assessment_label"]
        ] for item in snapshot["learnings"]])
    if "methodology" in snapshot.get("selected_sections", DEFAULT_REPORT_SECTIONS):
        sheet("Methodology", ["Aspect", "Record"], [[key.replace("_", " ").capitalize(), value] for key, value in snapshot["methodology"].items()] + [["Limitation", value] for value in snapshot["limitations"]])
    if comparison:
        sheet("Comparison", ["Metric", "Baseline", "Final", "Absolute change", "Relative change"], [[
            item["label"], item["baseline_value"], item["final_value"], item["absolute_change"], item["relative_change"]
        ] for item in snapshot["comparison"]])
    if len(snapshot.get("evolution") or ()) >= 3:
        evolution_rows = []
        for item in snapshot["evolution"]:
            for metric in item["metrics"]:
                evolution_rows.append([item["run"]["id"], _display_date(item["run"]["period_start"]), metric["label"], metric["value"], metric["known_count"], metric["total_count"], metric["coverage"]])
        sheet("Evolution", ["Run", "Start", "Metric", "Median", "Known", "Total", "Coverage"], evolution_rows)
    if snapshot.get("milestones"):
        sheet("Milestones", ["Date", "Type", "Description"], [[_display_date(item["at"]), item["kind"], item["title"]] for item in snapshot["milestones"]])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def reports_for_project(project: AnalyticsProject) -> tuple[InstagramReport, ...]:
    return tuple(db.session.scalars(select(InstagramReport).where(
        InstagramReport.analytics_project_id == project.id
    ).order_by(InstagramReport.created_at.desc(), InstagramReport.id.desc())).all())


def reports_page_for_project(project: AnalyticsProject, page: object) -> PlatformPage[InstagramReport]:
    """Histórico paginado no banco, sem carregar relatórios antigos inteiros."""
    requested = requested_page(page)
    total = int(db.session.scalar(select(func.count()).select_from(InstagramReport).where(
        InstagramReport.analytics_project_id == project.id
    )) or 0)
    pages = max(1, (total + 19) // 20)
    current = min(requested, pages)
    rows = db.session.scalars(
        select(InstagramReport).where(InstagramReport.analytics_project_id == project.id)
        .order_by(InstagramReport.created_at.desc(), InstagramReport.id.desc())
        .limit(20).offset((current - 1) * 20)
    ).all()
    return PlatformPage(tuple(rows), total, current)


def report_for_project(project: AnalyticsProject, report_id: str) -> InstagramReport | None:
    return db.session.scalar(select(InstagramReport).where(
        InstagramReport.analytics_project_id == project.id, InstagramReport.id == report_id
    ))


def generate_report(project: AnalyticsProject, values: dict[str, object], *, actor_id: str) -> InstagramReport:
    if project.status != "active" or project.archived_at is not None:
        raise InstagramReportError("Projetos arquivados não podem gerar novos relatórios.")
    requested_start, requested_end = _parse_requested_period(values)
    runs = resolve_report_runs(project, requested_start, requested_end)
    if not runs:
        raise InstagramReportError("Não há execuções concluídas com período conhecido para o recorte informado.")
    title = (
        "Relatório Instagram — Todo o histórico disponível"
        if requested_start is None or requested_end is None
        else f"Relatório Instagram — {requested_start.strftime('%d/%m/%Y')} a {requested_end.strftime('%d/%m/%Y')}"
    )
    report = InstagramReport(
        analytics_project_id=project.id, created_by_user_id=actor_id, title=title,
        # Valores canônicos apenas para compatibilidade com o schema da Fase 9.
        report_type="executive", analysis_mode="comparative_evolution" if len(runs) >= 2 else "single_period",
        status="generating", baseline_run_id=runs[0].id, final_run_id=runs[-1].id,
        selected_sections_json=list(DEFAULT_REPORT_SECTIONS), formats_json=["xlsx"],
        generator_version=INSTAGRAM_REPORT_VERSION,
    )
    try:
        # O estado transitório é persistido antes da geração de arquivo. Assim,
        # uma exceção posterior pode ser auditada como falha sem expor artefato.
        db.session.add(report)
        db.session.commit()
    except Exception as error:
        db.session.rollback()
        raise InstagramReportPersistenceError("Não foi possível iniciar a geração do relatório.") from error
    try:
        snapshot = build_report_snapshot(
            project,
            runs,
            requested_start=requested_start,
            requested_end=requested_end,
            requested_period_key=(
                str(values["period_preset"]) if values.get("period_preset") else None
            ),
        )
        directory = artifact_directory(project.id, report.id, create=True)
        report.snapshot_json = snapshot
        (directory / "report.xlsx").write_bytes(_workbook(snapshot))
        report.xlsx_filename = "report.xlsx"
        report.status = "completed"
        report.generated_at = datetime.now(timezone.utc)
        db.session.commit()
    except Exception as error:
        db.session.rollback()
        if report.id:
            try:
                directory = artifact_directory(project.id, report.id)
                if directory.exists():
                    rmtree(directory)
            except (InstagramReportError, InstagramReportPersistenceError, OSError):
                current_app.logger.exception("Não foi possível limpar artefatos de relatório falho")
        failed = db.session.get(InstagramReport, report.id)
        if failed is not None:
            failed.status = "failed"
            failed.error_message = "A geração falhou antes de publicar os artefatos. Tente novamente com outro período."
            db.session.commit()
        raise InstagramReportPersistenceError("Não foi possível gerar o relatório. Nenhum relatório concluído foi criado.") from error
    return report


def delete_report(project: AnalyticsProject, report: InstagramReport, confirmation: str) -> None:
    if project.status != "active" or project.archived_at is not None:
        raise InstagramReportError("Projetos arquivados não podem excluir relatórios.")
    if report.analytics_project_id != project.id:
        raise InstagramReportError("O relatório não pertence a este projeto.")
    if confirmation.strip() != "deletar":
        raise InstagramReportError('Digite exatamente "deletar" para confirmar a exclusão.')
    directory = artifact_directory(project.id, report.id)
    staged: Path | None = None
    try:
        if directory.exists():
            parent = directory.parent
            staged = parent / f".deleting-{uuid4()}"
            if parent.is_symlink() or staged.exists() or staged.is_symlink():
                raise InstagramReportPersistenceError("Diretório privado de relatórios inválido.")
            os.replace(directory, staged)
        db.session.delete(report)
        db.session.commit()
    except Exception as error:
        db.session.rollback()
        if staged is not None and staged.exists() and not directory.exists():
            os.replace(staged, directory)
        if isinstance(error, InstagramReportPersistenceError):
            raise
        raise InstagramReportPersistenceError("Não foi possível excluir o relatório.") from error
    if staged is not None:
        try:
            rmtree(staged)
        except OSError:
            # O registro já foi removido; o diretório isolado não é exposto por
            # rota alguma e o log permite uma limpeza operacional posterior.
            current_app.logger.exception("Artefatos de relatório aguardam limpeza: %s", staged)


__all__ = [
    "ANALYSIS_MODES", "ANALYSIS_MODE_LABELS", "DEFAULT_ANALYTICAL_SECTIONS",
    "DEFAULT_EXECUTIVE_SECTIONS", "DEFAULT_REPORT_SECTIONS", "INSTAGRAM_REPORT_VERSION", "InstagramReportError",
    "InstagramReportPersistenceError", "REPORT_FORMATS", "REPORT_TYPE_LABELS", "REPORT_TYPES",
    "SECTION_LABELS", "artifact_path", "delete_report", "eligible_report_runs", "generate_report",
    "report_chart_data", "report_download_name", "report_for_project", "report_summary_cards",
    "reports_for_project", "resolve_report_runs",
]
