"""Etapa 3B: âncoras canônicas, vínculos contextuais e busca por ocorrência."""

import hashlib
import io
import re
import tempfile
import unittest
from xml.etree import ElementTree
from pathlib import Path
from unittest.mock import patch

import pymupdf
from openpyxl import load_workbook
from sqlalchemy import select

from platform_helpers import create_user, csrf_from, isolated_platform, login
from platform_core.analyses import create_analysis, preserve_documents
from platform_core.extensions import db
from platform_core.models import (
    AnalysisDocument, Project, QualitativeCode, QualitativeCoding,
    QualitativeExcerpt, QualitativeMemo, UserToolOverride,
)
from platform_core.qualitative_corpus import read_qualitative_page
from platform_core.qualitative_annotations import SelectionConflict, validated_selection
from platform_core.qualitative_layout import build_native_layout, read_qualitative_layout
from platform_core.qualitative_routes import _run_corpus_job
from platform_core.scraping_types import QUALITATIVE, QUALITATIVE_TOOL


class QualitativeContextTests(unittest.TestCase):
    def setUp(self):
        self.scope = isolated_platform()
        self.app = self.scope.__enter__()
        self.user = create_user()
        self.client = self.app.test_client()
        db.session.add(UserToolOverride(user_id=self.user.id, tool_id=QUALITATIVE_TOOL, decision="allow"))
        db.session.commit()
        login(self.client)
        self.csrf = csrf_from(self.client.get("/"))
        self.project = Project(owner_user_id=self.user.id, name="Leitura", scrape_type=QUALITATIVE,
                               description="")
        db.session.add(self.project)
        db.session.commit()
        self.analysis = create_analysis(user_id=self.user.id, project_id=self.project.id,
                                        tool_id=QUALITATIVE_TOOL, tool_version="manual-v1",
                                        name="Leitura", parameters={})
        source_pdf = pymupdf.open()
        page = source_pdf.new_page()
        page.insert_text((72, 72), "Documento Documento pesquisa racial")
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "texto.pdf"
            path.write_bytes(source_pdf.tobytes())
            preserve_documents(self.analysis, [(path, "texto.pdf")])
        source_pdf.close()

        def extract(path):
            with pymupdf.open(path) as document:
                for number, item in enumerate(document, start=1):
                    yield {"pagina_pdf": number, "texto": item.get_text("text"), "ocr_utilizado": False}

        with patch("platform_core.qualitative_corpus.pdf_extractor.extrair_paginas", side_effect=extract):
            _run_corpus_job(self.app, self.analysis.id)
        db.session.expire_all()
        self.document = db.session.scalar(select(AnalysisDocument).where(
            AnalysisDocument.analysis_id == self.analysis.id))
        self.page = read_qualitative_page(self.analysis, self.document.id, 1)
        self.base_url = (f"/analise-qualitativa/bases/{self.analysis.id}/documentos/"
                         f"{self.document.id}/paginas/1")

    def tearDown(self):
        self.scope.__exit__(None, None, None)

    def selection(self, quote="Documento", *, start=None):
        start = self.page["text"].index(quote) if start is None else start
        return {"start": start, "end": start + len(quote), "selected_text": quote,
                "page_hash": self.page["sha256"]}

    def post(self, action, extra=None, *, url=None, token=True, selection=None):
        headers = {"X-CSRFToken": self.csrf} if token else {}
        return self.client.post(url or f"{self.base_url}/contexto", json={
            **(selection or self.selection()), "action": action, **(extra or {}),
        }, headers=headers)

    def test_apply_multiple_codes_is_idempotent_and_excerpt_has_stable_uuid(self):
        codes = [QualitativeCode(analysis_id=self.analysis.id, name=name,
                                 created_by_user_id=self.user.id) for name in ("Tema A", "Tema B")]
        db.session.add_all(codes)
        db.session.commit()
        response = self.post("apply_codes", {"code_ids": [code.id for code in codes]})
        self.assertEqual(response.status_code, 201)
        excerpt_id = response.json["excerpt_id"]
        self.assertEqual(len(db.session.scalars(select(QualitativeExcerpt)).all()), 1)
        self.assertEqual(len(db.session.scalars(select(QualitativeCoding)).all()), 2)
        self.assertEqual(self.post("apply_codes", {"code_ids": [codes[0].id, codes[0].id]})
                         .json["excerpt_id"], excerpt_id)
        self.assertEqual(len(db.session.scalars(select(QualitativeCoding)).all()), 2)
        second = self.page["text"].index("Documento", 1)
        another = self.post("apply_codes", {"code_ids": [codes[0].id]},
                            selection=self.selection(start=second))
        self.assertNotEqual(another.json["excerpt_id"], excerpt_id)
        self.assertEqual(db.session.get(QualitativeExcerpt, excerpt_id).start_offset, 0)
        listed = self.client.get(f"{self.base_url}/trechos")
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(len(listed.json["excerpts"]), 2)

    def test_contextual_code_in_vivo_and_memo_keep_general_memos_separate(self):
        created = self.post("create_code", {"name": "Identidade", "description": "Nota"})
        self.assertEqual(created.status_code, 201)
        self.assertEqual(len(db.session.scalars(select(QualitativeCoding)).all()), 1)
        self.assertEqual(created.json["records"]["codes"][0]["name"], "Identidade")
        self.assertEqual(self.post("create_code", {"name": "identidade"}).status_code, 409)
        second = self.page["text"].index("Documento", 1)
        in_vivo = self.post("in_vivo", selection=self.selection(start=second))
        self.assertEqual(in_vivo.status_code, 201)
        self.assertIn("Documento", [item["name"] for item in in_vivo.json["records"]["codes"]])
        self.assertEqual(self.post("in_vivo").status_code, 201)
        self.assertEqual(len(db.session.scalars(select(QualitativeCode).where(
            QualitativeCode.name == "Documento")).all()), 1)
        contextual = self.post("create_memo", {"text": "Memo do trecho"})
        self.assertEqual(contextual.status_code, 201)
        memo = db.session.scalar(select(QualitativeMemo).where(QualitativeMemo.text == "Memo do trecho"))
        self.assertEqual(memo.excerpt_id, created.json["excerpt_id"])
        self.assertEqual(contextual.json["records"]["memos"][0]["context"], "trecho")
        general = self.client.post(f"/analise-qualitativa/bases/{self.analysis.id}/memos",
                                   json={"text": "Memo geral"}, headers={"X-CSRFToken": self.csrf})
        self.assertEqual(general.status_code, 201)
        self.assertIn("geral", [item["context"] for item in general.json["memos"]])
        edited = self.client.patch(f"/analise-qualitativa/bases/{self.analysis.id}/memos/{memo.id}",
                                   json={"text": "Revisado"}, headers={"X-CSRFToken": self.csrf})
        self.assertEqual(edited.status_code, 200)
        self.assertEqual(db.session.get(QualitativeMemo, memo.id).excerpt_id, created.json["excerpt_id"])

    def test_invalid_offsets_hash_text_long_name_cross_document_and_csrf(self):
        for selection, status in [
            ({**self.selection(), "start": -1}, 400),
            ({**self.selection(), "selected_text": "Outro"}, 400),
            ({**self.selection(), "page_hash": "0" * 64}, 409),
            ({**self.selection(), "start": 0, "end": 0}, 400),
            ({**self.selection(), "start": 0, "end": 1, "selected_text": " "}, 400),
        ]:
            self.assertEqual(self.post("create_memo", {"text": "Teste"}, selection=selection).status_code, status)
        self.assertEqual(self.post("create_memo", {"text": "Teste"}, token=False).status_code, 400)
        long_text = "x" * 161
        with patch("platform_core.qualitative_annotations.read_qualitative_page",
                   return_value={"text": long_text, "sha256": self.page["sha256"]}):
            response = self.post("in_vivo", selection={"start": 0, "end": 161,
                "selected_text": long_text, "page_hash": self.page["sha256"]})
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.json["suggested_name"], long_text)
        self.assertEqual(len(db.session.scalars(select(QualitativeExcerpt)).all()), 0)
        self.assertEqual(self.client.get(f"{self.base_url}/trechos").status_code, 200)
        foreign = create_user("Outro", "outro-contexto@example.org")
        other_project = Project(owner_user_id=foreign.id, name="Outro", scrape_type=QUALITATIVE)
        db.session.add(other_project)
        db.session.commit()
        other_analysis = create_analysis(user_id=foreign.id, project_id=other_project.id,
                                         tool_id=QUALITATIVE_TOOL, tool_version="manual-v1",
                                         name="Outro", parameters={})
        other_document = AnalysisDocument(analysis_id=other_analysis.id,
                                          original_name="alheio.pdf", stored_name="alheio.pdf")
        db.session.add(other_document)
        db.session.commit()
        crossed = self.base_url.replace(self.document.id, other_document.id)
        self.assertEqual(self.post("create_memo", {"text": "Teste"}, url=f"{crossed}/contexto").status_code, 404)
        self.assertEqual(self.client.get(f"{crossed}/trechos").status_code, 404)
        foreign_code = QualitativeCode(analysis_id=other_analysis.id, name="Alheio",
                                       created_by_user_id=foreign.id)
        db.session.add(foreign_code)
        db.session.commit()
        self.assertEqual(self.post("apply_codes", {"code_ids": [foreign_code.id]}).status_code, 404)
        foreign_excerpt = QualitativeExcerpt(
            analysis_id=other_analysis.id, document_id=other_document.id, page_number=1,
            start_offset=0, end_offset=9, quoted_text="Documento",
            page_text_hash=self.page["sha256"], created_by_user_id=foreign.id,
        )
        db.session.add(foreign_excerpt)
        db.session.commit()
        self.assertEqual(self.client.patch(
            f"/analise-qualitativa/bases/{self.analysis.id}/trechos/{foreign_excerpt.id}",
            json={"start": 0, "end": 9, "page_hash": self.page["sha256"]},
            headers={"X-CSRFToken": self.csrf},
        ).status_code, 404)
        db.session.add(UserToolOverride(user_id=foreign.id, tool_id=QUALITATIVE_TOOL, decision="allow"))
        db.session.commit()
        self.client.post('/logout', data={"csrf_token": self.csrf})
        login(self.client, email="outro-contexto@example.org")
        self.assertEqual(self.client.get(f"{self.base_url}/trechos").status_code, 404)
        self.assertEqual(self.client.post(f"{self.base_url}/contexto", json={
            **self.selection(), "action": "create_memo", "text": "Tentativa",
        }, headers={"X-CSRFToken": csrf_from(self.client.get('/'))}).status_code, 404)

    def test_unicode_codepoint_offsets_and_composite_rollback(self):
        synthetic = "Raça 😊 e memória"
        with patch("platform_core.qualitative_annotations.read_qualitative_page",
                   return_value={"text": synthetic, "sha256": "a" * 64}):
            _, start, end, quote = validated_selection(self.analysis, self.document.id, 1,
                {"start": 5, "end": 6, "selected_text": "😊", "page_hash": "a" * 64})
        self.assertEqual((start, end, quote), (5, 6, "😊"))
        with patch("platform_core.qualitative_routes.get_or_create_excerpt",
                   side_effect=SelectionConflict("Falha de teste")):
            response = self.post("create_code", {"name": "Não persistir"})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(db.session.scalars(select(QualitativeCode)).all(), [])
        self.assertEqual(db.session.scalars(select(QualitativeExcerpt)).all(), [])

    def test_search_and_layout_use_exact_offsets_without_reprocessing(self):
        url = self.base_url.rsplit("/paginas/1", 1)[0]
        with patch("platform_core.qualitative_corpus.pdf_extractor.extrair_paginas",
                   side_effect=AssertionError("Não reprocessar")):
            found = self.client.get(f"{url}/buscar", query_string={"q": "Documento"})
            regex = self.client.get(f"{url}/buscar", query_string={"q": "Documento", "grep": "1"})
            layout = self.client.get(f"{self.base_url}/layout")
        self.assertEqual(found.status_code, 200)
        self.assertEqual(regex.status_code, 200)
        self.assertEqual(len(found.json["results"]), 2)
        self.assertNotEqual(found.json["results"][0]["start_offset"],
                            found.json["results"][1]["start_offset"])
        self.assertEqual(found.json["results"][0]["page_text_hash"], self.page["sha256"])
        self.assertTrue(layout.json["layout_available"])
        self.assertTrue(all(item["end"] - item["start"] == 1 for item in layout.json["items"]))

    def test_native_layout_preserves_codepoint_positions_across_lines(self):
        pdf = pymupdf.open()
        try:
            page = pdf.new_page()
            page.insert_text((72, 72), "Primeira linha\nSegunda linha")
            text = page.get_text("text")
            digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
            layout = build_native_layout(page, text, 1, digest)
            self.assertIsNotNone(layout)
            self.assertEqual(layout["page_text_hash"], digest)
            first = next(item for item in layout["items"] if item["start"] == text.index("Primeira"))
            second = next(item for item in layout["items"] if item["start"] == text.index("Segunda"))
            self.assertEqual(first["text"], "P")
            self.assertEqual(second["text"], "S")
            self.assertGreater(second["bbox"][1], first["bbox"][1])
        finally:
            pdf.close()

    def test_interface_has_canonical_selection_and_accessible_context_menu(self):
        html = self.client.get(self.base_url).get_data(as_text=True)
        for label in ("Aplicar código", "Código in vivo", "Criar memo", "Buscar ou criar código"):
            self.assertIn(label, html)
        self.assertNotIn('data-context-action="create_code"', html)
        self.assertIn('data-context-create', html)
        self.assertIn('data-context-menu role="menu"', html)
        self.assertIn('data-page-data-url-template=', html)
        self.assertIn('data-excerpts-url-template=', html)
        self.assertIn('data-context-url-template=', html)
        self.assertIn('Vinculado a', (Path(__file__).parents[1] /
                                     'templates/platform/_qualitative_records.html').read_text(encoding="utf-8"))
        viewer = (Path(__file__).parents[1] / 'static/js/qualitative_viewer.js').read_text(encoding="utf-8")
        self.assertIn('Array.from(container.textContent.slice', viewer)
        self.assertIn('selectionBoundary(range.startContainer', viewer)
        self.assertIn('selectionBoundary(range.endContainer', viewer)
        self.assertIn('if (!anchor || !anchor.range.getClientRects().length) return', viewer)
        self.assertIn('event.preventDefault();', viewer)
        self.assertIn("event.key === 'F10' && event.shiftKey", viewer)
        self.assertIn('closeContextMenu();', viewer)
        self.assertIn('platform-qualitative-excerpt-overlay', viewer)

    def test_inline_code_reuses_normalized_name_and_is_atomic(self):
        first = self.post("create_and_apply", {"name": "  Raça  "})
        self.assertEqual(first.status_code, 201)
        reused = self.post("create_and_apply", {"name": "raça"})
        self.assertEqual(reused.status_code, 201)
        self.assertEqual(first.json["excerpt_id"], reused.json["excerpt_id"])
        self.assertEqual(len(db.session.scalars(select(QualitativeCode)).all()), 1)
        self.assertEqual(len(db.session.scalars(select(QualitativeCoding)).all()), 1)
        self.assertEqual(first.json["page"]["excerpts"][0]["codes"][0]["name"], "Raça")
        with patch("platform_core.qualitative_routes.get_or_create_excerpt",
                   side_effect=SelectionConflict("Falha de teste")):
            response = self.post("create_and_apply", {"name": "Não persistir"})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(len(db.session.scalars(select(QualitativeCode)).all()), 1)

    def test_edit_bounds_keeps_identity_codings_and_memos_and_detects_conflicts(self):
        created = self.post("create_and_apply", {"name": "Identidade"})
        excerpt_id = created.json["excerpt_id"]
        self.assertEqual(self.post("create_memo", {"text": "Contexto"}).status_code, 201)
        url = f"/analise-qualitativa/bases/{self.analysis.id}/trechos/{excerpt_id}"
        desired_end = len("Documento Documento")
        changed = self.client.patch(url, json={"start": 0, "end": desired_end,
                                               "page_hash": self.page["sha256"]},
                                    headers={"X-CSRFToken": self.csrf})
        self.assertEqual(changed.status_code, 200)
        self.assertEqual(self.client.patch(url, json={"start": 0, "end": desired_end,
            "page_hash": self.page["sha256"]}).status_code, 400)
        excerpt = db.session.get(QualitativeExcerpt, excerpt_id)
        self.assertEqual(excerpt.id, excerpt_id)
        self.assertEqual(excerpt.quoted_text, "Documento Documento")
        self.assertEqual(excerpt.document_id, self.document.id)
        self.assertEqual(excerpt.page_number, 1)
        self.assertEqual(len(db.session.scalars(select(QualitativeCoding)).all()), 1)
        self.assertEqual(db.session.scalar(select(QualitativeMemo)).excerpt_id, excerpt_id)
        for payload, expected in [
            ({"start": 2, "end": 2, "page_hash": self.page["sha256"]}, 400),
            ({"start": -1, "end": 4, "page_hash": self.page["sha256"]}, 400),
            ({"start": 0, "end": self.page["char_count"] + 1,
              "page_hash": self.page["sha256"]}, 400),
            ({"start": 0, "end": 9, "page_hash": "0" * 64}, 409),
        ]:
            self.assertEqual(self.client.patch(url, json=payload,
                headers={"X-CSRFToken": self.csrf}).status_code, expected)
        self.assertEqual(self.client.patch(url, json={"start": 0, "end": 9,
            "page_hash": self.page["sha256"]}).status_code, 400)
        second_start = self.page["text"].index("Documento", 1)
        second = self.post("in_vivo", selection=self.selection(start=second_start))
        self.assertEqual(second.status_code, 201)
        second_url = f"/analise-qualitativa/bases/{self.analysis.id}/trechos/{second.json['excerpt_id']}"
        self.assertEqual(self.client.patch(second_url, json={"start": 0, "end": desired_end,
            "page_hash": self.page["sha256"]}, headers={"X-CSRFToken": self.csrf}).status_code, 409)

    def test_report_groups_current_excerpts_and_deep_links_without_pdf_processing(self):
        codes = [QualitativeCode(analysis_id=self.analysis.id, name=name,
            created_by_user_id=self.user.id) for name in ("Tema A", "Tema B")]
        db.session.add_all(codes)
        db.session.commit()
        first = self.post("apply_codes", {"code_ids": [item.id for item in codes]})
        self.assertEqual(first.status_code, 201)
        self.post("create_memo", {"text": "Memo associado"})
        report_url = f"/analise-qualitativa/bases/{self.analysis.id}/relatorio-codificacao"
        with patch("platform_core.qualitative_corpus.pdf_extractor.extrair_paginas",
                   side_effect=AssertionError("Não reprocessar")):
            report = self.client.get(report_url)
        html = report.get_data(as_text=True)
        self.assertEqual(report.status_code, 200)
        self.assertIn("Tema A", html)
        self.assertIn("Tema B", html)
        self.assertIn("1 trecho", html)
        self.assertIn("texto.pdf", html)
        self.assertIn("1 memo", html)
        self.assertIn(f"excerpt={first.json['excerpt_id']}", html)
        articles = re.findall(r'<article class="platform-qualitative-report-excerpt">.*?</article>',
                              html, re.DOTALL)
        self.assertEqual(len(articles), 2)
        for markup in articles:
            article = ElementTree.fromstring(markup)
            link = article.find('div/a')
            self.assertEqual(link.text, "texto.pdf · p. 1")
            self.assertEqual(link.attrib['href'], f"{self.base_url}?excerpt={first.json['excerpt_id']}")
            self.assertEqual(len(list(link)), 0)
            self.assertEqual(article.find('q').text, "Documento")
            self.assertIn("1 memo", article.find('div/span').text)
        self.assertIn('Exportar Excel</a>', html)
        new_end = len("Documento Documento")
        self.client.patch(f"/analise-qualitativa/bases/{self.analysis.id}/trechos/{first.json['excerpt_id']}",
            json={"start": 0, "end": new_end, "page_hash": self.page["sha256"]},
            headers={"X-CSRFToken": self.csrf})
        self.assertIn("Documento Documento", self.client.get(report_url).get_data(as_text=True))
        self.assertEqual(self.client.get(self.base_url, query_string={
            "excerpt": first.json["excerpt_id"]}).status_code, 200)
        self.assertEqual(self.client.get(self.base_url, query_string={"excerpt": "invalido"}).status_code, 200)
        self.assertEqual(self.client.get(self.base_url, query_string={"excerpt": str(__import__('uuid').uuid4())}).status_code, 200)

    def test_code_excerpt_counts_and_dynamic_excel_export(self):
        codes = [QualitativeCode(analysis_id=self.analysis.id, name=name,
            created_by_user_id=self.user.id) for name in ("Raça", "Ideologia", "Sem trechos")]
        db.session.add_all(codes)
        db.session.commit()
        payload = self.client.get(self.base_url).get_data(as_text=True)
        self.assertIn("Sem trechos (0)", payload)
        first = self.post("apply_codes", {"code_ids": [codes[0].id, codes[1].id]})
        self.assertEqual(first.status_code, 201)
        self.assertEqual({item["name"]: item["excerpt_count"] for item in first.json["records"]["codes"]},
                         {"Raça": 1, "Ideologia": 1, "Sem trechos": 0})
        repeated = self.post("apply_codes", {"code_ids": [codes[0].id]})
        self.assertEqual(next(item["excerpt_count"] for item in repeated.json["records"]["codes"]
                              if item["id"] == codes[0].id), 1)
        self.post("create_memo", {"text": "Primeira observação"})
        self.post("create_memo", {"text": "Segunda observação"})
        second_start = self.page["text"].index("pesquisa")
        second = self.post("apply_codes", {"code_ids": [codes[0].id]},
                           selection=self.selection(quote="pesquisa", start=second_start))
        self.assertEqual(next(item["excerpt_count"] for item in second.json["records"]["codes"]
                              if item["id"] == codes[0].id), 2)
        url = f"/analise-qualitativa/bases/{self.analysis.id}/relatorio-codificacao.xlsx"
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml.sheet", response.content_type)
        self.assertIn("relatorio-codificacao-", response.headers["Content-Disposition"])
        sheet = load_workbook(io.BytesIO(response.data))["Codificação"]
        self.assertEqual([cell.value for cell in sheet[1]],
                         ["Código", "Documento", "Página", "Trecho codificado", "Memo contextual"])
        exported = [tuple(cell.value for cell in row) for row in sheet.iter_rows(min_row=2)]
        self.assertEqual(len(exported), 3)
        self.assertEqual(sorted(row[0] for row in exported), ["Ideologia", "Raça", "Raça"])
        self.assertTrue(all(row[1] == "texto.pdf" and row[2] == 1 for row in exported))
        self.assertEqual(sum("Primeira observação --- Segunda observação" == row[4]
                             for row in exported), 2)
        changed_end = len("Documento Documento")
        self.client.patch(f"/analise-qualitativa/bases/{self.analysis.id}/trechos/{first.json['excerpt_id']}",
            json={"start": 0, "end": changed_end, "page_hash": self.page["sha256"]},
            headers={"X-CSRFToken": self.csrf})
        edited = load_workbook(io.BytesIO(self.client.get(url).data))["Codificação"]
        self.assertEqual(sum(row[3].value == "Documento Documento" for row in edited.iter_rows(min_row=2)), 2)

    def test_report_excel_flattens_line_breaks_only_in_export_and_disables_wrap(self):
        quote = "Primeira linha com ç e ação\nSegunda linha 😀\nTerceira linha"
        digest = hashlib.sha256(quote.encode('utf-8')).hexdigest()
        # Âncora canônica multilinhas: o relatório não precisa reler o corpus.
        with patch('platform_core.qualitative_annotations.read_qualitative_page',
                   return_value={"text": quote, "sha256": digest}):
            selection = {"start": 0, "end": len(quote), "selected_text": quote, "page_hash": digest}
            first = self.post('create_and_apply', {"name": "Tema A"}, selection=selection)
            second = self.post('create_and_apply', {"name": "Tema B"}, selection=selection)
        self.assertEqual(first.status_code, 201)
        self.assertEqual(second.status_code, 201)
        excerpt = db.session.get(QualitativeExcerpt, first.json['excerpt_id'])
        memo_text = "=Observação\r\ncom quebra\risolada\ne acentuação"
        memo = QualitativeMemo(analysis_id=self.analysis.id, excerpt_id=excerpt.id,
                                text=memo_text, created_by_user_id=self.user.id)
        db.session.add(memo)
        db.session.commit()
        before = (excerpt.quoted_text, excerpt.start_offset, excerpt.end_offset, excerpt.page_text_hash)
        url = f'/analise-qualitativa/bases/{self.analysis.id}/relatorio-codificacao'
        response = self.client.get(f'{url}.xlsx')
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.data))
        sheet = workbook['Codificação']
        self.assertEqual([cell.value for cell in sheet[1]],
                         ['Código', 'Documento', 'Página', 'Trecho codificado', 'Memo contextual'])
        rows = list(sheet.iter_rows(min_row=2))
        self.assertEqual(len(rows), 2)
        self.assertEqual({row[0].value for row in rows}, {'Tema A', 'Tema B'})
        for row in rows:
            self.assertEqual(row[3].value, ' '.join(quote.splitlines()))
            self.assertEqual(row[4].value, "'" + ' '.join(memo_text.splitlines()))
            for cell in row:
                self.assertIsNot(cell.alignment.wrap_text, True)
                self.assertEqual(cell.alignment.vertical, 'top')
                self.assertNotEqual(cell.data_type, 'f')
                self.assertNotRegex(str(cell.value), r'[\r\n]')
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertLessEqual(sheet.column_dimensions['D'].width, 80)
        self.assertTrue(all(dimension.height is None for dimension in sheet.row_dimensions.values()))
        db.session.expire_all()
        self.assertEqual((excerpt.quoted_text, excerpt.start_offset, excerpt.end_offset,
                          excerpt.page_text_hash), before)
        self.assertEqual(memo.text, memo_text)
        self.assertIn(quote, self.client.get(url).get_data(as_text=True))

    def test_report_excerpt_presentation_uses_full_width_and_selectable_natural_wrapping(self):
        css = (Path(__file__).parents[1] / 'static/css/platform.css').read_text(encoding='utf-8')
        excerpt_rule = re.search(r'\.platform-qualitative-report-excerpt \{([^}]+)', css)[1]
        quote_rule = re.search(r'\.platform-qualitative-report-excerpt q \{([^}]+)', css)[1]
        self.assertIn('grid-template-columns: minmax(0, 1fr)', excerpt_rule)
        for rule in ('width: 100%', 'white-space: normal', 'overflow-wrap: break-word',
                     'user-select: text', 'cursor: text'):
            self.assertIn(rule, quote_rule)
        self.assertNotIn('break-all', quote_rule)
        self.assertNotIn('.platform-qualitative-report-excerpt:hover', css)

    def test_report_empty_states_ownership_and_foreign_excerpt_link(self):
        report_url = f"/analise-qualitativa/bases/{self.analysis.id}/relatorio-codificacao"
        export_url = f"{report_url}.xlsx"
        self.assertIn("Nenhum código foi criado", self.client.get(report_url).get_data(as_text=True))
        self.assertEqual(self.client.get(export_url).status_code, 302)
        code = QualitativeCode(analysis_id=self.analysis.id, name="Sem trecho",
                               created_by_user_id=self.user.id)
        db.session.add(code)
        db.session.commit()
        self.assertIn("Ainda não há trechos codificados", self.client.get(report_url).get_data(as_text=True))
        other = create_user("Outro", "outro-report@example.org")
        db.session.add(UserToolOverride(user_id=other.id, tool_id=QUALITATIVE_TOOL, decision="allow"))
        db.session.commit()
        self.client.post('/logout', data={"csrf_token": self.csrf})
        self.assertIn(self.client.get(export_url).status_code, (302, 401))
        login(self.client, email="outro-report@example.org")
        self.assertEqual(self.client.get(report_url).status_code, 404)
        self.assertEqual(self.client.get(export_url).status_code, 404)
        self.assertEqual(self.client.patch(
            f"/analise-qualitativa/bases/{self.analysis.id}/trechos/{__import__('uuid').uuid4()}",
            json={"start": 0, "end": 9, "page_hash": self.page["sha256"]},
            headers={"X-CSRFToken": csrf_from(self.client.get('/'))}).status_code, 404)

    def test_report_reflects_code_and_document_deletion(self):
        report_url = f"/analise-qualitativa/bases/{self.analysis.id}/relatorio-codificacao"
        created = self.post("create_and_apply", {"name": "Código removível"})
        code_id = created.json["records"]["codes"][0]["id"]
        report_html = self.client.get(report_url).get_data(as_text=True)
        self.assertIn("Código removível", report_html)
        self.assertIn("(1 trecho)", report_html)
        removed = self.client.delete(f"/analise-qualitativa/bases/{self.analysis.id}/codigos/{code_id}",
                                     headers={"X-CSRFToken": self.csrf})
        self.assertEqual(removed.status_code, 200)
        self.assertNotIn("Código removível", self.client.get(report_url).get_data(as_text=True))
        self.assertEqual(self.post("create_and_apply", {"name": "Novo código"}).status_code, 201)
        deleted_document = self.client.post(
            f"/analise-qualitativa/bases/{self.analysis.id}/documentos/{self.document.id}/excluir",
            headers={"X-CSRFToken": self.csrf})
        self.assertEqual(deleted_document.status_code, 200)
        self.assertTrue(all(code["excerpt_count"] == 0 for code in deleted_document.json["records"]["codes"]))
        html = self.client.get(report_url).get_data(as_text=True)
        self.assertIn("Ainda não há trechos codificados", html)
        report_content = html.split('<dialog class="platform-tutorial-dialog"', 1)[0]
        self.assertNotIn("Documento", report_content)

    def test_reader_exposes_edit_margin_report_and_keeps_search_layer_separate(self):
        html = self.client.get(self.base_url).get_data(as_text=True)
        self.assertIn('data-edit-excerpt-url-template=', html)
        self.assertIn('data-margin-track', html)
        self.assertIn('>Relatório</a>', html)
        self.assertNotIn('Ver relatório de codificação', html)
        self.assertIn('data-context-create', html)
        self.assertNotIn('data-context-action="create_code"', html)
        viewer = (Path(__file__).parents[1] / 'static/js/qualitative_viewer.js').read_text(encoding='utf-8')
        css = (Path(__file__).parents[1] / 'static/css/platform.css').read_text(encoding='utf-8')
        for marker in ('pointerdown', 'pointermove', 'pointerup', 'pointercancel',
                       'nearestCanonicalBoundary', 'renderMargin', 'activateExcerpt',
                       'data-target-excerpt-id'):
            self.assertIn(marker, viewer if marker != 'data-target-excerpt-id' else html)
        self.assertIn('platform-qualitative-excerpt-handle', css)
        self.assertIn('platform-qualitative-margin-card', css)
        self.assertIn('platform-qualitative-search-overlay', css)


if __name__ == "__main__":
    unittest.main()
