"""Testa somente a apresentação compacta dos XLSX exportados."""

from __future__ import annotations

from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
import pandas as pd

from analyzer import v1, v2, v3


TERMOS = pd.DataFrame([{"Termo": "São Paulo"}])
DIAGNOSTICO = pd.DataFrame(
    [
        {
            "arquivo": "livro.pdf",
            "paginas_com_texto_extraivel": 1,
            "paginas_processadas_com_OCR": 0,
            "idioma_OCR": "não utilizado",
        }
    ]
)


def _registro_lexical() -> dict[str, object]:
    return {
        "ID livro": "livro",
        "Termo": "São Paulo",
        "Categoria do termo": "TERMO INFORMADO",
        "Parágrafo anterior": "Bloco anterior com texto longo.",
        "Parágrafo do termo": "São Paulo aparece neste trecho longo.\nO texto permanece integral.",
        "Parágrafo posterior": "Bloco posterior com texto longo.",
        "Unidade": "",
        "Capítulo": "",
        "Seção": "",
        "Subseção": "",
        "Página": 1,
        "Tipo da ocorrência": "parágrafo",
        "Contexto sociológico da ocorrência": "Desigualdades sociais e regionais",
        "Descrição da ocorrência": "Descrição longa preservada integralmente.",
        "Validação manual": "revisar",
        "_quantidade_no_registro": 1,
        "_metodo": "texto",
    }


def _registro_v3() -> dict[str, object]:
    return {
        "ID livro": "livro",
        "Consulta": "São Paulo",
        "Termo encontrado": "São Paulo",
        "Tipo de correspondência": "Lexical",
        "Similaridade semântica": None,
        "Página inicial": 1,
        "Página final": 1,
        "Bloco anterior": "Bloco anterior com texto longo.",
        "Trecho da ocorrência": "São Paulo aparece neste trecho longo.\nO texto permanece integral.",
        "Bloco posterior": "Bloco posterior com texto longo.",
        "Contexto sociológico": "Desigualdades sociais e regionais",
        "Descrição": "Descrição longa preservada integralmente.",
        "Validação manual": "",
        "Observações do pesquisador": "",
        "_quantidade_no_registro": 1,
        "_pagina_inicial_pdf": 1,
        "_pagina_final_pdf": 1,
        "_texto_normalizado": "sao paulo aparece neste trecho longo",
    }


class TesteApresentacaoXlsx(unittest.TestCase):
    def _verificar_planilha_compacta(self, caminho: Path, aba_principal: str) -> None:
        workbook = load_workbook(caminho)
        for aba in workbook.worksheets:
            for linha in aba.iter_rows():
                for celula in linha:
                    # openpyxl reabre o valor padrão ``False`` como ``None``;
                    # ambos representam ausência explícita de wrap no XLSX.
                    self.assertFalse(
                        celula.alignment.wrap_text,
                        f"{aba.title}!{celula.coordinate} deveria estar sem quebra automática",
                    )
            if aba.max_row >= 2:
                self.assertEqual(aba.row_dimensions[2].height, 15)

        principal = workbook[aba_principal]
        self.assertEqual(principal.cell(1, 1).value, "ID resultado")
        self.assertEqual(principal.cell(2, 1).number_format, "@")

    def test_v1_v2_v3_exportam_todas_as_abas_sem_quebra_automatica(self):
        with tempfile.TemporaryDirectory() as diretorio:
            pasta = Path(diretorio)
            df_lexical = pd.DataFrame([_registro_lexical()])
            caminho_v1 = pasta / "v1.xlsx"
            caminho_v2 = pasta / "v2.xlsx"
            caminho_v3 = pasta / "v3.xlsx"

            v1.salvar_excel_completo(caminho_v1, df_lexical, TERMOS, DIAGNOSTICO)
            v2.salvar_excel_completo(caminho_v2, df_lexical, TERMOS, DIAGNOSTICO)
            v3.salvar_excel_completo(
                caminho_v3,
                pd.DataFrame([_registro_v3()]),
                TERMOS,
                DIAGNOSTICO,
                {"incluir_lexical": True, "incluir_semantica": False},
            )

            self._verificar_planilha_compacta(caminho_v1, "Ocorrencias")
            self._verificar_planilha_compacta(caminho_v2, "Ocorrencias")
            self._verificar_planilha_compacta(caminho_v3, "Resultados")

            self.assertEqual(
                load_workbook(caminho_v1)["Ocorrencias"]["E2"].value,
                "São Paulo aparece neste trecho longo.\nO texto permanece integral.",
            )
            self.assertEqual(load_workbook(caminho_v2)["Ocorrencias"]["B2"].value, "livro")
            self.assertEqual(
                load_workbook(caminho_v3)["Resultados"]["J2"].value,
                "São Paulo aparece neste trecho longo.\nO texto permanece integral.",
            )


if __name__ == "__main__":
    unittest.main()
