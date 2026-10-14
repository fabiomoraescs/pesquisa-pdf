"""Contrato da Fase 9: relatórios Instagram reprodutíveis e privados."""

from __future__ import annotations

import tempfile
import unittest
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from flask_migrate import downgrade, upgrade
from openpyxl import load_workbook
from sqlalchemy import inspect, select

from platform_helpers import create_user, csrf_from, isolated_platform, login
from platform_core.extensions import db
from platform_core.instagram_reports import (
    InstagramReportError,
    artifact_path,
    delete_report,
    eligible_report_runs,
    generate_report,
    report_chart_data,
    report_download_name,
    report_summary_cards,
    resolve_report_runs,
)
from platform_core.models import (
    AnalyticsProject,
    AnalyticsRun,
    InstagramAccount,
    InstagramMedia,
    InstagramMediaObservation,
    InstagramReport,
    PlanTool,
)
from platform_core.services import ANALYTICS_INSTAGRAM_TOOL


UTC = timezone.utc


class FailingReportProvider:
    """Garante que Relatórios nunca acionem o provider de IA global."""

    def __init__(self):
        self.calls = 0

    def generate(self, **_kwargs):
        self.calls += 1
        raise AssertionError("Relatórios não devem chamar providers de IA.")


class InstagramReportsTests(unittest.TestCase):
    def setUp(self):
        self.scope = isolated_platform()
        self.app = self.scope.__enter__()
        self.client = self.app.test_client()
        self.user = create_user()
        self.other = create_user("Outra pessoa", "outra@example.org")
        db.session.add(PlanTool(plan_id="student", tool_id=ANALYTICS_INSTAGRAM_TOOL))
        self.project = AnalyticsProject(owner_user_id=self.user.id, name="Relatórios")
        self.other_project = AnalyticsProject(owner_user_id=self.other.id, name="Outro")
        db.session.add_all((self.project, self.other_project))
        db.session.flush()
        account = InstagramAccount(analytics_project_id=self.project.id, username="relatorios", username_normalized="relatorios")
        db.session.add(account)
        db.session.flush()
        self.media = []
        for index, kind in enumerate(("REELS", "REELS", "IMAGES")):
            item = InstagramMedia(instagram_account_id=account.id, external_id=f"post-{index}", media_type_raw=kind)
            db.session.add(item)
            self.media.append(item)
        db.session.flush()
        self.baseline = self.make_run("baseline", datetime(2026, 1, 1, tzinfo=UTC), datetime(2026, 1, 7, 23, 59, tzinfo=UTC), (100, 200, 300))
        self.middle = self.make_run("middle", datetime(2026, 1, 8, tzinfo=UTC), datetime(2026, 1, 14, 23, 59, tzinfo=UTC), (150, 250, 350))
        self.final = self.make_run("final", datetime(2026, 1, 15, tzinfo=UTC), datetime(2026, 1, 21, 23, 59, tzinfo=UTC), (200, 300, 400))
        db.session.commit()
        login(self.client)

    def tearDown(self):
        self.scope.__exit__(None, None, None)

    def make_run(self, source_hash, start, end, interactions):
        run = AnalyticsRun(
            analytics_project_id=self.project.id, source_kind="csv", status="completed",
            period_start=start, period_end=end, completed_at=end, processor_version="instagram-import-v1",
            source_hash=source_hash, record_count=3,
        )
        db.session.add(run)
        db.session.flush()
        for media, value in zip(self.media, interactions):
            db.session.add(InstagramMediaObservation(
                instagram_media_id=media.id, analytics_run_id=run.id, observed_at=end,
                interactions=value, saves=value // 2, reach=value * 10,
                impressions=value * 20, plays=value * 30,
            ))
        return run

    def make_meta_run(self, source_hash, start, end, interactions, *, preset="last_7_days", status="completed"):
        parameters = {"period_preset": preset, "presentation_timezone": "America/Maceio"}
        if start is not None and end is not None:
            parameters.update({"published_from": start.isoformat(), "published_to": end.isoformat()})
        run = AnalyticsRun(
            analytics_project_id=self.project.id, source_kind="meta_api", status=status,
            completed_at=end if status == "completed" else None, processor_version="instagram-meta-sync-v1",
            source_hash=source_hash, record_count=3, parameters_json=parameters,
        )
        db.session.add(run)
        db.session.flush()
        for media, value in zip(self.media, interactions):
            db.session.add(InstagramMediaObservation(
                instagram_media_id=media.id, analytics_run_id=run.id, observed_at=end or datetime(2026, 10, 14, tzinfo=UTC),
                interactions=value, saves=value // 2, reach=value * 10,
                impressions=value * 20, plays=value * 30,
            ))
        return run

    def values(self, **overrides):
        data = {"period_start": "2026-01-01", "period_end": "2026-01-21"}
        data.update(overrides)
        return data

    def legacy_test_comparison_progression_snapshot_and_exports_are_reproducible(self):
        report = generate_report(self.project, self.values(), actor_id=self.user.id)
        self.assertEqual(report.status, "completed")
        self.assertEqual([item["run"]["id"] for item in report.snapshot_json["evolution"]], [
            self.baseline.id, self.middle.id, self.final.id,
        ])
        interactions = next(item for item in report.snapshot_json["comparison"] if item["key"] == "interactions")
        self.assertEqual((interactions["baseline_value"], interactions["final_value"], interactions["absolute_change"]), (200, 300, 100))
        self.assertAlmostEqual(interactions["relative_change"], .5)
        before = report.snapshot_json
        db.session.add(InstagramMediaObservation(
            instagram_media_id=self.media[0].id, analytics_run_id=self.final.id,
            observed_at=self.final.period_end, interactions=9999,
        ))
        # A unicidade impede esse dado novo no mesmo run; a invariância é do JSON armazenado.
        db.session.rollback()
        self.assertEqual(db.session.get(InstagramReport, report.id).snapshot_json, before)
        pdf = artifact_path(self.project, report, "pdf")
        with pymupdf.open(pdf) as document:
            pdf_text = "".join(page.get_text() for page in document)
            self.assertIn("Análysis Instagram", pdf_text)
            self.assertIn("Síntese e indicadores", pdf_text)
            self.assertIn("Evolução observada", pdf_text)
        workbook = load_workbook(artifact_path(self.project, report, "xlsx"), read_only=True)
        try:
            self.assertIn("Comparison", workbook.sheetnames)
            self.assertIn("Evolution", workbook.sheetnames)
        finally:
            workbook.close()

    def legacy_test_chart_specs_are_snapshot_only_and_cover_composition_comparison_and_evolution(self):
        report = generate_report(self.project, self.values(), actor_id=self.user.id)
        charts = report_chart_data(report.snapshot_json)
        self.assertEqual(charts["media_type_distribution"]["kind"], "pie")
        self.assertEqual(sorted(charts["media_type_distribution"]["values"]), [1, 2])
        self.assertEqual(charts["media_type_performance"]["kind"], "bar")
        self.assertEqual(charts["comparison"]["kind"], "diverging_bar")
        self.assertEqual(charts["comparison"]["zero_value"], 0)
        self.assertEqual(charts["evolution"]["kind"], "line")
        stored_snapshot = report.snapshot_json
        db.session.add(InstagramMediaObservation(
            instagram_media_id=self.media[0].id, analytics_run_id=self.final.id,
            observed_at=self.final.period_end, interactions=9999,
        ))
        db.session.rollback()
        self.assertEqual(report_chart_data(stored_snapshot), charts)

    def legacy_test_summary_cards_are_formatted_from_the_frozen_snapshot(self):
        report = generate_report(self.project, self.values(report_type="executive"), actor_id=self.user.id)
        cards = report_summary_cards(report.snapshot_json)
        self.assertEqual([card["label"] for card in cards], ["Alcance", "Impressões", "Reproduções"])
        self.assertEqual([card["value_display"] for card in cards], ["3.000", "6.000", "9.000"])
        self.assertTrue(all(card["context"] == "Período final" for card in cards))
        self.assertEqual(report_summary_cards(deepcopy(report.snapshot_json)), cards)

    def legacy_test_diverging_chart_keeps_extreme_signed_values_in_independent_rows(self):
        snapshot = {
            "analysis_mode": "comparative_evolution",
            "comparison": [
                {"label": "Alcance", "relative_change": 6.129},
                {"label": "Impressões", "relative_change": -20.0},
                {"label": "Reproduções", "relative_change": 1.845},
                {"label": "Curtidas", "relative_change": -30.8017},
                {"label": "Comentários", "relative_change": -36.7347},
                {"label": "Compartilhamentos", "relative_change": -26.9841},
            ],
        }
        chart = report_chart_data(snapshot)["comparison"]
        self.assertEqual(chart["kind"], "diverging_bar")
        self.assertEqual(chart["zero_value"], 0)
        self.assertEqual(chart["labels"], [item["label"] for item in snapshot["comparison"]])
        self.assertEqual(chart["values"], [612.9, -2000.0, 184.5, -3080.17, -3673.47, -2698.41])

    def legacy_test_executive_is_concise_and_analytical_exposes_deeper_pdf_sections(self):
        executive = generate_report(self.project, self.values(
            title="Síntese executiva", report_type="executive",
            sections=("summary", "metrics", "findings", "recommendations", "actions", "learnings", "methodology"),
            formats=("pdf",),
        ), actor_id=self.user.id)
        analytical = generate_report(self.project, self.values(
            title="Leitura analítica", report_type="analytical",
            sections=("summary", "metrics", "contents", "comments", "findings", "recommendations", "actions", "evaluations", "learnings", "methodology"),
            formats=("pdf",),
        ), actor_id=self.user.id)
        with pymupdf.open(artifact_path(self.project, executive, "pdf")) as document:
            executive_text = "".join(page.get_text() for page in document)
        with pymupdf.open(artifact_path(self.project, analytical, "pdf")) as document:
            analytical_text = "".join(page.get_text() for page in document)
        self.assertIn("Resumo executivo", executive_text)
        self.assertIn("3.000", executive_text)
        self.assertIn("6.000", executive_text)
        self.assertIn("9.000", executive_text)
        self.assertNotIn("Conteúdos", executive_text)
        self.assertIn("3. Conteúdos", analytical_text)
        self.assertIn("14. Metodologia", analytical_text)
        self.assertIn("15. Rastreabilidade", analytical_text)

    def legacy_test_detail_page_is_a_visual_dashboard_with_downloads_not_an_html_export(self):
        report = generate_report(self.project, self.values(
            sections=("summary", "metrics", "contents", "comments", "findings", "recommendations", "actions", "evaluations", "learnings", "methodology"),
        ), actor_id=self.user.id)
        html = self.client.get(f"/analytics/instagram/projects/{self.project.id}/reports/{report.id}").get_data(as_text=True)
        self.assertIn("Indicadores principais do relatório", html)
        self.assertIn("Relatório completo", html)
        self.assertIn("Baixar PDF", html)
        self.assertIn("instagram-report-chart-data", html)
        self.assertIn("chart-card", html)
        self.assertNotIn("report-chart-svg", html)
        self.assertIn('class="col-12 col-md-4"', html)
        self.assertIn('<div class="metric-card"><span>Alcance</span><strong>3.000</strong>', html)
        self.assertIn('<small class="platform-muted">Período final</small>', html)
        self.assertIn("Fonte e reprodutibilidade", html)
        self.assertNotIn("Comparação de métricas e cobertura", html)
        self.assertNotIn("Achados e evidências registradas", html)
        self.assertNotIn("Método, limitações e rastreabilidade", html)

    def legacy_test_reports_do_not_offer_or_call_gemini(self):
        provider = FailingReportProvider()
        self.app.extensions["assistant_ai_provider"] = provider
        report = generate_report(self.project, self.values(), actor_id=self.user.id)
        self.assertEqual(provider.calls, 0)
        self.assertNotIn("ai_interpretation", report.snapshot_json)
        form = self.client.get(f"/analytics/instagram/projects/{self.project.id}/reports/new").get_data(as_text=True)
        self.assertNotIn("report-ai-interpretation", form)
        self.assertNotIn("Interpretação assistida por IA", form)
        self.assertNotIn("Gemini", form)

    def legacy_test_historical_ai_snapshot_is_ignored_without_calling_a_provider(self):
        provider = FailingReportProvider()
        self.app.extensions["assistant_ai_provider"] = provider
        report = generate_report(self.project, self.values(), actor_id=self.user.id)
        snapshot = deepcopy(report.snapshot_json)
        snapshot["ai_interpretation"] = {
            "enabled": True,
            "sections": {"primary_indicators": {"summary": "Síntese Gemini histórica."}},
        }
        report.snapshot_json = snapshot
        db.session.commit()
        html = self.client.get(f"/analytics/instagram/projects/{self.project.id}/reports/{report.id}").get_data(as_text=True)
        self.assertNotIn("Síntese Gemini histórica.", html)
        self.assertNotIn("assistência de IA", html)
        with pymupdf.open(stream=_pdf(snapshot, report.title), filetype="pdf") as document:
            text = "".join(page.get_text() for page in document)
        self.assertNotIn("Síntese Gemini histórica.", text)
        self.assertNotIn("assistência de IA", text)
        self.assertEqual(provider.calls, 0)

    def legacy_test_pdf_bar_layout_uses_content_aware_label_space(self):
        layout = _pdf_bar_layout(
            ["Alcance", "Impressões", "Reproduções", "Compartilhamentos"], width=595.0, margin=50.0,
        )
        self.assertEqual(layout["bar_x"] - layout["label_right"], layout["gap"])
        self.assertLess(layout["bar_x"], 160.0)
        self.assertGreater(layout["bar_end"], layout["bar_x"])

    def legacy_test_validation_rejects_overlapping_periods_and_missing_runs(self):
        with self.assertRaisesRegex(InstagramReportError, "terminar antes"):
            generate_report(self.project, self.values(final_run_id=self.baseline.id), actor_id=self.user.id)
        with self.assertRaisesRegex(InstagramReportError, "execução concluída"):
            generate_report(self.project, self.values(baseline_run_id="00000000-0000-0000-0000-000000000000"), actor_id=self.user.id)

    def legacy_test_single_period_uses_one_source_run_and_formula_like_content_is_safe_in_xlsx(self):
        self.media[0].caption_original = "=formula não executável"
        db.session.commit()
        report = generate_report(self.project, self.values(
            analysis_mode="single_period", baseline_run_id=self.final.id, final_run_id=self.final.id,
            formats=("xlsx",), sections=("summary", "metrics", "contents", "methodology"),
        ), actor_id=self.user.id)
        self.assertEqual([row["id"] for row in report.snapshot_json["source_runs"]], [self.final.id])
        workbook = load_workbook(artifact_path(self.project, report, "xlsx"), data_only=False)
        try:
            captions = [row[2].value for row in workbook["Contents"].iter_rows(min_row=2)]
            self.assertIn("'=formula não executável", captions)
        finally:
            workbook.close()

    def legacy_test_delete_only_removes_report_and_its_artifacts(self):
        report = generate_report(self.project, self.values(), actor_id=self.user.id)
        path = artifact_path(self.project, report, "pdf")
        delete_report(self.project, report, " deletar ")
        self.assertIsNone(db.session.get(InstagramReport, report.id))
        self.assertFalse(path.exists())
        self.assertIsNotNone(db.session.get(AnalyticsRun, self.baseline.id))
        with self.assertRaisesRegex(InstagramReportError, "exatamente"):
            delete_report(self.project, InstagramReport(analytics_project_id=self.project.id), "apagar")

    def legacy_test_routes_have_history_csrf_ownership_archive_and_download_contract(self):
        reports_url = f"/analytics/instagram/projects/{self.project.id}/reports"
        history = self.client.get(reports_url)
        self.assertIn("Ainda não há relatórios gerados para este projeto.", history.get_data(as_text=True))
        new_url = reports_url + "/new"
        form = self.client.get(new_url)
        values = self.values()
        posted = {key: value for key, value in values.items() if key not in {"formats", "sections"}}
        posted["formats"] = list(values["formats"])
        posted["sections"] = list(values["sections"])
        self.assertEqual(self.client.post(new_url, data=posted).status_code, 400)
        response = self.client.post(new_url, data={**posted, "csrf_token": csrf_from(form)})
        self.assertEqual(response.status_code, 302)
        report = db.session.scalar(select(InstagramReport))
        self.assertEqual(self.client.get(reports_url + f"/{report.id}/download/pdf").status_code, 200)
        self.client.post("/logout", data={"csrf_token": csrf_from(self.client.get("/"))})
        login(self.client, self.other.email)
        self.assertEqual(self.client.get(reports_url + f"/{report.id}").status_code, 404)
        self.client.post("/logout", data={"csrf_token": csrf_from(self.client.get("/"))})
        login(self.client)
        self.project.status, self.project.archived_at = "archived", datetime.now(UTC)
        db.session.commit()
        self.assertEqual(self.client.get(reports_url + f"/{report.id}").status_code, 200)
        self.assertEqual(self.client.get(new_url).status_code, 404)
        self.assertEqual(self.client.get(reports_url + f"/{report.id}/delete").status_code, 404)

    def legacy_test_form_fields_and_trash_action_follow_visual_contract(self):
        reports_url = f"/analytics/instagram/projects/{self.project.id}/reports"
        form = self.client.get(reports_url + "/new")
        form_html = form.get_data(as_text=True)
        self.assertIn('id="report-final-run-group" class="col-md-6" hidden', form_html)
        self.assertIn("Execução principal", form_html)
        self.assertNotIn("No período único, a mesma execução é usada como baseline e final.", form_html)
        self.assertIn("js/instagram_reports.js", form_html)
        values = self.values(final_run_id=self.baseline.id)
        posted = {key: value for key, value in values.items() if key not in {"formats", "sections"}}
        posted.update({"formats": list(values["formats"]), "sections": list(values["sections"]), "csrf_token": csrf_from(form)})
        comparative = self.client.post(reports_url + "/new", data=posted)
        comparative_html = comparative.get_data(as_text=True)
        self.assertEqual(comparative.status_code, 400)
        self.assertIn("Execução inicial", comparative_html)
        self.assertNotIn('id="report-final-run-group" class="col-md-6" hidden', comparative_html)

        report = generate_report(self.project, self.values(), actor_id=self.user.id)
        history = self.client.get(reports_url).get_data(as_text=True)
        detail = self.client.get(reports_url + f"/{report.id}").get_data(as_text=True)
        for page in (history, detail):
            self.assertIn('title="Excluir relatório"', page)
            self.assertIn('aria-label="Excluir relatório"', page)
            self.assertIn("M4 7h16", page)
        self.assertIn("3.000", detail)
        self.assertIn("6.000", detail)
        self.assertIn("9.000", detail)
        self.assertIn("Período final", detail)
        self.assertIn("Relatório completo", detail)
        self.assertIn("instagram-report-chart-data", detail)
        self.assertNotIn("report-chart-svg", detail)
        self.assertNotIn("Comparação de métricas e cobertura", detail)

    def test_period_resolves_only_completed_dated_runs_fully_contained_and_deduplicated(self):
        partial = self.make_run(
            "partial", datetime(2025, 12, 31, 23, 59, tzinfo=UTC),
            datetime(2026, 1, 2, 23, 59, tzinfo=UTC), (1, 2, 3),
        )
        replacement = self.make_run(
            "middle-retry", self.middle.period_start, self.middle.period_end, (151, 251, 351),
        )
        replacement.completed_at = datetime(2026, 1, 15, tzinfo=UTC)
        db.session.commit()
        runs = resolve_report_runs(self.project, datetime(2026, 1, 1, tzinfo=UTC).date(), datetime(2026, 1, 21, tzinfo=UTC).date())
        self.assertEqual([run.id for run in runs], [self.baseline.id, replacement.id, self.final.id])
        self.assertNotIn(partial.id, [run.id for run in runs])

    def test_meta_windows_all_history_and_unknown_legacy_follow_the_shared_resolution_contract(self):
        meta_last_seven = self.make_meta_run(
            "meta-last-seven", datetime(2026, 10, 8, 3, tzinfo=UTC), datetime(2026, 10, 14, 15, tzinfo=UTC), (10, 20, 30),
        )
        meta_last_thirty = self.make_meta_run(
            "meta-last-thirty", datetime(2026, 9, 11, 3, tzinfo=UTC), datetime(2026, 10, 10, 4, tzinfo=UTC), (11, 21, 31),
            preset="last_30_days",
        )
        partial = self.make_meta_run(
            "meta-partial", datetime(2026, 10, 7, 3, tzinfo=UTC), datetime(2026, 10, 14, 15, tzinfo=UTC), (12, 22, 32),
        )
        all_history = self.make_meta_run("meta-history", None, None, (13, 23, 33), preset="all_history")
        legacy = AnalyticsRun(
            analytics_project_id=self.project.id, source_kind="meta_api", status="completed", source_hash="meta-legacy",
            completed_at=datetime(2026, 10, 14, tzinfo=UTC), parameters_json={},
        )
        pending = self.make_meta_run(
            "meta-pending", datetime(2026, 10, 8, 3, tzinfo=UTC), datetime(2026, 10, 14, 15, tzinfo=UTC),
            (14, 24, 34), status="pending",
        )
        db.session.add(legacy)
        db.session.commit()

        week = resolve_report_runs(self.project, datetime(2026, 10, 8, tzinfo=UTC).date(), datetime(2026, 10, 14, tzinfo=UTC).date())
        self.assertEqual([run.id for run in week], [meta_last_seven.id])
        self.assertNotIn(meta_last_thirty.id, [run.id for run in week])
        self.assertNotIn(partial.id, [run.id for run in week])
        self.assertNotIn(legacy.id, [run.id for run in week])
        self.assertNotIn(pending.id, [run.id for run in week])

        history = resolve_report_runs(self.project, None, None)
        self.assertIn(all_history.id, [run.id for run in history])
        self.assertNotIn(legacy.id, [run.id for run in history])
        self.assertNotIn(pending.id, [run.id for run in history])
        report = generate_report(self.project, {"period_preset": "all"}, actor_id=self.user.id)
        self.assertEqual(report.snapshot_json["requested_period"], {"preset": "all"})
        snapshot = next(item for item in report.snapshot_json["runs_used"] if item["id"] == all_history.id)
        self.assertEqual((snapshot["period_kind"], snapshot["period_start"], snapshot["period_end"]), ("all_history", None, None))
        self.assertIsNone(db.session.get(AnalyticsRun, meta_last_seven.id).period_start)
        self.assertIsNone(db.session.get(AnalyticsRun, meta_last_seven.id).period_end)

    def test_one_two_and_three_observations_have_adaptive_snapshot_and_xlsx_sheets(self):
        one = generate_report(self.project, self.values(period_end="2026-01-07"), actor_id=self.user.id)
        self.assertEqual(one.formats_json, ["xlsx"])
        self.assertEqual(one.baseline_run_id, self.baseline.id)
        self.assertEqual(one.final_run_id, self.baseline.id)
        self.assertEqual(one.snapshot_json["comparison"], [])
        self.assertEqual(one.snapshot_json["evolution"], [])
        one_book = load_workbook(artifact_path(self.project, one, "xlsx"), read_only=True)
        try:
            self.assertNotIn("Comparison", one_book.sheetnames)
            self.assertNotIn("Evolution", one_book.sheetnames)
        finally:
            one_book.close()

        two = generate_report(self.project, self.values(period_end="2026-01-14"), actor_id=self.user.id)
        self.assertEqual(len(two.snapshot_json["runs_used"]), 2)
        self.assertTrue(two.snapshot_json["comparison"])
        self.assertEqual(two.snapshot_json["evolution"], [])
        two_book = load_workbook(artifact_path(self.project, two, "xlsx"), read_only=True)
        try:
            self.assertIn("Comparison", two_book.sheetnames)
            self.assertNotIn("Evolution", two_book.sheetnames)
        finally:
            two_book.close()

        three = generate_report(self.project, self.values(), actor_id=self.user.id)
        self.assertEqual([item["id"] for item in three.snapshot_json["runs_used"]], [self.baseline.id, self.middle.id, self.final.id])
        self.assertEqual(len(three.snapshot_json["evolution"]), 3)
        book = load_workbook(artifact_path(self.project, three, "xlsx"), read_only=True)
        try:
            self.assertIn("Comparison", book.sheetnames)
            self.assertIn("Evolution", book.sheetnames)
        finally:
            book.close()

    def test_no_eligible_run_rejects_generation_without_creating_a_report(self):
        with self.assertRaisesRegex(InstagramReportError, "período conhecido"):
            generate_report(self.project, self.values(period_start="2026-02-01", period_end="2026-02-02"), actor_id=self.user.id)
        self.assertIsNone(db.session.scalar(select(InstagramReport)))

    def test_snapshot_dates_and_xlsx_filename_are_reproducible_and_formula_safe(self):
        self.media[0].caption_original = "=formula não executável"
        db.session.commit()
        report = generate_report(self.project, self.values(), actor_id=self.user.id)
        self.assertEqual(report.snapshot_json["requested_period"], {"start": "2026-01-01", "end": "2026-01-21"})
        self.assertEqual(report_download_name(report, "xlsx"), "instagram_relatorio_2026-01-01_2026-01-21.xlsx")
        before = deepcopy(report.snapshot_json)
        workbook = load_workbook(artifact_path(self.project, report, "xlsx"), data_only=False)
        try:
            captions = [row[3].value for row in workbook["Contents"].iter_rows(min_row=2)]
            self.assertIn("'=formula não executável", captions)
        finally:
            workbook.close()
        self.assertEqual(db.session.get(InstagramReport, report.id).snapshot_json, before)

    def test_period_form_routes_and_history_hide_manual_controls_and_preserve_contracts(self):
        reports_url = f"/analytics/instagram/projects/{self.project.id}/reports"
        form = self.client.get(reports_url + "/new")
        form_html = form.get_data(as_text=True)
        self.assertIn("Período do relatório", form_html)
        self.assertIn('name="period_start"', form_html)
        self.assertIn('name="period_end"', form_html)
        self.assertNotIn('name="report_type"', form_html)
        self.assertNotIn('name="analysis_mode"', form_html)
        self.assertNotIn('name="baseline_run_id"', form_html)
        self.assertNotIn('id="format-pdf"', form_html)
        self.assertEqual(self.client.post(reports_url + "/new", data=self.values()).status_code, 400)
        response = self.client.post(reports_url + "/new", data={**self.values(), "csrf_token": csrf_from(form)})
        self.assertEqual(response.status_code, 302)
        report = db.session.scalar(select(InstagramReport))
        self.assertEqual(self.client.get(reports_url + f"/{report.id}/download/xlsx").status_code, 200)
        self.assertEqual(self.client.get(reports_url + f"/{report.id}/download/pdf").status_code, 404)
        history = self.client.get(reports_url).get_data(as_text=True)
        detail = self.client.get(reports_url + f"/{report.id}").get_data(as_text=True)
        self.assertIn("XLSX", history)
        self.assertNotIn("Executivo", history)
        self.assertNotIn("Analítico", history)
        self.assertNotIn("Tipo e modo", detail)
        self.assertNotIn("Baixar PDF", detail)
        self.assertIn("Baixar XLSX", detail)
        self.assertIn(f"/reports/{report.id}/download/xlsx", detail)
        self.assertIn("Execuções usadas", detail)
        self.assertNotIn("Observações usadas", detail)
        self.assertNotIn("report-downloads", detail)
        self.assertNotIn(">Download</span>", detail)
        workbook = load_workbook(artifact_path(self.project, report, "xlsx"), read_only=True)
        try:
            summary_rows = {
                row[0].value: row[1].value
                for row in workbook["Summary"].iter_rows(min_row=2)
            }
            self.assertEqual(summary_rows["Execuções usadas"], len(report.snapshot_json["runs_used"]))
            self.assertNotIn("Observações usadas", summary_rows)
        finally:
            workbook.close()
        self.client.post("/logout", data={"csrf_token": csrf_from(self.client.get("/"))})
        login(self.client, self.other.email)
        self.assertEqual(self.client.get(reports_url + f"/{report.id}").status_code, 404)
        self.client.post("/logout", data={"csrf_token": csrf_from(self.client.get("/"))})
        login(self.client)
        self.project.status, self.project.archived_at = "archived", datetime.now(UTC)
        db.session.commit()
        self.assertEqual(self.client.get(reports_url + f"/{report.id}").status_code, 200)
        self.assertEqual(self.client.get(reports_url + "/new").status_code, 404)


class InstagramReportsMigrationTests(unittest.TestCase):
    def test_report_migration_is_reversible_child_of_learning_head(self):
        from app import create_app

        migrations = str(Path(__file__).resolve().parents[1] / "migrations")
        with tempfile.TemporaryDirectory(prefix="pesquisapdf-reports-migration-") as root:
            database = Path(root) / "platform.sqlite3"
            app = create_app({"TESTING": True, "SQLALCHEMY_DATABASE_URI": f"sqlite:///{database.as_posix()}", "PLATFORM_DATA_DIR": root, "SECRET_KEY": "reports-migration-test"})
            with app.app_context():
                upgrade(directory=migrations, revision="e2f3a4b5c6d")
                self.assertNotIn("instagram_reports", inspect(db.engine).get_table_names())
                upgrade(directory=migrations, revision="f3a4b5c6d7e")
                self.assertIn("instagram_reports", inspect(db.engine).get_table_names())
                downgrade(directory=migrations, revision="e2f3a4b5c6d")
                self.assertNotIn("instagram_reports", inspect(db.engine).get_table_names())
                db.session.remove()
                db.engine.dispose()
