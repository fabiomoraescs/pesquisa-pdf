"""Contrato comparativo dos métodos de busca nas três ferramentas."""

import json
from io import BytesIO
from tempfile import TemporaryDirectory
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook
from sqlalchemy import select

import test_qualitative_automatic as automatic_fixture
import test_qualitative_context as context_fixture

from analyzer import v1, v3
from analyzer.common import criar_dashboard
from analyzer.search_matching import SearchPatternError
from historico_racial.entities import Entidade
from historico_racial.exporter import gerar_xlsx
from historico_racial.occurrences import BuscadorLexical
from platform_core.extensions import db
from platform_core.models import QualitativeCode
from platform_core.qualitative_annotations import page_excerpts
from platform_core.qualitative_expanded_search import search_lexical
from platform_core.qualitative_search import search_qualitative_document


class SearchMethodConsistencyTests(unittest.TestCase):
    setUp = context_fixture.QualitativeContextTests.setUp
    tearDown = context_fixture.QualitativeContextTests.tearDown
    fixture_pages = automatic_fixture.QualitativeAutomaticTests.fixture_pages
    auto_url = automatic_fixture.QualitativeAutomaticTests.auto_url
    automatic = automatic_fixture.QualitativeAutomaticTests.automatic
    codings = automatic_fixture.QualitativeAutomaticTests.codings

    def setUp(self):
        context_fixture.QualitativeContextTests.setUp(self)
        self.text = (
            "Raça, raças, racial, raciais, racismo, racializado e racializados. "
            "racional não entra."
        )
        self.fixture_pages([[self.text]])
        self.entity = Entidade("raca", "raça", ("raça",), "conceito", ("teste",))

    def _qualitative_words(self, result):
        return [self.text[item["start_offset"]:item["end_offset"]] for item in result["results"]]

    def test_literal_is_equivalent_without_family_expansion(self):
        qualitative = search_qualitative_document(self.analysis, self.document.id, "raça")
        terms = [self.text[start:end] for start, end, _ in v1.encontrar_ocorrencias(
            self.text, "raça", metodo="literal",
        )]
        structured = [item.forma_original_no_texto for item in BuscadorLexical(
            (self.entity,), metodo="literal",
        ).localizar(self.text)]
        self.assertEqual(self._qualitative_words(qualitative), ["Raça", "raça"])
        self.assertEqual(terms, self._qualitative_words(qualitative))
        self.assertEqual(structured, self._qualitative_words(qualitative))

    def test_lexical_family_is_shared_by_qualitative_terms_and_structured(self):
        expected = ["Raça", "raças", "racial", "raciais", "racismo", "racializado", "racializados"]
        qualitative = self._qualitative_words(search_lexical(self.analysis, self.document.id, "raça"))
        terms = [self.text[start:end] for start, end, _ in v1.encontrar_ocorrencias(self.text, "raça")]
        structured = [item.forma_original_no_texto for item in BuscadorLexical(
            (self.entity,), metodo="lexical",
        ).localizar(self.text)]
        self.assertEqual(qualitative, expected)
        self.assertEqual(terms, expected)
        self.assertEqual(structured, expected)

    def test_lexical_codes_preserve_the_concrete_match_and_query_origin(self):
        expected_codes = ["raça", "raças", "racial", "raciais", "racismo", "racializado", "racializados"]
        qualitative = self.automatic(q="raça", mode="lexical")
        self.assertEqual(qualitative.status_code, 201, qualitative.json)
        self.assertEqual({code.name for code in db.session.scalars(select(QualitativeCode))}, set(expected_codes))
        self.assertEqual({coding.source_query for coding in self.codings()}, {"raça"})
        labels = {
            code["name"]
            for excerpt in page_excerpts(self.analysis, self.document.id, 1)["excerpts"]
            for code in excerpt["codes"]
        }
        self.assertEqual(labels, set(expected_codes))

        terms = dict(v1.agrupar_formas_encontradas(self.text, "raça"))
        self.assertEqual(list(terms), expected_codes)
        self.assertTrue(all(terms[name] == 1 for name in expected_codes))

        blocks = [{
            "texto": self.text, "pagina": 1, "pagina_pdf": 1, "unidade": "",
            "capitulo": "", "secao": "", "subsecao": "", "metodo": "texto",
            "tipo_ocorrencia": "texto",
        }]
        with patch("analyzer.v1.extrair_pdf", return_value=(blocks, 1, 0, None)):
            term_records, term_diagnostic = v1.analisar_pdf(
                Path("controle.pdf"), [{"categoria": "TERMO", "termo": "raça"}],
                {"metodo_busca": "lexical"},
            )
        self.assertEqual([item["Consulta"] for item in term_records], ["raça"] * len(expected_codes))
        self.assertEqual([item["Termo"] for item in term_records], expected_codes)
        with TemporaryDirectory() as temporary:
            output = Path(temporary) / "resultado.xlsx"
            v1.salvar_excel_completo(
                output, pd.DataFrame(term_records),
                pd.DataFrame([{"Categoria do termo": "TERMO", "Termo": "raça"}]),
                pd.DataFrame([term_diagnostic]), {"metodo_busca": "lexical"},
            )
            workbook = load_workbook(output, read_only=True)
            try:
                worksheet = workbook["Ocorrencias"]
                headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
                query_column = headers.index("Consulta")
                term_column = headers.index("Termo")
                rows = list(worksheet.iter_rows(min_row=2, values_only=True))
                self.assertEqual([row[query_column] for row in rows], ["raça"] * len(expected_codes))
                self.assertEqual([row[term_column] for row in rows], expected_codes)
            finally:
                workbook.close()

        hybrid_records = v3._registros_lexicais(
            Path("controle.pdf"), blocks,
            [{"categoria": "TERMO", "termo": "raça"}],
        )
        self.assertEqual([item["Termo encontrado"] for item in hybrid_records], expected_codes)
        self.assertEqual({item["Consulta"] for item in hybrid_records}, {"raça"})

        structured = BuscadorLexical((self.entity,), metodo="lexical").localizar(self.text)
        self.assertEqual([item.variante_configurada for item in structured], ["raça"] * len(expected_codes))
        self.assertEqual([item.termo_encontrado for item in structured], expected_codes)
        structured_workbook = load_workbook(BytesIO(gerar_xlsx({
            "metodo_analise": "lexical", "documentos": [],
            "ocorrencias": [{
                "id_ocorrencia": f"occ-{index}", "id_documento": "doc-1", "arquivo_pdf": "controle.pdf",
                "pagina_pdf": 1, "categoria_busca": "conceito", "tipo_entidade": "conceito",
                "id_entidade": self.entity.id_entidade, "entidade_canonica": self.entity.forma_canonica,
                "variante_configurada": item.variante_configurada,
                "termo_encontrado": item.termo_encontrado,
                "forma_original_no_texto": item.forma_original_no_texto, "grupo": ["teste"],
                "trecho_anterior": "", "trecho_ocorrencia": self.text, "trecho_posterior": "",
                "contexto_completo": self.text,
            } for index, item in enumerate(structured, start=1)],
        }).getvalue()), read_only=True)
        try:
            structured_sheet = structured_workbook["OCORRENCIAS"]
            structured_headers = [cell.value for cell in next(structured_sheet.iter_rows(min_row=1, max_row=1))]
            variant_column = structured_headers.index("Variante configurada")
            found_column = structured_headers.index("Termo encontrado")
            structured_rows = list(structured_sheet.iter_rows(min_row=2, values_only=True))
            self.assertEqual([row[variant_column] for row in structured_rows], ["raça"] * len(expected_codes))
            self.assertEqual([row[found_column] for row in structured_rows], expected_codes)
        finally:
            structured_workbook.close()

    def test_lexical_export_keeps_all_64_occurrences_above_fifty(self):
        """O dashboard e o XLSX devem receber uma linha para cada span lexical."""
        blocks = [{
            "texto": "racismo racismo" if index < 14 else "racismo",
            "pagina": index + 1,
            "pagina_pdf": index + 1,
            "unidade": "",
            "capitulo": "",
            "secao": "",
            "subsecao": "",
            "metodo": "texto",
            "tipo_ocorrencia": "texto",
        } for index in range(50)]
        raw_matches = sum(
            len(v1.encontrar_ocorrencias(block["texto"], "raça", metodo="lexical"))
            for block in blocks
        )
        self.assertEqual(raw_matches, 64)

        with patch("analyzer.v1.extrair_pdf", return_value=(blocks, 50, 0, None)):
            records, diagnostic = v1.analisar_pdf(
                Path("controle.pdf"), [{"categoria": "TERMO", "termo": "raça"}],
                {"metodo_busca": "lexical"},
            )

        self.assertEqual(len(records), 64)
        self.assertEqual(sum(item["_quantidade_no_registro"] for item in records), 64)
        self.assertTrue(all(item["_quantidade_no_registro"] == 1 for item in records))
        self.assertEqual(
            [(item["_inicio_ocorrencia"], item["_fim_ocorrencia"]) for item in records[:2]],
            [(0, 7), (8, 15)],
        )

        occurrences = pd.DataFrame(records)
        dashboard = criar_dashboard({
            "versao": "v1",
            "ocorrencias": occurrences,
            "termos": [{"termo": "raça"}],
            "diagnosticos": pd.DataFrame([diagnostic]),
            "quantidade_pdfs": 1,
            "quantidade_termos": 1,
        })
        self.assertEqual(dashboard["indicadores"]["ocorrencias"], 64)

        with TemporaryDirectory() as temporary:
            output = Path(temporary) / "resultado.xlsx"
            v1.salvar_excel_completo(
                output,
                occurrences,
                pd.DataFrame([{"Categoria do termo": "TERMO", "Termo": "raça"}]),
                pd.DataFrame([diagnostic]),
                {"metodo_busca": "lexical"},
            )
            workbook = load_workbook(output, read_only=True)
            try:
                rows = list(workbook["Ocorrencias"].iter_rows(min_row=2, values_only=True))
                self.assertEqual(len(rows), 64)
            finally:
                workbook.close()

    def test_lexical_export_keeps_each_repeated_found_variant(self):
        text = " ".join(["raça racial racismo racializado"] * 3)
        blocks = [{
            "texto": text,
            "pagina": 1,
            "pagina_pdf": 1,
            "unidade": "",
            "capitulo": "",
            "secao": "",
            "subsecao": "",
            "metodo": "texto",
            "tipo_ocorrencia": "texto",
        }]
        with patch("analyzer.v1.extrair_pdf", return_value=(blocks, 1, 0, None)):
            records, diagnostic = v1.analisar_pdf(
                Path("variantes.pdf"), [{"categoria": "TERMO", "termo": "raça"}],
                {"metodo_busca": "lexical"},
            )

        expected_terms = ["raça", "racial", "racismo", "racializado"] * 3
        self.assertEqual([item["Termo"] for item in records], expected_terms)
        self.assertEqual(len(records), 12)
        self.assertEqual(len({(item["_inicio_ocorrencia"], item["_fim_ocorrencia"]) for item in records}), 12)

        with TemporaryDirectory() as temporary:
            output = Path(temporary) / "variantes.xlsx"
            v1.salvar_excel_completo(
                output,
                pd.DataFrame(records),
                pd.DataFrame([{"Categoria do termo": "TERMO", "Termo": "raça"}]),
                pd.DataFrame([diagnostic]),
                {"metodo_busca": "lexical"},
            )
            workbook = load_workbook(output, read_only=True)
            try:
                worksheet = workbook["Ocorrencias"]
                headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
                term_column = headers.index("Termo")
                exported_terms = [
                    row[term_column]
                    for row in worksheet.iter_rows(min_row=2, values_only=True)
                ]
                self.assertEqual(exported_terms, expected_terms)
            finally:
                workbook.close()

    def test_regex_is_literal_only_and_reports_invalid_patterns(self):
        query = r"(?i)\braças?\b"
        qualitative = self._qualitative_words(search_qualitative_document(
            self.analysis, self.document.id, query, grep=True,
        ))
        terms = [self.text[start:end] for start, end, _ in v1.encontrar_ocorrencias(
            self.text, query, metodo="literal", usar_regex=True,
        )]
        structured = [item.forma_original_no_texto for item in BuscadorLexical(
            (Entidade("raca", "raça", (query,), "conceito", ("teste",)),),
            metodo="literal", usar_regex=True,
        ).localizar(self.text)]
        self.assertEqual(qualitative, ["Raça", "raças"])
        self.assertEqual(terms, qualitative)
        self.assertEqual(structured, qualitative)
        with self.assertRaises(SearchPatternError):
            v1.encontrar_ocorrencias(self.text, "(", metodo="literal", usar_regex=True)
        with self.assertRaises(ValueError):
            v1.encontrar_ocorrencias(self.text, "raça", usar_regex=True)

    def test_hybrid_keeps_the_shared_lexical_stage_without_loading_a_model(self):
        blocks = [{
            "texto": self.text, "pagina": 1, "pagina_pdf": 1, "unidade": "",
            "capitulo": "", "secao": "", "subsecao": "", "metodo": "texto",
            "tipo_ocorrencia": "texto",
        }]
        records = v3._registros_lexicais(
            Path("controle.pdf"), blocks, [{"categoria": "TERMO", "termo": "raça"}],
        )
        self.assertEqual([item["Termo encontrado"] for item in records],
                         ["raça", "raças", "racial", "raciais", "racismo", "racializado", "racializados"])
        self.assertEqual(sum(item["_quantidade_no_registro"] for item in records), 7)
        with (patch("analyzer.v3.carregar_modelo_semantico", side_effect=AssertionError("não carregar")),
              patch("analyzer.v3.v1.extrair_pdf", return_value=(blocks, 1, 0, None))):
            v3.analisar_pdf(Path("controle.pdf"), [], {"incluir_lexical": True, "incluir_semantica": False})


class SearchMethodFrontendAndTutorialTests(unittest.TestCase):
    def test_forms_offer_valid_method_states_and_result_keeps_exports_out_of_view(self):
        root = Path(__file__).resolve().parents[1]
        terms_template = (root / "templates" / "index.html").read_text(encoding="utf-8")
        structured_template = (root / "templates" / "historico_racial" / "index.html").read_text(encoding="utf-8")
        structured_js = (root / "static" / "js" / "historico_racial.js").read_text(encoding="utf-8")
        result_template = (root / "templates" / "resultado.html").read_text(encoding="utf-8")
        structured_result_template = (root / "templates" / "historico_racial" / "resultado.html").read_text(encoding="utf-8")
        stylesheet = (root / "static" / "css" / "platform.css").read_text(encoding="utf-8")
        app_source = (root / "app.py").read_text(encoding="utf-8")
        structured_route_source = (root / "historico_racial" / "routes.py").read_text(encoding="utf-8")
        actions = (root / "templates" / "platform" / "_analysis_actions.html").read_text(encoding="utf-8")
        for template, field in ((terms_template, "versao"), (structured_template, "metodo_analise")):
            self.assertIn('class="platform-method-toggles"', template)
            self.assertIn(f'name="{field}" value="literal"', template)
            self.assertIn(f'name="{field}" value="v1"' if field == "versao" else f'name="{field}" value="lexical"', template)
            self.assertIn(f'name="{field}" value="v3"' if field == "versao" else f'name="{field}" value="hibrido"', template)
            self.assertIn('name="usar_regex"', template)
            self.assertLess(template.index('value="literal"'), template.index('value="v1"' if field == "versao" else 'value="lexical"'))
            self.assertLess(template.index('value="v1"' if field == "versao" else 'value="lexical"'), template.index('value="v3"' if field == "versao" else 'value="hibrido"'))
            self.assertLess(template.index('value="v3"' if field == "versao" else 'value="hibrido"'), template.index('name="usar_regex"'))
            self.assertIn('class="platform-qualitative-switch"', template)
        self.assertIn("regex.disabled = !literal", terms_template)
        self.assertIn("if (!literal) regex.checked = false", terms_template)
        self.assertIn("regex.disabled = !literal", structured_js)
        self.assertIn("if (!literal) regex.checked = false", structured_js)
        self.assertNotIn("Planilhas geradas", result_template)
        self.assertIn("url_for('analyses.excel'", actions)
        self.assertIn('.platform-method-toggles { display: flex;', stylesheet)
        self.assertNotIn('col-xl-7', result_template)
        self.assertNotIn('col-xl-5', result_template)
        self.assertNotIn('col-xl-7', structured_result_template)
        self.assertNotIn('col-xl-5', structured_result_template)
        self.assertIn('request.form.get("versao", "literal")', app_source)
        self.assertIn('request.form.get("metodo_analise", "literal")', structured_route_source)

    def test_tutorials_document_literal_regex_and_variant_scope(self):
        root = Path(__file__).resolve().parents[1]
        registry = json.loads((root / "resources" / "platform_tutorials.json").read_text(encoding="utf-8"))
        tutorials = {item["key"]: item for item in registry["tutorials"]}
        terms = " ".join(
            [tutorials["term_search"]["summary"]]
            + [" ".join(section["body"]) for section in tutorials["term_search"]["sections"]]
        )
        structured = " ".join(
            [tutorials["structured_search"]["summary"]]
            + [" ".join(section["body"]) for section in tutorials["structured_search"]["sections"]]
        )
        for text in (terms, structured):
            self.assertIn("Literal", text)
            self.assertIn("Regex", text)
            self.assertIn("Lexical", text)
            self.assertIn("Híbrido", text)
        self.assertIn("raça”, “raças”, “racial”, “raciais”, “racismo”, “racializado” e “racializados”", terms)
        self.assertIn("Grupo organiza", structured)
        self.assertIn("Variante", structured)


if __name__ == "__main__":
    unittest.main()
