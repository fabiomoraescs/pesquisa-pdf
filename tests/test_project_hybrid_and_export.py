"""Regressões do modo de projeto e da exportação, sem baixar o modelo semântico."""

import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

import pymupdf
from openpyxl import load_workbook

from historico_racial.entities import listar_entidades
from historico_racial.exporter import gerar_xlsx
from historico_racial.processor import ArquivoPDF, processar_documentos
from historico_racial.semantic import combinar_semantica
from historico_racial.routes import JOBS_LOCK, PROGRESSOS_HR, RESULTADOS_HR
from platform_helpers import create_project, create_user, csrf_from, isolated_platform, login
from platform_core.extensions import db
from platform_core.models import Analysis


def pdf_curto(texto: str) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 80), texto)
    data = doc.tobytes()
    doc.close()
    return data


class ProjectMethodTests(unittest.TestCase):
    def test_new_structured_jobs_enable_morphology_without_changing_versions(self):
        with isolated_platform() as app:
            create_user()
            with app.test_client() as client:
                login(client)
                project_id = create_project(client)
                path = f"/analise-documental/projetos/{project_id}"
                for method in ("lexical", "hibrido"):
                    with self.subTest(method=method):
                        token = csrf_from(client.get(path))
                        with patch("historico_racial.routes.EXECUTOR_HR.submit") as submit:
                            response = client.post(f"{path}/analisar", data={
                                "csrf_token": token, "metodo_analise": method,
                                "pdfs": (io.BytesIO(pdf_curto("As trabalhadoras foram citadas.")), "teste.pdf"),
                            }, content_type="multipart/form-data")
                        self.assertEqual(response.status_code, 202)
                        analysis = db.session.get(Analysis, response.json["job_id"])
                        self.assertEqual(analysis.tool_version, method)
                        self.assertTrue(analysis.parameters_json["morfologia_automatica"])
                        self.assertTrue(submit.call_args.kwargs["incluir_morfologia"])
                        submit.call_args.args[3].cleanup()
                        with JOBS_LOCK:
                            PROGRESSOS_HR.pop(response.json["job_id"], None)

    def test_hibrido_sem_limiar_explicito_usa_default_compartilhado(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "controle.pdf"
            path.write_bytes(pdf_curto("Du Bois foi citado em debate publico."))
            with patch("historico_racial.semantic.combinar_semantica", return_value=[]) as semantic:
                result = processar_documentos(
                    [ArquivoPDF(path, path.name, "arquivo-1")], metodo_analise="hibrido"
                )
        self.assertEqual(result["limiar_semantico"], 0.50)
        self.assertEqual(semantic.call_args.args[-2], 0.50)

    def test_lexical_padrao_nao_importa_motor_semantico(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "controle.pdf"
            path.write_bytes(pdf_curto("Du Bois foi citado em debate publico."))
            with patch("historico_racial.semantic.combinar_semantica") as semantic:
                result = processar_documentos([ArquivoPDF(path, path.name, "arquivo-1")])
            semantic.assert_not_called()
        self.assertEqual(result["metodo_analise"], "lexical")
        self.assertIsNone(result["limiar_semantico"])
        self.assertTrue(all(item["metodo_localizacao"] == "lexical" for item in result["ocorrencias"]))

    def test_hibrido_preserva_lexical_e_chama_adaptador(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "controle.pdf"
            path.write_bytes(pdf_curto("Du Bois foi citado em debate publico."))
            file = ArquivoPDF(path, path.name, "arquivo-1")
            baseline = processar_documentos([file])
            events = []
            with patch("historico_racial.semantic.combinar_semantica", return_value=[]) as semantic:
                result = processar_documentos([file], events.append, metodo_analise="hibrido", limiar_semantico=0.70)
            semantic.assert_called_once()
        projection = lambda rows: [(item["id_entidade"], item["pagina_pdf"], item["forma_original_no_texto"], item["trecho_ocorrencia"]) for item in rows]
        self.assertEqual(projection(baseline["ocorrencias"]), projection(result["ocorrencias"]))
        self.assertEqual(result["metodo_analise"], "hibrido")
        self.assertEqual(result["limiar_semantico"], 0.70)
        values = [event["percentual"] for event in events if event.get("percentual") is not None]
        self.assertEqual(values, sorted(values))

    def test_progresso_hibrido_por_etapas_reais_e_multiplos_pdfs(self):
        with tempfile.TemporaryDirectory() as root:
            files = []
            for number in (1, 2):
                path = Path(root) / f"controle-{number}.pdf"
                path.write_bytes(pdf_curto("Du Bois foi citado em debate publico."))
                files.append(ArquivoPDF(path, path.name, f"arquivo-{number}"))
            events = []

            def engine(*args):
                callback = args[-1]
                callback({"fase": "preparando_blocos_semanticos", "etapa": "Preparando blocos…", "bloco_atual": 1, "blocos_total": 1})
                callback({"fase": "preparando_modelo_semantico", "etapa": "Preparando modelo semântico…"})
                callback({"fase": "codificando_blocos_semanticos", "etapa": "Analisando correspondências semânticas…", "bloco_atual": 1, "blocos_total": 1})
                callback({"fase": "comparando_semantica", "etapa": "Analisando correspondências semânticas…", "consulta_atual": 1, "consultas_total": 1})
                return []

            with patch("historico_racial.semantic.combinar_semantica", side_effect=engine):
                result = processar_documentos(files, events.append, metodo_analise="hibrido", limiar_semantico=0.7)
        percentages = [event["percentual"] for event in events if event.get("percentual") is not None]
        self.assertEqual(percentages, sorted(percentages))
        self.assertLessEqual(max(percentages), 99)
        self.assertEqual(result["total_pdfs"], 2)
        self.assertEqual({event["arquivo_indice"] for event in events}, {1, 2})

    def test_motor_v3_adaptado_preserva_score_e_semantica_nao_e_lexical(self):
        entity = next(item for item in listar_entidades() if item.id_entidade == "du_bois")
        pages = [{"pagina_pdf": 1, "blocos": ["Du Bois foi citado em debate publico."]}]
        record = {
            "Consulta": entity.forma_canonica, "Trecho da ocorrência": "Du Bois foi citado em debate publico.",
            "Bloco anterior": "", "Bloco posterior": "", "Similaridade semântica": 0.82,
            "_pagina_inicial_pdf": 1, "_pagina_final_pdf": 1,
        }
        with patch("analyzer.v3._registros_semanticos", return_value=[record]) as engine:
            semantic = combinar_semantica(Path("livro.pdf"), pages, (entity,), [], "doc-1", "livro.pdf", "project-1", {}, 0.70)
        self.assertEqual(engine.call_args.args[3], 0.70)
        self.assertEqual(semantic[0]["tipo_correspondencia"], "Semântica")
        self.assertEqual(semantic[0]["similaridade_semantica"], 0.82)
        self.assertEqual(semantic[0]["termo_encontrado"], "")

    def test_planilha_quatro_abas_mesmo_esquema(self):
        for method in ("lexical", "hibrido"):
            result = {
                "metodo_analise": method, "limiar_semantico": 0.7 if method == "hibrido" else None,
                "modelo_semantico": "modelo" if method == "hibrido" else None,
                "vocabulario_version": "v1.0", "vocabulario_hash": "abc", "library_names": ["Relações raciais"],
                "data_processamento": "2026-09-23T00:00:00+00:00",
                "documentos": [{"id_project": "project", "id_documento": "doc", "arquivo_pdf": "a.pdf", "id_arquivo": "file",
                                "pagina_pdf_inicio": 1, "pagina_pdf_fim": 1, "ocr_utilizado": False}],
                "ocorrencias": [{"id_ocorrencia": "occ", "id_documento": "doc", "arquivo_pdf": "a.pdf", "pagina_pdf": 1,
                                 "categoria_busca": "entidades", "tipo_entidade": "autor", "id_entidade": "du_bois",
                                 "entidade_canonica": "W. E. B. Du Bois", "termo_encontrado": "Du Bois",
                                 "forma_original_no_texto": "Du Bois", "grupo": ["C"], "trecho_anterior": "",
                                 "trecho_ocorrencia": "Du Bois foi citado.", "trecho_posterior": "",
                                 "contexto_completo": "Du Bois foi citado."}],
            }
            book = load_workbook(io.BytesIO(gerar_xlsx(result).getvalue()))
            self.assertEqual(book.sheetnames, ["DOCUMENTOS", "OCORRENCIAS", "CODIFICACAO", "COOCORRENCIAS"])
            self.assertEqual(book["DOCUMENTOS"]["O2"].value, "Híbrido" if method == "hibrido" else "Lexical")
            self.assertEqual(book["DOCUMENTOS"]["P2"].value, 0.7 if method == "hibrido" else None)
            self.assertEqual(book["OCORRENCIAS"]["A2"].value, "occ")
            self.assertEqual(book["CODIFICACAO"]["A2"].value, "occ")
            self.assertEqual(book["COOCORRENCIAS"].max_row, 1)

    def test_download_de_resultado_pertence_ao_projeto(self):
        with isolated_platform() as app:
            user = create_user()
            with app.test_client() as client:
                login(client)
                project_id = create_project(client)
                job_id = str(uuid4())
                data = {
                    "project_id": project_id, "owner_user_id": user.id,
                    "documentos": [], "ocorrencias": [], "metodo_analise": "lexical",
                }
                with JOBS_LOCK:
                    RESULTADOS_HR[job_id] = data
                try:
                    url = f"/analise-documental/projetos/{project_id}/resultado/{job_id}/xlsx"
                    response = client.get(url)
                    self.assertEqual(response.status_code, 200)
                    self.assertTrue(response.data.startswith(b"PK"))
                    self.assertEqual(load_workbook(io.BytesIO(response.data)).sheetnames,
                                     ["DOCUMENTOS", "OCORRENCIAS", "CODIFICACAO", "COOCORRENCIAS"])
                finally:
                    with JOBS_LOCK:
                        RESULTADOS_HR.pop(job_id, None)


if __name__ == "__main__":
    unittest.main()
