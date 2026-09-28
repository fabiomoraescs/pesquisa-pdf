"""Renomeação por identidade estável, inclusive após autocodificação."""

import io
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from sqlalchemy import select

from platform_helpers import create_user, csrf_from, login
import test_qualitative_automatic as automatic_fixture
import test_qualitative_context as context_fixture
from test_qualitative_expanded_search import FakeModel
from platform_core.extensions import db
from platform_core.models import (QualitativeCode, QualitativeCoding, QualitativeExcerpt,
                                  QualitativeMemo, QualitativeRejection, UserToolOverride)
from platform_core.qualitative_annotations import page_excerpts
from platform_core.scraping_types import QUALITATIVE_TOOL


class QualitativeCodeRenameTests(unittest.TestCase):
    setUp = context_fixture.QualitativeContextTests.setUp
    tearDown = context_fixture.QualitativeContextTests.tearDown
    post = context_fixture.QualitativeContextTests.post
    selection = context_fixture.QualitativeContextTests.selection
    automatic = automatic_fixture.QualitativeAutomaticTests.automatic
    remove = automatic_fixture.QualitativeAutomaticTests.remove
    codings = automatic_fixture.QualitativeAutomaticTests.codings
    fixture_pages = automatic_fixture.QualitativeAutomaticTests.fixture_pages
    auto_url = automatic_fixture.QualitativeAutomaticTests.auto_url

    def code_url(self, identifier):
        return f"/analise-qualitativa/bases/{self.analysis.id}/codigos/{identifier}"

    def rename(self, identifier, name, *, token=True):
        return self.client.patch(self.code_url(identifier), json={"name": name},
                                 headers={"X-CSRFToken": self.csrf} if token else {})

    def test_manual_code_rename_propagates_to_three_excerpts_report_and_xlsx(self):
        created = self.client.post(f"/analise-qualitativa/bases/{self.analysis.id}/codigos",
            json={"name": "Racismo", "description": "Memo conceitual"},
            headers={"X-CSRFToken": self.csrf})
        identifier = created.json["codes"][0]["id"]
        other = self.client.post(f"/analise-qualitativa/bases/{self.analysis.id}/codigos",
            json={"name": "Outro"}, headers={"X-CSRFToken": self.csrf}).json["codes"][1]["id"]
        text = self.page["text"]
        for quote, start in (("Documento", text.index("Documento")),
                             ("Documento", text.index("Documento", 1)),
                             ("pesquisa", text.index("pesquisa"))):
            selection = self.selection(quote, start=start)
            self.assertEqual(self.post("apply_codes", {"code_ids": [identifier]}, selection=selection)
                             .status_code, 201)
        coding_ids = [coding.id for coding in self.codings()]
        excerpt_ids = [coding.excerpt_id for coding in self.codings()]
        db.session.add(QualitativeMemo(analysis_id=self.analysis.id, code_id=identifier,
                                       text="Memo preservado", created_by_user_id=self.user.id))
        db.session.commit()
        for name in ("Racismo estrutural", "Classificação racial", "Relações étnico-raciais"):
            response = self.rename(identifier, f"  {name}  ")
            self.assertEqual(response.status_code, 200, response.json)
            record = next(item for item in response.json["codes"] if item["id"] == identifier)
            self.assertEqual((record["name"], record["excerpt_count"]), (name, 3))
            self.assertEqual(db.session.get(QualitativeCode, identifier).description, "Memo conceitual")
            self.assertEqual(db.session.get(QualitativeCode, other).name, "Outro")
            self.assertEqual([item.id for item in self.codings()], coding_ids)
            self.assertEqual([item.excerpt_id for item in self.codings()], excerpt_ids)
            excerpts = page_excerpts(self.analysis, self.document.id, 1)["excerpts"]
            self.assertEqual([item["codes"][0]["name"] for item in excerpts], [name] * 3)
            self.assertEqual({item["codes"][0]["id"] for item in excerpts}, {identifier})
            report = self.client.get(f"/analise-qualitativa/bases/{self.analysis.id}/relatorio-codificacao")
            self.assertIn(name, report.get_data(as_text=True))
            workbook = load_workbook(io.BytesIO(self.client.get(
                f"/analise-qualitativa/bases/{self.analysis.id}/relatorio-codificacao.xlsx").data))
            self.assertEqual([workbook.active.cell(row, 1).value for row in range(2, 5)], [name] * 3)
        self.assertEqual(db.session.query(QualitativeExcerpt).count(), 3)
        self.assertEqual(db.session.scalar(select(QualitativeMemo.code_id)), identifier)

    def test_name_conflict_invalid_name_csrf_and_deleted_code(self):
        codes_url = f"/analise-qualitativa/bases/{self.analysis.id}/codigos"
        first = self.client.post(codes_url, json={"name": "Raça"},
            headers={"X-CSRFToken": self.csrf}).json["codes"][0]["id"]
        second = next(item["id"] for item in self.client.post(codes_url,
            json={"name": "Racismo"}, headers={"X-CSRFToken": self.csrf}).json["codes"]
            if item["name"] == "Racismo")
        self.assertEqual(self.rename(first, "  rAcIsMo  ").status_code, 409)
        self.assertEqual(self.rename(first, "   ").status_code, 400)
        self.assertEqual(self.rename(first, "Novo", token=False).status_code, 400)
        self.assertEqual(db.session.get(QualitativeCode, first).name, "Raça")
        self.assertEqual(db.session.get(QualitativeCode, second).name, "Racismo")
        self.assertEqual(self.client.delete(self.code_url(first), headers={"X-CSRFToken": self.csrf})
                         .status_code, 200)
        self.assertEqual(self.rename(first, "Tardio").status_code, 404)

    def test_other_user_cannot_rename_code_by_direct_id(self):
        created = self.client.post(f"/analise-qualitativa/bases/{self.analysis.id}/codigos",
            json={"name": "Privado"}, headers={"X-CSRFToken": self.csrf})
        identifier = created.json["codes"][0]["id"]
        other = create_user(name="Outro", email="outro@example.org")
        db.session.add(UserToolOverride(user_id=other.id, tool_id=QUALITATIVE_TOOL, decision="allow"))
        db.session.commit()
        self.client.post("/logout", data={"csrf_token": self.csrf})
        login(self.client, email="outro@example.org")
        token = csrf_from(self.client.get("/"))
        response = self.client.patch(self.code_url(identifier), json={"name": "Intruso"},
                                     headers={"X-CSRFToken": token})
        self.assertIn(response.status_code, (403, 404))
        self.assertEqual(db.session.get(QualitativeCode, identifier).name, "Privado")

    def test_literal_rename_repeat_uses_prior_code_even_if_query_name_is_reused(self):
        self.fixture_pages([["raça e sociedade. raça e escola. raça e trabalho."]])
        first = self.automatic(q="raça", mode="literal")
        identifier = first.json["terms"][0]["code_id"]
        self.assertEqual(first.json["summary"]["created"], 3)
        original_sources = [coding.source_query for coding in self.codings()]
        self.assertEqual(self.rename(identifier, "Relações raciais").status_code, 200)
        other = self.client.post(f"/analise-qualitativa/bases/{self.analysis.id}/codigos",
            json={"name": "raça"}, headers={"X-CSRFToken": self.csrf})
        self.assertEqual(other.status_code, 201)
        repeated = self.automatic(q="raça", mode="literal")
        self.assertEqual((repeated.json["summary"]["created"],
                          repeated.json["summary"]["existing"]), (0, 3))
        self.assertEqual(repeated.json["terms"][0]["code_id"], identifier)
        self.assertEqual({coding.code_id for coding in self.codings()}, {identifier})
        self.assertEqual([coding.source_query for coding in self.codings()], original_sources)
        self.assertEqual(db.session.query(QualitativeCode).count(), 2)

    def test_lexical_rename_repeat_preserves_identity(self):
        self.fixture_pages([["professores professoras professor"]])
        first = self.automatic(q="professor", mode="lexical")
        identifier = first.json["terms"][0]["code_id"]
        self.assertEqual(self.rename(identifier, "Docência").status_code, 200)
        repeated = self.automatic(q="professor", mode="lexical")
        self.assertEqual((repeated.json["summary"]["created"],
                          repeated.json["summary"]["existing"]), (0, 3))
        self.assertEqual(repeated.json["terms"][0]["code_id"], identifier)
        self.assertEqual(db.session.query(QualitativeCode).count(), 1)

    def test_semantic_rename_repeat_preserves_identity(self):
        self.fixture_pages([["Barreiras institucionais persistem."]])
        model = FakeModel(lambda _: [1, 0])
        with patch("platform_core.qualitative_expanded_search._model", return_value=model):
            first = self.automatic(q="iniquidade", mode="semantic")
            identifier = first.json["terms"][0]["code_id"]
            self.assertEqual(self.rename(identifier, "Desigualdade institucional").status_code, 200)
            repeated = self.automatic(q="iniquidade", mode="semantic")
        self.assertEqual((repeated.json["summary"]["created"],
                          repeated.json["summary"]["existing"]), (0, 1))
        self.assertEqual(repeated.json["terms"][0]["code_id"], identifier)
        self.assertEqual(self.codings()[0].source_query, "iniquidade")
        self.assertEqual(db.session.query(QualitativeCode).count(), 1)

    def test_regex_rename_keeps_distinct_match_codes_and_query(self):
        self.fixture_pages([["raça racismo"]])
        first = self.automatic(q=r"ra\w+", mode="literal", grep=True)
        identifier = next(item.id for item in db.session.scalars(select(QualitativeCode))
                          if item.name == "raça")
        self.assertEqual(self.rename(identifier, "Questão racial").status_code, 200)
        repeated = self.automatic(q=r"ra\w+", mode="literal", grep=True)
        self.assertEqual((repeated.json["summary"]["created"],
                          repeated.json["summary"]["existing"]), (0, 2))
        self.assertEqual({coding.code_id for coding in self.codings()},
                         {code.id for code in db.session.scalars(select(QualitativeCode))})
        self.assertEqual({coding.source_query for coding in self.codings()}, {r"ra\w+"})
        self.assertEqual(first.json["search"]["total"], 2)

    def test_multiple_terms_rename_keeps_each_term_on_its_original_code(self):
        self.fixture_pages([["raça racismo"]])
        first = self.automatic(q="raça;racismo", mode="literal", multiple_terms=True)
        codes = {item["term"]: item["code_id"] for item in first.json["terms"]}
        self.assertNotEqual(codes["raça"], codes["racismo"])
        self.assertEqual(self.rename(codes["raça"], "Relações raciais").status_code, 200)
        repeated = self.automatic(q="raça;racismo", mode="literal", multiple_terms=True)
        self.assertEqual((repeated.json["summary"]["created"],
                          repeated.json["summary"]["existing"]), (0, 2))
        self.assertEqual({item["term"]: item["code_id"] for item in repeated.json["terms"]}, codes)
        self.assertEqual(db.session.query(QualitativeCode).count(), 2)

    def test_contextual_rejection_remains_bound_to_renamed_code(self):
        self.fixture_pages([["raça"]])
        first = self.automatic(q="raça", mode="literal", contextual_rejection_enabled=True)
        identifier = first.json["terms"][0]["code_id"]
        self.assertEqual(self.remove(self.codings()[0]).status_code, 200)
        rejection = db.session.scalar(select(QualitativeRejection))
        self.assertEqual((rejection.code_id, rejection.query), (identifier, "raça"))
        self.assertEqual(self.rename(identifier, "Relações raciais").status_code, 200)
        repeated = self.automatic(q="raça", mode="literal")
        self.assertEqual((repeated.json["summary"]["created"],
                          repeated.json["summary"]["rejected"]), (0, 1))
        self.assertEqual(repeated.json["terms"][0]["code_id"], identifier)
        self.assertEqual(db.session.query(QualitativeCode).count(), 1)
        self.assertEqual(db.session.query(QualitativeCoding).count(), 0)


if __name__ == "__main__":
    unittest.main()
