"""Fase 7: comentários, recepção e seu histórico verificável."""

from __future__ import annotations

import csv
import io
import re
import tempfile
import unittest
from datetime import datetime, timezone
from importlib import import_module
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import sqlalchemy as sa
from flask_migrate import downgrade, upgrade
from alembic.migration import MigrationContext
from alembic.operations import Operations
from sqlalchemy import delete, func, select, text
from sqlalchemy import inspect
from sqlalchemy.exc import IntegrityError

from platform_helpers import create_user, csrf_from, isolated_platform, login
from platform_core.extensions import db
from platform_core.instagram_comment_import import (
    COMMENT_CANONICAL_HEADERS,
    COMMENT_IMPORT_FIELDS,
    InstagramImportError,
    parse_instagram_comment_import,
)
from platform_core.instagram_comment_storage import persist_instagram_comment_import
from platform_core.instagram_comments import build_instagram_comments_dashboard, classify_comment, effective_assessment
from platform_core.instagram_lifecycle import delete_instagram_project_permanently
from platform_core.models import (
    AnalyticsProject,
    AnalyticsRun,
    InstagramAccount,
    InstagramComment,
    InstagramCommentAssessment,
    InstagramCommentObservation,
    InstagramCommentTextVersion,
    InstagramMedia,
    InstagramMediaObservation,
    PlanTool,
    User,
)
from platform_core.services import ANALYTICS_INSTAGRAM_TOOL


BASE = {
    "media_external_id": "post-1",
    "comment_external_id": "comment-1",
    "comment_text": "Gostei muito.",
    "published_at": "2026-10-02T14:30:00-03:00",
    "observed_at": "2026-10-03T09:00:00-03:00",
    "parent_comment_external_id": "",
    "author_username": "leitora",
    "author_external_id": "user-1",
    "like_count": "3",
    "replies_count": "0",
}


def csv_bytes(*rows: dict[str, str], headers: tuple[str, ...] = COMMENT_CANONICAL_HEADERS) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=headers, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode("utf-8")


def xlsx_bytes(*rows: dict[str, str]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(COMMENT_CANONICAL_HEADERS)
    for row in rows:
        sheet.append([row.get(header, "") for header in COMMENT_CANONICAL_HEADERS])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class InstagramCommentsTests(unittest.TestCase):
    def setUp(self):
        self.scope = isolated_platform()
        self.app = self.scope.__enter__()
        self.client = self.app.test_client()
        self.user = create_user()
        self.other = create_user("Outra", "outra@example.org")
        db.session.add(PlanTool(plan_id="research_quali", tool_id=ANALYTICS_INSTAGRAM_TOOL))
        self.project = AnalyticsProject(owner_user_id=self.user.id, name="Recepção Instagram")
        self.account = InstagramAccount(
            analytics_project=self.project, username="revista", username_normalized="revista", display_name="Revista",
        )
        self.media = InstagramMedia(
            instagram_account=self.account, external_id="post-1", caption_original="Publicação um",
            published_at=datetime(2026, 10, 1, 12, tzinfo=timezone.utc),
        )
        self.other_project = AnalyticsProject(owner_user_id=self.other.id, name="Outro projeto")
        db.session.add_all((self.project, self.account, self.media, self.other_project))
        db.session.commit()
        login(self.client)

    def tearDown(self):
        self.scope.__exit__(None, None, None)

    def preview(self, data: bytes, filename="comentarios.csv"):
        page = self.client.get(f"/analytics/instagram/projects/{self.project.id}/comments/imports/new")
        return self.client.post(
            f"/analytics/instagram/projects/{self.project.id}/comments/imports/preview",
            data={"csrf_token": csrf_from(page), "dataset": (io.BytesIO(data), filename)},
        )

    def confirm(self, response):
        token = re.search(r'name="preview_token" value="([^"]+)"', response.get_data(as_text=True)).group(1)
        return self.client.post(
            f"/analytics/instagram/projects/{self.project.id}/comments/imports/confirm",
            data={"csrf_token": csrf_from(response), "preview_token": token},
        )

    def import_rows(self, *rows: dict[str, str]):
        response = self.preview(csv_bytes(*rows))
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        completed = self.confirm(response)
        self.assertEqual(completed.status_code, 200, completed.get_data(as_text=True))
        return completed

    def test_csv_import_preserves_original_versions_observations_and_replies(self):
        reply = dict(BASE, comment_external_id="comment-2", comment_text="Também gostei.", parent_comment_external_id="comment-1", like_count="")
        self.import_rows(BASE, reply)
        edited = dict(BASE, comment_text="Gostei muito do artigo.", observed_at="2026-10-04T09:00:00-03:00", like_count="8")
        self.import_rows(edited)
        comments = db.session.scalars(select(InstagramComment).order_by(InstagramComment.external_id)).all()
        self.assertEqual(len(comments), 2)
        original = comments[0]
        child = comments[1]
        self.assertEqual(original.original_text, "Gostei muito.")
        self.assertEqual(child.parent_comment_id, original.id)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramCommentTextVersion).where(
            InstagramCommentTextVersion.instagram_comment_id == original.id
        )), 2)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramCommentObservation).where(
            InstagramCommentObservation.instagram_comment_id == original.id
        )), 2)
        self.assertEqual(sorted(db.session.scalars(select(InstagramCommentObservation.like_count).where(
            InstagramCommentObservation.instagram_comment_id == original.id
        )).all()), [3, 8])
        latest = max(original.text_versions, key=lambda value: (value.first_observed_at, value.id))
        self.assertEqual(latest.text, "Gostei muito do artigo.")
        with self.assertRaises(ValueError):
            original.original_text = "Não permitido"
            db.session.commit()
        db.session.rollback()

    def test_comment_without_available_text_has_a_safe_presentation_and_cannot_be_classified(self):
        comment = InstagramComment(
            instagram_media=self.media,
            external_id="comment-without-text",
            published_at=datetime(2026, 10, 2, tzinfo=timezone.utc),
            original_text=None,
        )
        db.session.add(comment)
        db.session.commit()

        listing = self.client.get(f"/analytics/instagram/projects/{self.project.id}/comments?period=all")
        detail = self.client.get(
            f"/analytics/instagram/projects/{self.project.id}/comments/{comment.id}"
        )
        self.assertEqual(listing.status_code, 200)
        self.assertEqual(detail.status_code, 200)
        self.assertIn("Texto não disponibilizado pela API.", listing.get_data(as_text=True))
        self.assertIn("A API ainda não disponibilizou o texto deste comentário.", detail.get_data(as_text=True))
        self.assertIn("Texto ainda indisponível", detail.get_data(as_text=True))
        with self.assertRaisesRegex(ValueError, "não possui uma versão textual classificável"):
            classify_comment(comment, self.user.id, {"sentiment": "positive"})

    def test_original_text_allows_only_a_proven_first_fill_after_reload_or_expire(self):
        comment = InstagramComment(
            instagram_media=self.media,
            external_id="comment-original-text-immutable",
            published_at=datetime(2026, 10, 2, tzinfo=timezone.utc),
            original_text=None,
        )
        db.session.add(comment)
        db.session.commit()

        # The sole permitted transition is the first real text after the API
        # previously reported a valid comment identity without its body.
        comment.original_text = "Primeiro texto disponibilizado."
        db.session.commit()
        self.assertEqual(comment.original_text, "Primeiro texto disponibilizado.")

        # Repeating the same value is idempotent and does not manufacture a
        # second mutation in SQLAlchemy history.
        comment.original_text = "Primeiro texto disponibilizado."
        db.session.commit()
        self.assertEqual(comment.original_text, "Primeiro texto disponibilizado.")

        for forbidden_value in ("Texto substituído.", None):
            with self.subTest(forbidden_value=forbidden_value):
                comment.original_text = forbidden_value
                with self.assertRaisesRegex(ValueError, "não pode ser alterado"):
                    db.session.commit()
                db.session.rollback()
                db.session.refresh(comment)
                self.assertEqual(comment.original_text, "Primeiro texto disponibilizado.")

        # active_history=True must load the old value before assignment even
        # when this attribute has been explicitly expired.
        db.session.expire(comment, ["original_text"])
        comment.original_text = "Texto sobrescrito após expire."
        with self.assertRaisesRegex(ValueError, "não pode ser alterado"):
            db.session.commit()
        db.session.rollback()

        # The same protection holds when the row is loaded in a brand-new ORM
        # session, where no prior in-memory history exists.
        comment_id = comment.id
        db.session.remove()
        reloaded = db.session.get(InstagramComment, comment_id)
        self.assertEqual(reloaded.original_text, "Primeiro texto disponibilizado.")
        reloaded.original_text = "Texto sobrescrito em nova sessão."
        with self.assertRaisesRegex(ValueError, "não pode ser alterado"):
            db.session.commit()
        db.session.rollback()
        self.assertEqual(
            db.session.get(InstagramComment, comment_id).original_text,
            "Primeiro texto disponibilizado.",
        )

    def test_xlsx_schema_template_and_parser_are_canonical(self):
        page = self.client.get(f"/analytics/instagram/projects/{self.project.id}/comments/imports/new")
        self.assertIn("Importar comentários", page.get_data(as_text=True))
        response = self.client.get(f"/analytics/instagram/projects/{self.project.id}/comments/imports/modelo")
        self.assertEqual(response.status_code, 200)
        self.assertIn("modelo_importacao_comentarios_instagram_analysis.xlsx", response.headers["Content-Disposition"])
        workbook = load_workbook(io.BytesIO(response.data))
        try:
            sheet = workbook.active
            self.assertEqual(tuple(cell.value for cell in sheet[1]), COMMENT_CANONICAL_HEADERS)
            self.assertEqual(COMMENT_CANONICAL_HEADERS, tuple(field.name for field in COMMENT_IMPORT_FIELDS))
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.auto_filter.ref, "A1:J1")
            self.assertTrue(all(sheet.cell(1, column).comment for column in range(1, 11)))
            self.assertTrue(all(sheet.cell(2, column).value is None for column in range(1, 11)))
            self.assertFalse(any(cell.data_type == "f" for row in sheet.iter_rows(max_row=2) for cell in row))
        finally:
            workbook.close()
        response = self.preview(xlsx_bytes(BASE), "comentarios.xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.confirm(response).status_code, 200)

    def test_parser_rejects_missing_required_invalid_counts_and_duplicate_conflicts(self):
        parsed = parse_instagram_comment_import("comments.csv", csv_bytes(BASE))
        self.assertEqual(parsed.rows[0].published_at, datetime(2026, 10, 2, 17, 30, tzinfo=timezone.utc))
        self.assertEqual(parsed.rows[0].observed_at, datetime(2026, 10, 3, 12, 0, tzinfo=timezone.utc))
        with self.assertRaisesRegex(InstagramImportError, "Coluna obrigatória ausente: comment_text"):
            parse_instagram_comment_import("comments.csv", csv_bytes(BASE, headers=tuple(
                header for header in COMMENT_CANONICAL_HEADERS if header != "comment_text"
            )))
        with self.assertRaisesRegex(InstagramImportError, "não pode ser negativo"):
            parse_instagram_comment_import("comments.csv", csv_bytes(dict(BASE, like_count="-1")))
        with self.assertRaisesRegex(InstagramImportError, "número inteiro"):
            parse_instagram_comment_import("comments.csv", csv_bytes(dict(BASE, replies_count="1.5")))
        with self.assertRaisesRegex(InstagramImportError, "resposta de si mesmo"):
            parse_instagram_comment_import("comments.csv", csv_bytes(dict(BASE, parent_comment_external_id="comment-1")))
        with self.assertRaisesRegex(InstagramImportError, "conflita"):
            parse_instagram_comment_import("comments.csv", csv_bytes(BASE, dict(BASE, comment_text="Outro texto")))
        dataset = parse_instagram_comment_import("comments.csv", csv_bytes(BASE, BASE))
        self.assertEqual(len(dataset.rows), 1)
        self.assertTrue(dataset.warnings)

    def test_unknown_media_and_cross_media_parent_are_rejected_before_persistence(self):
        unknown = self.preview(csv_bytes(dict(BASE, media_external_id="missing")))
        self.assertEqual(unknown.status_code, 400)
        self.assertIn("não pertence a este projeto", unknown.get_data(as_text=True))
        other_media = InstagramMedia(
            instagram_account=self.account, external_id="post-2", caption_original="Outra",
            published_at=datetime(2026, 10, 1, tzinfo=timezone.utc),
        )
        db.session.add(other_media)
        db.session.commit()
        bad_child = dict(BASE, comment_external_id="comment-2", media_external_id="post-2", parent_comment_external_id="comment-1")
        response = self.preview(csv_bytes(BASE, bad_child))
        self.assertEqual(response.status_code, 400)
        self.assertIn("outra publicação", response.get_data(as_text=True))
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 0)

    def test_confirmation_is_atomic_when_commit_fails(self):
        dataset = parse_instagram_comment_import("comments.csv", csv_bytes(BASE))
        with patch("platform_core.instagram_comment_storage.db.session.commit", side_effect=RuntimeError("controlled")):
            with self.assertRaises(RuntimeError):
                persist_instagram_comment_import(self.project, dataset)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 0)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramComment)), 0)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramCommentObservation)), 0)

    def test_comment_constraints_keep_identity_counts_and_confidence_valid(self):
        self.import_rows(BASE)
        comment = db.session.scalar(select(InstagramComment))
        duplicate = InstagramComment(
            instagram_media=self.media, external_id=comment.external_id,
            published_at=datetime(2026, 10, 2, tzinfo=timezone.utc), original_text="Duplicado",
        )
        db.session.add(duplicate)
        with self.assertRaises(IntegrityError):
            db.session.commit()
        db.session.rollback()
        run = db.session.scalar(select(AnalyticsRun))
        db.session.add(InstagramCommentObservation(
            instagram_comment=comment, analytics_run=run,
            observed_at=datetime(2026, 10, 4, tzinfo=timezone.utc), like_count=-1,
        ))
        with self.assertRaises(IntegrityError):
            db.session.commit()
        db.session.rollback()
        db.session.add(InstagramCommentAssessment(
            instagram_comment=comment, text_version=comment.text_versions[0], source="manual",
            reviewer_user_id=self.user.id, confidence=1.1,
        ))
        with self.assertRaises(IntegrityError):
            db.session.commit()
        db.session.rollback()

    def test_manual_classification_review_and_effective_state_are_historical(self):
        self.import_rows(BASE)
        comment = db.session.scalar(select(InstagramComment))
        classify_comment(comment, self.user.id, {
            "sentiment": "negative", "intensity": "low", "intent": "criticism", "topic": "tom",
            "topic_normalized": "tom", "needs_response": "yes", "risk": "medium", "rationale": "Primeira leitura.",
        })
        db.session.commit()
        first = effective_assessment(comment.assessments)
        self.assertEqual(first.source, "manual")
        classify_comment(comment, self.user.id, {
            "sentiment": "mixed", "intensity": "high", "intent": "question", "topic": "conteúdo",
            "topic_normalized": "conteudo", "needs_response": "no", "risk": "low", "rationale": "Revisado.",
        })
        db.session.commit()
        assessments = db.session.scalars(select(InstagramCommentAssessment).where(
            InstagramCommentAssessment.instagram_comment_id == comment.id
        )).all()
        effective = effective_assessment(assessments)
        self.assertEqual(len(assessments), 2)
        self.assertIsNotNone(first.superseded_at)
        self.assertEqual(effective.source, "human_review")
        self.assertEqual(effective.sentiment, "mixed")
        self.assertEqual(effective.reviewer_user_id, self.user.id)
        self.assertEqual(effective.instagram_comment_text_version_id, comment.text_versions[0].id)

    def test_dashboard_uses_imported_comments_not_media_comments_count_and_reports_coverage(self):
        media_run = AnalyticsRun(
            analytics_project=self.project, source_kind="csv", status="completed", record_count=1,
            started_at=datetime(2026, 10, 3, tzinfo=timezone.utc),
            completed_at=datetime(2026, 10, 3, tzinfo=timezone.utc),
        )
        db.session.add(media_run)
        db.session.flush()
        db.session.add(InstagramMediaObservation(
            instagram_media=self.media, analytics_run=media_run,
            observed_at=datetime(2026, 10, 3, tzinfo=timezone.utc), comments_count=80,
        ))
        for index in range(100):
            comment = InstagramComment(
                instagram_media=self.media, external_id=f"c-{index}", published_at=datetime(2026, 10, 2, tzinfo=timezone.utc),
                original_text=f"Comentário {index}",
            )
            version = InstagramCommentTextVersion(
                instagram_comment=comment, text=f"Comentário {index}", content_hash=f"{index:064x}",
                first_observed_at=datetime(2026, 10, 3, tzinfo=timezone.utc), is_original=True,
            )
            db.session.add_all((comment, version))
            if index < 40:
                db.session.add(InstagramCommentAssessment(
                    instagram_comment=comment, text_version=version, source="manual", reviewer_user_id=self.user.id,
                    sentiment="positive", needs_response=False, risk="none",
                ))
        db.session.commit()
        dashboard = build_instagram_comments_dashboard(self.project, period="all", sort="recent")
        self.assertEqual(dashboard.metrics.imported, 100)
        self.assertEqual(dashboard.metrics.classified, 40)
        self.assertEqual(dashboard.metrics.coverage_percentage, 40.0)
        self.assertEqual(dashboard.chart_data["sentiment"]["values"], [40, 60])
        html = self.client.get(f"/analytics/instagram/projects/{self.project.id}/comments?period=all").get_data(as_text=True)
        self.assertIn("40 / 100", html)
        self.assertIn("40% classificados", html)
        self.assertIn("Comentários coletados", html)
        self.assertNotIn("Comentários importados", html)
        self.assertNotIn("80 / 100", html)

    def test_comments_presentation_is_neutral_about_collection_source_and_keeps_replies(self):
        reply = dict(
            BASE,
            comment_external_id="comment-reply",
            comment_text="Resposta vinculada.",
            parent_comment_external_id="comment-1",
        )
        self.import_rows(BASE, reply)
        root = db.session.scalar(select(InstagramComment).where(InstagramComment.external_id == "comment-1"))

        listing = self.client.get(
            f"/analytics/instagram/projects/{self.project.id}/comments?period=all"
        ).get_data(as_text=True)
        self.assertIn("Comentários coletados, observações históricas e classificação rastreável.", listing)
        self.assertNotIn("Comentários importados", listing)

        detail = self.client.get(
            f"/analytics/instagram/projects/{self.project.id}/comments/{root.id}"
        ).get_data(as_text=True)
        self.assertIn('<h2 id="instagram-comment-replies" class="h5">Respostas</h2>', detail)
        self.assertNotIn("Respostas importadas", detail)
        self.assertIn("Resposta vinculada.", detail)

    def test_routes_ownership_csrf_archived_read_only_and_partial_filters(self):
        self.import_rows(BASE)
        comment = db.session.scalar(select(InstagramComment))
        page = self.client.get(f"/analytics/instagram/projects/{self.project.id}/comments")
        html = page.get_data(as_text=True)
        self.assertIn("Comentários e recepção", html)
        self.assertIn("data-partial-target=\"#instagram-comments-content\"", html)
        self.assertIn("platform-themed-table", html)
        self.assertIn("Comentários", html)
        self.assertNotIn(
            f'<a class="btn btn-outline-primary btn-sm" href="/analytics/instagram/projects/{self.project.id}/dashboard">Dashboard</a>',
            html,
        )
        self.assertIn(
            'class="platform-filter-control-row platform-filter-control-row--start platform-filter-control-row--comment-filters"',
            html,
        )
        self.assertNotIn('platform-filter-control-group', html)
        # O seletor compartilhado mantém o rótulo do filtro compacto e acessível;
        # os rótulos de data existem somente dentro do bloco personalizado oculto.
        self.assertIn('<label class="visually-hidden" for="instagram-comments-period-select">', html)
        self.assertIn('data-instagram-temporal-period-custom hidden', html)
        expected_filters = (
            ('instagram-comments-period-select', 'Período', 'Período: Últimos 30 dias'),
            ('comment-status', 'Status', 'Status: Todos'),
            ('comment-sentiment', 'Sentimento', 'Sentimento: Todos'),
            ('comment-needs-response', 'Resposta', 'Resposta: Todas'),
            ('comment-risk', 'Risco', 'Risco: Todos'),
            ('comment-sort', 'Ordenar', 'Ordenar: Recentes'),
        )
        for control_id, accessible_name, selected_label in expected_filters:
            self.assertIn(f'<label class="visually-hidden" for="{control_id}">{accessible_name}</label>', html)
            self.assertIn(f'id="{control_id}"', html)
            self.assertIn(f'>{selected_label}</option>', html)
        styles = (Path(__file__).resolve().parents[1] / "static" / "css" / "platform.css").read_text(encoding="utf-8")
        self.assertIn("field-sizing: content", styles)
        self.assertNotIn("instagram-comment-filter-scale", styles)
        self.assertNotIn("--platform-filter-select-width", styles)
        self.assertIn(".platform-filter-control-row--comment-filters { width: 100%; gap: .55rem .8rem; }", styles)
        self.assertIn(".platform-filter-control-row--comment-filters .platform-filter-select { flex: 0 0 100%; width: 100%; max-width: none; field-sizing: fixed; }", styles)
        self.assertIn('id="comment-needs-response" name="needs_response" data-filter-width="long"', html)
        self.assertIn('id="comment-sort" name="sort" data-filter-width="long"', html)
        filtered = self.client.get(
            f"/analytics/instagram/projects/{self.project.id}/comments?period=7d&status=classified&sentiment=negative&needs_response=yes&risk=high&sort=likes"
        ).get_data(as_text=True)
        self.assertIn('<option value="7d" selected>Período: Últimos 7 dias</option>', filtered)
        self.assertIn('<option value="classified" selected>Status: Classificados</option>', filtered)
        self.assertIn('<option value="negative" selected>Sentimento: Negativo</option>', filtered)
        self.assertIn('<option value="yes" selected>Resposta: Sim</option>', filtered)
        self.assertIn('<option value="high" selected>Risco: Alto</option>', filtered)
        self.assertIn('<option value="likes" selected>Ordenar: Curtidos</option>', filtered)
        long_labels = self.client.get(
            f"/analytics/instagram/projects/{self.project.id}/comments?period=all&status=unclassified&sentiment=negative&needs_response=yes&risk=all&sort=risk"
        ).get_data(as_text=True)
        self.assertIn('<option value="all" selected>Período: Todo o histórico disponível</option>', long_labels)
        self.assertIn('<option value="unclassified" selected>Status: Pendentes</option>', long_labels)
        self.assertIn('<option value="yes" selected>Resposta: Sim</option>', long_labels)
        self.assertIn('<option value="risk" selected>Ordenar: Risco</option>', long_labels)
        tab_order = (
            html.index(f'href="/analytics/instagram/projects/{self.project.id}/dashboard">Dashboard</a>'),
            html.index(f'href="/analytics/instagram/projects/{self.project.id}/taxonomies">Taxonomias</a>'),
            html.index(f'href="/analytics/instagram/projects/{self.project.id}/classifications">Classificação de conteúdos</a>'),
            html.index(f'href="/analytics/instagram/projects/{self.project.id}/imports/new">Fonte de dados</a>'),
            html.index(f'href="/analytics/instagram/projects/{self.project.id}/comments" aria-current="page">Comentários</a>'),
            html.index(f'href="/analytics/instagram/projects/{self.project.id}/insights">Insights</a>'),
        )
        self.assertEqual(tab_order, tuple(sorted(tab_order)))
        self.assertEqual(self.client.post(
            f"/analytics/instagram/projects/{self.project.id}/comments/{comment.id}", data={}
        ).status_code, 400)
        detail = self.client.get(f"/analytics/instagram/projects/{self.project.id}/comments/{comment.id}")
        self.assertEqual(detail.status_code, 200)
        self.assertIn("Texto original", detail.get_data(as_text=True))
        self.client.post("/logout", data={"csrf_token": csrf_from(self.client.get("/"))})
        login(self.client, self.other.email)
        self.assertEqual(self.client.get(f"/analytics/instagram/projects/{self.project.id}/comments").status_code, 404)
        self.assertEqual(self.client.get(
            f"/analytics/instagram/projects/{self.project.id}/comments/imports/modelo"
        ).status_code, 404)
        self.client.post("/logout", data={"csrf_token": csrf_from(self.client.get("/"))})
        login(self.client)
        self.project.status = "archived"
        self.project.archived_at = datetime.now(timezone.utc)
        db.session.commit()
        archived = self.client.get(f"/analytics/instagram/projects/{self.project.id}/comments")
        self.assertEqual(archived.status_code, 200)
        self.assertIn("Projeto arquivado", self.client.get(
            f"/analytics/instagram/projects/{self.project.id}/comments/{comment.id}"
        ).get_data(as_text=True))
        self.assertEqual(self.client.get(
            f"/analytics/instagram/projects/{self.project.id}/comments/imports/new"
        ).status_code, 404)
        self.assertEqual(self.client.post(
            f"/analytics/instagram/projects/{self.project.id}/comments/{comment.id}",
            data={"csrf_token": csrf_from(archived), "sentiment": "positive"},
        ).status_code, 404)

    def test_permanent_deletion_removes_comments_and_preview_directory(self):
        preview = self.preview(csv_bytes(BASE))
        self.assertEqual(preview.status_code, 200)
        self.confirm(preview)
        comment = db.session.scalar(select(InstagramComment))
        classify_comment(comment, self.user.id, {"sentiment": "positive", "needs_response": "no"})
        self.project.status = "archived"
        self.project.archived_at = datetime.now(timezone.utc)
        db.session.commit()
        delete_instagram_project_permanently(self.project, self.user, "deletar")
        self.assertIsNone(db.session.get(AnalyticsProject, self.project.id))
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramComment)), 0)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramCommentTextVersion)), 0)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramCommentObservation)), 0)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramCommentAssessment)), 0)


class InstagramCommentsMigrationTests(unittest.TestCase):
    def test_text_unavailable_migration_is_reversible_when_no_unavailable_text_exists(self):
        from app import create_app

        migrations = str(Path(__file__).resolve().parents[1] / "migrations")
        with tempfile.TemporaryDirectory(prefix="pesquisapdf-comment-text-migration-") as root:
            database = Path(root) / "platform.sqlite3"
            app = create_app({
                "TESTING": True,
                "SQLALCHEMY_DATABASE_URI": f"sqlite:///{database.as_posix()}",
                "PLATFORM_DATA_DIR": root,
                "SECRET_KEY": "comments-text-migration-test-only",
            })
            with app.app_context():
                upgrade(directory=migrations, revision="d2e1f0a9b8c7")
                with db.engine.connect() as connection:
                    self.assertEqual(connection.exec_driver_sql("PRAGMA foreign_keys").scalar_one(), 1)
                before = {column["name"]: column for column in inspect(db.engine).get_columns("instagram_comments")}
                self.assertFalse(before["original_text"]["nullable"])

                # Seed a complete pre-existing comment lineage before the
                # nullable change: root/reply links, a text version and its
                # human assessment must all survive both schema operations.
                user = User(name="Migration user", email="migration@example.org")
                user.set_password("migration-test-only")
                db.session.add(user)
                db.session.flush()
                project = AnalyticsProject(owner_user_id=user.id, name="Projeto de migration")
                account = InstagramAccount(
                    analytics_project=project,
                    username="migration_account",
                    username_normalized="migration_account",
                )
                media = InstagramMedia(
                    instagram_account=account,
                    external_id="migration-media",
                    caption_original="Publicação preservada",
                    published_at=datetime(2026, 10, 1, tzinfo=timezone.utc),
                )
                root = InstagramComment(
                    instagram_media=media,
                    external_id="migration-root",
                    published_at=datetime(2026, 10, 2, tzinfo=timezone.utc),
                    original_text="Texto original preservado.",
                )
                reply = InstagramComment(
                    instagram_media=media,
                    external_id="migration-reply",
                    parent=root,
                    published_at=datetime(2026, 10, 2, 1, tzinfo=timezone.utc),
                    original_text="Resposta preservada.",
                )
                text_version = InstagramCommentTextVersion(
                    instagram_comment=root,
                    text="Texto original preservado.",
                    content_hash="a" * 64,
                    first_observed_at=datetime(2026, 10, 2, tzinfo=timezone.utc),
                    is_original=True,
                )
                assessment = InstagramCommentAssessment(
                    instagram_comment=root,
                    text_version=text_version,
                    source="manual",
                    reviewer_user_id=user.id,
                    sentiment="positive",
                    needs_response=False,
                    risk="none",
                )
                db.session.add_all((project, account, media, root, reply, text_version, assessment))
                db.session.commit()
                root_id, reply_id = root.id, reply.id
                version_id, assessment_id = text_version.id, assessment.id

                upgrade(directory=migrations, revision="e3f4a5b6c7d8")
                after_upgrade = {
                    column["name"]: column for column in inspect(db.engine).get_columns("instagram_comments")
                }
                self.assertTrue(after_upgrade["original_text"]["nullable"])
                self.assertIn(
                    "ck_instagram_comment_original_text",
                    {item["name"] for item in inspect(db.engine).get_check_constraints("instagram_comments")},
                )
                with db.engine.connect() as connection:
                    self.assertEqual(connection.exec_driver_sql("PRAGMA foreign_keys").scalar_one(), 1)
                    self.assertEqual(
                        connection.exec_driver_sql("SELECT version_num FROM alembic_version").scalar_one(),
                        "e3f4a5b6c7d8",
                    )

                # Existing canonical rows and their linked history remain
                # intact after SQLite's batch table rebuild.
                db.session.remove()
                preserved_root = db.session.get(InstagramComment, root_id)
                preserved_reply = db.session.get(InstagramComment, reply_id)
                preserved_version = db.session.get(InstagramCommentTextVersion, version_id)
                preserved_assessment = db.session.get(InstagramCommentAssessment, assessment_id)
                self.assertEqual(preserved_root.original_text, "Texto original preservado.")
                self.assertEqual(preserved_reply.parent_comment_id, root_id)
                self.assertEqual(preserved_version.instagram_comment_id, root_id)
                self.assertEqual(preserved_assessment.instagram_comment_text_version_id, version_id)
                self.assertEqual(db.session.execute(text("PRAGMA foreign_key_check")).all(), [])

                # A populated NULL is a deliberate data state after this
                # revision. Downgrade must fail rather than inventing or
                # destroying content in order to restore NOT NULL.
                unavailable = InstagramComment(
                    instagram_media_id=preserved_root.instagram_media_id,
                    external_id="migration-text-unavailable",
                    published_at=datetime(2026, 10, 3, tzinfo=timezone.utc),
                    original_text=None,
                )
                db.session.add(unavailable)
                db.session.commit()
                unavailable_id = unavailable.id
                # flask_migrate translates a migration RuntimeError into the
                # CLI-compatible SystemExit(1); that is the explicit block.
                with self.assertRaises(SystemExit) as blocked_downgrade:
                    downgrade(directory=migrations, revision="d2e1f0a9b8c7")
                self.assertEqual(blocked_downgrade.exception.code, 1)
                db.session.remove()
                with db.engine.connect() as connection:
                    self.assertEqual(connection.exec_driver_sql("PRAGMA foreign_keys").scalar_one(), 1)
                    self.assertEqual(
                        connection.exec_driver_sql("SELECT version_num FROM alembic_version").scalar_one(),
                        "e3f4a5b6c7d8",
                    )
                self.assertEqual(db.session.get(InstagramComment, root_id).original_text, "Texto original preservado.")
                self.assertEqual(db.session.get(InstagramComment, reply_id).parent_comment_id, root_id)
                self.assertIsNotNone(db.session.get(InstagramCommentAssessment, assessment_id))

                # Removing only the synthetic unavailable row permits a safe
                # rollback while preserving all pre-existing linked records.
                db.session.delete(db.session.get(InstagramComment, unavailable_id))
                db.session.commit()
                downgrade(directory=migrations, revision="d2e1f0a9b8c7")
                after_downgrade = {
                    column["name"]: column for column in inspect(db.engine).get_columns("instagram_comments")
                }
                self.assertFalse(after_downgrade["original_text"]["nullable"])
                with db.engine.connect() as connection:
                    self.assertEqual(connection.exec_driver_sql("PRAGMA foreign_keys").scalar_one(), 1)
                    self.assertEqual(
                        connection.exec_driver_sql("SELECT version_num FROM alembic_version").scalar_one(),
                        "d2e1f0a9b8c7",
                    )
                db.session.remove()
                self.assertEqual(db.session.get(InstagramComment, root_id).original_text, "Texto original preservado.")
                self.assertEqual(db.session.get(InstagramComment, reply_id).parent_comment_id, root_id)
                self.assertEqual(db.session.get(InstagramCommentTextVersion, version_id).instagram_comment_id, root_id)
                self.assertEqual(
                    db.session.get(InstagramCommentAssessment, assessment_id).instagram_comment_text_version_id,
                    version_id,
                )
                self.assertEqual(db.session.execute(text("PRAGMA foreign_key_check")).all(), [])
                db.session.remove()
                db.engine.dispose()

    def test_sqlite_fk_guard_rejects_an_active_transaction_without_touching_schema(self):
        migration = import_module(
            "migrations.versions.e3f4a5b6c7d8_permite_texto_indisponivel_em_comentarios"
        )
        engine = sa.create_engine("sqlite://")
        statements: list[str] = []

        def record_sql(_connection, _cursor, statement, _parameters, _context, _executemany):
            statements.append(statement)

        sa.event.listen(engine, "before_cursor_execute", record_sql)
        try:
            with engine.connect() as connection:
                self.assertEqual(connection.exec_driver_sql("PRAGMA foreign_keys").scalar_one(), 1)
                connection.exec_driver_sql("CREATE TABLE migration_fk_guard_probe (id INTEGER PRIMARY KEY)")
                connection.commit()
                transaction = connection.begin()
                try:
                    connection.exec_driver_sql("INSERT INTO migration_fk_guard_probe (id) VALUES (1)")
                    self.assertTrue(migration._sqlite_driver_transaction_active(connection))
                    with Operations.context(MigrationContext.configure(connection)):
                        with self.assertRaisesRegex(RuntimeError, "transação SQLite ativa"):
                            migration._alter_original_text_nullability(
                                nullable=True,
                                constraint_sql="original_text IS NULL OR length(trim(original_text)) > 0",
                            )
                    self.assertEqual(connection.exec_driver_sql("PRAGMA foreign_keys").scalar_one(), 1)
                    self.assertTrue(migration._sqlite_driver_transaction_active(connection))
                    self.assertFalse(any("FOREIGN_KEYS=OFF" in statement.upper().replace(" ", "") for statement in statements))
                    self.assertFalse(any("DROP TABLE INSTAGRAM_COMMENTS" in statement.upper() for statement in statements))
                finally:
                    transaction.rollback()
                self.assertEqual(connection.exec_driver_sql("PRAGMA foreign_keys").scalar_one(), 1)
                self.assertEqual(connection.exec_driver_sql("PRAGMA foreign_key_check").all(), [])
        finally:
            sa.event.remove(engine, "before_cursor_execute", record_sql)
            engine.dispose()

    def test_sqlite_fk_state_is_restored_after_a_controlled_batch_failure(self):
        migration = import_module(
            "migrations.versions.e3f4a5b6c7d8_permite_texto_indisponivel_em_comentarios"
        )
        engine = sa.create_engine("sqlite://")
        statements: list[str] = []

        def record_sql(_connection, _cursor, statement, _parameters, _context, _executemany):
            statements.append(statement)

        sa.event.listen(engine, "before_cursor_execute", record_sql)
        try:
            with engine.connect() as connection:
                self.assertEqual(connection.exec_driver_sql("PRAGMA foreign_keys").scalar_one(), 1)
                with Operations.context(MigrationContext.configure(connection)):
                    with patch.object(
                        migration.op,
                        "batch_alter_table",
                        side_effect=RuntimeError("batch controlado"),
                    ):
                        with self.assertRaisesRegex(RuntimeError, "batch controlado"):
                            migration._alter_original_text_nullability(
                                nullable=True,
                                constraint_sql="original_text IS NULL OR length(trim(original_text)) > 0",
                            )
                normalized = [statement.upper().replace(" ", "") for statement in statements]
                self.assertLess(
                    normalized.index("PRAGMAFOREIGN_KEYS=OFF"),
                    normalized.index("PRAGMAFOREIGN_KEYS=ON"),
                )
                self.assertEqual(connection.exec_driver_sql("PRAGMA foreign_keys").scalar_one(), 1)
                self.assertEqual(connection.exec_driver_sql("PRAGMA foreign_key_check").all(), [])
        finally:
            sa.event.remove(engine, "before_cursor_execute", record_sql)
            engine.dispose()

    def test_comments_migration_is_a_reversible_child_of_phase_six_head(self):
        from app import create_app

        migrations = str(Path(__file__).resolve().parents[1] / "migrations")
        with tempfile.TemporaryDirectory(prefix="pesquisapdf-comment-migration-") as root:
            database = Path(root) / "platform.sqlite3"
            app = create_app({
                "TESTING": True,
                "SQLALCHEMY_DATABASE_URI": f"sqlite:///{database.as_posix()}",
                "PLATFORM_DATA_DIR": root,
                "SECRET_KEY": "comments-migration-test-only",
            })
            with app.app_context():
                upgrade(directory=migrations, revision="e5f6a7b8c9d0")
                self.assertNotIn("instagram_comments", inspect(db.engine).get_table_names())
                upgrade(directory=migrations, revision="f6a7b8c9d0e1")
                inspector = inspect(db.engine)
                expected = {
                    "instagram_comments", "instagram_comment_text_versions",
                    "instagram_comment_observations", "instagram_comment_assessments",
                }
                self.assertTrue(expected.issubset(inspector.get_table_names()))
                self.assertIn(
                    "uq_instagram_comment_media_external_id",
                    {item["name"] for item in inspector.get_unique_constraints("instagram_comments")},
                )
                downgrade(directory=migrations, revision="e5f6a7b8c9d0")
                self.assertTrue(expected.isdisjoint(inspect(db.engine).get_table_names()))
                db.session.remove()
                db.engine.dispose()


if __name__ == "__main__":
    unittest.main()
