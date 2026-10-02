"""Importação CSV/XLSX do Análysis Instagram sem efeitos parciais."""

from __future__ import annotations

import io
import json
import re
import unittest
from datetime import datetime
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, func, select

from platform_helpers import create_user, csrf_from, isolated_platform, login
from platform_core.extensions import db
from platform_core.instagram_import import (
    ACCOUNT_METRIC_FIELDS,
    CANONICAL_HEADERS,
    IMPORT_FIELDS,
    MAX_IMPORT_BYTES,
    MAX_IMPORT_ROWS,
    InstagramImportError,
    parse_instagram_import,
)
from platform_core.instagram_previews import InstagramImportPreviews, previews
from platform_core.instagram_storage import persist_instagram_import
from platform_core.models import (
    AnalyticsProject,
    AnalyticsRun,
    InstagramAccount,
    InstagramAccountSnapshot,
    InstagramMedia,
    InstagramMediaObservation,
    PlanTool,
)
from platform_core.services import ANALYTICS_INSTAGRAM_TOOL


BASE_ROW = {
    "account_username": "@RevistaX",
    "media_external_id": "post-1",
    "published_at": "2026-09-03T12:00:00Z",
    "observed_at": "2026-10-02T14:30:00Z",
    "account_display_name": "Revista X",
    "account_external_id": "account-123",
    "media_type": "REEL",
    "permalink": "https://instagram.example/post-1",
    "caption": "Texto original.",
    "period_start": "2026-09-01T00:00:00Z",
    "period_end": "2026-09-30T23:59:59Z",
    "reach": "1000",
    "impressions": "1200",
    "plays": "900",
    "likes": "50",
    "comments_count": "2",
    "shares": "3",
    "saves": "4",
    "interactions": "59",
    "follows_generated": "1",
    "followers_count": "8200",
    "new_followers": "30",
    "account_reach": "5000",
    "profile_views": "200",
    "website_clicks": "7",
}


def csv_bytes(*rows: dict[str, str], headers: tuple[str, ...] = CANONICAL_HEADERS) -> bytes:
    output = io.StringIO(newline="")
    import csv
    writer = csv.DictWriter(output, fieldnames=headers, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode("utf-8")


def xlsx_bytes(*rows: dict[str, str]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Importação"
    sheet.append(CANONICAL_HEADERS)
    for row in rows:
        sheet.append([row.get(header, "") for header in CANONICAL_HEADERS])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class InstagramImportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.scope = isolated_platform()
        cls.app = cls.scope.__enter__()
        cls.client = cls.app.test_client()
        cls.user = create_user()
        db.session.add(PlanTool(plan_id="student", tool_id=ANALYTICS_INSTAGRAM_TOOL))
        db.session.commit()
        login(cls.client)

    @classmethod
    def tearDownClass(cls):
        cls.scope.__exit__(None, None, None)

    def setUp(self):
        db.session.rollback()
        for model in (
            InstagramAccountSnapshot,
            InstagramMediaObservation,
            InstagramMedia,
            InstagramAccount,
            AnalyticsRun,
            AnalyticsProject,
        ):
            db.session.execute(delete(model))
        db.session.commit()

    def create_project(self, name: str = "Instagram Revista Latitude") -> AnalyticsProject:
        response = self.client.get("/analytics/instagram/projects/new")
        created = self.client.post("/analytics/instagram/projects", data={
            "csrf_token": csrf_from(response), "name": name, "description": "Acompanhamento editorial.",
        })
        self.assertEqual(created.status_code, 302)
        project_id = created.headers["Location"].rsplit("/", 1)[1]
        return db.session.get(AnalyticsProject, project_id)

    def preview(self, project: AnalyticsProject, data: bytes, filename: str = "coleta.csv"):
        page = self.client.get(f"/analytics/instagram/projects/{project.id}/imports/new")
        return self.client.post(
            f"/analytics/instagram/projects/{project.id}/imports/preview",
            data={"csrf_token": csrf_from(page), "dataset": (io.BytesIO(data), filename)},
        )

    def confirm(self, project: AnalyticsProject, preview_response):
        html = preview_response.get_data(as_text=True)
        token = re.search(r'name="preview_token" value="([^"]+)"', html).group(1)
        return self.client.post(f"/analytics/instagram/projects/{project.id}/imports/confirm", data={
            "csrf_token": csrf_from(preview_response), "preview_token": token,
        })

    def preview_token(self, response):
        return re.search(r'name="preview_token" value="([^"]+)"', response.get_data(as_text=True)).group(1)

    def preview_path(self, project, response):
        return previews._directory(self.user.id, project.id) / f"{self.preview_token(response)}.json"

    def test_existing_pages_follow_shared_project_and_import_layout(self):
        project = self.create_project()
        listing = self.client.get("/analytics/instagram").get_data(as_text=True)
        self.assertIn('id="instagram-projects-list" data-view-container', listing)
        self.assertIn('class="platform-view-list"', listing)
        self.assertIn('aria-label="Visualização em cartões"', listing)
        self.assertIn(f'Abrir projeto {project.name}', listing)
        self.assertNotIn("Acompanhe desempenho, recepção e evolução", listing)
        new_project = self.client.get("/analytics/instagram/projects/new").get_data(as_text=True)
        self.assertIn('class="platform-panel"', new_project)
        self.assertIn('>Novo projeto</button>', new_project)
        detail = self.client.get(f"/analytics/instagram/projects/{project.id}").get_data(as_text=True)
        self.assertIn("Nenhuma importação registrada neste projeto.", detail)
        self.assertIn("Importar dados", detail)
        upload = self.client.get(f"/analytics/instagram/projects/{project.id}/imports/new").get_data(as_text=True)
        self.assertIn("Modelo para importação", upload)
        self.assertIn("Baixar modelo Excel", upload)
        self.assertIn("Importar arquivo", upload)
        self.assertIn('class="form-control file-upload-input"', upload)
        preview = self.preview(project, csv_bytes(BASE_ROW))
        self.assertIn('<h1 class="h4 mb-1">Prévia da importação</h1>', preview.get_data(as_text=True))
        self.assertIn("Amostra das publicações validadas", preview.get_data(as_text=True))
        completed = self.confirm(project, preview)
        self.assertEqual(completed.status_code, 200)
        self.assertIn("Resumo da importação", completed.get_data(as_text=True))
        self.assertIn("Histórico de importações", self.client.get(
            f"/analytics/instagram/projects/{project.id}"
        ).get_data(as_text=True))
        pages = (listing, new_project, detail, upload,
                 preview.get_data(as_text=True), completed.get_data(as_text=True))
        for html in pages:
            main = html.split('<main class="platform-main', 1)[1].split('</main>', 1)[0]
            self.assertNotIn('class="platform-page-heading"', main)
            self.assertNotIn('class="display-', main)
            self.assertNotIn('nesta raspagem', main.casefold())
            self.assertIn('class="app-content-container"', main)

    def test_import_period_is_stored_in_utc_and_presented_in_configured_timezone(self):
        self.app.config["PRESENTATION_TIMEZONE"] = "America/Maceio"
        project = self.create_project()
        row = dict(BASE_ROW, period_start="2026-09-01T00:00:00-03:00",
                   period_end="2026-09-30T23:59:59-03:00")
        preview = self.preview(project, csv_bytes(row))
        self.assertEqual(preview.status_code, 200)
        preview_html = preview.get_data(as_text=True)
        self.assertIn("01/09/2026 00:00 a 30/09/2026 23:59", preview_html)
        self.assertIn("03/09/2026 09:00", preview_html)  # publicação, não observação
        self.assertIn("02/10/2026 11:30", preview_html)  # observação, não publicação
        completed = self.confirm(project, preview)
        self.assertEqual(completed.status_code, 200)
        run = db.session.scalar(select(AnalyticsRun).where(AnalyticsRun.analytics_project_id == project.id))
        self.assertEqual(run.period_start.replace(tzinfo=None), datetime(2026, 9, 1, 3, 0))
        self.assertEqual(run.period_end.replace(tzinfo=None), datetime(2026, 10, 1, 2, 59, 59))
        run.completed_at = datetime(2026, 10, 2, 11, 2)
        db.session.commit()
        detail = self.client.get(f"/analytics/instagram/projects/{project.id}").get_data(as_text=True)
        self.assertIn("01/09/2026 00:00 a 30/09/2026 23:59", detail)
        self.assertEqual(detail.count("02/10/2026 08:02"), 2)  # última importação e histórico
        self.assertNotIn("01/09/2026 03:00 a 01/10/2026 02:59", detail)

    def test_missing_period_still_says_not_informed(self):
        project = self.create_project()
        row = dict(BASE_ROW, period_start="", period_end="")
        preview = self.preview(project, csv_bytes(row))
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(self.confirm(project, preview).status_code, 200)
        detail = self.client.get(f"/analytics/instagram/projects/{project.id}").get_data(as_text=True)
        self.assertIn("<td>Não informado</td>", detail)

    def test_official_excel_model_download_matches_import_schema(self):
        project = self.create_project()
        page = self.client.get(f"/analytics/instagram/projects/{project.id}/imports/new")
        self.assertIn(f'/analytics/instagram/projects/{project.id}/imports/modelo', page.get_data(as_text=True))
        response = self.client.get(f"/analytics/instagram/projects/{project.id}/imports/modelo")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("modelo_importacao_instagram_analysis.xlsx", response.headers["Content-Disposition"])
        self.assertLessEqual(len(response.data), MAX_IMPORT_BYTES)
        workbook = load_workbook(io.BytesIO(response.data))
        try:
            self.assertEqual(workbook.sheetnames, ["Dados Instagram"])
            sheet = workbook.active
            self.assertEqual(tuple(cell.value for cell in sheet[1]), CANONICAL_HEADERS)
            self.assertEqual(len(CANONICAL_HEADERS), 25)
            self.assertEqual(CANONICAL_HEADERS, tuple(field.name for field in IMPORT_FIELDS))
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.auto_filter.ref, "A1:Y1")
            self.assertTrue(all(sheet.cell(1, column).comment for column in range(1, 26)))
            self.assertFalse(any(cell.data_type == "f" for row in sheet.iter_rows(max_row=2) for cell in row))
            self.assertTrue(all(sheet.cell(2, column).value is None for column in range(1, len(CANONICAL_HEADERS) + 1)))
            datetime_headers = ("published_at", "observed_at", "period_start", "period_end")
            for header in datetime_headers:
                column = CANONICAL_HEADERS.index(header) + 1
                with self.subTest(header=header):
                    self.assertEqual(sheet.cell(2, column).number_format, "@")
                    self.assertEqual(sheet.cell(1000, column).number_format, "@")
                    self.assertEqual(sheet.cell(MAX_IMPORT_ROWS + 1, column).number_format, "@")
                    note = sheet.cell(1, column).comment.text
                    self.assertIn("ISO 8601", note)
                    self.assertIn("timezone", note)
                    self.assertIn("não converta", note)
                    self.assertIn("-03:00", note)
            for column, header in enumerate(CANONICAL_HEADERS, start=1):
                sheet.cell(2, column, BASE_ROW.get(header))
            iso_values = {
                "published_at": "2026-09-20T18:30:00-03:00",
                "observed_at": "2026-10-02T09:00:00-03:00",
                "period_start": "2026-09-01T00:00:00-03:00",
                "period_end": "2026-09-30T23:59:59-03:00",
            }
            for header, value in iso_values.items():
                sheet.cell(2, CANONICAL_HEADERS.index(header) + 1, value)
            filled = io.BytesIO()
            workbook.save(filled)
        finally:
            workbook.close()
            response.close()
        reopened = load_workbook(io.BytesIO(filled.getvalue()), read_only=True)
        try:
            for header, value in iso_values.items():
                self.assertEqual(reopened.active.cell(2, CANONICAL_HEADERS.index(header) + 1).value, value)
        finally:
            reopened.close()
        dataset = parse_instagram_import("modelo_importacao_instagram_analysis.xlsx", filled.getvalue())
        self.assertEqual(len(dataset.rows), 1)
        self.assertEqual(dataset.rows[0].media_metrics["reach"], 1000)
        self.assertEqual(dataset.rows[0].published_at.isoformat(), "2026-09-20T21:30:00+00:00")
        self.assertEqual(dataset.rows[0].observed_at.isoformat(), "2026-10-02T12:00:00+00:00")
        self.assertEqual(dataset.rows[0].period_start.isoformat(), "2026-09-01T03:00:00+00:00")
        self.assertEqual(dataset.rows[0].period_end.isoformat(), "2026-10-01T02:59:59+00:00")

    def test_import_instructions_share_the_canonical_fields(self):
        project = self.create_project()
        page = self.client.get(f"/analytics/instagram/projects/{project.id}/imports/new")
        html = page.get_data(as_text=True)
        self.assertIn("Ver instruções de preenchimento", html)
        for field in IMPORT_FIELDS:
            self.assertIn(f"<code>{field.name}</code>", html)
        self.assertEqual(html.count("<th scope=\"row\"><code>"), len(IMPORT_FIELDS))

    def test_excel_model_requires_login_access_and_project_ownership(self):
        project = self.create_project()
        url = f"/analytics/instagram/projects/{project.id}/imports/modelo"
        other = create_user("Outra", "excel-other@example.org")
        try:
            self.client.post("/logout", data={"csrf_token": csrf_from(self.client.get("/"))})
            self.assertEqual(self.client.get(url).status_code, 302)
            login(self.client, other.email)
            self.assertEqual(self.client.get(url).status_code, 404)
        finally:
            self.client.post("/logout", data={"csrf_token": csrf_from(self.client.get("/"))})
            login(self.client, self.user.email)

    def test_preview_is_server_side_json_and_survives_a_new_store_instance(self):
        project = self.create_project()
        response = self.preview(project, csv_bytes(BASE_ROW), "coleta.csv")
        path = self.preview_path(project, response)
        self.assertTrue(path.exists())
        envelope = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual(envelope["user_id"], self.user.id)
        self.assertEqual(envelope["project_id"], project.id)
        self.assertEqual(envelope["preview"]["dataset"]["rows"][0]["media_metrics"]["reach"], 1000)
        self.assertNotIn("_items", vars(previews))
        # A rota usa outro objeto recém-criado, como em outro worker/processo.
        with patch("platform_core.instagram_routes.previews", InstagramImportPreviews()):
            result = self.confirm(project, response)
        self.assertEqual(result.status_code, 200)
        self.assertFalse(path.exists())
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 1)

    def test_expired_preview_is_removed_and_cannot_create_a_run(self):
        project = self.create_project()
        response = self.preview(project, csv_bytes(BASE_ROW))
        path = self.preview_path(project, response)
        expiration = json.loads(path.read_text(encoding="utf-8"))["expires_at"]
        with patch("platform_core.instagram_previews.now_timestamp", return_value=expiration + 1):
            result = self.confirm(project, response)
        self.assertEqual(result.status_code, 302)
        self.assertFalse(path.exists())
        with self.client.session_transaction() as session:
            self.assertIn("A prévia expirou. Envie o arquivo novamente.", str(session.get("_flashes")))
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 0)

    def test_new_preview_opportunistically_removes_expired_sibling(self):
        project = self.create_project()
        first = self.preview(project, csv_bytes(BASE_ROW))
        first_path = self.preview_path(project, first)
        envelope = json.loads(first_path.read_text(encoding="utf-8"))
        self.assertEqual(envelope["expires_at"] - envelope["created_at"], 15 * 60)
        with patch("platform_core.instagram_previews.now_timestamp", return_value=envelope["expires_at"] + 1):
            second = self.preview(project, csv_bytes(dict(BASE_ROW, media_external_id="post-2")))
        self.assertEqual(second.status_code, 200)
        self.assertFalse(first_path.exists())
        self.assertTrue(self.preview_path(project, second).exists())

    def test_user_and_project_binding_and_invalid_tokens(self):
        project = self.create_project()
        other_project = self.create_project("Outro projeto")
        response = self.preview(project, csv_bytes(BASE_ROW))
        token = self.preview_token(response)
        other = create_user("Outro usuário", "outro-instagram@example.org")
        self.assertIsNone(previews.get(token, other.id, project.id))
        self.assertIsNone(previews.claim(token, other.id, project.id))
        self.assertIsNone(previews.get(token, self.user.id, other_project.id))
        self.assertIsNone(previews.claim(token, self.user.id, other_project.id))
        for invalid in ("../" + token, "x" * 43, "", "../../platform.sqlite3"):
            result = self.client.post(f"/analytics/instagram/projects/{project.id}/imports/confirm", data={
                "csrf_token": csrf_from(response), "preview_token": invalid,
            })
            self.assertEqual(result.status_code, 302)
        self.assertTrue(self.preview_path(project, response).exists())
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 0)

    def test_failure_rolls_back_and_releases_preview_for_retry(self):
        project = self.create_project()
        response = self.preview(project, csv_bytes(BASE_ROW))
        with patch.object(db.session, "flush", side_effect=RuntimeError("falha transitória")):
            failure = self.confirm(project, response)
        self.assertEqual(failure.status_code, 500)
        self.assertTrue(self.preview_path(project, response).exists())
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 0)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramMedia)), 0)
        self.assertEqual(self.confirm(project, response).status_code, 200)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 1)
        self.assertFalse(self.preview_path(project, response).exists())

    def test_claim_blocks_simultaneous_and_replayed_confirmation(self):
        project = self.create_project()
        response = self.preview(project, csv_bytes(BASE_ROW))
        token = self.preview_token(response)
        claim = previews.claim(token, self.user.id, project.id)
        self.assertIsNotNone(claim)
        try:
            concurrent = self.confirm(project, response)
            self.assertEqual(concurrent.status_code, 302)
            self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 0)
        finally:
            claim.release()
        self.assertEqual(self.confirm(project, response).status_code, 200)
        self.assertEqual(self.confirm(project, response).status_code, 302)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 1)

    def test_cancel_is_idempotent_and_removes_persistent_artifact(self):
        project = self.create_project()
        response = self.preview(project, csv_bytes(BASE_ROW))
        token = self.preview_token(response)
        path = self.preview_path(project, response)
        for _ in range(2):
            result = self.client.post(f"/analytics/instagram/projects/{project.id}/imports/cancel", data={
                "csrf_token": csrf_from(response), "preview_token": token,
            })
            self.assertEqual(result.status_code, 302)
        self.assertFalse(path.exists())
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 0)

    def test_client_cannot_replace_server_side_dataset_at_confirmation(self):
        project = self.create_project()
        response = self.preview(project, csv_bytes(BASE_ROW))
        confirmed = self.client.post(f"/analytics/instagram/projects/{project.id}/imports/confirm", data={
            "csrf_token": csrf_from(response), "preview_token": self.preview_token(response),
            "reach": "999999", "media_external_id": "adulterado", "caption": "adulterada",
        })
        self.assertEqual(confirmed.status_code, 200)
        self.assertEqual(db.session.scalar(select(InstagramMediaObservation.reach)), 1000)
        self.assertEqual(db.session.scalar(select(InstagramMedia.external_id)), "post-1")

    def test_csv_preview_then_confirmation_creates_project_account_media_run_observations_and_snapshot(self):
        project = self.create_project()
        response = self.preview(project, csv_bytes(BASE_ROW))
        self.assertEqual(response.status_code, 200)
        self.assertIn("Prévia da importação", response.get_data(as_text=True))
        self.assertIn("Amostra das publicações validadas", response.get_data(as_text=True))
        self.assertIn("post-1", response.get_data(as_text=True))
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 0)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramMedia)), 0)

        confirmed = self.confirm(project, response)
        self.assertEqual(confirmed.status_code, 200)
        self.assertIn("Importação concluída", confirmed.get_data(as_text=True))
        account = db.session.scalar(select(InstagramAccount))
        self.assertEqual(account.username_normalized, "revistax")
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 1)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramMedia)), 1)
        observation = db.session.scalar(select(InstagramMediaObservation))
        self.assertEqual(observation.reach, 1000)
        self.assertEqual(observation.observed_at.isoformat(), "2026-10-02T14:30:00")
        self.assertEqual(observation.period_end.isoformat(), "2026-09-30T23:59:59")
        self.assertEqual(db.session.scalar(select(InstagramAccountSnapshot).where(
            InstagramAccountSnapshot.instagram_account_id == account.id
        )).followers_count, 8200)
        detail = self.client.get(f"/analytics/instagram/projects/{project.id}").get_data(as_text=True)
        self.assertIn("Histórico de importações", detail)
        self.assertIn("Publicações observadas", detail)

    def test_xlsx_validates_and_persists_the_same_canonical_schema(self):
        project = self.create_project()
        response = self.preview(project, xlsx_bytes(BASE_ROW), "coleta.xlsx")
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        confirmed = self.confirm(project, response)
        self.assertEqual(confirmed.status_code, 200)
        run = db.session.scalar(select(AnalyticsRun))
        self.assertEqual(run.source_kind, "xlsx")
        self.assertEqual(run.processor_version, "instagram-import-v1")
        self.assertEqual(run.parameters_json["schema_version"], "instagram-import-v1")
        self.assertEqual(run.period_start.isoformat(), "2026-09-01T00:00:00")
        self.assertEqual(run.period_end.isoformat(), "2026-09-30T23:59:59")

    def test_second_import_reuses_media_and_keeps_both_historical_observations(self):
        project = self.create_project()
        first_preview = self.preview(project, csv_bytes(BASE_ROW))
        self.assertEqual(self.confirm(project, first_preview).status_code, 200)
        second = dict(BASE_ROW, reach="1500", likes="73", observed_at="2026-10-10T14:30:00Z")
        second_preview = self.preview(project, csv_bytes(second), "coleta-2.csv")
        self.assertEqual(second_preview.status_code, 200)
        self.assertIn("0 / 1", second_preview.get_data(as_text=True))
        self.assertEqual(self.confirm(project, second_preview).status_code, 200)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramMedia)), 1)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 2)
        values = db.session.scalars(select(InstagramMediaObservation.reach).order_by(
            InstagramMediaObservation.observed_at
        )).all()
        self.assertEqual(values, [1000, 1500])
        detail = self.client.get(f"/analytics/instagram/projects/{project.id}").get_data(as_text=True)
        self.assertEqual(detail.count("<td>CSV</td>"), 2)

    def test_missing_metrics_remain_null_and_zero_is_preserved(self):
        project = self.create_project()
        row = dict(BASE_ROW, reach="", likes="0")
        for field in ACCOUNT_METRIC_FIELDS:
            row[field] = ""
        response = self.preview(project, csv_bytes(row))
        self.assertEqual(self.confirm(project, response).status_code, 200)
        observation = db.session.scalar(select(InstagramMediaObservation))
        self.assertIsNone(observation.reach)
        self.assertEqual(observation.likes, 0)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramAccountSnapshot)), 0)

    def test_invalid_later_row_prevents_preview_and_any_persistence(self):
        project = self.create_project()
        invalid = dict(BASE_ROW, media_external_id="post-2", reach="-5")
        response = self.preview(project, csv_bytes(BASE_ROW, invalid))
        self.assertEqual(response.status_code, 400)
        self.assertIn("não pode ser negativo", response.get_data(as_text=True))
        for model in (InstagramAccount, InstagramMedia, AnalyticsRun, InstagramMediaObservation):
            self.assertEqual(db.session.scalar(select(func.count()).select_from(model)), 0)

    def test_normalizes_usernames_and_rejects_project_account_divergence(self):
        project = self.create_project()
        first = self.preview(project, csv_bytes(BASE_ROW))
        self.assertEqual(self.confirm(project, first).status_code, 200)
        matching = self.preview(project, csv_bytes(dict(BASE_ROW, observed_at="2026-10-10T14:30:00Z", account_username="revistax")))
        self.assertEqual(matching.status_code, 200)
        divergent = self.preview(project, csv_bytes(dict(BASE_ROW, account_username="@Outra")))
        self.assertEqual(divergent.status_code, 400)
        self.assertIn("associado a @RevistaX", divergent.get_data(as_text=True))

    def test_repeated_hash_warns_but_does_not_block_new_historical_run(self):
        project = self.create_project()
        data = csv_bytes(BASE_ROW)
        first = self.preview(project, data)
        self.assertEqual(self.confirm(project, first).status_code, 200)
        repeated = self.preview(project, data)
        self.assertEqual(repeated.status_code, 200)
        self.assertIn("parece já ter sido importado", repeated.get_data(as_text=True))
        self.assertEqual(self.confirm(project, repeated).status_code, 200)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 2)

    def test_persistence_rolls_back_every_new_record_when_flush_fails(self):
        project = self.create_project()
        dataset = parse_instagram_import("coleta.csv", csv_bytes(BASE_ROW))
        with patch.object(db.session, "flush", side_effect=RuntimeError("falha controlada")):
            with self.assertRaisesRegex(RuntimeError, "falha controlada"):
                persist_instagram_import(project, dataset)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 0)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramAccount)), 0)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramMedia)), 0)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramMediaObservation)), 0)

    def test_failed_second_import_after_flush_preserves_first_run(self):
        project = self.create_project()
        first = self.preview(project, csv_bytes(BASE_ROW))
        self.assertEqual(self.confirm(project, first).status_code, 200)
        changed = dict(BASE_ROW, reach="1500", observed_at="2026-10-10T14:30:00Z")
        dataset = parse_instagram_import("coleta-2.csv", csv_bytes(changed))
        with patch.object(db.session, "commit", side_effect=RuntimeError("falha após flush")):
            with self.assertRaisesRegex(RuntimeError, "falha após flush"):
                persist_instagram_import(project, dataset)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 1)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramMedia)), 1)
        self.assertEqual(db.session.scalars(select(InstagramMediaObservation.reach)).all(), [1000])
        self.assertEqual(db.session.scalar(select(func.count()).select_from(InstagramAccountSnapshot)), 1)

    def test_cancel_discards_preview_without_persisting(self):
        project = self.create_project()
        response = self.preview(project, csv_bytes(BASE_ROW))
        self.assertEqual(response.status_code, 200)
        token = re.search(r'name="preview_token" value="([^"]+)"', response.get_data(as_text=True)).group(1)
        csrf = csrf_from(response)
        cancelled = self.client.post(f"/analytics/instagram/projects/{project.id}/imports/cancel", data={
            "csrf_token": csrf, "preview_token": token,
        })
        self.assertEqual(cancelled.status_code, 302)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 0)
        rejected = self.client.post(f"/analytics/instagram/projects/{project.id}/imports/confirm", data={
            "csrf_token": csrf, "preview_token": token,
        })
        self.assertEqual(rejected.status_code, 302)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 0)

    def test_preview_token_cannot_be_used_in_another_owned_project(self):
        first_project = self.create_project("Primeiro")
        second_project = self.create_project("Segundo")
        response = self.preview(first_project, csv_bytes(BASE_ROW))
        token = re.search(r'name="preview_token" value="([^"]+)"', response.get_data(as_text=True)).group(1)
        rejected = self.client.post(f"/analytics/instagram/projects/{second_project.id}/imports/confirm", data={
            "csrf_token": csrf_from(response), "preview_token": token,
        })
        self.assertEqual(rejected.status_code, 302)
        self.assertEqual(db.session.scalar(select(func.count()).select_from(AnalyticsRun)), 0)

    def test_routes_enforce_ownership_and_csrf(self):
        project = self.create_project()
        other = create_user("Outra", "outra@example.org")
        other_project = AnalyticsProject(owner_user_id=other.id, module_key="instagram", name="Projeto de outra pessoa")
        db.session.add(other_project)
        db.session.commit()
        self.assertEqual(self.client.get(f"/analytics/instagram/projects/{other_project.id}").status_code, 404)
        self.assertEqual(self.client.get(f"/analytics/instagram/projects/{other_project.id}/imports/new").status_code, 404)
        self.assertEqual(self.client.get(f"/analytics/instagram/projects/{other_project.id}/imports/modelo").status_code, 404)
        own_preview = self.preview(project, csv_bytes(BASE_ROW))
        token = re.search(r'name="preview_token" value="([^"]+)"', own_preview.get_data(as_text=True)).group(1)
        self.assertEqual(self.client.post(
            f"/analytics/instagram/projects/{other_project.id}/imports/confirm",
            data={"csrf_token": csrf_from(own_preview), "preview_token": token},
        ).status_code, 404)
        self.assertEqual(self.client.post(
            f"/analytics/instagram/projects/{project.id}/imports/preview",
            data={"dataset": (io.BytesIO(csv_bytes(BASE_ROW)), "coleta.csv")},
        ).status_code, 400)
        self.assertEqual(self.client.post(
            "/logout", data={"csrf_token": csrf_from(self.client.get("/"))}
        ).status_code, 302)
        try:
            login(self.client, other.email)
            csrf = csrf_from(self.client.get("/"))
            self.assertEqual(self.client.get(f"/analytics/instagram/projects/{project.id}").status_code, 404)
            self.assertEqual(self.client.get(f"/analytics/instagram/projects/{project.id}/imports/modelo").status_code, 404)
            self.assertEqual(self.client.post(
                f"/analytics/instagram/projects/{project.id}/imports/confirm",
                data={"csrf_token": csrf, "preview_token": token},
            ).status_code, 404)
            self.assertEqual(self.client.post(
                f"/analytics/instagram/projects/{project.id}/imports/preview",
                data={"csrf_token": csrf, "dataset": (io.BytesIO(csv_bytes(BASE_ROW)), "coleta.csv")},
            ).status_code, 404)
        finally:
            self.client.post("/logout", data={"csrf_token": csrf_from(self.client.get("/"))})
            login(self.client, self.user.email)

class InstagramImportParserTests(unittest.TestCase):
    def test_xlsx_dates_without_explicit_timezone_and_excel_native_dates_are_rejected(self):
        invalid_values = ("2026-10-02T09:00:00", 46300, datetime(2026, 10, 2, 9))
        for value in invalid_values:
            with self.subTest(value=value), self.assertRaisesRegex(InstagramImportError, "ISO 8601 com timezone"):
                parse_instagram_import("coleta.xlsx", xlsx_bytes(dict(BASE_ROW, observed_at=value)))

    def test_parser_rejects_invalid_metrics_required_fields_dates_periods_and_two_accounts(self):
        with self.assertRaisesRegex(InstagramImportError, "não pode ser negativo"):
            parse_instagram_import("coleta.csv", csv_bytes(dict(BASE_ROW, reach="-1")))
        with self.assertRaisesRegex(InstagramImportError, "número inteiro"):
            parse_instagram_import("coleta.csv", csv_bytes(dict(BASE_ROW, likes="12.5")))
        without_required = tuple(header for header in CANONICAL_HEADERS if header != "media_external_id")
        with self.assertRaisesRegex(InstagramImportError, "Coluna obrigatória ausente: media_external_id"):
            parse_instagram_import("coleta.csv", csv_bytes(BASE_ROW, headers=without_required))
        with self.assertRaisesRegex(InstagramImportError, "ISO 8601"):
            parse_instagram_import("coleta.csv", csv_bytes(dict(BASE_ROW, observed_at="2026-10-02 14:30")))
        with self.assertRaisesRegex(InstagramImportError, "informados juntos"):
            parse_instagram_import("coleta.csv", csv_bytes(dict(BASE_ROW, period_end="")))
        with self.assertRaisesRegex(InstagramImportError, "mais de uma conta"):
            parse_instagram_import("coleta.csv", csv_bytes(BASE_ROW, dict(BASE_ROW, media_external_id="post-2", account_username="@Outra")))

    def test_duplicate_identical_rows_are_consolidated_and_conflicts_are_rejected(self):
        duplicated = parse_instagram_import("coleta.csv", csv_bytes(BASE_ROW, BASE_ROW))
        self.assertEqual(len(duplicated.rows), 1)
        self.assertEqual(len(duplicated.warnings), 1)
        with self.assertRaisesRegex(InstagramImportError, "conflita"):
            parse_instagram_import("coleta.csv", csv_bytes(BASE_ROW, dict(BASE_ROW, reach="1001")))
        with self.assertRaisesRegex(InstagramImportError, "métricas da conta"):
            parse_instagram_import("coleta.csv", csv_bytes(BASE_ROW, dict(BASE_ROW, followers_count="8300")))

    def test_invalid_extensions_size_and_xlsx_formula_are_rejected_before_preview(self):
        with self.assertRaisesRegex(InstagramImportError, "somente arquivos"):
            parse_instagram_import("coleta.xlsm", b"conteudo")
        with self.assertRaisesRegex(InstagramImportError, "até 2 MB"):
            parse_instagram_import("coleta.csv", b"x" * (MAX_IMPORT_BYTES + 1))
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(CANONICAL_HEADERS)
        sheet.append([BASE_ROW[header] for header in CANONICAL_HEADERS])
        sheet.cell(2, CANONICAL_HEADERS.index("reach") + 1).value = "=1+1"
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        with self.assertRaisesRegex(InstagramImportError, "fórmulas não são permitidas"):
            parse_instagram_import("coleta.xlsx", output.getvalue())
        with io.BytesIO() as malicious:
            with ZipFile(io.BytesIO(xlsx_bytes(BASE_ROW))) as source, ZipFile(malicious, "w") as target:
                for member in source.infolist():
                    target.writestr(member, source.read(member))
                target.writestr("xl/vbaProject.bin", b"not-a-macro")
            with self.assertRaisesRegex(InstagramImportError, "macros"):
                parse_instagram_import("coleta.xlsx", malicious.getvalue())
        with io.BytesIO() as linked:
            with ZipFile(io.BytesIO(xlsx_bytes(BASE_ROW))) as source, ZipFile(linked, "w") as target:
                for member in source.infolist():
                    target.writestr(member, source.read(member))
                target.writestr("custom/_rels/external.xml.rels", (
                    b'<Relationships><Relationship TargetMode="External" '
                    b'Target="https://example.invalid/data" /></Relationships>'
                ))
            with self.assertRaisesRegex(InstagramImportError, "links externos"):
                parse_instagram_import("coleta.xlsx", linked.getvalue())


if __name__ == "__main__":
    unittest.main()
