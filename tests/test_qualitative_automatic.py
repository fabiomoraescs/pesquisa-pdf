"""Codificação Literal/Regex usa o corpus e os mesmos vínculos da edição manual."""
import hashlib
import io
import json
import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

from flask_migrate import upgrade
from openpyxl import load_workbook
from sqlalchemy import select, text

import test_qualitative_context as context_fixture
from platform_helpers import create_user, csrf_from, isolated_platform, login
from platform_core.analyses import analysis_dir, create_analysis, delete_analysis
from platform_core.extensions import db
from platform_core.models import (Analysis, AnalysisDocument, Project, QualitativeCode, QualitativeCoding,
                                  QualitativeExcerpt, QualitativeMemo, QualitativeRejection, UserToolOverride)
from platform_core.qualitative_corpus import load_qualitative_manifest, qualitative_corpus_dir, qualitative_manifest_path, read_qualitative_page
from platform_core.qualitative_annotations import get_or_create_excerpt
from platform_core.qualitative_automatic import remove_coding
from platform_core.qualitative_search import QualitativeSearchError
from platform_core.scraping_types import QUALITATIVE_TOOL


class QualitativeAutomaticTests(unittest.TestCase):
    # A mesma fixture PDF/corpus já usada nas regressões de codificação manual.
    setUp = context_fixture.QualitativeContextTests.setUp
    tearDown = context_fixture.QualitativeContextTests.tearDown
    selection = context_fixture.QualitativeContextTests.selection
    post = context_fixture.QualitativeContextTests.post

    @property
    def auto_url(self):
        return f"/analise-qualitativa/bases/{self.analysis.id}/documentos/{self.document.id}/codificar"

    def automatic(self, **changes):
        payload = {"q": "Documento", "scope": "document", "mode": "literal",
                   "grep": False, "case_sensitive": False, **changes}
        return self.client.post(self.auto_url, json=payload, headers={"X-CSRFToken": self.csrf})

    def remove(self, coding, *, token=True, analysis_id=None):
        return self.client.delete(f"/analise-qualitativa/bases/{analysis_id or self.analysis.id}/codificacoes/{coding.id}",
                                  headers={"X-CSRFToken": self.csrf} if token else {})

    def codings(self):
        return db.session.scalars(select(QualitativeCoding).order_by(QualitativeCoding.id)).all()

    def fixture_pages(self, documents):
        """Amplia só a fixture persistida: Unicode e múltiplas páginas sem OCR real."""
        manifest = load_qualitative_manifest(self.analysis)
        first = manifest["documents"][0]
        manifest["documents"] = []
        for index, pages in enumerate(documents):
            if index == 0:
                identifier, stored, original = self.document.id, self.document.stored_name, self.document.original_name
            else:
                document = AnalysisDocument(analysis_id=self.analysis.id, stored_name=f"fixture-{index}.pdf",
                                            original_name=f"Documento {index}.pdf")
                db.session.add(document)
                db.session.flush()
                identifier, stored, original = document.id, document.stored_name, document.original_name
            entries = []
            for number, content in enumerate(pages, 1):
                relative = f"{identifier}/pages/{number:06d}.txt"
                path = qualitative_corpus_dir(self.analysis.id) / relative
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(content.encode("utf-8"))
                entries.append({"page_number": number, "file": relative,
                    "sha256": hashlib.sha256(content.encode("utf-8")).hexdigest(), "char_count": len(content),
                    "extraction_method": "text", "layout_available": False})
            manifest["documents"].append({**first, "document_id": identifier, "original_name": original,
                                          "stored_name": stored, "page_count": len(pages), "pages": entries})
        self.analysis.document_count = len(documents)
        db.session.commit()
        qualitative_manifest_path(self.analysis.id).write_text(json.dumps(manifest), encoding="utf-8")

    def test_normal_search_is_read_only_and_toggle_default_is_off(self):
        html = self.client.get(self.base_url).get_data(as_text=True)
        from html.parser import HTMLParser
        class Inputs(HTMLParser):
            def __init__(self):
                super().__init__(); self.items = []; self.stack = []
            def handle_starttag(self, tag, attrs):
                attrs = dict(attrs)
                attrs["_tag"] = tag
                if tag in {"input", "fieldset", "span", "select", "button", "div"}:
                    self.items.append({**attrs, "_ancestors": [i.get("class", "") for i in self.stack],
                                       "_parents": [i["_tag"] for i in self.stack]})
                if tag not in {"input", "meta", "link", "br", "hr", "img", "source", "wbr"}:
                    self.stack.append(attrs)
            def handle_endtag(self, tag):
                if self.stack: self.stack.pop()
        parser = Inputs(); parser.feed(html)
        toggle = next(i for i in parser.items if "data-automatic-toggle" in i)
        self.assertNotIn("checked", toggle)
        options = next(i for i in parser.items if "data-automatic-options" in i)
        self.assertIn("hidden", options)
        row_class = "platform-qualitative-search-options platform-qualitative-search-control-row"
        self.assertIn(row_class, toggle["_ancestors"])
        rejection = next(i for i in parser.items if "data-contextual-rejection" in i)
        self.assertNotIn("checked", rejection)
        self.assertIn("disabled", rejection)
        for name in ("grep", "case", "automatic", "multiple_terms", "contextual_rejection_enabled"):
            self.assertIn(row_class, next(i for i in parser.items if i.get("name") == name)["_ancestors"])
        self.assertNotIn(row_class, options["_ancestors"])
        separators = [i for i in parser.items if i.get("class") == "platform-qualitative-mode-separator"]
        self.assertEqual(separators, [])
        multiple = next(i for i in parser.items if "data-multiple-terms" in i)
        self.assertNotIn("checked", multiple)
        self.assertIn("disabled", multiple)
        self.assertEqual(rejection["role"], "switch")
        self.assertIn("platform-qualitative-switch", rejection["class"])
        settings_class = "platform-qualitative-search-options platform-qualitative-automatic-settings"
        settings = next(i for i in parser.items if "data-automatic-settings" in i)
        self.assertIn("hidden", settings)
        for item in (multiple, rejection):
            self.assertIn(row_class, item["_ancestors"])
            self.assertIn(settings_class, item["_ancestors"])
            self.assertNotIn(options["class"], item["_ancestors"])
        self.assertLess(parser.items.index(multiple), parser.items.index(rejection))
        modes = {i["value"]: i for i in parser.items if i.get("name") == "automatic_mode"}
        self.assertNotIn("disabled", modes["literal"])
        self.assertNotIn("disabled", modes["lexical"])
        self.assertNotIn("disabled", modes["semantic"])
        self.assertEqual(list(modes), ["literal", "lexical", "semantic"])
        self.assertEqual([name for name, mode in modes.items() if "checked" in mode], ["literal"])
        for mode in modes.values():
            self.assertEqual(mode["type"], "radio")
            self.assertEqual(mode["name"], "automatic_mode")
            self.assertNotIn("role", mode)  # aparência de switch não muda semântica de radio
            self.assertIn(options["class"], mode["_ancestors"])
            self.assertIn("platform-qualitative-search-options platform-qualitative-automatic-modes", mode["_ancestors"])
        scope = next(i for i in parser.items if i.get("name") == "automatic_scope")
        self.assertIn(options["class"], scope["_ancestors"])
        self.assertNotIn(settings_class, scope["_ancestors"])
        self.assertLess(parser.items.index(rejection), parser.items.index(scope))
        switches = [i for i in parser.items if i.get("_tag") == "input" and "platform-qualitative-switch" in i.get("class", "")]
        self.assertEqual(len(switches), 8)
        self.assertEqual([i["name"] for i in switches[:5]],
                         ["grep", "case", "automatic", "multiple_terms", "contextual_rejection_enabled"])
        for item in switches:
            self.assertIn("label", item["_parents"])
            self.assertEqual(item["type"], "radio" if item["name"] == "automatic_mode" else "checkbox")
        self.assertEqual(len([i for i in parser.items if i.get("id") == "qualitative-query"]), 1)
        for mode in modes.values():
            self.assertLess(parser.items.index(rejection), parser.items.index(mode))
            self.assertLess(parser.items.index(mode), parser.items.index(scope))
        info_buttons = [i for i in parser.items if "data-platform-tutorial-open" in i]
        self.assertEqual(len(info_buttons), 1)
        info_button = info_buttons[0]
        self.assertEqual(info_button["type"], "button")
        self.assertEqual(info_button["aria-controls"], "platform-tutorial-dialog")
        self.assertEqual(info_button["aria-label"], "Como funciona a Análise quali-dados")
        self.assertNotIn("disabled", info_button)
        self.assertNotIn("label", info_button["_parents"])
        self.assertEqual(html.count('id="platform-tutorial-dialog"'), 1)
        url = self.auto_url.replace("/codificar", "/buscar")
        self.assertEqual(self.client.get(url, query_string={"q": "Documento"}).json["total"], 2)
        self.assertEqual(self.codings(), [])
        self.assertEqual(db.session.query(QualitativeExcerpt).count(), 0)
        self.assertEqual(db.session.query(QualitativeRejection).count(), 0)

    def test_default_off_removal_can_be_created_again_and_preference_is_not_retroactive(self):
        self.automatic()  # campo omitido: default seguro OFF também no endpoint
        coding = self.codings()[0]
        identifier = coding.id
        self.assertFalse(coding.contextual_rejection_enabled)
        self.automatic(contextual_rejection_enabled=True)  # existentes não mudam de política
        db.session.expire_all()
        coding = db.session.get(QualitativeCoding, identifier)
        self.assertFalse(coding.contextual_rejection_enabled)
        self.assertEqual(self.remove(coding).status_code, 200)
        self.assertEqual(db.session.query(QualitativeRejection).count(), 0)
        result = self.automatic()
        self.assertEqual(result.json["summary"]["created"], 1)
        self.assertEqual(result.json["summary"]["existing"], 1)
        self.assertTrue(all(not c.contextual_rejection_enabled for c in self.codings()))

    def test_on_preference_survives_later_off_operation_and_database_reload(self):
        self.automatic(contextual_rejection_enabled=True)
        identifier = self.codings()[0].id
        self.automatic(contextual_rejection_enabled=False)
        db.session.expire_all()
        coding = db.session.get(QualitativeCoding, identifier)
        self.assertTrue(coding.contextual_rejection_enabled)
        # O DELETE não recebe a preferência atual da interface.
        self.assertEqual(self.remove(coding).status_code, 200)
        self.assertEqual(db.session.query(QualitativeRejection).count(), 1)
        result = self.automatic(contextual_rejection_enabled=False)
        self.assertEqual(result.json["summary"]["rejected"], 1)
        self.assertEqual(result.json["summary"]["created"], 0)

    def test_literal_creates_real_codings_exact_offsets_and_is_idempotent_without_pdf(self):
        with patch("platform_core.qualitative_corpus.pdf_extractor.extrair_paginas", side_effect=AssertionError("OCR/PDF")):
            response = self.automatic()
            self.assertEqual(response.status_code, 201)
            self.assertEqual(response.json["summary"]["created"], 2)
            self.assertEqual(response.json["records"]["codes"][0]["name"], "Documento")
            ids = [c.id for c in self.codings()]
            for coding in self.codings():
                excerpt = db.session.get(QualitativeExcerpt, coding.excerpt_id)
                self.assertEqual(coding.origin, "automatic_literal")
                self.assertEqual(coding.source_query, "Documento")
                self.assertEqual(excerpt.quoted_text, self.page["text"][excerpt.start_offset:excerpt.end_offset])
                self.assertEqual(excerpt.page_text_hash, self.page["sha256"])
                self.assertEqual((coding.source_start, coding.source_end), (excerpt.start_offset, excerpt.end_offset))
            again = self.automatic()
            self.assertEqual(again.json["summary"]["created"], 0)
            self.assertEqual(again.json["summary"]["existing"], 2)
            self.assertEqual([c.id for c in self.codings()], ids)
            self.assertEqual(db.session.query(QualitativeExcerpt).count(), 2)
            page = self.client.get(f"{self.base_url}/trechos").json
            self.assertEqual(len(page["excerpts"]), 2)
            self.assertTrue(all(e["codes"][0]["coding_id"] in ids for e in page["excerpts"]))

    def test_one_occurrence_phrase_case_sensitive_and_regex_use_same_search(self):
        for query, case, grep, expected in (
            ("pesquisa racial", False, False, 1), ("documento", False, False, 2),
            ("documento", True, False, 0), ("Documento", True, False, 2),
            (r"^Docu", True, True, 1),
        ):
            with self.subTest(query=query, case=case, grep=grep):
                response = self.automatic(q=query, case_sensitive=case, grep=grep)
                self.assertEqual(response.status_code, 201)
                self.assertEqual(response.json["search"]["total"], expected)
                self.assertEqual(response.json["summary"]["created"] + response.json["summary"]["existing"], expected)
        regex_coding = db.session.scalar(select(QualitativeCoding).where(QualitativeCoding.origin == "automatic_regex"))
        self.assertEqual(regex_coding.source_query, r"^Docu")

    def test_invalid_regex_timeout_and_unsupported_modes_do_not_write(self):
        for changes in ({"q": "(", "grep": True}, {"mode": "unknown"},
                        {"mode": "lexical", "grep": True}, {"mode": "semantic", "grep": True},
                        {"q": " "}, {"scope": []}, {"grep": "false"}, {"q": "a" * 201},
                        {"contextual_rejection_enabled": "false"}, {"contextual_rejection_enabled": 1},
                        {"contextual_rejection_enabled": None}):
            with self.subTest(changes=changes):
                self.assertEqual(self.automatic(**changes).status_code, 400)
                self.assertEqual(db.session.query(QualitativeCode).count(), 0)
        with patch("platform_core.qualitative_search.SEARCH_TIMEOUT_SECONDS", 0):
            response = self.automatic(grep=True)
            self.assertEqual(response.status_code, 400)
            self.assertIn("não foi concluída", response.json["error"])
        self.assertEqual(self.codings(), [])
        self.assertEqual(db.session.query(QualitativeExcerpt).count(), 0)

    def test_existing_excerpt_and_manual_coding_are_reused_without_changing_origin(self):
        manual = self.post("create_and_apply", {"name": "Documento"})
        excerpt_id = manual.json["excerpt_id"]
        code_id = manual.json["records"]["codes"][0]["id"]
        response = self.automatic(contextual_rejection_enabled=True)
        self.assertEqual(response.json["records"]["codes"][0]["id"], code_id)
        self.assertEqual(response.json["summary"]["existing"], 1)
        self.assertEqual(response.json["summary"]["created"], 1)
        coding = db.session.scalar(select(QualitativeCoding).where(QualitativeCoding.excerpt_id == excerpt_id))
        self.assertEqual(coding.origin, "manual")
        self.assertFalse(coding.contextual_rejection_enabled)
        second = self.automatic(q=r"Documen[t]o", grep=True)
        self.assertEqual(second.json["summary"]["created"], 0)
        self.assertEqual(second.json["summary"]["existing"], 2)
        self.assertEqual(db.session.query(QualitativeExcerpt).count(), 2)

    def test_removing_one_manual_association_preserves_neighbors_and_last_excerpt(self):
        first = self.post("create_and_apply", {"name": "Primeiro"})
        self.post("create_and_apply", {"name": "Segundo"})
        excerpt_id = first.json["excerpt_id"]
        codings = self.codings()
        response = self.remove(codings[0])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json["page"]["excerpts"][0]["codes"]), 1)
        self.assertIsNotNone(db.session.get(QualitativeCoding, codings[1].id))
        self.assertIsNotNone(db.session.get(QualitativeExcerpt, excerpt_id))
        last = self.remove(codings[1])
        self.assertEqual(last.json["page"]["excerpts"], [])
        self.assertIsNotNone(db.session.get(QualitativeExcerpt, excerpt_id))
        self.assertEqual(db.session.query(QualitativeRejection).count(), 0)
        reapplied = self.post("apply_codes", {"code_ids": [codings[0].code_id]})
        self.assertEqual(reapplied.json["excerpt_id"], excerpt_id)
        self.assertEqual(len(self.codings()), 1)

    def test_last_coding_removal_keeps_contextual_memo_and_independent_codes(self):
        self.post("create_and_apply", {"name": "Tema"})
        self.post("create_memo", {"text": "Não apagar esta interpretação"})
        coding = self.codings()[0]
        response = self.remove(coding)
        self.assertEqual(response.status_code, 200)
        excerpt = response.json["page"]["excerpts"][0]
        self.assertEqual(excerpt["codes"], [])
        self.assertEqual(excerpt["memo_count"], 1)
        self.assertEqual(db.session.query(QualitativeMemo).one().text, "Não apagar esta interpretação")
        self.assertEqual(db.session.query(QualitativeCode).count(), 1)

    def test_rejection_blocks_same_operation_but_not_new_code_query_mode_case_or_manual(self):
        self.automatic(contextual_rejection_enabled=True)
        coding = self.codings()[0]
        excerpt_id, code_id = coding.excerpt_id, coding.code_id
        self.assertEqual(self.remove(coding).status_code, 200)
        rejection = db.session.query(QualitativeRejection).one()
        self.assertEqual((rejection.query, rejection.code_id, rejection.origin), ("Documento", code_id, "automatic_literal"))
        repeated = self.automatic()
        self.assertEqual(repeated.json["summary"], {"created": 0, "existing": 1, "rejected": 1, "invalid": 0, "documents": 0})
        self.assertEqual(self.automatic(q=r"Document[o]", grep=True,
                                       contextual_rejection_enabled=True).json["summary"]["created"], 1)
        other_query = self.automatic(q="documento", contextual_rejection_enabled=True)
        self.assertEqual(other_query.json["summary"]["created"], 0)
        self.assertEqual(other_query.json["summary"]["rejected"], 1)
        self.remove(db.session.scalar(select(QualitativeCoding).where(
            QualitativeCoding.origin == "automatic_regex", QualitativeCoding.code_id == code_id)))
        regex_result = self.automatic(q="Documento", grep=True, contextual_rejection_enabled=True)
        self.assertEqual(regex_result.json["summary"]["created"], 1)
        self.remove(db.session.scalar(select(QualitativeCoding).where(
            QualitativeCoding.origin == "automatic_regex", QualitativeCoding.code_id == code_id)))
        self.assertEqual(self.automatic(case_sensitive=True).json["summary"]["created"], 1)
        self.remove(db.session.scalar(select(QualitativeCoding).where(QualitativeCoding.source_case_sensitive.is_(True))))
        excerpt = db.session.get(QualitativeExcerpt, excerpt_id)
        manual = self.post("apply_codes", {"code_ids": [code_id]}, selection={
            "start": excerpt.start_offset, "end": excerpt.end_offset,
            "selected_text": excerpt.quoted_text, "page_hash": excerpt.page_text_hash})
        self.assertEqual(manual.status_code, 201)
        self.assertEqual(db.session.scalar(select(QualitativeCoding).where(
            QualitativeCoding.excerpt_id == excerpt_id, QualitativeCoding.code_id == code_id)).origin, "manual")

    def test_editing_auto_excerpt_keeps_identity_origin_and_original_rejection_anchor(self):
        self.automatic(q="pesquisa", contextual_rejection_enabled=True)
        coding = self.codings()[0]
        excerpt = db.session.get(QualitativeExcerpt, coding.excerpt_id)
        original = (excerpt.start_offset, excerpt.end_offset)
        self.post("create_memo", {"text": "Nota persistente"}, selection=self.selection("pesquisa"))
        response = self.client.patch(f"/analise-qualitativa/bases/{self.analysis.id}/trechos/{excerpt.id}",
            json={"start": original[0], "end": original[1] + len(" racial"), "page_hash": self.page["sha256"]},
            headers={"X-CSRFToken": self.csrf})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.automatic(q="pesquisa").json["summary"]["existing"], 1)
        self.assertEqual(db.session.query(QualitativeExcerpt).count(), 1)
        self.assertEqual(self.remove(coding).status_code, 200)
        rejection = db.session.query(QualitativeRejection).one()
        self.assertEqual((rejection.start_offset, rejection.end_offset), original)
        self.assertEqual(self.automatic(q="pesquisa").json["summary"]["rejected"], 1)
        self.assertEqual(db.session.query(QualitativeMemo).count(), 1)

    def test_project_scope_multiple_documents_pages_and_unicode(self):
        content = "A😀ç Documento. Documento racial.\n"
        self.fixture_pages([[content, content], [content]])
        local = self.automatic(q="Documento")
        self.assertEqual(local.json["summary"]["created"], 4)
        project = self.automatic(scope="project")
        self.assertEqual(project.json["summary"]["created"], 2)
        self.assertEqual(project.json["summary"]["existing"], 4)
        first = db.session.scalar(select(QualitativeExcerpt).where(QualitativeExcerpt.start_offset == 4))
        self.assertIsNotNone(first)  # emoji conta um code point, não dois UTF-16
        emoji = self.automatic(q="😀ç", scope="project")
        self.assertEqual(emoji.json["summary"]["created"], 3)
        self.assertTrue(all(c.origin == "automatic_literal" for c in self.codings()))

    def test_all_250_document_and_637_project_matches_are_searched_and_coded(self):
        phrase = "Ação 😀"
        self.fixture_pages([[(phrase + "\n") * 125] * 2,
                            [(phrase + "\n") * 200, (phrase + "\n") * 187]])
        search_url = self.auto_url.replace("/codificar", "/buscar")
        found = self.client.get(search_url, query_string={"q": phrase}).json
        self.assertEqual(found["total"], 250)
        self.assertEqual(len(found["results"]), 250)
        self.assertEqual({m["page_number"] for m in found["results"]}, {1, 2})
        current = self.automatic(q=phrase)
        self.assertEqual(current.status_code, 201)
        self.assertEqual(current.json["summary"]["created"], 250)
        project = self.automatic(q=phrase, scope="project")
        self.assertEqual(project.status_code, 201)
        self.assertEqual(project.json["search"]["total"], 637)
        self.assertEqual(project.json["summary"]["created"], 387)
        self.assertEqual(project.json["summary"]["existing"], 250)
        self.assertEqual(db.session.query(QualitativeExcerpt).count(), 637)
        again = self.automatic(q=phrase, scope="project")
        self.assertEqual(again.json["summary"]["created"], 0)
        self.assertEqual(again.json["summary"]["existing"], 637)
        regex_result = self.automatic(q=r"Ação\s+😀", grep=True, scope="project")
        self.assertEqual(regex_result.status_code, 201)
        self.assertEqual(regex_result.json["summary"]["created"], 0)
        self.assertEqual(regex_result.json["summary"]["existing"], 637)
        self.assertEqual(db.session.query(QualitativeExcerpt).count(), 637)
        self.assertEqual(db.session.query(QualitativeCoding).count(), 637)
        for match in project.json["search"]["results"]:
            self.assertEqual(match["end_offset"] - match["start_offset"], len(phrase))

    def test_consultative_literal_and_regex_return_all_3000_without_writing(self):
        self.fixture_pages([["Ação 😀\n" * 1500, "Ação 😀\n" * 1500]])
        url = self.auto_url.replace("/codificar", "/buscar")
        # A proteção temporal de Regex não se aplica à consulta literal.
        with patch("platform_core.qualitative_search.SEARCH_TIMEOUT_SECONDS", 0):
            literal = self.client.get(url, query_string={"q": "ação 😀"})
        regex_result = self.client.get(url, query_string={"q": r"ação\s+😀", "grep": "1"})
        for response in (literal, regex_result):
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json["total"], 3000)
            self.assertEqual(len(response.json["results"]), 3000)
            self.assertEqual(response.json["results"][-1]["page_number"], 2)
            self.assertEqual(response.json["results"][-1]["start_offset"], 1499 * len("Ação 😀\n"))
            self.assertNotIn("truncated", response.json)
        self.assertEqual(db.session.query(QualitativeExcerpt).count(), 0)
        self.assertEqual(self.codings(), [])

    def test_300_matches_include_100_existing_50_rejected_and_150_new(self):
        self.fixture_pages([["raça\n" * 300]])
        page = read_qualitative_page(self.analysis, self.document.id, 1)
        code = QualitativeCode(analysis_id=self.analysis.id, name="Raça", created_by_user_id=self.user.id)
        db.session.add(code)
        db.session.flush()
        for index in range(150):
            start, end = index * 5, index * 5 + 4
            excerpt = get_or_create_excerpt(self.analysis, self.document.id, 1, start, end,
                                           "raça", page["sha256"], self.user.id)
            coding = QualitativeCoding(analysis_id=self.analysis.id, excerpt_id=excerpt.id, code_id=code.id,
                created_by_user_id=self.user.id, origin="automatic_literal", source_query="raça",
                source_case_sensitive=False, source_start=start, source_end=end, source_page_hash=page["sha256"],
                contextual_rejection_enabled=True)
            db.session.add(coding)
            db.session.flush()
            if index >= 100:
                remove_coding(self.analysis, coding, self.user.id)
        db.session.commit()
        response = self.automatic(q="raça")
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json["search"]["total"], 300)
        self.assertEqual(response.json["summary"], {
            "created": 150, "existing": 100, "rejected": 50, "invalid": 0, "documents": 1})
        self.assertEqual(db.session.query(QualitativeCoding).count(), 250)
        self.assertEqual(db.session.query(QualitativeRejection).count(), 50)
        self.assertEqual(self.automatic(q="raça").json["summary"], {
            "created": 0, "existing": 250, "rejected": 50, "invalid": 0, "documents": 0})

    def test_regex_timeout_after_matches_does_not_return_or_save_partial_results(self):
        self.fixture_pages([["raça\n" * 250, "a" * 20000 + "!"]])
        query = r"raça|(a+)+$"
        with patch("platform_core.qualitative_search.SEARCH_TIMEOUT_SECONDS", 0.02):
            response = self.automatic(q=query, grep=True)
            normal = self.client.get(self.auto_url.replace("/codificar", "/buscar"),
                                     query_string={"q": query, "grep": "1"})
        for result in (response, normal):
            self.assertEqual(result.status_code, 400)
            self.assertIn("não foi concluída", result.json["error"])
            self.assertNotIn("results", result.json)
        self.assertEqual(db.session.query(QualitativeCode).count(), 0)
        self.assertEqual(db.session.query(QualitativeExcerpt).count(), 0)
        self.assertEqual(self.codings(), [])

    def test_invalid_long_or_whitespace_matches_are_reported_without_excerpts(self):
        self.fixture_pages([["a" * 10001 + "\n   \n"]])
        response = self.automatic(q="a+| +", grep=True)
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json["summary"]["invalid"], 2)
        self.assertEqual(response.json["summary"]["created"], 0)
        self.assertEqual(db.session.query(QualitativeExcerpt).count(), 0)

    def test_report_export_and_counts_follow_individual_deletion(self):
        self.automatic()
        report = f"/analise-qualitativa/bases/{self.analysis.id}/relatorio-codificacao"
        self.assertEqual(self.client.get(report).status_code, 200)
        workbook = load_workbook(io.BytesIO(self.client.get(report + ".xlsx").data))
        self.assertEqual(workbook.active.max_row, 3)
        response = self.remove(self.codings()[0])
        self.assertEqual(response.json["records"]["codes"][0]["excerpt_count"], 1)
        workbook = load_workbook(io.BytesIO(self.client.get(report + ".xlsx").data))
        self.assertEqual(workbook.active.max_row, 2)
        self.assertEqual(workbook.active.cell(2, 4).value, "Documento")

    def test_auth_csrf_tool_project_and_foreign_coding_guards(self):
        self.automatic()
        coding = self.codings()[0]
        self.assertEqual(self.remove(coding, token=False).status_code, 400)
        self.assertEqual(self.client.post(self.auto_url, json={"q": "Documento"}).status_code, 400)
        self.assertEqual(self.client.get(self.auto_url).status_code, 405)
        other = create_user("Outra pessoa", "automatic-other@example.org")
        other_project = Project(owner_user_id=other.id, name="Alheio", scrape_type="qualitative")
        db.session.add(other_project); db.session.commit()
        other_analysis = create_analysis(user_id=other.id, project_id=other_project.id,
            tool_id=QUALITATIVE_TOOL, tool_version="manual-v1", parameters={}, name="Alheio")
        foreign_code = QualitativeCode(analysis_id=other_analysis.id, name="Não acessar", created_by_user_id=other.id)
        db.session.add(foreign_code); db.session.commit()
        self.assertEqual(self.automatic(code_id=foreign_code.id).status_code, 400)
        self.assertEqual(self.automatic(name="Nome arbitrário").status_code, 400)
        self.assertEqual(self.remove(coding, analysis_id=other_analysis.id).status_code, 404)
        self.project.owner_user_id = other.id; db.session.commit()
        self.assertEqual(self.remove(coding).status_code, 404)
        self.assertEqual(self.automatic().status_code, 404)
        self.project.owner_user_id = self.user.id; db.session.commit()
        self.project.status = "archived"; db.session.commit()
        self.assertIn(self.automatic().status_code, (403, 404))
        self.assertIn(self.remove(coding).status_code, (403, 404))
        self.project.status = "active"
        db.session.get(UserToolOverride, (self.user.id, QUALITATIVE_TOOL)).decision = "deny"
        db.session.commit()
        self.assertIn(self.remove(coding).status_code, (403, 404))
        self.assertIn(self.automatic().status_code, (403, 404))
        self.assertEqual(db.session.query(QualitativeRejection).count(), 0)

    def test_other_project_and_user_do_not_inherit_rejections(self):
        self.automatic(contextual_rejection_enabled=True)
        self.remove(self.codings()[0])
        other = create_user("Outra pessoa", "separate@example.org")
        project = Project(owner_user_id=other.id, name="Pesquisa independente", scrape_type="qualitative")
        db.session.add(project); db.session.commit()
        own_ids = (self.analysis.id, self.document.id)
        self.analysis = create_analysis(user_id=other.id, project_id=project.id,
            tool_id=QUALITATIVE_TOOL, tool_version="manual-v1", parameters={}, name="Independente")
        self.document = AnalysisDocument(analysis_id=self.analysis.id, original_name="Mesmo texto.pdf", stored_name="mesmo.pdf")
        db.session.add_all([self.document, UserToolOverride(user_id=other.id, tool_id=QUALITATIVE_TOOL, decision="allow")])
        self.analysis.status = "concluida"; self.analysis.document_count = 1
        db.session.commit()
        old_manifest = load_qualitative_manifest(db.session.get(Analysis, own_ids[0]))
        new_manifest = {**old_manifest, "analysis_id": self.analysis.id, "documents": []}
        original = old_manifest["documents"][0]
        text_content = self.page["text"]
        relative = f"{self.document.id}/pages/000001.txt"
        page_file = qualitative_corpus_dir(self.analysis.id) / relative
        page_file.parent.mkdir(parents=True); page_file.write_bytes(text_content.encode("utf-8"))
        new_manifest["documents"] = [{**original, "document_id": self.document.id, "original_name": self.document.original_name,
            "stored_name": self.document.stored_name, "pages": [{**original["pages"][0], "file": relative, "layout_available": False}]}]
        qualitative_manifest_path(self.analysis.id).write_text(json.dumps(new_manifest), encoding="utf-8")
        self.client.post("/logout", data={"csrf_token": self.csrf})
        login(self.client, other.email); self.csrf = csrf_from(self.client.get("/"))
        self.assertEqual(self.automatic().json["summary"]["created"], 2)

    def test_rejections_removed_by_code_document_and_analysis_lifecycle(self):
        self.automatic(contextual_rejection_enabled=True); self.remove(self.codings()[0])
        code = db.session.query(QualitativeCode).one()
        self.assertEqual(self.client.delete(f"/analise-qualitativa/bases/{self.analysis.id}/codigos/{code.id}",
                                          headers={"X-CSRFToken": self.csrf}).status_code, 200)
        self.assertEqual(db.session.query(QualitativeRejection).count(), 0)
        self.automatic(contextual_rejection_enabled=True); self.remove(self.codings()[0])
        response = self.client.post(self.auto_url.replace("/codificar", "/excluir"),
            json={"confirm": True}, headers={"X-CSRFToken": self.csrf})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(db.session.query(QualitativeRejection).count(), 0)
        self.assertEqual(db.session.query(QualitativeCoding).count(), 0)
        delete_analysis(self.analysis)
        self.assertIsNone(db.session.get(Analysis, self.analysis.id))

    def test_corruption_and_upload_in_progress_do_not_create_partial_codings(self):
        lock = analysis_dir(self.analysis.id) / ".qualitative_upload.lock"
        lock.touch()
        self.assertEqual(self.automatic().status_code, 409)
        lock.unlink()
        manifest = load_qualitative_manifest(self.analysis)
        path = qualitative_corpus_dir(self.analysis.id) / manifest["documents"][0]["pages"][0]["file"]
        path.write_text("corrompido", encoding="utf-8")
        self.assertEqual(self.automatic().status_code, 409)
        self.assertEqual(db.session.query(QualitativeCode).count(), 0)
        self.assertEqual(self.codings(), [])

    def test_analysis_deletion_cleans_rejections_without_touching_other_analysis(self):
        self.automatic(contextual_rejection_enabled=True)
        self.remove(self.codings()[0])
        other = create_analysis(user_id=self.user.id, project_id=self.project.id,
            tool_id=QUALITATIVE_TOOL, tool_version="manual-v1", parameters={}, name="Preservar")
        other_code = QualitativeCode(analysis_id=other.id, name="Preservar", created_by_user_id=self.user.id)
        db.session.add(other_code)
        db.session.commit()
        self.assertEqual(db.session.query(QualitativeRejection).count(), 1)
        delete_analysis(self.analysis)
        self.assertEqual(db.session.query(QualitativeRejection).count(), 0)
        self.assertEqual(db.session.query(QualitativeCoding).count(), 0)
        self.assertIsNotNone(db.session.get(Analysis, other.id))
        self.assertIsNotNone(db.session.get(QualitativeCode, other_code.id))

    def test_failed_removal_rolls_back_coding_and_rejection_together(self):
        self.automatic(contextual_rejection_enabled=True)
        coding = self.codings()[0]
        coding_id = coding.id
        manifest = load_qualitative_manifest(self.analysis)
        path = qualitative_corpus_dir(self.analysis.id) / manifest["documents"][0]["pages"][0]["file"]
        path.write_bytes(b"corrompido")
        self.assertEqual(self.remove(coding).status_code, 409)
        self.assertIsNotNone(db.session.get(QualitativeCoding, coding_id))
        self.assertEqual(db.session.query(QualitativeRejection).count(), 0)


class AutomaticMigrationTests(unittest.TestCase):
    def test_upgrade_preserves_legacy_manual_and_assisted_data(self):
        from app import create_app
        from platform_core.services import seed_platform
        with tempfile.TemporaryDirectory() as directory:
            application = create_app({"TESTING": True, "SECRET_KEY": "migration-only",
                "SQLALCHEMY_DATABASE_URI": f"sqlite:///{Path(directory).as_posix()}/test.sqlite3", "PLATFORM_DATA_DIR": directory})
            with application.app_context(), ExitStack() as cleanup:
                cleanup.callback(db.engine.dispose)
                cleanup.callback(db.session.remove)
                migrations = str(Path(__file__).resolve().parents[1] / "migrations")
                from platform_helpers import seed_legacy_platform
                seed_legacy_platform(migrations, "a81d6e3f902b")
                # O modelo atual já inclui colunas posteriores a este schema legado.
                def legacy_code(name, analysis_id, user_id):
                    identifier = str(uuid4())
                    db.session.execute(text("INSERT INTO qualitative_codes "
                        "(id,analysis_id,name,normalized_name,description,active,created_by_user_id,created_at,updated_at) "
                        "VALUES (:id,:analysis,:name,:normalized,'',1,:user,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)"),
                        {"id": identifier, "analysis": analysis_id, "name": name,
                         "normalized": name, "user": user_id})
                    return identifier
                user = create_user()
                project = Project(owner_user_id=user.id, name="Anterior", scrape_type="qualitative")
                db.session.add(project); db.session.commit()
                analysis = create_analysis(user_id=user.id, project_id=project.id, tool_id=QUALITATIVE_TOOL,
                                           tool_version="manual-v1", parameters={}, name="Anterior")
                document = AnalysisDocument(analysis_id=analysis.id, original_name="original.pdf", stored_name="original.pdf")
                db.session.add(document); db.session.flush()
                excerpt = QualitativeExcerpt(analysis_id=analysis.id, document_id=document.id, page_number=1,
                    start_offset=0, end_offset=1, quoted_text="A", page_text_hash="a" * 64, created_by_user_id=user.id)
                db.session.add(excerpt); db.session.flush()
                ids = []
                for origin in ("manual", "assisted"):
                    code_id = legacy_code(origin, analysis.id, user.id)
                    identifier = str(uuid4()); ids.append(identifier)
                    db.session.execute(text("INSERT INTO qualitative_codings (id,analysis_id,excerpt_id,code_id,created_by_user_id,origin,created_at) "
                        "VALUES (:id,:analysis,:excerpt,:code,:user,:origin,CURRENT_TIMESTAMP)"),
                        {"id": identifier, "analysis": analysis.id, "excerpt": excerpt.id, "code": code_id, "user": user.id, "origin": origin})
                analysis_id, user_id, excerpt_id, document_id = analysis.id, user.id, excerpt.id, document.id
                db.session.commit(); db.session.remove()
                upgrade(directory=migrations, revision="b62f8a10d947")
                old_auto_ids = []
                for origin in ("automatic_literal", "automatic_regex"):
                    code_id = legacy_code(origin, analysis_id, user_id)
                    identifier = str(uuid4()); old_auto_ids.append(identifier)
                    db.session.execute(text("INSERT INTO qualitative_codings "
                        "(id,analysis_id,excerpt_id,code_id,created_by_user_id,origin,created_at,source_query,"
                        "source_case_sensitive,source_start,source_end,source_page_hash) VALUES "
                        "(:id,:analysis,:excerpt,:code,:user,:origin,CURRENT_TIMESTAMP,'A',0,0,1,:hash)"),
                        {"id": identifier, "analysis": analysis_id, "excerpt": excerpt_id, "code": code_id,
                         "user": user_id, "origin": origin, "hash": "a" * 64})
                rejection = QualitativeRejection(analysis_id=analysis_id, document_id=document_id,
                    code_id=code_id, page_number=1, start_offset=0, end_offset=1, page_text_hash="a" * 64,
                    origin="automatic_regex", query="Outro padrão", case_sensitive=True, created_by_user_id=user_id)
                db.session.add(rejection); db.session.flush()
                rejection_id = rejection.id
                db.session.commit(); db.session.remove()
                upgrade(directory=migrations, revision="head")
                for identifier, origin in zip(ids, ("manual", "assisted")):
                    coding = db.session.get(QualitativeCoding, identifier)
                    self.assertEqual(coding.origin, origin)
                    self.assertIsNone(coding.source_query)
                    self.assertFalse(coding.contextual_rejection_enabled)
                for identifier, origin in zip(old_auto_ids, ("automatic_literal", "automatic_regex")):
                    coding = db.session.get(QualitativeCoding, identifier)
                    self.assertEqual(coding.origin, origin)
                    self.assertTrue(coding.contextual_rejection_enabled)
                    self.assertEqual(coding.source_query, "A")
                    self.assertEqual(coding.source_page_hash, "a" * 64)
                # Default de banco novo também é OFF (sem depender do default ORM).
                # Remover o vínculo antigo libera o par sem mudar código/trecho.
                saved = db.session.get(QualitativeCoding, old_auto_ids[0])
                values = dict(analysis=saved.analysis_id, excerpt=saved.excerpt_id, code=saved.code_id,
                              user=saved.created_by_user_id, hash=saved.source_page_hash)
                db.session.delete(saved); db.session.flush()
                db.session.execute(text("INSERT INTO qualitative_codings "
                    "(id,analysis_id,excerpt_id,code_id,created_by_user_id,origin,created_at,source_query,"
                    "source_case_sensitive,source_start,source_end,source_page_hash) VALUES "
                    "('new-default',:analysis,:excerpt,:code,:user,'automatic_literal',CURRENT_TIMESTAMP,'A',0,0,1,:hash)"), values)
                db.session.commit()
                self.assertFalse(db.session.get(QualitativeCoding, "new-default").contextual_rejection_enabled)
                self.assertEqual(db.session.query(QualitativeRejection).count(), 1)
                self.assertEqual(db.session.get(QualitativeRejection, rejection_id).query, "Outro padrão")
                self.assertEqual(db.session.execute(text("PRAGMA foreign_key_check")).all(), [])
                db.session.remove(); db.engine.dispose()
