"""Análises persistem fora do cache e respeitam plano, projeto e proprietário."""

import io
import errno
import tempfile
import unittest
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from unittest.mock import patch

import pandas as pd
import pymupdf
from openpyxl import Workbook, load_workbook
from sqlalchemy import inspect, select, text
from flask_migrate import upgrade

import app as legacy
from historico_racial.routes import EXECUTOR_HR, RESULTADOS_HR
from platform_helpers import create_project, create_user, csrf_from, isolated_platform, login
from platform_core.analyses import analysis_dir, create_analysis, delete_analysis, load_result
from platform_core.extensions import db
from platform_core.models import Analysis, PlanTool, Project, Tool
from platform_core.services import can_use_tool, seed_platform
from platform_core.project_lifecycle import archive, delete_archived


def small_pdf():
    document = pymupdf.open()
    document.new_page().insert_text((50, 80), "Du Bois discutiu a populacao negra.")
    payload = document.tobytes()
    document.close()
    return payload


def fake_standard(pdfs, _terms, folder, version, _config, progress_callback=None):
    folder.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.active["A1"] = "resultado"
    workbook.save(folder / "resultado.xlsx")
    if progress_callback:
        progress_callback({"fase": "paginas", "etapa": "Analisando páginas…", "pagina_atual": 1, "paginas_total": 1})
    return {
        "versao": version, "erros": [], "arquivos": [{"nome": "resultado.xlsx", "rotulo": "Excel", "consolidado": True}],
        "quantidade_pdfs": len(pdfs), "quantidade_termos": 1,
        "ocorrencias": pd.DataFrame([{"ID livro": "teste", "Termo": "Du Bois"}]),
        "diagnosticos": pd.DataFrame(), "termos": [{"termo": "Du Bois", "categoria": "TERMO INFORMADO"}],
    }


class PersistentAnalysisTests(unittest.TestCase):
    def setUp(self):
        self.scope = isolated_platform()
        self.app = self.scope.__enter__()
        self.client = self.app.test_client()

    def tearDown(self):
        self.scope.__exit__(None, None, None)

    def _standard(self, *, duplicate=None):
        token = csrf_from(self.client.get("/raspagem-livre"))
        data = {"csrf_token": token, "versao": "v1", "termos": "Du Bois"}
        if duplicate:
            data["duplicate_id"] = duplicate
        else:
            data["pdfs"] = (io.BytesIO(small_pdf()), "fonte.pdf")
        with patch.object(legacy, "executar_analises", side_effect=fake_standard), patch.object(legacy, "criar_dashboard", return_value={}):
            response = self.client.post("/", data=data, content_type="multipart/form-data",
                                        headers={"X-Requested-With": "XMLHttpRequest"})
            self.assertEqual(response.status_code, 202, response.get_data(as_text=True))
            legacy.EXECUTOR_ANALISES.submit(lambda: None).result(timeout=30)
        return response.json["job_id"]

    def _name_base(self, identifier, name="Base de teste"):
        pending = self.client.get(f"/analises/{identifier}")
        self.assertEqual(pending.status_code, 302)
        self.assertIn("/nomear", pending.headers["Location"])
        token = csrf_from(self.client.get(pending.headers["Location"]))
        saved = self.client.post(pending.headers["Location"], data={
            "csrf_token": token, "name": name,
        })
        self.assertEqual(saved.status_code, 302)
        self.assertTrue(db.session.get(Analysis, identifier).name_confirmed)

    def test_delete_moves_upload_across_filesystems(self):
        user = create_user()
        analysis = create_analysis(user_id=user.id, project_id=None,
                                   tool_id="pdf_scraper", tool_version="v1", parameters={})
        with tempfile.TemporaryDirectory() as application_root, patch.object(self.app, "root_path", application_root):
            upload = Path(application_root) / "uploads" / analysis.id
            upload.mkdir(parents=True)
            (upload / "documento.pdf").write_bytes(b"PDF de teste")
            original_rename = Path.rename

            def rename_with_cross_device(source, destination):
                if source == upload:
                    raise OSError(errno.EXDEV, "filesystems diferentes")
                return original_rename(source, destination)

            with patch.object(Path, "rename", rename_with_cross_device):
                delete_analysis(analysis)
            self.assertIsNone(db.session.get(Analysis, analysis.id))
            self.assertFalse(upload.exists())
            self.assertFalse(analysis_dir(analysis.id).exists())

    def test_delete_rolls_back_cross_device_moves_after_later_failure(self):
        user = create_user()
        analysis = create_analysis(user_id=user.id, project_id=None,
                                   tool_id="pdf_scraper", tool_version="v1", parameters={})
        with tempfile.TemporaryDirectory() as application_root, patch.object(self.app, "root_path", application_root):
            upload = Path(application_root) / "uploads" / analysis.id
            output = Path(application_root) / "outputs" / analysis.id
            upload.mkdir(parents=True)
            output.mkdir(parents=True)
            (upload / "documento.pdf").write_bytes(b"PDF de teste")
            (output / "resultado.xlsx").write_bytes(b"XLSX de teste")
            quarantine_upload = analysis_dir(analysis.id).with_name(analysis.id + ".removing") / "upload"
            original_rename = Path.rename

            def rename_with_failure(source, destination):
                if source in (upload, quarantine_upload):
                    raise OSError(errno.EXDEV, "filesystems diferentes")
                if source == output:
                    raise OSError(errno.EACCES, "falha de I/O posterior")
                return original_rename(source, destination)

            with patch.object(Path, "rename", rename_with_failure):
                with self.assertRaises(OSError):
                    delete_analysis(analysis)
            self.assertIsNotNone(db.session.get(Analysis, analysis.id))
            self.assertTrue((upload / "documento.pdf").is_file())
            self.assertTrue((output / "resultado.xlsx").is_file())
            self.assertTrue(analysis_dir(analysis.id).exists())
            self.assertFalse(quarantine_upload.parent.exists())

    def test_standard_persists_reopens_duplicates_moves_and_deletes(self):
        user = create_user()
        login(self.client)
        token = csrf_from(self.client.get("/projetos/livres/novo"))
        project_response = self.client.post("/projetos/livres/novo", data={
            "csrf_token": token, "name": "Pesquisa livre", "description": "",
        })
        self.assertEqual(project_response.status_code, 302)
        project_id = parse_qs(urlparse(project_response.headers["Location"]).query)["project_id"][0]
        first = self._standard()
        record = db.session.get(Analysis, first)
        self.assertEqual((record.status, record.document_count, record.result_count), ("concluida", 1, 1))
        self.assertTrue((analysis_dir(first) / "resultado.xlsx").is_file())
        with legacy.ANALISES_LOCK:
            legacy.ANALISES.pop(first, None)
        self._name_base(first, "Racismo")
        self.assertEqual(self.client.get(f"/analises/{first}").status_code, 200)
        self.assertEqual(self.client.get(f"/analises/{first}/excel/resultado.xlsx").status_code, 200)
        self.assertIn(first, self.client.get("/analises").get_data(as_text=True))
        self.client.post("/logout", data={"csrf_token": csrf_from(self.client.get("/analises"))})
        login(self.client)
        self.assertEqual(self.client.get(f"/analises/{first}").status_code, 200)
        self.assertEqual(self.client.get(f"/analises/{first}/duplicar").status_code, 302)
        copy = self._standard(duplicate=first)
        self.assertNotEqual(first, copy)
        self.assertEqual(db.session.get(Analysis, copy).result_count, 1)
        self._name_base(copy, "Nordeste")
        token = csrf_from(self.client.get(f"/analises/{first}"))
        moved = self.client.post(f"/analises/{first}/mover", data={"csrf_token": token, "project_id": project_id})
        self.assertEqual(moved.status_code, 302)
        self.assertEqual(db.session.get(Analysis, first).project_id, project_id)
        self.assertNotIn(first, self.client.get("/analises").get_data(as_text=True))
        self.assertIn(first, self.client.get(f"/analises/projeto/{project_id}").get_data(as_text=True))
        deleted = self.client.post(f"/analises/{first}/excluir", data={"csrf_token": token, "confirm": "yes"})
        self.assertEqual(deleted.status_code, 302)
        self.assertIsNone(db.session.get(Analysis, first))
        self.assertFalse(analysis_dir(first).exists())
        self.assertFalse((legacy.UPLOAD_DIR / first).exists())
        self.assertFalse((legacy.OUTPUT_DIR / first).exists())
        self.assertIsNotNone(db.session.get(Project, project_id))

    def test_real_v1_output_serializes_without_changing_analysis(self):
        create_user()
        login(self.client)
        token = csrf_from(self.client.get("/raspagem-livre"))
        response = self.client.post("/raspagem-livre", data={
            "csrf_token": token, "versao": "v1", "termos": "Du Bois",
            "pdfs": (io.BytesIO(small_pdf()), "real.pdf"),
        }, content_type="multipart/form-data", headers={"X-Requested-With": "XMLHttpRequest"})
        self.assertEqual(response.status_code, 202, response.get_data(as_text=True))
        legacy.EXECUTOR_ANALISES.submit(lambda: None).result(timeout=30)
        identifier = response.json["job_id"]
        record = db.session.get(Analysis, identifier)
        self.assertEqual(record.status, "concluida", record.error_message)
        self.assertTrue((analysis_dir(identifier) / "result.json").is_file())
        self.assertEqual(load_result(record)["ocorrencias"],
                         legacy.ANALISES[identifier]["ocorrencias"].to_dict("records"))
        self._name_base(identifier)
        self.assertEqual(self.client.get(f"/analises/{identifier}").status_code, 200)

    def test_duplicate_v3_shows_original_threshold_in_slider_and_label(self):
        user = create_user()
        login(self.client)
        analysis = create_analysis(
            user_id=user.id, project_id=None, tool_id="pdf_scraper", tool_version="v3",
            parameters={"termos": [{"termo": "Du Bois"}], "configuracoes_v3": {
                "incluir_lexical": True, "incluir_semantica": True, "limiar_semantico": 0.75,
            }},
        )
        analysis.status = "concluida"
        db.session.commit()
        page = self.client.get(f"/raspagem-livre?duplicate={analysis.id}")
        self.assertEqual(page.status_code, 200)
        html = page.get_data(as_text=True)
        self.assertIn('id="valor-limiar" for="limiar-semantico">0,75', html)
        self.assertIn('id="limiar-semantico" name="limiar_semantico"', html)
        self.assertIn('value="0.75"', html)

    def test_project_execution_persists_and_is_private(self):
        owner = create_user()
        login(self.client)
        project_id = create_project(self.client)
        token = csrf_from(self.client.get(f"/analise-documental/projetos/{project_id}"))
        response = self.client.post(f"/analise-documental/projetos/{project_id}/analisar", data={
            "csrf_token": token, "pdfs": (io.BytesIO(small_pdf()), "fonte.pdf"),
        }, content_type="multipart/form-data", headers={"X-Requested-With": "XMLHttpRequest"})
        self.assertEqual(response.status_code, 202)
        EXECUTOR_HR.submit(lambda: None).result(timeout=30)
        identifier = response.json["job_id"]
        self.assertEqual(db.session.get(Analysis, identifier).status, "concluida")
        self.assertTrue((analysis_dir(identifier) / "resultado.xlsx").is_file())
        from historico_racial.exporter import gerar_xlsx
        saved = load_workbook(analysis_dir(identifier) / "resultado.xlsx", read_only=True)
        legacy_export = load_workbook(gerar_xlsx(RESULTADOS_HR[identifier]), read_only=True)
        self.assertEqual(saved.sheetnames, legacy_export.sheetnames)
        for name in saved.sheetnames:
            self.assertEqual(list(saved[name].values), list(legacy_export[name].values))
        saved.close()
        legacy_export.close()
        RESULTADOS_HR.pop(identifier, None)
        self._name_base(identifier)
        page = self.client.get(f"/analises/{identifier}")
        self.assertEqual(page.status_code, 200, page.get_data(as_text=True)[:500])
        self.assertIn("Gráficos", page.get_data(as_text=True))
        self.assertIn(identifier, self.client.get(f"/analises/projeto/{project_id}").get_data(as_text=True))
        repeat_page = self.client.get(f"/analises/{identifier}/duplicar")
        self.assertEqual(repeat_page.status_code, 302)
        self.assertIn("duplicate=", repeat_page.headers["Location"])
        repeat = self.client.post(f"/analise-documental/projetos/{project_id}/analisar", data={
            "csrf_token": token, "duplicate_id": identifier, "metodo_analise": "lexical",
        }, headers={"X-Requested-With": "XMLHttpRequest"})
        self.assertEqual(repeat.status_code, 202, repeat.get_data(as_text=True))
        EXECUTOR_HR.submit(lambda: None).result(timeout=30)
        self.assertNotEqual(identifier, repeat.json["job_id"])
        self.assertEqual(db.session.get(Analysis, repeat.json["job_id"]).status, "concluida")
        archive(db.session.get(Project, project_id), owner)
        delete_archived([db.session.get(Project, project_id)], owner, "deletar")
        self.assertIsNone(db.session.get(Project, project_id))
        self.assertIsNone(db.session.get(Analysis, identifier))
        self.assertFalse(analysis_dir(identifier).exists())
        self.client.post("/logout", data={"csrf_token": token})
        create_user("Outro", "outro@example.org")
        login(self.client, "outro@example.org")
        self.assertEqual(self.client.get(f"/analises/{identifier}").status_code, 404)
        self.assertEqual(self.client.get(f"/analises/{identifier}/excel/resultado.xlsx").status_code, 404)

    def test_access_matrix_is_persistent_and_enforced(self):
        student = create_user()
        admin = create_user("Admin", "admin@example.org", "admin", "institutional")
        login(self.client, "admin@example.org")
        token = csrf_from(self.client.get("/admin/definir-acessos"))
        current = [f"{row.plan_id}|{row.tool_id}" for row in db.session.scalars(select(PlanTool))]
        current.remove("student|pdf_scraper")
        response = self.client.post("/admin/definir-acessos", data={"csrf_token": token, "access": current})
        self.assertEqual(response.status_code, 302)
        seed_platform()
        self.assertFalse(can_use_tool(student, "pdf_scraper"))
        self.client.post("/logout", data={"csrf_token": token})
        login(self.client)
        self.assertEqual(self.client.get("/").headers["Location"], "/perfil")
        self.assertEqual(self.client.get("/raspagem-livre").status_code, 403)
        self.assertEqual(self.client.get("/analises").status_code, 403)
        self.assertEqual(self.client.get("/projetos").status_code, 200)
        self.client.post("/logout", data={"csrf_token": csrf_from(self.client.get("/projetos"))})
        login(self.client, "admin@example.org")
        token = csrf_from(self.client.get("/admin/definir-acessos"))
        restored = current + ["student|pdf_scraper"]
        self.assertEqual(self.client.post("/admin/definir-acessos", data={
            "csrf_token": token, "access": restored,
        }).status_code, 302)
        self.client.post("/logout", data={"csrf_token": token})
        login(self.client)
        self.assertEqual(self.client.get("/").headers["Location"], "/perfil")
        self.assertEqual(self.client.get("/raspagem-livre").status_code, 200)

    def test_new_tool_route_is_guarded_by_dynamic_plan_matrix(self):
        self.app.add_url_rule("/ferramenta-nova", "future_tool", lambda: "ok")
        student = create_user()
        db.session.add(Tool(id="future_tool", name="Ferramenta nova", route="/ferramenta-nova", active=True))
        db.session.commit()
        login(self.client)
        self.assertEqual(self.client.get("/ferramenta-nova").status_code, 403)
        db.session.add(PlanTool(plan_id="student", tool_id="future_tool"))
        db.session.commit()
        self.assertEqual(self.client.get("/ferramenta-nova").status_code, 200)

    def test_project_owner_can_open_analysis_started_by_admin(self):
        owner = create_user()
        login(self.client)
        project_id = create_project(self.client)
        self.client.post("/logout", data={"csrf_token": csrf_from(self.client.get("/projetos"))})
        create_user("Admin", "admin@example.org", "admin", "institutional")
        login(self.client, "admin@example.org")
        token = csrf_from(self.client.get(f"/analise-documental/projetos/{project_id}"))
        response = self.client.post(f"/analise-documental/projetos/{project_id}/analisar", data={
            "csrf_token": token, "pdfs": (io.BytesIO(small_pdf()), "admin.pdf"),
        }, content_type="multipart/form-data", headers={"X-Requested-With": "XMLHttpRequest"})
        self.assertEqual(response.status_code, 202)
        EXECUTOR_HR.submit(lambda: None).result(timeout=30)
        identifier = response.json["job_id"]
        self.assertNotEqual(db.session.get(Analysis, identifier).user_id, owner.id)
        self.client.post("/logout", data={"csrf_token": token})
        login(self.client)
        self._name_base(identifier, "Base criada pelo administrador")
        self.assertEqual(self.client.get(f"/analises/{identifier}").status_code, 200)
        self.assertIn(identifier, self.client.get(f"/analises/projeto/{project_id}").get_data(as_text=True))

    def test_navigation_has_indirect_admin_projects_and_access_matrix(self):
        owner = create_user()
        login(self.client)
        project_id = create_project(self.client)
        self.client.post("/logout", data={"csrf_token": csrf_from(self.client.get("/projetos"))})
        create_user("Admin", "admin@example.org", "admin", "institutional")
        login(self.client, "admin@example.org")
        page = self.client.get("/admin/definir-acessos")
        self.assertEqual(page.status_code, 200)
        self.assertIn("Definir acessos por plano", page.get_data(as_text=True))
        self.assertNotIn("href=\"/admin/projetos\"", page.get_data(as_text=True))
        self.assertIn("href=\"/admin/definir-acessos\"", page.get_data(as_text=True))
        user_page = self.client.get(f"/admin/usuarios/{owner.id}")
        self.assertIn(f"/admin/projetos/{project_id}", user_page.get_data(as_text=True))
        response = self.client.post(f"/admin/projetos/{project_id}/estado", data={
            "csrf_token": csrf_from(user_page), "confirm": "yes", "status": "blocked",
        })
        self.assertIn(f"/admin/usuarios/{owner.id}", response.headers["Location"])


class AnalysisMigrationTests(unittest.TestCase):
    def test_upgrade_adds_analysis_tables_on_existing_schema(self):
        from app import create_app
        with tempfile.TemporaryDirectory(prefix="pesquisapdf-analysis-migration-") as root:
            db_path = Path(root) / "migration.sqlite3"
            application = create_app({
                "TESTING": True, "SQLALCHEMY_DATABASE_URI": f"sqlite:///{db_path.as_posix()}",
                "PLATFORM_DATA_DIR": root, "SECRET_KEY": "migration-test-only",
            })
            migrations = str(Path(__file__).resolve().parent.parent / "migrations")
            with application.app_context():
                upgrade(directory=migrations, revision="f7e2b9c4d601")
                upgrade(directory=migrations, revision="d83f4b6a19c2")
                with db.engine.begin() as connection:
                    connection.execute(text("""INSERT INTO users
                        (id, name, email, password_hash, must_change_password, role, status,
                         student_verification_status, created_at, updated_at)
                        VALUES ('old-user', 'Antigo', 'antigo@example.org', 'hash', 0, 'user', 'active',
                                'not_required', '2026-01-01', '2026-01-01')"""))
                    connection.execute(text("""INSERT INTO tools
                        (id, name, description, route, active)
                        VALUES ('document_analysis', 'Ferramenta antiga', '', '/analise-documental', 1)"""))
                    connection.execute(text("""INSERT INTO projects
                        (id, owner_user_id, name, description, status, created_at, updated_at)
                        VALUES ('old-project', 'old-user', 'Projeto existente', '', 'active',
                                '2026-01-01', '2026-01-01')"""))
                    for identifier, name in (("old-auto", "Análise 24/09/2026 00:34"),
                                             ("old-custom", "Racismo no PNLD")):
                        connection.execute(text("""INSERT INTO analyses
                            (id, user_id, project_id, name, source_type, tool_id, tool_version,
                             status, created_at, document_count, result_count, parameters_json,
                             excel_files_json)
                            VALUES (:id, 'old-user', 'old-project', :name, 'project',
                                    'document_analysis', 'lexical', 'concluida', '2026-01-01',
                                    1, 2, '{}', '[]')"""), {"id": identifier, "name": name})
                upgrade(directory=migrations, revision="head")
                tables = set(inspect(db.engine).get_table_names())
                self.assertTrue({"users", "projects", "plan_tools", "analyses", "analysis_documents"} <= tables)
                with db.engine.connect() as connection:
                    self.assertEqual(connection.scalar(text(
                        "SELECT scrape_type FROM projects WHERE id = 'old-project'")), "systematic")
                    names = dict(connection.execute(text(
                        "SELECT id, name_confirmed FROM analyses WHERE id IN ('old-auto', 'old-custom')")).all())
                    self.assertEqual(names, {"old-auto": 0, "old-custom": 1})
                db.session.remove()
                db.engine.dispose()


if __name__ == "__main__":
    unittest.main()
