"""Testes controlados da V3 sem depender de PDFs enviados pelo usuário."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import app as aplicacao_web
from analyzer import v3
from analyzer import v1, v2
from analyzer.common import criar_dashboard, executar_analises, montar_termos


class ModeloControlado:
    """Embeddings determinísticos para testar a lógica sem rede/modelo externo."""

    def encode(self, textos, **_kwargs):
        vetores = []
        for texto in textos:
            normalizado = v3.normalizar(str(texto))
            if normalizado == "seca" or "climatic" in normalizado:
                vetores.append([1.0, 0.0])
            elif normalizado == "rio grande do norte" or "rio grande do norte" in normalizado:
                vetores.append([0.0, 1.0])
            else:
                vetores.append([0.0, 0.0])
        return np.asarray(vetores, dtype=float)


class ModeloQueRegistraConsultas:
    """Modelo mínimo para confirmar quais consultas entram na etapa semântica."""

    def __init__(self):
        self.chamadas: list[list[str]] = []

    def encode(self, textos, **_kwargs):
        textos = [str(texto) for texto in textos]
        self.chamadas.append(textos)
        return np.tile(np.asarray([[1.0, 0.0]]), (len(textos), 1))


def blocos_controlados():
    trecho_semantico = " ".join(
        [
            "A comunidade rural enfrentava desafios climaticos, perda de renda, "
            "migração e dificuldades para garantir condições de vida dignas"
        ]
        * 42
    ) + "."
    trecho_lexical = " ".join(
        [
            "No Rio Grande do Norte trabalhadores rurais discutiam cultura, "
            "território, desigualdade regional e políticas públicas"
        ]
        * 42
    ) + "."
    campos = {"unidade": "", "capitulo": "", "secao": "", "subsecao": "", "metodo": "texto", "tipo_ocorrencia": "parágrafo"}
    return [
        {**campos, "texto": trecho_semantico, "pagina": 1, "pagina_pdf": 1},
        {**campos, "texto": trecho_lexical, "pagina": 2, "pagina_pdf": 2},
    ]


class TesteV3(unittest.TestCase):
    def setUp(self):
        self.termos = [
            {"termo": "Rio Grande do Norte", "categoria": "TERMO INFORMADO"},
            {"termo": "seca", "categoria": "TERMO INFORMADO"},
        ]
        self.configuracao = {
            "incluir_lexical": True,
            "incluir_semantica": True,
            "limiar_semantico": 0.70,
        }

    def _extrair_pdf(self, _caminho):
        return blocos_controlados(), 2, 0, None

    @staticmethod
    def _parametros(workbook):
        aba = workbook["Parâmetros da análise"]
        return {
            linha[0]: linha[1]
            for linha in aba.iter_rows(min_row=2, values_only=True)
            if linha[0]
        }

    @patch("analyzer.v3.carregar_modelo_semantico", return_value=ModeloControlado())
    @patch("analyzer.v3.v1.extrair_pdf")
    def test_busca_hibrida_diferencia_evidencias(self, extrair_pdf, _modelo):
        extrair_pdf.side_effect = self._extrair_pdf
        registros, diagnostico = v3.analisar_pdf(
            Path("livro_um.pdf"), self.termos, self.configuracao
        )

        tipos = {registro["Consulta"]: registro["Tipo de correspondência"] for registro in registros}
        self.assertEqual(tipos["Rio Grande do Norte"], "Lexical + semântica")
        self.assertEqual(tipos["seca"], "Semântica")
        self.assertEqual(diagnostico["paginas_processadas_com_OCR"], 0)
        self.assertTrue(all("Categoria do termo" not in registro for registro in registros))

    @patch("analyzer.v3.carregar_modelo_semantico", return_value=ModeloControlado())
    @patch("analyzer.v3.v1.extrair_pdf")
    def test_excel_v3_e_consolidado(self, extrair_pdf, _modelo):
        extrair_pdf.side_effect = self._extrair_pdf
        with TemporaryDirectory() as diretorio:
            saida = Path(diretorio)
            resultado = executar_analises(
                [Path("livro_um.pdf"), Path("livro_dois.pdf")],
                self.termos,
                saida,
                "v3",
                self.configuracao,
            )
            nomes = {arquivo["nome"] for arquivo in resultado["arquivos"]}
            self.assertEqual(
                nomes,
                {"livro_um_v3.xlsx", "livro_dois_v3.xlsx", "resultado_todos_pdfs_v3.xlsx"},
            )
            workbook = load_workbook(saida / "livro_um_v3.xlsx")
            self.assertEqual(
                workbook.sheetnames,
                [
                    "Resultados",
                    "Resumo lexical",
                    "Resumo semântico",
                    "Resumo por consulta",
                    "Termos pesquisados",
                    "Diagnostico OCR",
                    "Parâmetros da análise",
                ],
            )
            cabecalhos = [celula.value for celula in workbook["Resultados"][1]]
            self.assertEqual(cabecalhos, v3.COLUNAS_RESULTADOS)
            self.assertNotIn("Categoria do termo", cabecalhos)
            for nome_aba in workbook.sheetnames:
                if nome_aba != "Resultados":
                    self.assertNotIn(
                        "ID resultado",
                        [celula.value for celula in workbook[nome_aba][1]],
                    )
            ids = [
                workbook["Resultados"].cell(linha, 1).value
                for linha in range(2, workbook["Resultados"].max_row + 1)
            ]
            self.assertEqual(len(ids), len(set(ids)))
            self.assertEqual(
                ids,
                [f"V3-001-{numero:06d}" for numero in range(1, len(ids) + 1)],
            )
            self.assertTrue(
                all(
                    workbook["Resultados"].cell(linha, 1).data_type == "s"
                    and workbook["Resultados"].cell(linha, 1).number_format == "@"
                    for linha in range(2, workbook["Resultados"].max_row + 1)
                )
            )
            coluna_validacao = cabecalhos.index("Validação manual") + 1
            coluna_observacoes = cabecalhos.index("Observações do pesquisador") + 1
            for linha in range(2, workbook["Resultados"].max_row + 1):
                self.assertIn(workbook["Resultados"].cell(linha, coluna_validacao).value, (None, ""))
                self.assertIn(workbook["Resultados"].cell(linha, coluna_observacoes).value, (None, ""))

            validacoes = workbook["Resultados"].data_validations.dataValidation
            self.assertTrue(any(item.formula1 == '"incluir,excluir,revisar"' for item in validacoes))
            self.assertEqual(workbook["Resultados"].freeze_panes, "A2")
            self.assertTrue(workbook["Resultados"].auto_filter.ref.startswith("A1:"))
            coluna_trecho = cabecalhos.index("Trecho da ocorrência") + 1
            self.assertGreater(
                workbook["Resultados"].column_dimensions[
                    workbook["Resultados"].cell(1, coluna_trecho).column_letter
                ].width,
                60,
            )
            parametros = self._parametros(workbook)
            self.assertEqual(parametros["Método de varredura"], "V3: Busca híbrida (lexical + semântica)")
            self.assertEqual(parametros["Busca lexical/morfológica"], "Sim")
            self.assertEqual(parametros["Busca semântica"], "Sim")
            self.assertEqual(parametros["Limiar semântico"], "0.70")
            self.assertEqual(
                parametros["Modelo de embeddings"],
                "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2",
            )
            self.assertIn("candidatos à análise", parametros["Definição operacional"])
            self.assertIn("não representam ocorrências textuais", parametros["Observações metodológicas"])
            self.assertEqual(workbook["Parâmetros da análise"].freeze_panes, "A2")
            self.assertGreater(workbook["Parâmetros da análise"].column_dimensions["B"].width, 90)

            consolidado = load_workbook(saida / "resultado_todos_pdfs_v3.xlsx")
            ids_consolidados = [
                consolidado["Resultados"].cell(linha, 1).value
                for linha in range(2, consolidado["Resultados"].max_row + 1)
            ]
            self.assertEqual(len(ids_consolidados), len(set(ids_consolidados)))
            self.assertTrue(all(identificador.startswith("V3-") for identificador in ids_consolidados))
            self.assertEqual(
                [identificador.split("-")[1] for identificador in ids_consolidados],
                ["001", "001", "002", "002"],
            )
            self.assertEqual(
                self._parametros(consolidado)["Abrangência da análise"],
                "Análise consolidada de múltiplos PDFs",
            )

            dashboard = criar_dashboard(resultado)
            self.assertEqual(dashboard["indicadores"]["livros"], 2)
            self.assertIn("sankey", dashboard)

    def test_v1_e_v2_preservam_planilhas_e_recebem_parametros_de_auditoria(self):
        diagnosticos = pd.DataFrame(
            [
                {
                    "arquivo": "livro_teste.pdf",
                    "paginas_com_texto_extraivel": 1,
                    "paginas_processadas_com_OCR": 0,
                    "idioma_OCR": "não utilizado",
                }
            ]
        )
        termos = pd.DataFrame(
            [{"Termo": "Rio de Janeiro", "Categoria do termo": "TERMO INFORMADO"}]
        )
        cabecalhos_v2 = [
            "ID resultado",
            "ID livro",
            "Termo",
            "Unidade",
            "Capítulo",
            "Seção",
            "Subseção",
            "Página",
            "Tipo de ocorrência",
            "Contexto sociológico",
            "Descrição",
        ]
        with TemporaryDirectory() as diretorio:
            diretorio_saida = Path(diretorio)
            for versao, analisador, metodo in (
                ("v1", v1, "V1: Busca lexical detalhada"),
                ("v2", v2, "V2: Busca lexical (planilha simplificada)"),
            ):
                arquivo = diretorio_saida / f"{versao}.xlsx"
                analisador.salvar_excel_completo(
                    arquivo, analisador.dataframe_vazio(), termos, diagnosticos
                )
                workbook = load_workbook(arquivo)
                self.assertIn("Parâmetros da análise", workbook.sheetnames)
                self.assertEqual(self._parametros(workbook)["Método de varredura"], metodo)
                self.assertEqual(self._parametros(workbook)["Busca lexical/morfológica"], "Sim")
                self.assertEqual(self._parametros(workbook)["Busca semântica"], "Não")
                cabecalhos = [celula.value for celula in workbook["Ocorrencias"][1]]
                if versao == "v1":
                    self.assertIn("Validação manual", cabecalhos)
                    self.assertEqual(cabecalhos.count("Validação manual"), 1)
                    self.assertEqual(cabecalhos[-1], "Observações do pesquisador")
                    self.assertTrue(
                        any(
                            item.formula1 == '"incluir,excluir,revisar"'
                            for item in workbook["Ocorrencias"].data_validations.dataValidation
                        )
                    )
                else:
                    self.assertEqual(
                        cabecalhos, cabecalhos_v2
                    )

    @staticmethod
    def _ocorrencias_lexicais_para_exportacao(caminho: Path):
        base = {
            "ID livro": caminho.stem,
            "Categoria do termo": "TERMO INFORMADO",
            "Parágrafo anterior": "Contexto anterior.",
            "Parágrafo posterior": "Contexto posterior.",
            "Unidade": "",
            "Capítulo": "",
            "Seção": "",
            "Subseção": "",
            "Página": 1,
            "Tipo da ocorrência": "Lexical",
            "Contexto sociológico da ocorrência": "Cultura e identidade",
            "Descrição da ocorrência": "Descrição automática preservada.",
            "Validação manual": "revisar",
            "_quantidade_no_registro": 1,
            "_metodo": "texto",
        }
        return [
            {
                **base,
                "Termo": "São Paulo",
                "Parágrafo do termo": "São Paulo aparece neste primeiro bloco.",
            },
            {
                **base,
                "Termo": "Rio de Janeiro",
                "Parágrafo do termo": "Rio de Janeiro aparece neste segundo bloco.",
            },
        ]

    def test_ids_v1_e_v2_sao_textuais_e_unicos_no_consolidado(self):
        termos = [
            {"termo": "São Paulo", "categoria": "TERMO INFORMADO"},
            {"termo": "Rio de Janeiro", "categoria": "TERMO INFORMADO"},
        ]
        cabecalhos_v1 = [
            "ID resultado",
            "ID livro",
            "Termo",
            "Parágrafo anterior",
            "Parágrafo do termo",
            "Parágrafo posterior",
            "Unidade",
            "Capítulo",
            "Seção",
            "Subseção",
            "Página",
            "Tipo da ocorrência",
            "Contexto sociológico da ocorrência",
            "Descrição da ocorrência",
            "Validação manual",
            "Observações do pesquisador",
        ]

        def analisar_controlado(caminho, _termos):
            return (
                self._ocorrencias_lexicais_para_exportacao(caminho),
                {
                    "arquivo": caminho.name,
                    "paginas_com_texto_extraivel": 1,
                    "paginas_processadas_com_OCR": 0,
                    "idioma_OCR": "não utilizado",
                },
            )

        with TemporaryDirectory() as diretorio:
            pasta_saida = Path(diretorio)
            for versao, analisador in (("v1", v1), ("v2", v2)):
                with patch.object(analisador, "analisar_pdf", side_effect=analisar_controlado):
                    resultado = executar_analises(
                        [Path("primeiro.pdf"), Path("segundo.pdf")],
                        termos,
                        pasta_saida / versao,
                        versao,
                    )

                nome_primeiro = (
                    f"resultado_primeiro_{versao}.xlsx"
                )
                individual = load_workbook(pasta_saida / versao / nome_primeiro)
                aba = individual["Ocorrencias"]
                cabecalhos = [celula.value for celula in aba[1]]
                self.assertEqual(cabecalhos[0], "ID resultado")
                for nome_aba in individual.sheetnames:
                    if nome_aba != "Ocorrencias":
                        self.assertNotIn(
                            "ID resultado",
                            [celula.value for celula in individual[nome_aba][1]],
                        )
                ids_individuais = [aba.cell(linha, 1).value for linha in range(2, aba.max_row + 1)]
                self.assertEqual(
                    ids_individuais,
                    [f"{versao.upper()}-001-000001", f"{versao.upper()}-001-000002"],
                )
                self.assertTrue(all(aba.cell(linha, 1).data_type == "s" for linha in range(2, aba.max_row + 1)))
                self.assertTrue(all(aba.cell(linha, 1).number_format == "@" for linha in range(2, aba.max_row + 1)))

                consolidado = load_workbook(
                    pasta_saida / versao / f"resultado_todos_pdfs_{versao}.xlsx"
                )["Ocorrencias"]
                ids_consolidados = [
                    consolidado.cell(linha, 1).value
                    for linha in range(2, consolidado.max_row + 1)
                ]
                self.assertEqual(len(ids_consolidados), len(set(ids_consolidados)))
                self.assertEqual(
                    ids_consolidados,
                    [
                        f"{versao.upper()}-001-000001",
                        f"{versao.upper()}-001-000002",
                        f"{versao.upper()}-002-000001",
                        f"{versao.upper()}-002-000002",
                    ],
                )
                self.assertEqual(criar_dashboard(resultado)["indicadores"]["ocorrencias"], 4)

                if versao == "v2":
                    campos_manuais = ["Unidade", "Capítulo", "Seção", "Subseção", "Tipo de ocorrência"]
                    for campo in campos_manuais:
                        coluna = cabecalhos.index(campo) + 1
                        self.assertTrue(
                            all(aba.cell(linha, coluna).value in (None, "") for linha in range(2, aba.max_row + 1))
                        )
                else:
                    self.assertEqual(cabecalhos, cabecalhos_v1)

    def test_v3_lexical_exporta_blocos_originais_adjacentes(self):
        anterior = "O crescimento urbano modificou profundamente a organização das cidades."
        central = (
            "A população negra encontra-se desproporcionalmente concentrada em áreas "
            "periféricas com menor acesso a serviços públicos."
        )
        posterior = "As desigualdades territoriais também se expressam no acesso ao saneamento."
        campos = {
            "unidade": "",
            "capitulo": "",
            "secao": "",
            "subsecao": "",
            "metodo": "texto",
            "tipo_ocorrencia": "parágrafo",
        }
        blocos = [
            {**campos, "texto": anterior, "pagina": 1, "pagina_pdf": 1},
            {**campos, "texto": central, "pagina": 1, "pagina_pdf": 1},
            {**campos, "texto": posterior, "pagina": 1, "pagina_pdf": 1},
        ]
        registros = v3._registros_lexicais(
            Path("contexto.pdf"),
            blocos,
            [{"termo": "população negra", "categoria": "TERMO INFORMADO"}],
        )
        self.assertEqual(len(registros), 1)
        registro = registros[0]
        self.assertEqual(registro["Bloco anterior"], anterior)
        self.assertEqual(registro["Trecho da ocorrência"], central)
        self.assertEqual(registro["Bloco posterior"], posterior)

        # Verificação de ponta a ponta: os três blocos chegam ao Excel na ordem
        # correta e preservam o texto original extraído.
        with TemporaryDirectory() as diretorio:
            arquivo = Path(diretorio) / "contexto_v3.xlsx"
            v3.salvar_excel_completo(
                arquivo,
                pd.DataFrame(registros),
                pd.DataFrame([{"Termo": "população negra"}]),
                pd.DataFrame(
                    [
                        {
                            "arquivo": "contexto.pdf",
                            "paginas_com_texto_extraivel": 1,
                            "paginas_processadas_com_OCR": 0,
                            "idioma_OCR": "não utilizado",
                        }
                    ]
                ),
                {"incluir_lexical": True, "incluir_semantica": False},
            )
            planilha = load_workbook(arquivo)["Resultados"]
            cabecalhos = [celula.value for celula in planilha[1]]
            self.assertEqual(
                planilha.cell(2, cabecalhos.index("Bloco anterior") + 1).value,
                anterior,
            )
            self.assertEqual(
                planilha.cell(2, cabecalhos.index("Trecho da ocorrência") + 1).value,
                central,
            )
            self.assertEqual(
                planilha.cell(2, cabecalhos.index("Bloco posterior") + 1).value,
                posterior,
            )

        primeiro = v3._registros_lexicais(
            Path("contexto.pdf"), blocos, [{"termo": "crescimento urbano"}]
        )[0]
        ultimo = v3._registros_lexicais(
            Path("contexto.pdf"), blocos, [{"termo": "desigualdades territoriais"}]
        )[0]
        self.assertEqual(primeiro["Bloco anterior"], "")
        self.assertEqual(ultimo["Bloco posterior"], "")

    def test_parametros_v3_registram_modalidades_efetivamente_utilizadas(self):
        diagnostico = pd.DataFrame(
            [
                {
                    "arquivo": "lexical.pdf",
                    "paginas_com_texto_extraivel": 1,
                    "paginas_processadas_com_OCR": 0,
                    "idioma_OCR": "não utilizado",
                    "busca_lexical_ativada": True,
                    "busca_semantica_ativada": False,
                    "limiar_semantico": 0.70,
                }
            ]
        )
        with TemporaryDirectory() as diretorio:
            arquivo = Path(diretorio) / "lexical_v3.xlsx"
            v3.salvar_excel_completo(
                arquivo,
                v3.dataframe_vazio(),
                pd.DataFrame([{"Termo": "São Paulo"}]),
                diagnostico,
                {"incluir_lexical": True, "incluir_semantica": False, "limiar_semantico": 0.70},
            )
            parametros = self._parametros(load_workbook(arquivo))
            self.assertEqual(parametros["Busca lexical/morfológica"], "Sim")
            self.assertEqual(parametros["Busca semântica"], "Não")
            self.assertEqual(parametros["Limiar semântico"], "Não aplicável")
            self.assertEqual(
                parametros["Modelo de embeddings"],
                "Não utilizado (busca semântica desativada)",
            )

    @patch("analyzer.v3.carregar_modelo_semantico", return_value=ModeloControlado())
    def test_v3_semantica_usa_contexto_original_sem_repetir_overlap(self, _modelo):
        anterior = "Bloco anterior real, extraído antes do trecho recuperado."
        central = "Bloco central real que trata de condições climaticas e migração."
        posterior = "Bloco posterior real, extraído depois do trecho recuperado."
        blocos = [{"texto": texto} for texto in (anterior, central, posterior)]
        trechos = [
            {
                "texto": "Trecho semântico recuperado sobre condições climaticas e migração.",
                "pagina_inicial": 2,
                "pagina_final": 2,
                "pagina_inicial_pdf": 2,
                "pagina_final_pdf": 2,
                "indice_bloco_inicial": 1,
                "indice_bloco_final": 1,
            }
        ]
        registros = v3._registros_semanticos(
            Path("semantico.pdf"),
            trechos,
            [{"termo": "seca", "categoria": "TERMO INFORMADO"}],
            0.70,
            blocos,
        )
        self.assertEqual(len(registros), 1)
        registro = registros[0]
        self.assertEqual(registro["Tipo de correspondência"], "Semântica")
        self.assertEqual(registro["Termo encontrado"], "")
        self.assertEqual(registro["Trecho da ocorrência"], trechos[0]["texto"])
        self.assertEqual(registro["Bloco anterior"], anterior)
        self.assertEqual(registro["Bloco posterior"], posterior)
        self.assertNotEqual(registro["Bloco anterior"], registro["Trecho da ocorrência"])
        self.assertNotEqual(registro["Bloco posterior"], registro["Trecho da ocorrência"])

    def test_termos_compostos_combinados_sem_duplicacao(self):
        termos = montar_termos(
            "Rio Grande do Norte; João Pessoa",
            "rio grande do norte\nRio de Janeiro",
            "v3",
        )
        self.assertEqual(
            [item["termo"] for item in termos],
            ["Rio Grande do Norte", "João Pessoa", "Rio de Janeiro"],
        )

    def test_termos_compostos_preservam_espacos_e_separacao_por_ponto_e_virgula(self):
        consultas_esperadas = [
            "São Paulo",
            "Rio de Janeiro",
            "Rio Grande do Norte",
            "João Pessoa",
            "racismo ambiental",
            "monopólio da violência legítima",
        ]
        termos = montar_termos("; ".join(consultas_esperadas), "", "v1")

        self.assertEqual([item["termo"] for item in termos], consultas_esperadas)
        self.assertEqual(len(termos), len(consultas_esperadas))
        self.assertNotIn("São", [item["termo"] for item in termos])
        self.assertNotIn("Paulo", [item["termo"] for item in termos])

    def test_tres_consultas_compostas_geram_exatamente_tres_itens(self):
        termos = montar_termos("São Paulo; Rio de Janeiro; João Pessoa", "", "v3")
        self.assertEqual(
            [item["termo"] for item in termos],
            ["São Paulo", "Rio de Janeiro", "João Pessoa"],
        )

    def test_v1_e_v2_encontram_expressao_composta_com_normalizacao(self):
        texto = "São Paulo, SÃO PAULO e Sao Paulo discutem racismo ambiental."
        for analisador in (v1, v2):
            self.assertEqual(analisador.contar_ocorrencias(texto, "São Paulo"), 3)
            self.assertEqual(analisador.contar_ocorrencias(texto, "racismo ambiental"), 1)

    @patch("analyzer.v3.v1.extrair_pdf")
    def test_v3_lexical_encontra_expressao_composta_sem_dividi_la(self, extrair_pdf):
        extrair_pdf.return_value = (
            [
                {
                    "texto": "São Paulo, SÃO PAULO e Sao Paulo aparecem no mesmo parágrafo.",
                    "pagina": 1,
                    "pagina_pdf": 1,
                }
            ],
            1,
            0,
            None,
        )
        termos = montar_termos("São Paulo", "", "v3")
        registros, _ = v3.analisar_pdf(
            Path("sao_paulo.pdf"),
            termos,
            {"incluir_lexical": True, "incluir_semantica": False},
        )

        self.assertEqual(len(registros), 1)
        self.assertEqual(registros[0]["Consulta"], "São Paulo")
        self.assertEqual(registros[0]["_quantidade_no_registro"], 3)
        self.assertEqual(registros[0]["Tipo de correspondência"], "Lexical")

    @patch("analyzer.v3.carregar_modelo_semantico")
    def test_v3_semantica_codifica_cada_expressao_composta_por_inteiro(self, carregar_modelo):
        modelo = ModeloQueRegistraConsultas()
        carregar_modelo.return_value = modelo
        consultas_esperadas = [
            "racismo ambiental",
            "representações do Nordeste brasileiro",
            "migração provocada pela seca",
            "desigualdade no acesso ao saneamento básico",
        ]
        termos = montar_termos("; ".join(consultas_esperadas), "", "v3")
        trechos = [
            {
                "texto": "Trecho usado apenas para confirmar a entrada do modelo semântico.",
                "pagina_inicial": 1,
                "pagina_final": 1,
                "pagina_inicial_pdf": 1,
                "pagina_final_pdf": 1,
            }
        ]

        v3._registros_semanticos(Path("livro.pdf"), trechos, termos, 0.70)

        self.assertEqual(modelo.chamadas[1], consultas_esperadas)
        self.assertNotIn("racismo", modelo.chamadas[1])
        self.assertNotIn("ambiental", modelo.chamadas[1])

    def test_estrutura_de_resultados_vazia(self):
        vazio = v3.dataframe_vazio()
        self.assertEqual(vazio.columns.tolist(), v3.COLUNAS_INTERNAS)
        self.assertTrue(vazio.empty)

    @patch("analyzer.v3.v1.extrair_pdf")
    @patch("analyzer.v3.carregar_modelo_semantico")
    def test_modo_lexical_isolado_nao_carrega_modelo(self, modelo, extrair_pdf):
        extrair_pdf.side_effect = self._extrair_pdf
        registros, _ = v3.analisar_pdf(
            Path("livro_lexical.pdf"),
            self.termos,
            {"incluir_lexical": True, "incluir_semantica": False},
        )
        self.assertTrue(registros)
        self.assertTrue(
            all(
                registro["Tipo de correspondência"] in {"Lexical", "Morfológica"}
                for registro in registros
            )
        )
        self.assertTrue(
            all(registro["Similaridade semântica"] is None for registro in registros)
        )
        modelo.assert_not_called()

    @patch("analyzer.v3.carregar_modelo_semantico", return_value=ModeloControlado())
    @patch("analyzer.v3.v1.extrair_pdf")
    def test_modo_semantico_isolado_nao_cria_ocorrencias_lexicais(self, extrair_pdf, modelo):
        extrair_pdf.side_effect = self._extrair_pdf
        registros, _ = v3.analisar_pdf(
            Path("livro_semantico.pdf"),
            self.termos,
            {"incluir_lexical": False, "incluir_semantica": True, "limiar_semantico": 0.70},
        )
        self.assertTrue(registros)
        self.assertTrue(all(registro["Tipo de correspondência"] == "Semântica" for registro in registros))
        self.assertTrue(all(registro["Termo encontrado"] == "" for registro in registros))
        self.assertTrue(all(registro["_quantidade_no_registro"] == 1 for registro in registros))
        modelo.assert_called_once()

    def test_configuracao_recusa_duas_modalidades_desativadas(self):
        with self.assertRaises(ValueError):
            v3._configuracao({"incluir_lexical": False, "incluir_semantica": False})

    @patch("analyzer.v3.carregar_modelo_semantico", return_value=ModeloControlado())
    @patch("analyzer.v3.v1.extrair_pdf")
    def test_paginas_flask_e_regressao_de_estruturas(self, extrair_pdf, _modelo):
        extrair_pdf.side_effect = self._extrair_pdf
        with TemporaryDirectory() as diretorio:
            resultado = executar_analises(
                [Path("livro_um.pdf")],
                self.termos,
                Path(diretorio),
                "v3",
                self.configuracao,
            )
            resultado["dashboard"] = criar_dashboard(resultado)
            self.assertNotIn("destaques_semanticos", resultado["dashboard"])
            resultado["pasta_saida"] = Path(diretorio)
            identificador = "teste-v3"
            aplicacao_web.ANALISES[identificador] = resultado
            try:
                with aplicacao_web.app.test_client() as cliente:
                    inicio = cliente.get("/")
                    pagina_v3 = cliente.get(f"/resultado/{identificador}")
                    self.assertEqual(inicio.status_code, 200)
                    self.assertIn(b"V3: Busca h", inicio.data)
                    self.assertIn(b"Fabio Monteiro de Moraes</a> - 2026 | Vers", inicio.data)
                    html_inicio = inicio.get_data(as_text=True)
                    self.assertIn("Incluir busca lexical e morfológica", html_inicio)
                    self.assertIn("Incluir busca semântica", html_inicio)
                    self.assertIn("Escolha o método de varredura", html_inicio)
                    self.assertIn(
                        "Separe os termos por ponto e vírgula ou coloque um termo por linha. Termos compostos devem permanecer inteiros.",
                        html_inicio,
                    )
                    self.assertIn(
                        "Use ponto e vírgula ou coloque um termo por linha.",
                        html_inicio,
                    )
                    self.assertIn(
                        "Exemplo: Nordeste; São Paulo; Rio de Janeiro",
                        html_inicio,
                    )
                    self.assertIn("Qual método escolher?", html_inicio)
                    self.assertIn("controle-limiar-semantico", html_inicio)
                    self.assertIn(
                        "Dependendo do método de varredura escolhido, a busca pode considerar ocorrências lexicais e morfológicas ou também recuperar trechos por similaridade semântica.",
                        html_inicio,
                    )
                    self.assertIn(
                        "Por isso, uma correspondência semântica não é o mesmo que uma ocorrência textual.",
                        html_inicio,
                    )
                    self.assertIn(
                        "Quando apenas a busca lexical e morfológica estiver ativada, a V3 utiliza a mesma lógica básica de localização lexical da V1.",
                        html_inicio,
                    )
                    self.assertIn(
                        "Os arquivos Excel registram também os parâmetros utilizados na análise, permitindo identificar posteriormente como a varredura foi realizada.",
                        html_inicio,
                    )
                    self.assertIn(
                        "Para auxiliar a validação qualitativa, a planilha V3 apresenta o trecho recuperado juntamente com o bloco textual anterior e posterior, quando disponíveis.",
                        html_inicio,
                    )
                    self.assertIn("http://lattes.cnpq.br/0075160400322127", html_inicio)
                    self.assertIn('target="_blank" rel="noopener noreferrer"', html_inicio)
                    self.assertEqual(pagina_v3.status_code, 200)
                    self.assertIn(b"Fabio Monteiro de Moraes</a> - 2026 | Vers", pagina_v3.data)
                    self.assertNotIn(b"Principais correspond", pagina_v3.data)
                    self.assertNotIn(b"destaques-v3", pagina_v3.data)
                    self.assertIn(b"dashboard-v3.js", pagina_v3.data)
            finally:
                aplicacao_web.ANALISES.pop(identificador, None)

        # A V3 não alterou os dataframes que alimentam as estruturas V1/V2.
        self.assertIn("Parágrafo do termo", v1.dataframe_vazio().columns)
        self.assertIn("Tipo da ocorrência", v2.dataframe_vazio().columns)
        self.assertNotIn("Resultados", v1.dataframe_vazio().columns)
        self.assertNotIn("Resultados", v2.dataframe_vazio().columns)

    def test_dashboard_lexical_v1_e_v2_continua_selecionando_o_script_original(self):
        diagnosticos = pd.DataFrame(
            columns=[
                "arquivo",
                "paginas_com_texto_extraivel",
                "paginas_processadas_com_OCR",
                "idioma_OCR",
            ]
        )
        for versao, analisador in (("v1", v1), ("v2", v2)):
            resultado = {
                "arquivos": [],
                "erros": [],
                "quantidade_pdfs": 0,
                "quantidade_termos": 0,
                "ocorrencias": analisador.dataframe_vazio(),
                "diagnosticos": diagnosticos,
                "termos": [],
                "versao": versao,
                "pasta_saida": Path("."),
            }
            resultado["dashboard"] = criar_dashboard(resultado)
            identificador = f"teste-{versao}"
            aplicacao_web.ANALISES[identificador] = resultado
            try:
                with aplicacao_web.app.test_client() as cliente:
                    pagina = cliente.get(f"/resultado/{identificador}")
                    self.assertEqual(pagina.status_code, 200)
                    self.assertIn(b"Fabio Monteiro de Moraes</a> - 2026 | Vers", pagina.data)
                    self.assertIn(b"dashboard.js", pagina.data)
                    self.assertNotIn(b"dashboard-v3.js", pagina.data)
            finally:
                aplicacao_web.ANALISES.pop(identificador, None)

    @patch("app.executar_analises")
    def test_backend_rejeita_virgula_e_ponto_no_campo_manual(self, executar):
        entradas_invalidas = (
            "Nordeste, Recife, sertão",
            "Nordeste. Recife. sertão",
            "Nordeste, Recife; sertão",
            "Nordeste. Recife; sertão",
            "São Paulo, Rio de Janeiro, João Pessoa",
            "São Paulo. Rio de Janeiro. João Pessoa",
        )
        with aplicacao_web.app.test_client() as cliente:
            for termos in entradas_invalidas:
                resposta = cliente.post(
                    "/",
                    data={
                        "versao": "v1",
                        "termos": termos,
                        "pdfs": (BytesIO(b"%PDF-teste"), "livro.pdf"),
                    },
                    content_type="multipart/form-data",
                )
                self.assertEqual(resposta.status_code, 200)
                self.assertIn(
                    "Use ponto e vírgula (;) para separar os termos de pesquisa.".encode(),
                    resposta.data,
                )

            resposta_valida = cliente.post(
                "/",
                data={"versao": "v1", "termos": "Nordeste; Recife; sertão"},
                content_type="multipart/form-data",
            )
            self.assertEqual(resposta_valida.status_code, 200)
            self.assertIn(b"Envie ao menos um arquivo PDF.", resposta_valida.data)
            self.assertNotIn(
                '<div class="alert alert-danger" role="alert">Use ponto e vírgula (;) para separar os termos de pesquisa.</div>',
                resposta_valida.get_data(as_text=True),
            )

        executar.assert_not_called()


if __name__ == "__main__":
    unittest.main()
