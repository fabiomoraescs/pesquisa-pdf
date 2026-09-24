"""Importação XLSX e navegação da plataforma sem tocar nos analisadores legados."""

import io
import re
import time
import unittest
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from platform_helpers import create_project, create_user, csrf_from, isolated_platform, login
from platform_core.extensions import db
from platform_core.library_spreadsheets import SpreadsheetImportError, parse_library_xlsx
from platform_core.models import AuditLog, Project, ProjectLibrary, VocabularyLibrary
from historico_racial.routes import JOBS_LOCK, PROGRESSOS_HR


def workbook_bytes(*, name="Biblioteca de teste", groups=None, entities=None, variants=None):
    workbook = Workbook()
    workbook.active.title = "BIBLIOTECA"
    workbook.active.append(["nome", "descricao"])
    workbook.active.append([name, "Vocabulário de teste"])
    sheets = {
        "GRUPOS": (["id_grupo", "nome", "descricao", "ativo"],
                   groups if groups is not None else [["educacao", "Educação", "", "Sim"], ["trabalho", "Trabalho", "", "Sim"]]),
        "ENTIDADES": (["entity_key", "forma_canonica", "tipo_entidade", "grupos", "ativo", "tradicao_intelectual", "pais_regiao", "observacoes"],
                     entities if entities is not None else [["escola_publica", "Escola pública", "conceito", "educacao;trabalho", "Sim", "", "", ""]]),
        "VARIANTES": (["entity_key", "variante", "ativo"],
                     variants if variants is not None else [["escola_publica", "Escola pública", "Sim"], ["escola_publica", "escolas públicas", "Sim"]]),
    }
    for title, (header, rows) in sheets.items():
        sheet = workbook.create_sheet(title)
        sheet.append(header)
        for row in rows:
            sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class ImportParserTests(unittest.TestCase):
    def test_official_template_is_valid_xlsx(self):
        path = Path(__file__).resolve().parent.parent / "resources" / "modelo_biblioteca.xlsx"
        workbook = load_workbook(io.BytesIO(path.read_bytes()), read_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["INSTRUÇÕES", "BIBLIOTECA", "GRUPOS", "ENTIDADES", "VARIANTES"])
            self.assertEqual([cell.value for cell in next(workbook["ENTIDADES"].iter_rows())],
                             ["identificador_entidade", "forma_canonica", "tipo_entidade", "grupos", "ativo", "tradicao_intelectual", "pais_regiao", "observacoes"])
            count_rows = lambda sheet: sum(1 for row in sheet.iter_rows(values_only=True) if any(value is not None for value in row))
            self.assertEqual(count_rows(workbook["GRUPOS"]), 9)
            self.assertEqual(count_rows(workbook["ENTIDADES"]), 76)
            self.assertEqual(count_rows(workbook["VARIANTES"]), 118)
        finally:
            workbook.close()

    def test_valid_multigroup_and_optional_metadata(self):
        imported = parse_library_xlsx(workbook_bytes())
        self.assertEqual(imported.counts, {"grupos": 2, "entidades": 1, "variantes": 2})
        self.assertEqual(imported.snapshot["entidades"][0]["grupo"], ["educacao", "trabalho"])
        self.assertEqual(imported.snapshot["entidades"][0]["tradicao_intelectual"], "")
        optional_type = [["escola_publica", "Escola pública", "", "educacao", "Sim", "", "", ""]]
        self.assertEqual(parse_library_xlsx(workbook_bytes(entities=optional_type)).snapshot["entidades"][0]["tipo_entidade"], "outro")

    def test_duplicate_group_rejected(self):
        with self.assertRaisesRegex(SpreadsheetImportError, "GRUPOS, linha 3: identificador de grupo duplicado"):
            parse_library_xlsx(workbook_bytes(groups=[["educacao", "Educação", "", "Sim"], ["educacao", "Outra", "", "Sim"]]))

    def test_duplicate_entity_key_rejected(self):
        entities = [["escola_publica", "Escola pública", "conceito", "educacao", "Sim", "", "", ""],
                    ["escola_publica", "Rede pública", "conceito", "trabalho", "Sim", "", "", ""]]
        with self.assertRaisesRegex(SpreadsheetImportError, "ENTIDADES, linha 3: identificador de entidade duplicado"):
            parse_library_xlsx(workbook_bytes(entities=entities))

    def test_portuguese_headers_and_legacy_compatibility(self):
        old = workbook_bytes()
        self.assertEqual(parse_library_xlsx(old).counts["entidades"], 1)
        workbook = load_workbook(io.BytesIO(old))
        workbook["ENTIDADES"]["A1"] = "identificador_entidade"
        workbook["VARIANTES"]["A1"] = "identificador_entidade"
        output = io.BytesIO()
        workbook.save(output)
        self.assertEqual(parse_library_xlsx(output.getvalue()).counts["variantes"], 2)

    def test_missing_group_rejected(self):
        entities = [["escola_publica", "Escola pública", "conceito", "inexistente", "Sim", "", "", ""]]
        with self.assertRaisesRegex(SpreadsheetImportError, "ENTIDADES, linha 2: o grupo 'inexistente' não existe"):
            parse_library_xlsx(workbook_bytes(entities=entities))

    def test_duplicate_variant_in_same_entity_rejected(self):
        variants = [["escola_publica", "Escola pública", "Sim"], ["escola_publica", "escola pública", "Sim"]]
        with self.assertRaisesRegex(SpreadsheetImportError, "VARIANTES, linha 3: variante duplicada"):
            parse_library_xlsx(workbook_bytes(variants=variants))

    def test_variant_conflict_between_entities_rejected(self):
        entities = [["escola_publica", "Escola pública", "conceito", "educacao", "Sim", "", "", ""],
                    ["rede_publica", "Rede pública", "conceito", "trabalho", "Sim", "", "", ""]]
        variants = [["escola_publica", "Escola pública", "Sim"], ["rede_publica", "Rede pública", "Sim"],
                    ["rede_publica", "escola publica", "Sim"]]
        with self.assertRaisesRegex(SpreadsheetImportError, "VARIANTES, linha 4: a variante .* já pertence a outra entidade"):
            parse_library_xlsx(workbook_bytes(entities=entities, variants=variants))

    def test_formula_rejected(self):
        data = workbook_bytes(name="=2+2")
        # openpyxl stores the leading '=' as a formula in the workbook fixture.
        with self.assertRaisesRegex(SpreadsheetImportError, "fórmulas não são permitidas"):
            parse_library_xlsx(data)

    def test_invalid_state_rejected(self):
        with self.assertRaisesRegex(SpreadsheetImportError, "use Sim ou Não"):
            parse_library_xlsx(workbook_bytes(groups=[["educacao", "Educação", "", "talvez"]]))


class ImportRouteTests(unittest.TestCase):
    def setUp(self):
        self.scope = isolated_platform()
        self.app = self.scope.__enter__()
        self.client = self.app.test_client()

    def tearDown(self):
        self.scope.__exit__(None, None, None)

    def _admin(self):
        create_user("Administrador", "admin@example.org", role="admin")
        login(self.client, "admin@example.org")

    def _upload(self, payload, filename="nova.xlsx"):
        token = csrf_from(self.client.get("/admin/bibliotecas/importar"))
        return self.client.post("/admin/bibliotecas/importar", data={
            "csrf_token": token, "spreadsheet": (io.BytesIO(payload), filename),
        }, content_type="multipart/form-data")

    def test_admin_template_download_and_user_restriction(self):
        create_user()
        login(self.client)
        self.assertEqual(self.client.get("/admin/bibliotecas/modelo").status_code, 403)
        self.assertEqual(self.client.get("/admin/bibliotecas/importar").status_code, 403)
        self.client.post("/logout", data={"csrf_token": csrf_from(self.client.get("/perfil"))})
        self._admin()
        response = self.client.get("/admin/bibliotecas/modelo")
        try:
            self.assertEqual(response.status_code, 200)
            self.assertTrue(response.data.startswith(b"PK"))
        finally:
            response.close()

    def test_preview_then_atomic_draft_then_publication(self):
        self._admin()
        # O projeto anterior deve manter o snapshot e os vínculos originais.
        project_id = create_project(self.client)
        original_libraries = db.session.execute(
            select(ProjectLibrary.library_id, ProjectLibrary.source_hash).where(ProjectLibrary.project_id == project_id)
        ).all()
        preview = self._upload(workbook_bytes())
        self.assertEqual(preview.status_code, 200)
        html = preview.get_data(as_text=True)
        self.assertIn("Pré-visualização da importação", html)
        self.assertIn("2 grupos · 1 entidades · 2 variantes", html)
        self.assertIsNone(db.session.get(VocabularyLibrary, "biblioteca_de_teste"))
        token = re.search(r'name="preview_token" value="([^"]+)"', html).group(1)
        confirmed = self.client.post("/admin/bibliotecas/importar/confirmar", data={
            "csrf_token": csrf_from(preview), "preview_token": token,
        })
        self.assertEqual(confirmed.status_code, 302)
        library = db.session.get(VocabularyLibrary, "biblioteca_de_teste")
        self.assertEqual(library.status, "draft")
        self.assertFalse(library.active)
        self.assertEqual(library.snapshot_json["entidades"][0]["grupo"], ["educacao", "trabalho"])
        self.assertNotIn("biblioteca_de_teste", self.client.get("/projetos/novo").get_data(as_text=True))
        self.assertEqual(original_libraries, db.session.execute(
            select(ProjectLibrary.library_id, ProjectLibrary.source_hash).where(ProjectLibrary.project_id == project_id)
        ).all())
        publication = self.client.post(f"/admin/bibliotecas/{library.id}/publicar", data={
            "csrf_token": csrf_from(self.client.get(f"/admin/bibliotecas/{library.id}")),
            "base_hash": library.content_hash, "confirm": "yes",
        })
        self.assertEqual(publication.status_code, 302)
        self.assertIn("biblioteca_de_teste", self.client.get("/projetos/novo").get_data(as_text=True))
        self.assertEqual(original_libraries, db.session.execute(
            select(ProjectLibrary.library_id, ProjectLibrary.source_hash).where(ProjectLibrary.project_id == project_id)
        ).all())
        official = db.session.get(VocabularyLibrary, "relacoes_raciais")
        self.assertEqual(official.counts_json, {"grupos": 8, "entidades": 75, "variantes": 117})

    def test_filled_official_model_imports_as_separate_draft(self):
        self._admin()
        path = Path(__file__).resolve().parent.parent / "resources" / "modelo_biblioteca.xlsx"
        untouched = self._upload(path.read_bytes(), "modelo_biblioteca.xlsx")
        self.assertEqual(untouched.status_code, 400)
        self.assertIn("altere o nome da cópia de referência", untouched.get_data(as_text=True))
        workbook = load_workbook(io.BytesIO(path.read_bytes()))
        workbook["BIBLIOTECA"]["A2"] = "Biblioteca de pesquisa racial"
        edited = io.BytesIO()
        workbook.save(edited)
        workbook.close()
        response = self._upload(edited.getvalue(), "biblioteca_editada.xlsx")
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn("8 grupos · 75 entidades · 117 variantes", html)
        token = re.search(r'name="preview_token" value="([^"]+)"', html).group(1)
        confirmation = self.client.post("/admin/bibliotecas/importar/confirmar", data={
            "csrf_token": csrf_from(response), "preview_token": token,
        })
        self.assertEqual(confirmation.status_code, 302)
        imported = db.session.get(VocabularyLibrary, "biblioteca_de_pesquisa_racial")
        self.assertIsNotNone(imported)
        self.assertEqual(imported.status, "draft")
        self.assertEqual(imported.counts_json, {"grupos": 8, "entidades": 75, "variantes": 117})
        self.assertEqual(db.session.get(VocabularyLibrary, "relacoes_raciais").status, "published")

    def test_invalid_spreadsheet_creates_nothing(self):
        self._admin()
        response = self._upload(workbook_bytes(groups=[["educacao", "Educação", "", "Sim"],
                                                      ["educacao", "Duplicado", "", "Sim"]]))
        self.assertEqual(response.status_code, 400)
        self.assertIn("GRUPOS, linha 3", response.get_data(as_text=True))
        self.assertIsNone(db.session.get(VocabularyLibrary, "biblioteca_de_teste"))

    def test_library_deletion_protects_default_and_project_links(self):
        self._admin()
        default = db.session.get(VocabularyLibrary, "relacoes_raciais")
        response = self.client.get(f"/admin/bibliotecas/{default.id}/excluir")
        self.assertEqual(response.status_code, 302)
        self.assertIn("A biblioteca padrão do sistema não pode ser excluída.",
                      self.client.get(response.headers["Location"]).get_data(as_text=True))
        self.assertIsNotNone(db.session.get(VocabularyLibrary, default.id))
        preview = self._upload(workbook_bytes())
        token = re.search(r'name="preview_token" value="([^"]+)"', preview.get_data(as_text=True)).group(1)
        self.client.post("/admin/bibliotecas/importar/confirmar", data={
            "csrf_token": csrf_from(preview), "preview_token": token,
        })
        library = db.session.get(VocabularyLibrary, "biblioteca_de_teste")
        url = f"/admin/bibliotecas/{library.id}/excluir"
        bad = self.client.post(url, data={"csrf_token": csrf_from(self.client.get(url)),
                                          "confirmation": "errado", "base_hash": library.content_hash})
        self.assertEqual(bad.status_code, 400)
        self.assertIsNotNone(db.session.get(VocabularyLibrary, library.id))
        project_id = create_project(self.client, libraries=("relacoes_raciais",))
        db.session.add(ProjectLibrary(project_id=project_id, library_id=library.id,
                                      source_hash=library.content_hash, source_version=library.version))
        db.session.commit()
        linked_response = self.client.get(url)
        self.assertEqual(linked_response.status_code, 302)
        self.assertIn("Esta biblioteca não pode ser excluída porque está vinculada",
                      self.client.get(linked_response.headers["Location"]).get_data(as_text=True))
        self.assertIsNotNone(db.session.get(VocabularyLibrary, library.id))
        db.session.delete(db.session.get(ProjectLibrary, (project_id, library.id)))
        db.session.commit()
        response = self.client.post(url, data={"csrf_token": csrf_from(self.client.get(url)),
                                               "confirmation": library.name, "base_hash": library.content_hash})
        self.assertEqual(response.status_code, 302)
        self.assertIsNone(db.session.get(VocabularyLibrary, library.id))
        self.assertIsNotNone(db.session.scalar(select(AuditLog).where(
            AuditLog.action == "library_deleted", AuditLog.target_id == library.id)))

    def test_navigation_and_shared_progress_markup(self):
        self._admin()
        projects = self.client.get("/projetos").get_data(as_text=True)
        self.assertIn('>Novo projeto</a>', projects.split('class="platform-nav"')[1].split('</nav>')[0])
        self.assertIn('>Novo projeto</a>', projects)
        self.assertNotIn('+ Novo projeto', projects)
        self.assertIn(
            'class="platform-section-title" aria-label="Nome da plataforma"><span class="platform-brand-mark"><strong>Análysis</strong><small>ferramentas para pesquisa</small></span>',
            projects,
        )
        self.assertIn('class="platform-header-context">Projetos', projects)
        profile = self.client.get("/perfil").get_data(as_text=True)
        self.assertIn(
            'class="platform-section-title" aria-label="Nome da plataforma"><span class="platform-brand-mark"><strong>Análysis</strong><small>ferramentas para pesquisa</small></span>',
            profile,
        )
        self.assertIn('class="platform-header-context">Perfil', profile)
        self.assertNotIn('class="platform-page-heading"', profile)
        admin = self.client.get("/admin").get_data(as_text=True)
        self.assertIn('id="platform-admin-toggle"', admin)
        self.assertIn('aria-expanded="true"', admin)
        self.assertNotIn('class="platform-admin-nav"', admin)
        for name in ("Visão geral", "Usuários", "Ferramentas", "Planos", "Definir acessos", "Bibliotecas", "Auditoria"):
            self.assertIn(name, admin)
        legacy = self.client.get("/raspagem-livre").get_data(as_text=True)
        self.assertIn('id="platform-sidebar"', legacy)
        self.assertIn('id="platform-menu-toggle"', legacy)
        self.assertIn(
            'class="platform-section-title" aria-label="Nome da plataforma"><span class="platform-brand-mark"><strong>Análysis</strong><small>ferramentas para pesquisa</small></span>',
            legacy,
        )
        self.assertIn('class="platform-header-context">Raspagem livre', legacy)
        self.assertNotIn('aria-label="Conta e projetos"', legacy)
        self.assertIn('id="overlay-processamento"', legacy)

    def test_document_progress_has_real_fields_and_eta(self):
        self._admin()
        project_id = create_project(self.client)
        page = self.client.get(f"/analise-documental/projetos/{project_id}").get_data(as_text=True)
        for marker in ('id="hr-overlay-processamento"', 'id="hr-progresso-processamento"',
                       'id="hr-detalhe-processamento"', 'id="hr-tempo-restante-processamento"'):
            self.assertIn(marker, page)
        job_id = str(uuid4())
        user_id = db.session.scalar(select(Project.owner_user_id).where(Project.id == project_id))
        with JOBS_LOCK:
            PROGRESSOS_HR[job_id] = {
                "status": "processando", "etapa": "Identificando ocorrências…", "percentual": 35,
                "arquivo_atual": "teste.pdf", "arquivo_indice": 1, "arquivos_total": 1,
                "pagina_atual": 7, "paginas_total": 20, "tempo_inicio": time.monotonic() - 20,
                "resultado_url": None, "erro": None, "project_id": project_id,
                "owner_user_id": user_id, "vocabulario_version": "v1.0", "vocabulario_hash": "hash",
            }
        try:
            response = self.client.get(f"/analise-documental/projetos/{project_id}/api/progresso/{job_id}")
            self.assertEqual(response.status_code, 200)
            data = response.get_json()
            self.assertEqual(data["percentual"], 35)
            self.assertEqual(data["pagina_atual"], 7)
            self.assertEqual(data["etapa"], "Identificando ocorrências…")
            self.assertIsInstance(data["eta_segundos"], int)
        finally:
            with JOBS_LOCK:
                PROGRESSOS_HR.pop(job_id, None)


if __name__ == "__main__":
    unittest.main()
