"""Cor global é propriedade do código, nunca da ocorrência."""

import io
import re
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from sqlalchemy import select

from platform_helpers import create_user, csrf_from, login
import test_qualitative_automatic as automatic_fixture
import test_qualitative_context as context_fixture
from test_qualitative_expanded_search import FakeModel
from platform_core.extensions import db
from platform_core.models import (QualitativeCode, QualitativeCoding, QualitativeExcerpt,
                                  QualitativeMemo, QualitativeRejection, UserToolOverride)
from platform_core.qualitative_annotations import page_excerpts
from platform_core.qualitative_corpus import read_qualitative_page
from platform_core.qualitative_colors import (LEGACY_CODE_COLORS, QUALITATIVE_CODE_PALETTE,
                                              code_color_style)
from platform_core.scraping_types import QUALITATIVE_TOOL


class QualitativeCodeColorTests(unittest.TestCase):
    setUp = context_fixture.QualitativeContextTests.setUp
    tearDown = context_fixture.QualitativeContextTests.tearDown
    post = context_fixture.QualitativeContextTests.post
    selection = context_fixture.QualitativeContextTests.selection
    automatic = automatic_fixture.QualitativeAutomaticTests.automatic
    remove = automatic_fixture.QualitativeAutomaticTests.remove
    codings = automatic_fixture.QualitativeAutomaticTests.codings
    fixture_pages = automatic_fixture.QualitativeAutomaticTests.fixture_pages
    auto_url = automatic_fixture.QualitativeAutomaticTests.auto_url

    def code_url(self, identifier):
        return f"/analise-qualitativa/bases/{self.analysis.id}/codigos/{identifier}"

    def create_code(self, name):
        response = self.client.post(f"/analise-qualitativa/bases/{self.analysis.id}/codigos",
                                    json={"name": name}, headers={"X-CSRFToken": self.csrf})
        self.assertEqual(response.status_code, 201, response.json)
        return next(item["id"] for item in response.json["codes"] if item["name"] == name)

    def patch_code(self, identifier, changes):
        return self.client.patch(self.code_url(identifier), json=changes,
                                 headers={"X-CSRFToken": self.csrf})

    def test_ten_labels_two_codes_and_new_coding_use_current_color_without_changing_links(self):
        self.fixture_pages([[" ".join(["raça"] * 10)]])
        page = read_qualitative_page(self.analysis, self.document.id, 1)
        first = self.create_code("Raça")
        second = self.create_code("Outro")
        self.assertEqual(self.patch_code(second, {"color": "green"}).status_code, 200)
        positions = [match.start() for match in re.finditer("raça", page["text"])]
        for start in positions[:9]:
            response = self.post("apply_codes", {"code_ids": [first]}, selection={
                "start": start, "end": start + 4, "selected_text": "raça", "page_hash": page["sha256"]})
            self.assertEqual(response.status_code, 201, response.json)
        self.assertEqual(self.post("apply_codes", {"code_ids": [second]}, selection={
            "start": positions[0], "end": positions[0] + 4,
            "selected_text": "raça", "page_hash": page["sha256"]}).status_code, 201)
        first_excerpt = db.session.scalar(select(QualitativeExcerpt).order_by(QualitativeExcerpt.start_offset))
        memo = QualitativeMemo(analysis_id=self.analysis.id, excerpt_id=first_excerpt.id,
                               text="Memo preservado", created_by_user_id=self.user.id)
        db.session.add(memo)
        db.session.commit()
        before_codings = [(item.id, item.code_id, item.excerpt_id, item.origin,
                           item.source_query) for item in self.codings()]
        before_excerpts = [item.id for item in db.session.scalars(
            select(QualitativeExcerpt).order_by(QualitativeExcerpt.start_offset))]
        changed = self.patch_code(first, {"color": "yellow"})
        self.assertEqual(changed.status_code, 200, changed.json)
        item = next(item for item in changed.json["codes"] if item["id"] == first)
        self.assertEqual((item["name"], item["color"], item["color_hex"], item["excerpt_count"]),
                         ("Raça", "yellow", "#FDE68A", 9))
        self.assertEqual(db.session.get(QualitativeCode, first).color, "yellow")
        self.assertEqual([(item.id, item.code_id, item.excerpt_id, item.origin, item.source_query)
                          for item in self.codings()], before_codings)
        self.assertEqual([item.id for item in db.session.scalars(
            select(QualitativeExcerpt).order_by(QualitativeExcerpt.start_offset))], before_excerpts)
        self.assertEqual(db.session.get(QualitativeMemo, memo.id).excerpt_id, first_excerpt.id)
        excerpts = page_excerpts(self.analysis, self.document.id, 1)["excerpts"]
        self.assertEqual(sum(code["color"] == "yellow" for excerpt in excerpts
                             for code in excerpt["codes"] if code["id"] == first), 9)
        self.assertEqual(next(code["color"] for code in excerpts[0]["codes"]
                              if code["id"] == second), "green")
        self.assertEqual(self.post("apply_codes", {"code_ids": [first]}, selection={
            "start": positions[9], "end": positions[9] + 4,
            "selected_text": "raça", "page_hash": page["sha256"]}).status_code, 201)
        excerpts = page_excerpts(self.analysis, self.document.id, 1)["excerpts"]
        self.assertEqual(sum(code["id"] == first and code["color"] == "yellow"
                             for excerpt in excerpts for code in excerpt["codes"]), 10)
        self.assertEqual(sum(code["id"] == second and code["color"] == "green"
                             for excerpt in excerpts for code in excerpt["codes"]), 1)
        self.assertIn("--qualitative-code-color: #FDE68A", self.client.get(self.base_url).get_data(as_text=True))
        self.client.post("/logout", data={"csrf_token": self.csrf})
        login(self.client)
        self.assertIn("--qualitative-code-color: #FDE68A", self.client.get(self.base_url).get_data(as_text=True))

    def test_invalid_colors_and_ownership_are_rejected(self):
        identifier = self.create_code("Privado")
        for invalid in ("javascript:alert(1)", "red; background:url(x)", "#XYZXYZ", "#2563EB", None, 42):
            with self.subTest(color=invalid):
                self.assertEqual(self.patch_code(identifier, {"color": invalid}).status_code, 400)
        self.assertIsNone(db.session.get(QualitativeCode, identifier).color)
        self.assertEqual(self.client.patch(self.code_url(identifier), json={"color": "red"}).status_code, 400)
        other = create_user(name="Outro", email="color-other@example.org")
        db.session.add(UserToolOverride(user_id=other.id, tool_id=QUALITATIVE_TOOL, decision="allow"))
        db.session.commit()
        self.client.post("/logout", data={"csrf_token": self.csrf})
        login(self.client, email="color-other@example.org")
        token = csrf_from(self.client.get("/"))
        self.assertIn(self.client.patch(self.code_url(identifier), json={"color": "red"},
                                        headers={"X-CSRFToken": token}).status_code, (403, 404))
        self.assertIsNone(db.session.get(QualitativeCode, identifier).color)

    def test_successive_changes_and_both_name_color_orders_keep_id(self):
        identifier = self.create_code("Raça")
        for color in ("blue", "yellow", "green", "purple", "blue"):
            self.assertEqual(self.patch_code(identifier, {"color": color}).json["codes"][0]["color"], color)
        self.assertEqual(self.patch_code(identifier, {"name": "Relações raciais"}).status_code, 200)
        self.assertEqual((db.session.get(QualitativeCode, identifier).name,
                          db.session.get(QualitativeCode, identifier).color), ("Relações raciais", "blue"))
        self.assertEqual(self.patch_code(identifier, {"color": "yellow"}).status_code, 200)
        self.assertEqual(self.patch_code(identifier, {"name": "Questão racial"}).status_code, 200)
        code = db.session.get(QualitativeCode, identifier)
        self.assertEqual((code.id, code.name, code.color), (identifier, "Questão racial", "yellow"))

    def _automatic_color_roundtrip(self, *, query, mode, text):
        self.fixture_pages([[text]])
        first = self.automatic(q=query, mode=mode)
        identifier = first.json["terms"][0]["code_id"]
        self.assertIsNotNone(identifier)
        before = [(item.id, item.code_id, item.excerpt_id, item.source_query) for item in self.codings()]
        self.assertEqual(self.patch_code(identifier, {"color": "purple"}).status_code, 200)
        repeated = self.automatic(q=query, mode=mode)
        self.assertEqual(repeated.json["terms"][0]["code_id"], identifier)
        self.assertEqual(repeated.json["summary"]["created"], 0)
        self.assertEqual([(item.id, item.code_id, item.excerpt_id, item.source_query)
                          for item in self.codings()], before)
        self.assertEqual(db.session.get(QualitativeCode, identifier).color, "purple")
        self.assertEqual(db.session.query(QualitativeCode).count(), 1)

    def test_literal_automatic_after_color(self):
        self._automatic_color_roundtrip(query="raça", mode="literal", text="raça raça")

    def test_lexical_automatic_after_color(self):
        self._automatic_color_roundtrip(query="professor", mode="lexical",
                                        text="professor professores professoras")

    def test_semantic_automatic_after_color(self):
        with patch("platform_core.qualitative_expanded_search._model", return_value=FakeModel(lambda _: [1, 0])):
            self._automatic_color_roundtrip(query="iniquidade", mode="semantic",
                                            text="Barreiras institucionais persistem.")

    def test_contextual_rejection_stays_bound_to_code_after_color_change(self):
        self.fixture_pages([["raça"]])
        first = self.automatic(q="raça", mode="literal", contextual_rejection_enabled=True)
        identifier = first.json["terms"][0]["code_id"]
        self.assertEqual(self.remove(self.codings()[0]).status_code, 200)
        rejection = db.session.scalar(select(QualitativeRejection))
        self.assertEqual(self.patch_code(identifier, {"color": "red"}).status_code, 200)
        repeated = self.automatic(q="raça", mode="literal")
        self.assertEqual((repeated.json["summary"]["created"],
                          repeated.json["summary"]["rejected"]), (0, 1))
        self.assertEqual((rejection.code_id, rejection.query), (identifier, "raça"))

    def test_pastel_palette_and_legacy_hex_are_resolved_centrally(self):
        self.assertEqual(len(QUALITATIVE_CODE_PALETTE), 9)
        self.assertEqual(code_color_style(None), code_color_style("blue"))
        for old_hex, identifier in LEGACY_CODE_COLORS.items():
            with self.subTest(identifier=identifier):
                style = code_color_style(old_hex.lower())
                self.assertEqual(style["color"], identifier)
                self.assertEqual(style["color_hex"], QUALITATIVE_CODE_PALETTE[identifier][1])
                self.assertEqual(style["color_text"], "#000000")
                self.assertNotEqual(style["color_hex"].upper(), old_hex)
        self.assertEqual(code_color_style("not-a-color"), code_color_style(None))

    def test_report_and_excel_use_current_color_only_on_code_cells(self):
        first = self.create_code("Raça")
        second = self.create_code("Racismo")
        self.assertEqual(self.patch_code(second, {"color": "yellow"}).status_code, 200)
        applied = self.post("apply_codes", {"code_ids": [first, second]})
        self.assertEqual(applied.status_code, 201)
        report_url = f"/analise-qualitativa/bases/{self.analysis.id}/relatorio-codificacao"

        def exported():
            response = self.client.get(f"{report_url}.xlsx")
            self.assertEqual(response.status_code, 200)
            workbook = load_workbook(io.BytesIO(response.data))
            self.assertEqual(workbook.sheetnames, ["Codificação"])
            sheet = workbook.active
            self.assertEqual(sheet.max_column, 5)
            self.assertEqual(sheet.max_row, 3)
            self.assertEqual([cell.value for cell in sheet[1]],
                             ["Código", "Documento", "Página", "Trecho codificado", "Memo contextual"])
            self.assertTrue(all(cell.alignment.wrap_text is not True for row in sheet.iter_rows(min_row=2)
                                for cell in row))
            self.assertTrue(all(cell.fill.patternType is None for row in sheet.iter_rows(min_row=2)
                                for cell in row[1:]))
            return sheet

        report = self.client.get(report_url).get_data(as_text=True)
        self.assertIn('style="--qualitative-code-color: #93C5FD; --qualitative-code-text: #000000">Raça</span>', report)
        self.assertIn('style="--qualitative-code-color: #FDE68A; --qualitative-code-text: #000000">Racismo</span>', report)
        sheet = exported()
        cells = {sheet.cell(row, 1).value: sheet.cell(row, 1) for row in (2, 3)}
        self.assertEqual(cells["Raça"].fill.fgColor.rgb, "FF93C5FD")
        self.assertEqual(cells["Racismo"].fill.fgColor.rgb, "FFFDE68A")
        self.assertEqual(cells["Raça"].font.color.rgb, "FF000000")
        self.assertEqual(cells["Racismo"].font.color.rgb, "FF000000")
        self.assertTrue(all(cell.fill.patternType == "solid" for cell in cells.values()))
        other_values = sorted(tuple(cell.value for cell in row[1:]) for row in sheet.iter_rows(min_row=2))

        self.assertEqual(self.patch_code(first, {"name": "Relações raciais"}).status_code, 200)
        self.assertEqual(self.patch_code(first, {"color": "purple"}).status_code, 200)
        report = self.client.get(report_url).get_data(as_text=True)
        self.assertIn('style="--qualitative-code-color: #C4B5FD; --qualitative-code-text: #000000">Relações raciais</span>', report)
        sheet = exported()
        cells = {sheet.cell(row, 1).value: sheet.cell(row, 1) for row in (2, 3)}
        self.assertEqual(cells["Relações raciais"].fill.fgColor.rgb, "FFC4B5FD")
        self.assertEqual(cells["Racismo"].fill.fgColor.rgb, "FFFDE68A")
        self.assertEqual(sorted(tuple(cell.value for cell in row[1:]) for row in sheet.iter_rows(min_row=2)),
                         other_values)
        self.assertNotIn("Cor", [cell.value for cell in sheet[1]])
        self.assertEqual({coding.code_id for coding in self.codings()}, {first, second})

        db.session.get(QualitativeCode, first).color = "#2563eb"
        db.session.commit()
        self.assertEqual(next(code["color_hex"] for excerpt in page_excerpts(
            self.analysis, self.document.id, 1)["excerpts"] for code in excerpt["codes"]
            if code["id"] == first), "#93C5FD")
        self.assertIn("--qualitative-code-color: #93C5FD",
                      self.client.get(self.base_url).get_data(as_text=True))
        self.assertIn('style="--qualitative-code-color: #93C5FD; --qualitative-code-text: #000000">Relações raciais</span>',
                      self.client.get(report_url).get_data(as_text=True))
        sheet = exported()
        cells = {sheet.cell(row, 1).value: sheet.cell(row, 1) for row in (2, 3)}
        self.assertEqual(cells["Relações raciais"].fill.fgColor.rgb, "FF93C5FD")
        self.assertEqual(self.patch_code(first, {"color": "green"}).json["codes"][0]["color"], "green")
        self.assertEqual(db.session.get(QualitativeCode, first).color, "green")


if __name__ == "__main__":
    unittest.main()
