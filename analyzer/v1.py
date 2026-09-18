from pathlib import Path
from collections import Counter, OrderedDict
import re
import shutil
import unicodedata

import pandas as pd
import pymupdf
import pytesseract
from PIL import Image
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .auditoria import adicionar_aba_parametros, gerar_ids_resultado


# O módulo fica em ``analyzer/``; a pasta de trabalho continua sendo a raiz
# do projeto, como no script de referência.
PASTA = Path(__file__).resolve().parent.parent
ARQUIVO_TERMOS = PASTA / "termos.txt"

# OCR só entra em ação quando a página não oferece texto extraível suficiente.
MIN_CARACTERES_EXTRAIVEIS = 80
OCR_DPI = 250


def _emitir_progresso(progress_callback, **evento):
    """Emite somente metadados de trabalho já concluído, quando solicitado."""
    if progress_callback is not None:
        progress_callback(evento)

STOPWORDS = {
    "a", "à", "ao", "aos", "as", "às", "o", "os", "um", "uma", "uns", "umas",
    "de", "da", "das", "do", "dos", "em", "na", "nas", "no", "nos", "por",
    "para", "com", "sem", "sob", "sobre", "entre", "até", "e", "ou", "mas",
    "que", "se", "não", "como", "quando", "onde", "qual", "quais", "quem",
    "porque", "pois", "também", "ainda", "ser", "estar", "ter", "foi", "foram",
    "era", "eram", "são", "seu", "seus", "sua", "suas", "ele", "ela", "eles",
    "elas", "isso", "isto", "esse", "essa", "este", "esta", "esses", "essas",
    "mais", "menos", "muito", "muita", "muitos", "muitas", "pode", "podem",
    "cada", "todo", "toda", "todos", "todas"
}

CONTEXTOS_SOCIOLOGICOS = {
    "Cultura e identidade": [
        "cultura", "identidade", "representacao", "simbolo", "tradicao",
        "regionalismo", "patrimonio", "costume", "pertencimento"
    ],
    "Migração e mobilidade": [
        "migracao", "migrante", "imigracao", "emigracao", "retirante",
        "deslocamento", "mobilidade", "exodo", "fluxo migratorio"
    ],
    "Desigualdades sociais e regionais": [
        "desigualdade", "pobreza", "renda", "exclusao", "vulnerabilidade",
        "regiao", "regional", "concentracao de renda"
    ],
    "Trabalho e relações de produção": [
        "trabalho", "trabalhador", "emprego", "salario", "producao",
        "capital", "industria", "operario", "mercado de trabalho"
    ],
    "Questão agrária e mundo rural": [
        "agrario", "agricultura", "latifundio", "terra", "rural", "campones",
        "fazenda", "engenho", "pecuaria", "reforma agraria"
    ],
    "Urbanização e território": [
        "urbanizacao", "urbano", "cidade", "metropole", "territorio",
        "periferia", "segregacao", "espaco urbano"
    ],
    "Estado, poder e políticas públicas": [
        "estado", "governo", "politica publica", "poder", "instituicao",
        "administracao", "sudene", "dnocs"
    ],
    "Movimentos sociais e ação coletiva": [
        "movimento social", "mobilizacao", "protesto", "acao coletiva",
        "associacao", "sindicato", "resistencia"
    ],
    "Raça e relações étnico-raciais": [
        "raca", "racial", "racismo", "etnico", "etnia", "negro", "negra",
        "indigena", "branquitude", "discriminacao racial"
    ],
    "Gênero e sexualidade": [
        "genero", "mulher", "mulheres", "sexualidade", "feminismo",
        "masculinidade", "patriarcado"
    ],
    "Religião": [
        "religiao", "religioso", "igreja", "fe", "catolico", "evangelico",
        "romaria", "beato"
    ],
    "Violência e conflitos sociais": [
        "violencia", "conflito", "guerra", "crime", "repressao", "cangaco",
        "coronelismo", "canudos"
    ],
    "Educação": [
        "educacao", "escola", "ensino", "professor", "estudante",
        "universidade", "escolar"
    ],
    "Meio ambiente e sociedade": [
        "meio ambiente", "ambiental", "seca", "semiarido", "caatinga",
        "clima", "agua", "desertificacao"
    ],
    "População e demografia": [
        "populacao", "demografia", "natalidade", "mortalidade", "densidade",
        "crescimento populacional"
    ],
    "Desenvolvimento regional": [
        "desenvolvimento regional", "subdesenvolvimento", "sudene",
        "desigualdade regional", "industrializacao", "infraestrutura"
    ],
    "História e formação social": [
        "historia", "historico", "colonial", "escravidao", "republica",
        "formacao social", "canudos", "cangaco"
    ],
    "Economia": [
        "economia", "economico", "mercado", "renda", "producao",
        "consumo", "industria", "agricultura"
    ],
}


# ============================================================
# TEXTO E TERMOS
# ============================================================

def normalizar(texto):
    texto = unicodedata.normalize("NFKD", texto or "")
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return texto.casefold()


def limpar_texto(texto):
    if not texto:
        return ""
    texto = re.sub(r"(?<=\w)-\s*\n\s*(?=\w)", "", texto)
    texto = texto.replace("\r", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", texto).strip()


def tokenizar(texto):
    return re.findall(r"\b[a-z][a-z\-]{2,}\b", normalizar(texto))


def gerar_variacoes_palavra(palavra):
    """
    Gera variações flexionais simples de número e, quando a terminação
    permite uma inferência razoável, de gênero.

    Exemplos:
      baiano -> baiano, baiana, baianos, baianas
      nordestino -> nordestino, nordestina, nordestinos, nordestinas
      cearense -> cearense, cearenses
      jagunço -> jagunço, jagunços
      sertão -> sertão, sertões, sertãos, sertães

    A busca continua sem diferenciar maiúsculas/minúsculas ou acentos.
    """

    p = palavra.strip().casefold()

    if not p:
        return set()

    formas = {p}

    # --------------------------------------------------------
    # TENTA RECUPERAR O SINGULAR QUANDO A ENTRADA ESTÁ NO PLURAL
    # --------------------------------------------------------

    singulares = {p}

    if p.endswith(("ões", "ãos", "ães")) and len(p) > 3:
        singulares.add(p[:-3] + "ão")

    elif p.endswith("ns") and len(p) > 3:
        singulares.add(p[:-2] + "m")

    elif p.endswith("is") and len(p) > 3:
        # regional -> regionais
        singulares.add(p[:-2] + "l")

    elif p.endswith(("res", "zes")) and len(p) > 3:
        # trabalhador(es), juiz(es)
        singulares.add(p[:-2])

    elif p.endswith("ses") and len(p) > 4:
        singulares.add(p[:-1])

    elif p.endswith(("os", "as")) and len(p) > 3:
        singulares.add(p[:-1])

    elif p.endswith("es") and len(p) > 3:
        # cearenses -> cearense; professores -> professor
        candidato_sem_s = p[:-1]
        candidato_sem_es = p[:-2]

        singulares.add(candidato_sem_s)

        if candidato_sem_es.endswith(("r", "z", "s")):
            singulares.add(candidato_sem_es)

    elif p.endswith("s") and len(p) > 3:
        singulares.add(p[:-1])

    formas.update(singulares)

    # --------------------------------------------------------
    # GÊNERO: APENAS TERMINAÇÕES COMUNS E RELATIVAMENTE SEGURAS
    # --------------------------------------------------------

    pares_genero = [
        ("ano", "ana"),
        ("ino", "ina"),
        ("eiro", "eira"),
        ("oso", "osa"),
        ("ivo", "iva"),
        ("ico", "ica"),
        ("ário", "ária"),
        ("ario", "aria"),
        ("ório", "ória"),
        ("orio", "oria"),
        ("dor", "dora"),
        ("tor", "tora"),
    ]

    novas = set(formas)

    for forma in list(formas):

        for masc, fem in pares_genero:

            if forma.endswith(masc) and len(forma) > len(masc):
                novas.add(
                    forma[:-len(masc)] + fem
                )

            if forma.endswith(fem) and len(forma) > len(fem):
                novas.add(
                    forma[:-len(fem)] + masc
                )

        # Gentílicos/adjetivos em -ês / -esa.
        if forma.endswith("ês") and len(forma) > 3:
            novas.add(
                forma[:-2] + "esa"
            )

        if forma.endswith("esa") and len(forma) > 4:
            novas.add(
                forma[:-3] + "ês"
            )

    formas.update(novas)

    # --------------------------------------------------------
    # PLURALIZAÇÃO DAS FORMAS GERADAS
    # --------------------------------------------------------

    novas = set(formas)

    for forma in list(formas):

        # Não tenta pluralizar algo que já parece plural.
        if forma.endswith(("ões", "ãos", "ães")):
            continue

        if forma.endswith("ão"):
            base = forma[:-2]

            # As três possibilidades existem em português.
            # Manter as três evita perder ocorrências; a regex só contará
            # a forma caso ela realmente apareça no PDF.
            novas.update({
                base + "ões",
                base + "ãos",
                base + "ães",
            })

        elif forma.endswith("m"):
            novas.add(
                forma[:-1] + "ns"
            )

        elif forma.endswith("l"):
            novas.add(
                forma[:-1] + "is"
            )

        elif forma.endswith(("r", "z")):
            novas.add(
                forma + "es"
            )

        elif forma.endswith(("a", "e", "i", "o", "u", "á", "é", "í", "ó", "ú")):
            novas.add(
                forma + "s"
            )

    formas.update(novas)

    return {
        normalizar(forma)
        for forma in formas
        if forma.strip()
    }


def gerar_variacoes_termo(termo):
    """
    Para expressões com mais de uma palavra, flexiona apenas a última.
    Exemplo:
      migrante nordestino
      -> migrante nordestino / nordestina / nordestinos / nordestinas
    """

    termo = termo.strip()

    if not termo:
        return set()

    partes = termo.split()

    if len(partes) == 1:
        return gerar_variacoes_palavra(
            partes[0]
        )

    prefixo = " ".join(
        partes[:-1]
    )

    ultima = partes[-1]

    variantes_ultima = gerar_variacoes_palavra(
        ultima
    )

    return {
        normalizar(
            prefixo + " " + variante
        )
        for variante in variantes_ultima
    }


def criar_regex(termo):
    variantes = gerar_variacoes_termo(
        termo
    )

    if not variantes:
        variantes = {
            normalizar(termo)
        }

    # Mais longas primeiro, para evitar que uma forma curta
    # interfira na alternativa mais específica.
    alternativas = []

    for variante in sorted(
        variantes,
        key=len,
        reverse=True
    ):
        partes = variante.split()

        alternativas.append(
            r"\s+".join(
                re.escape(parte)
                for parte in partes
            )
        )

    padrao = "(?:" + "|".join(
        alternativas
    ) + ")"

    return re.compile(
        rf"(?<!\w){padrao}(?!\w)",
        re.IGNORECASE
    )


def contar_ocorrencias(texto, termo):
    """
    Conta o termo e suas variações automáticas de gênero/número.
    """
    return len(
        criar_regex(
            termo
        ).findall(
            normalizar(texto)
        )
    )


def carregar_termos_referencia():
    """
    termos.txt funciona apenas como dicionário opcional de categorias.
    A lista de termos a pesquisar SEMPRE é pedida ao usuário.
    """
    itens = []

    if not ARQUIVO_TERMOS.exists():
        return itens

    with ARQUIVO_TERMOS.open("r", encoding="utf-8-sig") as arq:
        for linha in arq:
            linha = linha.strip()

            if not linha or linha.startswith("#") or "|" not in linha:
                continue

            categoria, termo = linha.split("|", 1)
            categoria = categoria.strip()
            termo = termo.strip()

            if categoria and termo:
                itens.append({
                    "categoria": categoria,
                    "termo": termo
                })

    return itens


def solicitar_termos(termos_referencia):
    mapa_categorias = {}

    for item in termos_referencia:
        chave = normalizar(item["termo"]).strip()
        if chave not in mapa_categorias:
            mapa_categorias[chave] = item["categoria"]

    print("\nDigite a lista de termos que deseja pesquisar.")
    print(
        "A busca considera automaticamente variações simples "
        "de singular/plural e masculino/feminino."
    )
    print("Separe os termos por ponto e vírgula (;).")
    print("\nExemplo:")
    print("Nordeste; nordestino; Alagoas; sertão; cangaço")
    print("\nOpcionalmente, você pode informar uma categoria manual:")
    print("REGIÃO|Nordeste; HISTÓRIA|cangaço")

    while True:
        entrada = input("\nTermos: ").strip()

        if not entrada:
            print("Informe pelo menos um termo.")
            continue

        partes = [p.strip() for p in entrada.split(";") if p.strip()]
        termos = []
        vistos = set()

        for parte in partes:
            categoria_manual = None
            termo = parte

            if "|" in parte:
                possivel_categoria, possivel_termo = parte.split("|", 1)
                if possivel_categoria.strip() and possivel_termo.strip():
                    categoria_manual = possivel_categoria.strip()
                    termo = possivel_termo.strip()

            termo = termo.strip()
            if not termo:
                continue

            chave = normalizar(termo).strip()
            if chave in vistos:
                continue

            vistos.add(chave)

            categoria = (
                categoria_manual
                or mapa_categorias.get(chave)
                or "TERMO INFORMADO"
            )

            termos.append({
                "categoria": categoria,
                "termo": termo
            })

        if termos:
            print(f"\nForam informados {len(termos)} termo(s):")
            for item in termos:
                print(f"- {item['termo']} [{item['categoria']}]")
            return termos

        print("Nenhum termo válido foi identificado. Tente novamente.")


# ============================================================
# OCR
# ============================================================

def configurar_tesseract():
    candidatos = [
        shutil.which("tesseract"),
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]

    for candidato in candidatos:
        if candidato and Path(candidato).exists():
            pytesseract.pytesseract.tesseract_cmd = str(candidato)
            return str(candidato)

    return None


def escolher_idioma_ocr():
    try:
        langs = set(pytesseract.get_languages(config=""))
    except Exception:
        return "eng"

    if "por" in langs and "eng" in langs:
        return "por+eng"
    if "por" in langs:
        return "por"
    if "eng" in langs:
        return "eng"

    return next(iter(langs), "eng")


def pagina_precisa_ocr(pagina):
    texto = pagina.get_text("text") or ""
    letras = re.findall(r"[A-Za-zÀ-ÿ]", texto)
    return len(letras) < MIN_CARACTERES_EXTRAIVEIS


def ocr_pagina(pagina, idioma):
    escala = OCR_DPI / 72
    pix = pagina.get_pixmap(
        matrix=pymupdf.Matrix(escala, escala),
        alpha=False
    )
    imagem = Image.frombytes(
        "RGB",
        [pix.width, pix.height],
        pix.samples
    )

    dados = pytesseract.image_to_data(
        imagem,
        lang=idioma,
        config="--psm 3",
        output_type=pytesseract.Output.DICT
    )

    grupos = OrderedDict()
    n = len(dados["text"])

    for i in range(n):
        palavra = (dados["text"][i] or "").strip()
        if not palavra:
            continue

        try:
            conf = float(dados["conf"][i])
        except Exception:
            conf = -1

        if conf >= 0 and conf < 25:
            continue

        chave = (
            dados.get("block_num", [0] * n)[i],
            dados.get("par_num", [0] * n)[i],
        )

        if chave not in grupos:
            grupos[chave] = {
                "palavras": [],
                "left": [],
                "top": [],
                "right": [],
                "bottom": [],
            }

        left = dados.get("left", [0] * n)[i]
        top = dados.get("top", [0] * n)[i]
        width = dados.get("width", [0] * n)[i]
        height = dados.get("height", [0] * n)[i]

        grupos[chave]["palavras"].append(palavra)
        grupos[chave]["left"].append(left)
        grupos[chave]["top"].append(top)
        grupos[chave]["right"].append(left + width)
        grupos[chave]["bottom"].append(top + height)

    blocos = []

    for grupo in grupos.values():
        texto = limpar_texto(" ".join(grupo["palavras"]))

        if not texto:
            continue

        bbox = (
            min(grupo["left"]),
            min(grupo["top"]),
            max(grupo["right"]),
            max(grupo["bottom"]),
        )

        blocos.append({
            "texto": texto,
            "tamanho_max": None,
            "metodo": "OCR",
            "bbox": bbox,
            "image_bboxes": [],
        })

    return blocos


# ============================================================
# EXTRAÇÃO E ESTRUTURA DO PDF
# ============================================================

def estimar_fonte_corpo(documento, progress_callback=None):
    tamanhos = Counter()
    total_paginas = len(documento)
    passo_progresso = max(1, total_paginas // 100)

    for numero_pagina, pagina in enumerate(documento, start=1):
        if pagina_precisa_ocr(pagina):
            dados = None
        else:
            dados = pagina.get_text("dict")

            for bloco in dados.get("blocks", []):
                if bloco.get("type") != 0:
                    continue

                for linha in bloco.get("lines", []):
                    for span in linha.get("spans", []):
                        texto = (span.get("text") or "").strip()

                        if texto:
                            tamanho = round(span.get("size", 0), 1)
                            tamanhos[tamanho] += len(texto)

        if (
            numero_pagina == total_paginas
            or numero_pagina % passo_progresso == 0
        ):
            _emitir_progresso(
                progress_callback,
                fase="estimando_fonte",
                etapa="Lendo PDFs…",
                pagina_atual=numero_pagina,
                paginas_total=total_paginas,
            )

    return tamanhos.most_common(1)[0][0] if tamanhos else 11.0


def extrair_blocos_texto(pagina, fonte_corpo):
    dados = pagina.get_text("dict")

    image_bboxes = [
        tuple(bloco.get("bbox", [0, 0, 0, 0]))
        for bloco in dados.get("blocks", [])
        if bloco.get("type") == 1
    ]

    blocos = [
        b
        for b in dados.get("blocks", [])
        if b.get("type") == 0
    ]

    blocos.sort(
        key=lambda b: (
            b.get("bbox", [0, 0, 0, 0])[1],
            b.get("bbox", [0, 0, 0, 0])[0]
        )
    )

    saida = []

    for bloco in blocos:
        linhas = []
        tamanhos = []

        for linha in bloco.get("lines", []):
            partes = []

            for span in linha.get("spans", []):
                txt = span.get("text", "")

                if txt.strip():
                    partes.append(txt)
                    tamanhos.append(
                        span.get("size", fonte_corpo)
                    )

            if partes:
                linhas.append(" ".join(partes))

        texto = limpar_texto(" ".join(linhas))

        if not texto or re.fullmatch(r"\d+", texto):
            continue

        saida.append({
            "texto": texto,
            "tamanho_max": (
                max(tamanhos)
                if tamanhos
                else fonte_corpo
            ),
            "metodo": "texto",
            "bbox": tuple(
                bloco.get("bbox", [0, 0, 0, 0])
            ),
            "image_bboxes": image_bboxes,
        })

    return saida


def classificar_estrutura(
    texto,
    tamanho_fonte=None,
    fonte_corpo=11.0
):
    t = texto.strip()
    n = normalizar(t)

    if not t or len(t) > 240 or len(t.split()) > 35:
        return None

    if re.match(r"^unidade\s+([0-9ivxlcdm]+)\b", n):
        return "unidade"

    if re.match(r"^capitulo\s+([0-9ivxlcdm]+)\b", n):
        return "capitulo"

    if re.match(r"^(secao|seção)\s+", t, re.IGNORECASE):
        return "secao"

    if re.match(r"^(subsecao|subseção)\s+", t, re.IGNORECASE):
        return "subsecao"

    m = re.match(r"^(\d+(?:\.\d+)+)\s+", t)

    if m:
        pontos = m.group(1).count(".")
        return "subsecao" if pontos >= 2 else "secao"

    if re.match(r"^\d+\s+[A-Za-zÀ-ÿ]", t):
        if tamanho_fonte is None or tamanho_fonte >= fonte_corpo + 0.5:
            return "capitulo"

    if tamanho_fonte is not None:
        letras = [c for c in t if c.isalpha()]
        prop_maiusculas = (
            sum(c.isupper() for c in letras) / len(letras)
            if letras else 0
        )

        if (
            tamanho_fonte >= fonte_corpo + 3
            and len(t.split()) <= 15
        ):
            return "capitulo"

        if (
            tamanho_fonte >= fonte_corpo + 1.7
            and prop_maiusculas >= 0.75
            and len(t.split()) <= 18
        ):
            return "secao"

    return None


def sobreposicao_horizontal(b1, b2):
    x0 = max(b1[0], b2[0])
    x1 = min(b1[2], b2[2])

    if x1 <= x0:
        return 0.0

    largura1 = max(1, b1[2] - b1[0])
    largura2 = max(1, b2[2] - b2[0])

    return (x1 - x0) / min(largura1, largura2)


def proximo_de_imagem(bbox_texto, image_bboxes):
    if not bbox_texto or not image_bboxes:
        return False

    for img in image_bboxes:
        overlap = sobreposicao_horizontal(
            bbox_texto,
            img
        )

        distancia_abaixo = bbox_texto[1] - img[3]
        distancia_acima = img[1] - bbox_texto[3]

        distancia = min(
            abs(distancia_abaixo),
            abs(distancia_acima)
        )

        if overlap >= 0.25 and distancia <= 90:
            return True

    return False


def classificar_tipo_ocorrencia(
    texto,
    tipo_estrutura=None,
    bbox=None,
    image_bboxes=None
):
    """
    Valores principais:
      - parágrafo
      - legenda de fotografia/imagem
      - outro local
    """
    t = texto.strip()
    n = normalizar(t)
    palavras = t.split()

    if tipo_estrutura:
        return "outro local"

    padroes_legenda = [
        r"^(figura|fig\.|foto|fotografia|imagem)\s*\d*[\s:.\-–—]",
        r"^(fonte|credito|creditos|crédito|créditos)\s*[:.\-–—]",
        r"^©\s*",
    ]

    if any(
        re.match(padrao, n, re.IGNORECASE)
        for padrao in padroes_legenda
    ):
        return "legenda de fotografia/imagem"

    if (
        len(palavras) <= 80
        and proximo_de_imagem(
            bbox,
            image_bboxes or []
        )
    ):
        return "legenda de fotografia/imagem"

    # Corpo textual típico.
    if (
        len(palavras) >= 18
        or re.search(r"[.!?][”\"']?$", t)
    ):
        return "parágrafo"

    return "outro local"


def extrair_pdf(caminho, progress_callback=None):
    documento = pymupdf.open(caminho)
    total_paginas = len(documento)
    _emitir_progresso(
        progress_callback,
        fase="leitura_pdf",
        etapa="Lendo PDFs…",
        pagina_atual=0,
        paginas_total=total_paginas,
    )
    fonte_corpo = (
        estimar_fonte_corpo(documento)
        if progress_callback is None
        else estimar_fonte_corpo(documento, progress_callback=progress_callback)
    )

    tesseract = configurar_tesseract()
    idioma_ocr = (
        escolher_idioma_ocr()
        if tesseract
        else None
    )

    unidade = ""
    capitulo = ""
    secao = ""
    subsecao = ""

    registros = []
    paginas_ocr = 0
    paginas_texto = 0

    for numero_pagina, pagina in enumerate(
        documento,
        start=1
    ):
        usar_ocr = pagina_precisa_ocr(pagina)

        if usar_ocr:
            if not tesseract:
                documento.close()
                raise RuntimeError(
                    "Uma ou mais páginas precisam de OCR, "
                    "mas o Tesseract não foi encontrado."
                )

            blocos = ocr_pagina(
                pagina,
                idioma_ocr
            )
            paginas_ocr += 1

        else:
            blocos = extrair_blocos_texto(
                pagina,
                fonte_corpo
            )
            paginas_texto += 1

        try:
            rotulo = (
                pagina.get_label()
                or ""
            ).strip()
        except Exception:
            rotulo = ""

        pagina_saida = (
            rotulo
            if rotulo
            else numero_pagina
        )

        for bloco in blocos:
            texto = bloco["texto"]

            tipo_estrutura = classificar_estrutura(
                texto,
                bloco["tamanho_max"],
                fonte_corpo
            )

            if tipo_estrutura == "unidade":
                unidade = texto
                capitulo = ""
                secao = ""
                subsecao = ""

            elif tipo_estrutura == "capitulo":
                capitulo = texto
                secao = ""
                subsecao = ""

            elif tipo_estrutura == "secao":
                secao = texto
                subsecao = ""

            elif tipo_estrutura == "subsecao":
                subsecao = texto

            tipo_ocorrencia = classificar_tipo_ocorrencia(
                texto,
                tipo_estrutura=tipo_estrutura,
                bbox=bloco.get("bbox"),
                image_bboxes=bloco.get(
                    "image_bboxes",
                    []
                )
            )

            registros.append({
                "pagina": pagina_saida,
                "pagina_pdf": numero_pagina,
                "unidade": unidade,
                "capitulo": capitulo,
                "secao": secao,
                "subsecao": subsecao,
                "texto": texto,
                "metodo": bloco["metodo"],
                "tipo_estrutura": tipo_estrutura or "",
                "tipo_ocorrencia": tipo_ocorrencia,
            })

        _emitir_progresso(
            progress_callback,
            fase="ocr" if usar_ocr else "paginas",
            etapa="Executando OCR…" if usar_ocr else "Analisando páginas…",
            pagina_atual=numero_pagina,
            paginas_total=total_paginas,
        )

    documento.close()

    return (
        registros,
        paginas_texto,
        paginas_ocr,
        idioma_ocr
    )


# ============================================================
# CONTEXTO DA OCORRÊNCIA
# ============================================================

def localizar_paragrafo_anterior(blocos, indice):
    for i in range(indice - 1, -1, -1):
        if blocos[i].get("tipo_ocorrencia") == "parágrafo":
            return blocos[i]["texto"]
    return ""


def localizar_paragrafo_posterior(blocos, indice):
    for i in range(indice + 1, len(blocos)):
        if blocos[i].get("tipo_ocorrencia") == "parágrafo":
            return blocos[i]["texto"]
    return ""


def identificar_contexto_sociologico(texto):
    n = normalizar(texto)
    pontuacoes = {}

    for contexto, palavras in CONTEXTOS_SOCIOLOGICOS.items():
        score = 0

        for palavra in palavras:
            p = normalizar(palavra)

            if p in n:
                score += n.count(p)

        if score:
            pontuacoes[contexto] = score

    if not pontuacoes:
        return "Outros / revisar"

    melhor = max(
        pontuacoes.values()
    )

    empatados = [
        contexto
        for contexto, score
        in pontuacoes.items()
        if score == melhor
    ]

    return " | ".join(
        empatados[:2]
    )


def extrair_frase_com_termo(texto, termo):
    frases = re.split(
        r"(?<=[.!?])\s+",
        limpar_texto(texto)
    )

    for frase in frases:
        if contar_ocorrencias(
            frase,
            termo
        ):
            return frase.strip()

    return limpar_texto(texto)[:700].strip()


# Explicações padronizadas utilizadas na coluna "Descrição".
# Elas não informam o tipo da ocorrência. Servem para transformar
# a classificação de contexto sociológico em uma descrição analítica.
DESCRICOES_CONTEXTOS = {
    "Cultura e identidade": {
        "tema": (
            "O texto aborda processos de construção de identidades, "
            "representações, símbolos, tradições e formas de pertencimento."
        ),
        "foco": (
            "O foco sociológico está na relação entre cultura, identidade, "
            "representação e produção social de sentidos."
        ),
    },
    "Migração e mobilidade": {
        "tema": (
            "O texto discute deslocamentos populacionais, migrações, "
            "mobilidade e as condições sociais associadas a esses movimentos."
        ),
        "foco": (
            "O foco sociológico está nos processos de mobilidade, nas causas "
            "sociais dos deslocamentos e em seus efeitos sobre grupos e territórios."
        ),
    },
    "Desigualdades sociais e regionais": {
        "tema": (
            "O texto discute desigualdades sociais e territoriais, destacando "
            "diferenças de renda, acesso a recursos, oportunidades e condições de vida."
        ),
        "foco": (
            "O foco sociológico está na produção e reprodução das desigualdades "
            "e na forma como elas se distribuem entre grupos sociais e regiões."
        ),
    },
    "Trabalho e relações de produção": {
        "tema": (
            "O texto aborda trabalho, relações de produção, emprego, mercado "
            "e formas de organização da atividade econômica."
        ),
        "foco": (
            "O foco sociológico está nas relações de trabalho, na organização "
            "da produção e nas posições ocupadas por diferentes grupos sociais."
        ),
    },
    "Questão agrária e mundo rural": {
        "tema": (
            "O texto discute o mundo rural, a propriedade da terra, a produção "
            "agrícola e as relações sociais constituídas no espaço agrário."
        ),
        "foco": (
            "O foco sociológico está nas relações de propriedade, trabalho e poder "
            "no campo e nos conflitos ligados à organização do espaço rural."
        ),
    },
    "Urbanização e território": {
        "tema": (
            "O texto aborda processos de urbanização, organização do território, "
            "segregação espacial e diferentes formas de ocupação das cidades."
        ),
        "foco": (
            "O foco sociológico está nas relações entre espaço, território, "
            "urbanização e desigualdades socioespaciais."
        ),
    },
    "Estado, poder e políticas públicas": {
        "tema": (
            "O texto discute o Estado, suas instituições e diferentes formas "
            "de exercício do poder, da autoridade e da regulação da vida social."
        ),
        "foco": (
            "O foco sociológico está nas relações entre Estado, poder, autoridade, "
            "legitimidade e regulação institucional da vida social."
        ),
    },
    "Movimentos sociais e ação coletiva": {
        "tema": (
            "O texto aborda mobilização coletiva, movimentos sociais, formas de "
            "organização e ações desenvolvidas por grupos em torno de demandas comuns."
        ),
        "foco": (
            "O foco sociológico está na ação coletiva, na organização de interesses "
            "e nas formas de mobilização e conflito social."
        ),
    },
    "Raça e relações étnico-raciais": {
        "tema": (
            "O texto discute relações étnico-raciais, processos de racialização, "
            "identidades raciais e formas de desigualdade ou discriminação."
        ),
        "foco": (
            "O foco sociológico está nas relações raciais, nos processos de "
            "classificação social e nas desigualdades produzidas ou reproduzidas racialmente."
        ),
    },
    "Gênero e sexualidade": {
        "tema": (
            "O texto aborda relações de gênero, sexualidade, papéis sociais e "
            "desigualdades associadas às classificações e expectativas de gênero."
        ),
        "foco": (
            "O foco sociológico está na construção social do gênero e da sexualidade "
            "e nas relações de poder que atravessam essas classificações."
        ),
    },
    "Religião": {
        "tema": (
            "O texto discute crenças, práticas, instituições e formas de organização "
            "religiosa presentes na vida social."
        ),
        "foco": (
            "O foco sociológico está no papel social da religião, em suas instituições "
            "e nas formas de pertencimento e organização que produz."
        ),
    },
    "Violência e conflitos sociais": {
        "tema": (
            "O texto aborda conflitos, formas de violência, repressão, resistência "
            "e disputas que envolvem diferentes grupos e instituições."
        ),
        "foco": (
            "O foco sociológico está nas relações de poder, nos conflitos sociais "
            "e nas formas pelas quais a violência é produzida, legitimada ou contestada."
        ),
    },
    "Educação": {
        "tema": (
            "O texto discute instituições educacionais, processos de ensino e "
            "socialização e as relações entre educação e sociedade."
        ),
        "foco": (
            "O foco sociológico está na educação como instituição social e em suas "
            "relações com socialização, desigualdades e reprodução ou transformação social."
        ),
    },
    "Meio ambiente e sociedade": {
        "tema": (
            "O texto aborda relações entre sociedade e ambiente, condições climáticas, "
            "recursos naturais e impactos sociais associados ao território."
        ),
        "foco": (
            "O foco sociológico está na relação entre organização social, território, "
            "meio ambiente e distribuição desigual de riscos e recursos."
        ),
    },
    "População e demografia": {
        "tema": (
            "O texto discute características e transformações da população, como "
            "crescimento, distribuição, densidade e dinâmica demográfica."
        ),
        "foco": (
            "O foco sociológico está nas mudanças populacionais e em sua relação "
            "com a organização social e territorial."
        ),
    },
    "Desenvolvimento regional": {
        "tema": (
            "O texto aborda desenvolvimento econômico e regional, infraestrutura, "
            "industrialização e desigualdades entre diferentes partes do território."
        ),
        "foco": (
            "O foco sociológico está nas diferenças de desenvolvimento entre regiões "
            "e nas condições históricas, econômicas e políticas que as produzem."
        ),
    },
    "História e formação social": {
        "tema": (
            "O texto mobiliza processos históricos para explicar a formação de "
            "instituições, relações sociais, grupos e estruturas da sociedade."
        ),
        "foco": (
            "O foco sociológico está na relação entre processos históricos e a "
            "constituição das estruturas e relações sociais."
        ),
    },
    "Economia": {
        "tema": (
            "O texto aborda produção, mercado, renda, consumo e formas de organização "
            "da atividade econômica."
        ),
        "foco": (
            "O foco sociológico está nas relações entre economia e sociedade e na "
            "forma como produção, distribuição e consumo organizam relações sociais."
        ),
    },
}


TERMOS_GEOGRAFICOS_NORDESTE = {
    normalizar(x)
    for x in [
        "Alagoas", "Bahia", "Ceará", "Maranhão", "Paraíba",
        "Pernambuco", "Piauí", "Rio Grande do Norte", "Sergipe",
        "Maceió", "Salvador", "Fortaleza", "São Luís",
        "João Pessoa", "Recife", "Teresina", "Natal", "Aracaju",
        "Nordeste",
    ]
}


TERMOS_E_CONTEXTOS_DIRETOS = {
    "seca": {
        "Meio ambiente e sociedade",
        "Desenvolvimento regional",
        "Desigualdades sociais e regionais",
        "História e formação social",
    },
    "sertao": {
        "Questão agrária e mundo rural",
        "Desenvolvimento regional",
        "Desigualdades sociais e regionais",
        "Meio ambiente e sociedade",
        "História e formação social",
    },
    "jagunco": {
        "Violência e conflitos sociais",
        "História e formação social",
        "Questão agrária e mundo rural",
    },
    "cangaco": {
        "Violência e conflitos sociais",
        "História e formação social",
    },
    "palmares": {
        "História e formação social",
        "Raça e relações étnico-raciais",
        "Violência e conflitos sociais",
    },
}


def contexto_principal(contexto_sociologico):
    if not contexto_sociologico:
        return ""

    return contexto_sociologico.split("|")[0].strip()


def termo_funciona_como_localizacao(termo, frase):
    termo_n = normalizar(termo)
    frase_n = normalizar(frase)

    partes = [
        re.escape(p)
        for p in termo_n.split()
    ]
    padrao_termo = r"\s+".join(partes)

    padrao = (
        rf"\b(?:em|no|na|nos|nas|de|do|da|dos|das|"
        rf"para|por|desde|ate)\s+(?:o\s+|a\s+)?{padrao_termo}\b"
    )

    return bool(
        re.search(
            padrao,
            frase_n,
            re.IGNORECASE
        )
    )


def papel_do_termo(
    termo,
    frase,
    contexto_sociologico,
    contexto_amplo
):
    termo_n = normalizar(termo)
    contexto = contexto_principal(
        contexto_sociologico
    )

    quantidade = contar_ocorrencias(
        contexto_amplo,
        termo
    )

    # Termos cujo próprio significado é diretamente associado
    # a determinados contextos sociológicos.
    if (
        termo_n in TERMOS_E_CONTEXTOS_DIRETOS
        and contexto in TERMOS_E_CONTEXTOS_DIRETOS[termo_n]
    ):
        return (
            f"A menção a {termo} integra diretamente a discussão "
            "desenvolvida no trecho."
        ), "central"

    # Estados, capitais e Nordeste.
    if termo_n in TERMOS_GEOGRAFICOS_NORDESTE:
        if termo_funciona_como_localizacao(
            termo,
            frase
        ):
            return (
                f"A menção a {termo} funciona como referência espacial, "
                "situando o local associado ao exemplo ou situação apresentada."
            ), "contextual"

        contextos_territoriais = {
            "Migração e mobilidade",
            "Desigualdades sociais e regionais",
            "Questão agrária e mundo rural",
            "Urbanização e território",
            "Meio ambiente e sociedade",
            "Desenvolvimento regional",
            "População e demografia",
        }

        if contexto in contextos_territoriais:
            return (
                f"A menção a {termo} participa da contextualização territorial "
                "do fenômeno social discutido."
            ), "central"

        return (
            f"A menção a {termo} aparece como referência geográfica utilizada "
            "para situar o exemplo apresentado."
        ), "contextual"

    if quantidade >= 3:
        return (
            f"O termo {termo} aparece de forma recorrente e participa diretamente "
            "do desenvolvimento da discussão."
        ), "central"

    return (
        f"A menção a {termo} aparece como elemento específico utilizado para "
        "contextualizar ou exemplificar a discussão."
    ), "contextual"


def tema_especifico_do_trecho(
    contexto_sociologico,
    contexto_amplo
):
    n = normalizar(contexto_amplo)
    contexto = contexto_principal(
        contexto_sociologico
    )

    # Regra específica para uma formulação clássica de Max Weber.
    if (
        contexto == "Estado, poder e políticas públicas"
        and "weber" in n
        and "violencia" in n
        and (
            "legitima" in n
            or "legitimidade" in n
        )
    ):
        return (
            "O texto discute uma definição clássica de Estado em Max Weber, "
            "associada ao monopólio do uso legítimo da violência e à "
            "legitimidade da coerção exercida institucionalmente."
        )

    dados = DESCRICOES_CONTEXTOS.get(
        contexto
    )

    if dados:
        return dados["tema"]

    return (
        "O termo aparece inserido em uma discussão sociológica mais ampla, "
        "cujo sentido deve ser interpretado a partir do trecho e de sua relação "
        "com o restante da seção."
    )


def foco_sociologico_descricao(
    contexto_sociologico
):
    contexto = contexto_principal(
        contexto_sociologico
    )

    dados = DESCRICOES_CONTEXTOS.get(
        contexto
    )

    if dados:
        return dados["foco"]

    return (
        "O foco sociológico não pôde ser determinado com segurança "
        "pelas regras automáticas e deve ser conferido manualmente."
    )


def descrever_ocorrencia(
    termo,
    texto_ocorrencia,
    contexto_amplo,
    contexto_sociologico
):
    """
    Produz uma descrição analítica do papel da menção no contexto.

    A descrição:
      1. explica como o termo participa do trecho;
      2. sintetiza o assunto sociológico em que aparece;
      3. explicita o foco sociológico da passagem;
      4. não informa a classificação "tipo de ocorrência".
    """
    frase = extrair_frase_com_termo(
        texto_ocorrencia,
        termo
    )

    frase = limpar_texto(
        frase
    )

    papel, centralidade = papel_do_termo(
        termo,
        frase,
        contexto_sociologico,
        contexto_amplo
    )

    tema = tema_especifico_do_trecho(
        contexto_sociologico,
        contexto_amplo
    )

    foco = foco_sociologico_descricao(
        contexto_sociologico
    )

    partes = [
        papel,
        tema,
        foco,
    ]

    if centralidade == "contextual":
        partes.append(
            f"Nesse contexto, a referência a {termo} funciona principalmente "
            "como elemento de contextualização ou exemplificação, e não como "
            "o foco central da discussão sociológica."
        )

    return " ".join(
        limpar_texto(p)
        for p in partes
        if p
    )

def id_livro(caminho):
    return caminho.stem


def analisar_pdf(caminho, termos, progress_callback=None):
    print(f"\nAnalisando: {caminho.name}")

    (
        blocos,
        paginas_texto,
        paginas_ocr,
        idioma_ocr
    ) = (
        extrair_pdf(caminho)
        if progress_callback is None
        else extrair_pdf(caminho, progress_callback=progress_callback)
    )

    ocorrencias = []
    total_blocos = len(blocos)
    passo_progresso = max(1, total_blocos // 100)
    _emitir_progresso(
        progress_callback,
        fase="busca_lexical",
        etapa="Executando busca lexical…",
        bloco_atual=0,
        blocos_total=total_blocos,
    )

    for indice, bloco in enumerate(blocos):
        atual = bloco["texto"]

        paragrafo_anterior = localizar_paragrafo_anterior(
            blocos,
            indice
        )

        paragrafo_posterior = localizar_paragrafo_posterior(
            blocos,
            indice
        )

        contexto_amplo = " ".join([
            bloco.get("unidade", ""),
            bloco.get("capitulo", ""),
            bloco.get("secao", ""),
            bloco.get("subsecao", ""),
            paragrafo_anterior,
            atual,
            paragrafo_posterior,
        ])

        for item in termos:
            termo = item["termo"]
            categoria = item["categoria"]

            quantidade = contar_ocorrencias(
                atual,
                termo
            )

            if quantidade == 0:
                continue

            tipo_ocorrencia = bloco[
                "tipo_ocorrencia"
            ]

            contexto_soc = identificar_contexto_sociologico(
                contexto_amplo
            )

            descricao = descrever_ocorrencia(
                termo,
                atual,
                contexto_amplo,
                contexto_soc
            )

            ocorrencias.append({
                "ID livro": id_livro(caminho),
                "Termo": termo,
                "Categoria do termo": categoria,
                "Parágrafo anterior": paragrafo_anterior,
                "Parágrafo do termo": atual,
                "Parágrafo posterior": paragrafo_posterior,
                "Unidade": bloco["unidade"],
                "Capítulo": bloco["capitulo"],
                "Seção": bloco["secao"],
                "Subseção": bloco["subsecao"],
                "Página": bloco["pagina"],
                "Tipo da ocorrência": tipo_ocorrencia,
                "Contexto sociológico da ocorrência": contexto_soc,
                "Descrição da ocorrência": descricao,
                "Validação manual": "revisar",
                "_quantidade_no_registro": quantidade,
                "_metodo": bloco["metodo"],
            })

        if (
            indice == total_blocos - 1
            or (indice + 1) % passo_progresso == 0
        ):
            _emitir_progresso(
                progress_callback,
                fase="busca_lexical",
                etapa="Executando busca lexical…",
                bloco_atual=indice + 1,
                blocos_total=total_blocos,
            )

    diagnostico = {
        "arquivo": caminho.name,
        "paginas_com_texto_extraivel": paginas_texto,
        "paginas_processadas_com_OCR": paginas_ocr,
        "idioma_OCR": (
            idioma_ocr
            or "não utilizado"
        ),
    }

    return ocorrencias, diagnostico


# ============================================================
# RESUMOS
# ============================================================

def criar_resumos(df_ocorrencias):
    """
    Cria apenas resumos que não utilizam a coluna
    'Categoria do termo'.

    A categoria pode continuar existindo internamente no script,
    mas não é exportada para nenhuma aba do Excel.
    """
    if df_ocorrencias.empty:
        resumo_termo = pd.DataFrame(
            columns=[
                "ID livro",
                "Termo",
                "Total de ocorrências"
            ]
        )

        resumo_tipo = pd.DataFrame(
            columns=[
                "ID livro",
                "Tipo da ocorrência",
                "Total de ocorrências"
            ]
        )

        return (
            resumo_termo,
            resumo_tipo
        )

    tmp = df_ocorrencias.copy()

    resumo_termo = (
        tmp.groupby(
            [
                "ID livro",
                "Termo"
            ],
            as_index=False
        )["_quantidade_no_registro"]
        .sum()
        .rename(
            columns={
                "_quantidade_no_registro":
                    "Total de ocorrências"
            }
        )
    )

    resumo_tipo = (
        tmp.groupby(
            [
                "ID livro",
                "Tipo da ocorrência"
            ],
            as_index=False
        )["_quantidade_no_registro"]
        .sum()
        .rename(
            columns={
                "_quantidade_no_registro":
                    "Total de ocorrências"
            }
        )
    )

    return (
        resumo_termo,
        resumo_tipo
    )


# ============================================================
# EXCEL
# ============================================================

def formatar_excel(caminho):
    wb = load_workbook(caminho)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )
            cell.fill = PatternFill(
                "solid",
                fgColor="1F4E78"
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

        for col in range(
            1,
            ws.max_column + 1
        ):
            titulo = ws.cell(
                1,
                col
            ).value

            if titulo in {
                "Parágrafo anterior",
                "Parágrafo do termo",
                "Parágrafo posterior",
                "Descrição da ocorrência",
                "Observações do pesquisador",
            }:
                largura = 70

            elif titulo in {
                "Unidade",
                "Capítulo",
                "Seção",
                "Subseção",
                "Contexto sociológico da ocorrência"
            }:
                largura = 35

            elif titulo in {
                "ID resultado",
                "Termo",
                "Categoria do termo",
                "Tipo da ocorrência",
                "ID livro"
            }:
                largura = 28

            else:
                largura = 20

            ws.column_dimensions[
                get_column_letter(col)
            ].width = largura

        if "ID resultado" in {
            cell.value for cell in ws[1]
        }:
            coluna_id = next(
                cell.column for cell in ws[1] if cell.value == "ID resultado"
            )
            for linha in range(2, ws.max_row + 1):
                ws.cell(linha, coluna_id).number_format = "@"

    if "Ocorrencias" in wb.sheetnames:
        ws = wb["Ocorrencias"]

        headers = {
            cell.value: cell.column
            for cell in ws[1]
        }

        col_validacao = headers.get(
            "Validação manual"
        )

        if col_validacao:
            letra = get_column_letter(
                col_validacao
            )

            dv = DataValidation(
                type="list",
                formula1='"incluir,excluir,revisar"',
                allow_blank=False
            )

            dv.error = (
                "Escolha: incluir, excluir ou revisar."
            )
            dv.errorTitle = "Valor inválido"
            dv.prompt = (
                "Classifique esta ocorrência."
            )
            dv.promptTitle = (
                "Validação manual"
            )

            ws.add_data_validation(dv)

            ultima_linha = max(
                ws.max_row,
                2
            )

            dv.add(
                f"{letra}2:{letra}{ultima_linha}"
            )

            for linha in range(
                2,
                ultima_linha + 1
            ):
                if ws.cell(
                    linha,
                    col_validacao
                ).value in (
                    None,
                    ""
                ):
                    ws.cell(
                        linha,
                        col_validacao
                    ).value = "revisar"

    wb.save(caminho)


def salvar_excel_completo(
    arquivo_saida,
    df_completo,
    df_termos,
    df_diag,
    configuracoes=None,
    indice_livro=1,
):
    colunas_saida = [
        "ID resultado",
        "ID livro",
        "Termo",
        "Parágrafo anterior",
        "Parágrafo do termo",
        "Parágrafo posterior",
        "Unidade",
        "Capítulo",
        "Seção",
        "Subseção",
        "Página",
        "Tipo da ocorrência",
        "Contexto sociológico da ocorrência",
        "Descrição da ocorrência",
        "Validação manual",
        "Observações do pesquisador",
    ]

    if df_completo.empty:
        df_saida = pd.DataFrame(
            columns=colunas_saida
        )
    else:
        df_saida = df_completo.reindex(columns=colunas_saida).copy()
        df_saida["Observações do pesquisador"] = ""

    df_saida["ID resultado"] = gerar_ids_resultado(
        df_completo,
        "V1",
        indice_livro,
    )

    (
        resumo_termo,
        resumo_tipo
    ) = criar_resumos(df_completo)

    if "Termo" in df_termos.columns:
        df_termos_saida = (
            df_termos[["Termo"]]
            .drop_duplicates()
            .reset_index(drop=True)
        )
    else:
        df_termos_saida = pd.DataFrame(
            columns=["Termo"]
        )

    with pd.ExcelWriter(
        arquivo_saida,
        engine="openpyxl"
    ) as writer:

        df_saida.to_excel(
            writer,
            sheet_name="Ocorrencias",
            index=False
        )

        resumo_termo.to_excel(
            writer,
            sheet_name="Resumo por termo",
            index=False
        )

        resumo_tipo.to_excel(
            writer,
            sheet_name="Resumo por tipo",
            index=False
        )

        df_termos_saida.to_excel(
            writer,
            sheet_name="Termos pesquisados",
            index=False
        )

        df_diag.to_excel(
            writer,
            sheet_name="Diagnostico OCR",
            index=False
        )

    formatar_excel(
        arquivo_saida
    )
    adicionar_aba_parametros(
        arquivo_saida,
        "v1",
        df_termos,
        df_diag,
        configuracoes,
    )


# ============================================================
# PROGRAMA PRINCIPAL
# ============================================================

def dataframe_vazio():
    return pd.DataFrame(
        columns=[
            "ID livro",
            "Termo",
            "Categoria do termo",
            "Parágrafo anterior",
            "Parágrafo do termo",
            "Parágrafo posterior",
            "Unidade",
            "Capítulo",
            "Seção",
            "Subseção",
            "Página",
            "Tipo da ocorrência",
            "Contexto sociológico da ocorrência",
            "Descrição da ocorrência",
            "Validação manual",
            "_quantidade_no_registro",
            "_metodo",
        ]
    )


def main():
    print(
        "\n"
        + "=" * 76
    )
    print(
        "VARREDURA SOCIOLÓGICA DE TERMOS EM PDFs"
    )
    print(
        "=" * 76
    )

    termos_referencia = carregar_termos_referencia()

    pdfs = sorted(
        PASTA.rglob(
            "*.pdf"
        )
    )

    if not pdfs:
        print(
            "\nNenhum PDF foi encontrado nesta pasta."
        )
        input(
            "\nPressione ENTER para sair..."
        )
        return

    if termos_referencia:
        print(
            f"\nO arquivo termos.txt contém "
            f"{len(termos_referencia)} termos "
            f"de referência para categorias."
        )
    else:
        print(
            "\nNenhum termos.txt foi encontrado. "
            "Termos novos serão classificados "
            "como TERMO INFORMADO."
        )

    print(
        f"Foram encontrados {len(pdfs)} PDF(s):\n"
    )

    for numero, pdf in enumerate(
        pdfs,
        start=1
    ):
        print(
            f"{numero}. {pdf.name}"
        )

    print(
        "\n0. ANALISAR TODOS OS PDFs"
    )

    while True:
        escolha = input(
            "\nDigite o número do PDF que deseja analisar: "
        ).strip()

        try:
            escolha = int(
                escolha
            )
        except ValueError:
            print(
                "Digite somente um número."
            )
            continue

        if escolha == 0:
            selecionados = pdfs
            break

        if 1 <= escolha <= len(pdfs):
            selecionados = [
                pdfs[
                    escolha - 1
                ]
            ]
            break

        print(
            "Número inválido."
        )

    # SEMPRE pergunta a lista de termos.
    termos = solicitar_termos(
        termos_referencia
    )

    df_termos = pd.DataFrame(
        termos
    ).rename(
        columns={
            "categoria":
                "Categoria do termo",
            "termo":
                "Termo"
        }
    )

    todos_registros = []
    todos_diagnosticos = []
    arquivos_gerados = []

    for numero, pdf in enumerate(
        selecionados,
        start=1
    ):
        print(
            f"\n[{numero}/{len(selecionados)}]"
        )

        try:
            ocorrencias, diag = analisar_pdf(
                pdf,
                termos
            )

        except Exception as erro:
            print(
                f"\nERRO ao analisar "
                f"{pdf.name}:\n{erro}"
            )
            continue

        df_livro = (
            pd.DataFrame(
                ocorrencias
            )
            if ocorrencias
            else dataframe_vazio()
        )

        df_diag_livro = pd.DataFrame(
            [diag]
        )

        nome_seguro = re.sub(
            r'[<>:"/\\|?*]',
            "_",
            pdf.stem
        )

        arquivo_individual = (
            PASTA
            / f"resultado_{nome_seguro}_v1.xlsx"
        )

        salvar_excel_completo(
            arquivo_individual,
            df_livro,
            df_termos,
            df_diag_livro
        )

        arquivos_gerados.append(
            arquivo_individual
        )

        todos_registros.extend(
            ocorrencias
        )

        todos_diagnosticos.append(
            diag
        )

        print(
            f"Excel individual criado: "
            f"{arquivo_individual.name}"
        )

    if len(selecionados) > 1:
        df_todos = (
            pd.DataFrame(
                todos_registros
            )
            if todos_registros
            else dataframe_vazio()
        )

        df_diag_todos = pd.DataFrame(
            todos_diagnosticos
        )

        arquivo_consolidado = (
            PASTA
            / "resultado_todos_pdfs_v1.xlsx"
        )

        salvar_excel_completo(
            arquivo_consolidado,
            df_todos,
            df_termos,
            df_diag_todos
        )

        arquivos_gerados.append(
            arquivo_consolidado
        )

    print(
        "\n"
        + "=" * 76
    )
    print(
        "ANÁLISE CONCLUÍDA"
    )
    print(
        "=" * 76
    )

    print(
        "\nArquivos gerados:"
    )

    for arquivo in arquivos_gerados:
        print(
            f"- {arquivo.name}"
        )

    if len(selecionados) > 1:
        print(
            "\nFoi criado um Excel para cada PDF "
            "e também um Excel consolidado."
        )
    else:
        print(
            "\nFoi criado um Excel para o PDF selecionado."
        )

    input(
        "\nPressione ENTER para fechar..."
    )


if __name__ == "__main__":
    main()
