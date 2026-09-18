"""Ponte reutilizável para a lógica analítica preservada da V2.

O arquivo ``varredura_pdf_v2.py`` permanece como a referência íntegra da
versão V2. Esta ponte expõe suas funções para a aplicação web sem misturar
a análise ou a exportação V2 com a V1.
"""

from copy import copy

from openpyxl import load_workbook

from varredura_pdf_v2 import (
    analisar_pdf as _analisar_pdf_referencia,
    carregar_termos_referencia,
    configurar_tesseract,
    contar_ocorrencias,
    criar_regex,
    criar_resumos,
    dataframe_vazio,
    descrever_ocorrencia,
    escolher_idioma_ocr,
    formatar_excel,
    gerar_variacoes_palavra,
    gerar_variacoes_termo,
    identificar_contexto_sociologico,
    limpar_texto,
    normalizar,
    salvar_excel_completo as _salvar_excel_completo_v2,
)

from .auditoria import adicionar_aba_parametros, gerar_ids_resultado


def analisar_pdf(caminho, termos, progress_callback=None):
    """Expõe marcos observacionais dos mesmos loops da V2 de referência."""
    resultado = (
        _analisar_pdf_referencia(caminho, termos)
        if progress_callback is None
        else _analisar_pdf_referencia(
            caminho, termos, progress_callback=progress_callback
        )
    )
    if progress_callback is not None:
        progress_callback(
            {
                "fase": "analise_v2_concluida",
                "etapa": "Preparando resultados…",
            }
        )
    return resultado


def salvar_excel_completo(
    arquivo_saida,
    df_completo,
    df_termos,
    df_diag,
    configuracoes=None,
    indice_livro=1,
):
    """Acrescenta apenas a auditoria, preservando a exportação V2 de referência."""
    _salvar_excel_completo_v2(arquivo_saida, df_completo, df_termos, df_diag)
    identificadores = gerar_ids_resultado(df_completo, "V2", indice_livro)

    # A referência V2 permanece íntegra. O identificador é inserido na cópia
    # exportada, sem modificar sua análise, seus resumos ou os campos manuais.
    workbook = load_workbook(arquivo_saida)
    aba = workbook["Ocorrencias"]
    aba.insert_cols(1)
    modelo_cabecalho = aba.cell(1, 2)
    celula_cabecalho = aba.cell(1, 1, "ID resultado")
    celula_cabecalho.font = copy(modelo_cabecalho.font)
    celula_cabecalho.fill = copy(modelo_cabecalho.fill)
    celula_cabecalho.alignment = copy(modelo_cabecalho.alignment)
    celula_cabecalho.border = copy(modelo_cabecalho.border)
    celula_cabecalho.protection = copy(modelo_cabecalho.protection)

    for linha, identificador in enumerate(identificadores, start=2):
        celula = aba.cell(linha, 1, identificador)
        modelo = aba.cell(linha, 2)
        celula.font = copy(modelo.font)
        celula.fill = copy(modelo.fill)
        celula.alignment = copy(modelo.alignment)
        celula.border = copy(modelo.border)
        celula.protection = copy(modelo.protection)
        celula.number_format = "@"

    aba.column_dimensions["A"].width = 20
    aba.auto_filter.ref = aba.dimensions
    workbook.save(arquivo_saida)
    adicionar_aba_parametros(
        arquivo_saida,
        "v2",
        df_termos,
        df_diag,
        configuracoes,
    )

__all__ = [
    "analisar_pdf",
    "carregar_termos_referencia",
    "configurar_tesseract",
    "contar_ocorrencias",
    "criar_regex",
    "criar_resumos",
    "dataframe_vazio",
    "descrever_ocorrencia",
    "escolher_idioma_ocr",
    "formatar_excel",
    "gerar_variacoes_palavra",
    "gerar_variacoes_termo",
    "identificar_contexto_sociologico",
    "limpar_texto",
    "normalizar",
    "salvar_excel_completo",
]
