"""Analisador V3: busca híbrida lexical e semântica.

Esta versão reutiliza exclusivamente as funções estáveis de extração, OCR,
normalização, variações morfológicas e contexto sociológico da V1. A etapa de
embeddings é própria da V3 e só importa/carrega o modelo quando a busca
semântica foi solicitada.
"""

from __future__ import annotations

from functools import lru_cache
from pathlib import Path
import re
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .auditoria import adicionar_aba_parametros, gerar_ids_resultado
from . import v1


MODELO_SEMANTICO = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"
LIMIAR_PADRAO = 0.70
TAMANHO_ALVO_TRECHO = 400
TAMANHO_MINIMO_TRECHO = 300
TAMANHO_MAXIMO_TRECHO = 500
SOBREPOSICAO_TRECHO = 60

COLUNAS_RESULTADOS = [
    "ID resultado",
    "ID livro",
    "Consulta",
    "Termo encontrado",
    "Tipo de correspondência",
    "Similaridade semântica",
    "Página inicial",
    "Página final",
    "Bloco anterior",
    "Trecho da ocorrência",
    "Bloco posterior",
    "Contexto sociológico",
    "Descrição",
    "Validação manual",
    "Observações do pesquisador",
]

COLUNAS_INTERNAS = [
    *COLUNAS_RESULTADOS,
    "_quantidade_no_registro",
    "_pagina_inicial_pdf",
    "_pagina_final_pdf",
    "_texto_normalizado",
]


def _emitir_progresso(progress_callback, **evento: Any) -> None:
    """Publica apenas metadados de etapas já executadas, quando solicitado."""
    if progress_callback is not None:
        progress_callback(evento)


def normalizar(texto: str) -> str:
    """Mantém a mesma normalização usada pela V1 para termos e entradas."""
    return v1.normalizar(texto)


def carregar_termos_referencia() -> list[dict[str, str]]:
    """Usa o dicionário de referência opcional da V1 sem expô-lo no Excel."""
    return v1.carregar_termos_referencia()


def dataframe_vazio() -> pd.DataFrame:
    return pd.DataFrame(columns=COLUNAS_INTERNAS)


def _configuracao(configuracoes: dict[str, Any] | None) -> dict[str, Any]:
    configuracoes = configuracoes or {}
    lexical = bool(configuracoes.get("incluir_lexical", True))
    semantica = bool(configuracoes.get("incluir_semantica", True))
    if not lexical and not semantica:
        raise ValueError("Selecione ao menos uma modalidade de busca para a V3.")

    try:
        limiar = float(configuracoes.get("limiar_semantico", LIMIAR_PADRAO))
    except (TypeError, ValueError):
        limiar = LIMIAR_PADRAO
    limiar = min(0.90, max(0.50, limiar))
    return {
        "incluir_lexical": lexical,
        "incluir_semantica": semantica,
        "limiar_semantico": limiar,
    }


@lru_cache(maxsize=1)
def carregar_modelo_semantico():
    """Carrega o modelo multilíngue somente na primeira busca semântica V3."""
    try:
        from sentence_transformers import SentenceTransformer
    except ImportError as erro:
        raise RuntimeError(
            "A busca semântica V3 requer a dependência sentence-transformers. "
            "Instale as dependências do projeto antes de executar a análise."
        ) from erro

    try:
        return SentenceTransformer(MODELO_SEMANTICO)
    except Exception as erro:  # mensagem amigável para falha de cache/download local
        raise RuntimeError(
            "Não foi possível carregar o modelo semântico multilíngue da V3. "
            "Verifique a conexão na primeira execução ou o cache local do modelo."
        ) from erro


def modelo_semantico_carregado() -> bool:
    """Indica se o cache em memória foi utilizado nesta execução do processo."""
    return carregar_modelo_semantico.cache_info().currsize > 0


def _frases_do_bloco(
    bloco: dict[str, Any], indice_bloco: int
) -> list[dict[str, Any]]:
    texto = v1.limpar_texto(str(bloco.get("texto", "")))
    if not texto:
        return []

    # Não corta uma sentença no meio: quando a pontuação não é confiável,
    # mantém o bloco como uma unidade textual.
    frases = re.split(r"(?<=[.!?…])\s+(?=[A-ZÀ-Ý0-9])", texto)
    frases = [frase.strip() for frase in frases if frase.strip()]
    if not frases:
        frases = [texto]

    return [
        {
            "texto": frase,
            "pagina": bloco.get("pagina"),
            "pagina_pdf": int(bloco.get("pagina_pdf", 0) or 0),
            # Metadado de proveniência: não altera a janela nem os embeddings.
            "indice_bloco": indice_bloco,
        }
        for frase in frases
    ]


def criar_trechos_semanticos(
    blocos: list[dict[str, Any]], progress_callback=None
) -> list[dict[str, Any]]:
    """Agrupa o texto em janelas de 300–500 palavras com sobreposição segura."""
    unidades: list[dict[str, Any]] = []
    total_blocos = len(blocos)
    passo_progresso = max(1, total_blocos // 100)
    _emitir_progresso(
        progress_callback,
        fase="preparando_blocos_semanticos",
        etapa="Preparando blocos semânticos…",
        bloco_atual=0,
        blocos_total=total_blocos,
    )
    for indice_bloco, bloco in enumerate(blocos):
        unidades.extend(_frases_do_bloco(bloco, indice_bloco))
        if (
            indice_bloco == total_blocos - 1
            or (indice_bloco + 1) % passo_progresso == 0
        ):
            _emitir_progresso(
                progress_callback,
                fase="preparando_blocos_semanticos",
                etapa="Preparando blocos semânticos…",
                bloco_atual=indice_bloco + 1,
                blocos_total=total_blocos,
            )

    trechos: list[dict[str, Any]] = []
    inicio = 0
    while inicio < len(unidades):
        fim = inicio
        palavras = 0
        while fim < len(unidades):
            quantidade = len(unidades[fim]["texto"].split())
            if palavras and palavras + quantidade > TAMANHO_MAXIMO_TRECHO:
                break
            palavras += quantidade
            fim += 1
            if palavras >= TAMANHO_ALVO_TRECHO:
                break

        # Uma frase muito longa é mantida intacta, em vez de ser cortada.
        if fim == inicio:
            fim += 1
            palavras = len(unidades[inicio]["texto"].split())

        janela = unidades[inicio:fim]
        texto = " ".join(item["texto"] for item in janela).strip()
        if texto:
            trechos.append(
                {
                    "texto": texto,
                    "pagina_inicial": janela[0]["pagina"],
                    "pagina_final": janela[-1]["pagina"],
                    "pagina_inicial_pdf": janela[0]["pagina_pdf"],
                    "pagina_final_pdf": janela[-1]["pagina_pdf"],
                    "indice_bloco_inicial": janela[0]["indice_bloco"],
                    "indice_bloco_final": janela[-1]["indice_bloco"],
                    "palavras": palavras,
                }
            )

        if fim >= len(unidades):
            break

        # Retrocede até acumular a sobreposição desejada, sempre avançando.
        proximo_inicio = fim
        sobrepostas = 0
        while proximo_inicio > inicio and sobrepostas < SOBREPOSICAO_TRECHO:
            proximo_inicio -= 1
            sobrepostas += len(unidades[proximo_inicio]["texto"].split())
        inicio = max(inicio + 1, proximo_inicio)

    return trechos


def _texto_original_do_bloco(
    blocos: list[dict[str, Any]], indice: int | None
) -> str:
    """Retorna somente texto extraído/OCR, sem normalização para exportação."""
    if indice is None or not 0 <= indice < len(blocos):
        return ""
    return str(blocos[indice].get("texto", "")).strip()


def _blocos_adjacentes(
    blocos: list[dict[str, Any]] | None,
    indice_inicial: int | None,
    indice_final: int | None,
) -> tuple[str, str]:
    """Obtém contexto real da sequência original, nunca do overlap semântico."""
    if blocos is None or indice_inicial is None or indice_final is None:
        return "", ""
    return (
        _texto_original_do_bloco(blocos, indice_inicial - 1),
        _texto_original_do_bloco(blocos, indice_final + 1),
    )


def _contexto_amplo(blocos: list[dict[str, Any]], indice: int) -> str:
    bloco = blocos[indice]
    return " ".join(
        [
            bloco.get("unidade", ""),
            bloco.get("capitulo", ""),
            bloco.get("secao", ""),
            bloco.get("subsecao", ""),
            v1.localizar_paragrafo_anterior(blocos, indice),
            bloco.get("texto", ""),
            v1.localizar_paragrafo_posterior(blocos, indice),
        ]
    )


def _tipo_lexical(texto: str, consulta: str) -> tuple[str, str]:
    """Diferencia a forma literal da variação morfológica encontrada."""
    encontrado = v1.criar_regex(consulta).search(v1.normalizar(texto))
    termo_encontrado = encontrado.group(0) if encontrado else consulta
    tipo = (
        "Lexical"
        if termo_encontrado == v1.normalizar(consulta)
        else "Morfológica"
    )
    return tipo, termo_encontrado


def _registros_lexicais(
    caminho: Path,
    blocos: list[dict[str, Any]],
    termos: list[dict[str, str]],
    progress_callback=None,
) -> list[dict[str, Any]]:
    registros: list[dict[str, Any]] = []
    total_blocos = len(blocos)
    passo_progresso = max(1, total_blocos // 100)
    _emitir_progresso(
        progress_callback,
        fase="busca_lexical",
        etapa="Executando busca lexical…",
        bloco_atual=0,
        blocos_total=total_blocos,
    )
    for indice, bloco in enumerate(blocos):
        atual = str(bloco.get("texto", ""))
        bloco_anterior, bloco_posterior = _blocos_adjacentes(
            blocos, indice, indice
        )
        contexto_amplo = _contexto_amplo(blocos, indice)
        contexto = v1.identificar_contexto_sociologico(contexto_amplo)

        for item in termos:
            consulta = item["termo"]
            quantidade = v1.contar_ocorrencias(atual, consulta)
            if not quantidade:
                continue
            tipo, termo_encontrado = _tipo_lexical(atual, consulta)
            trecho = v1.extrair_frase_com_termo(atual, consulta)
            descricao = v1.descrever_ocorrencia(
                consulta, atual, contexto_amplo, contexto
            )
            registros.append(
                {
                    "ID livro": v1.id_livro(caminho),
                    "Consulta": consulta,
                    "Termo encontrado": termo_encontrado,
                    "Tipo de correspondência": tipo,
                    "Similaridade semântica": None,
                    "Página inicial": bloco.get("pagina"),
                    "Página final": bloco.get("pagina"),
                    "Bloco anterior": bloco_anterior,
                    # O Excel guarda o bloco original; ``trecho`` continua sendo
                    # usado somente pela comparação híbrida já existente.
                    "Trecho da ocorrência": atual,
                    "Bloco posterior": bloco_posterior,
                    "Contexto sociológico": contexto,
                    "Descrição": descricao,
                    "Validação manual": "",
                    "Observações do pesquisador": "",
                    "_quantidade_no_registro": int(quantidade),
                    "_pagina_inicial_pdf": int(bloco.get("pagina_pdf", 0) or 0),
                    "_pagina_final_pdf": int(bloco.get("pagina_pdf", 0) or 0),
                    "_texto_normalizado": v1.normalizar(trecho),
                }
            )
        if (
            indice == total_blocos - 1
            or (indice + 1) % passo_progresso == 0
        ):
            _emitir_progresso(
                progress_callback,
                fase="busca_lexical",
                etapa="Executando busca lexical…",
                bloco_atual=indice + 1,
                blocos_total=total_blocos,
            )
    return registros


def _similaridade_jaccard(texto_a: str, texto_b: str) -> float:
    conjunto_a = set(re.findall(r"\w+", v1.normalizar(texto_a)))
    conjunto_b = set(re.findall(r"\w+", v1.normalizar(texto_b)))
    if not conjunto_a or not conjunto_b:
        return 0.0
    return len(conjunto_a & conjunto_b) / len(conjunto_a | conjunto_b)


def _registro_tem_evidencia_lexical(
    registros: list[dict[str, Any]], candidato: dict[str, Any]
) -> dict[str, Any] | None:
    for registro in registros:
        if registro["Consulta"] != candidato["Consulta"]:
            continue
        if registro["ID livro"] != candidato["ID livro"]:
            continue
        mesma_pagina = (
            registro["_pagina_inicial_pdf"]
            <= candidato["_pagina_final_pdf"]
            and registro["_pagina_final_pdf"]
            >= candidato["_pagina_inicial_pdf"]
        )
        if mesma_pagina and _similaridade_jaccard(
            registro["_texto_normalizado"], candidato["_texto_normalizado"]
        ) >= 0.08:
            return registro
    return None


def _deduplicar_semanticos(candidatos: list[dict[str, Any]]) -> list[dict[str, Any]]:
    aceitos: list[dict[str, Any]] = []
    for candidato in sorted(
        candidatos, key=lambda item: float(item["Similaridade semântica"]), reverse=True
    ):
        duplicado = False
        for anterior in aceitos:
            if (
                anterior["Consulta"] == candidato["Consulta"]
                and anterior["ID livro"] == candidato["ID livro"]
                and abs(
                    anterior["_pagina_inicial_pdf"]
                    - candidato["_pagina_inicial_pdf"]
                ) <= 1
                and _similaridade_jaccard(
                    anterior["_texto_normalizado"], candidato["_texto_normalizado"]
                ) >= 0.75
            ):
                duplicado = True
                break
        if not duplicado:
            aceitos.append(candidato)
    return aceitos


def _descricao_semantica(consulta: str, contexto: str) -> str:
    return (
        f"O trecho foi recuperado por similaridade semântica com a consulta “{consulta}”. "
        f"A relação é uma aproximação de sentido, sem ocorrência lexical direta verificada. "
        f"O contexto sociológico automático identificado é “{contexto}” e requer validação manual."
    )


def _codificar_trechos_com_progresso(
    modelo: Any,
    textos: list[str],
    progress_callback,
) -> np.ndarray:
    """Codifica os mesmos lotes de 32 e publica somente lotes concluídos.

    A V3 já usa ``batch_size=32``. Quando há callback, esta função preserva
    tamanho e ordem desses lotes e apenas expõe cada lote que o modelo já
    concluiu. Sem callback, o caminho original de uma única chamada a
    ``encode`` continua sendo usado literalmente.
    """
    vetores: list[np.ndarray] = []
    total_trechos = len(textos)
    for inicio in range(0, total_trechos, 32):
        lote = textos[inicio : inicio + 32]
        vetores.append(
            np.asarray(
                modelo.encode(
                    lote,
                    batch_size=32,
                    show_progress_bar=False,
                    normalize_embeddings=True,
                )
            )
        )
        _emitir_progresso(
            progress_callback,
            fase="codificando_blocos_semanticos",
            etapa="Analisando correspondências semânticas…",
            bloco_atual=min(inicio + len(lote), total_trechos),
            blocos_total=total_trechos,
        )
    return np.concatenate(vetores, axis=0)


def _registros_semanticos(
    caminho: Path,
    trechos: list[dict[str, Any]],
    termos: list[dict[str, str]],
    limiar: float,
    blocos: list[dict[str, Any]] | None = None,
    progress_callback=None,
    progresso_por_lote: bool = False,
) -> list[dict[str, Any]]:
    if not trechos or not termos:
        return []

    _emitir_progresso(
        progress_callback,
        fase="preparando_modelo_semantico",
        etapa="Preparando modelo semântico…",
        preparando_modelo=True,
    )
    modelo = carregar_modelo_semantico()
    _emitir_progresso(
        progress_callback,
        fase="modelo_semantico_carregado",
        etapa="Modelo semântico carregado",
    )
    total_trechos = len(trechos)
    _emitir_progresso(
        progress_callback,
        fase="codificando_blocos_semanticos",
        etapa="Analisando correspondências semânticas…",
        bloco_atual=0,
        blocos_total=total_trechos,
    )
    textos_trechos = [item["texto"] for item in trechos]
    if progress_callback is not None and progresso_por_lote:
        vetores_trechos = _codificar_trechos_com_progresso(
            modelo, textos_trechos, progress_callback
        )
    else:
        # Mantém a chamada única já usada pelo fluxo híbrido validado.
        vetores_trechos = np.asarray(
            modelo.encode(
                textos_trechos,
                batch_size=32,
                show_progress_bar=False,
                normalize_embeddings=True,
            )
        )
        _emitir_progresso(
            progress_callback,
            fase="codificando_blocos_semanticos",
            etapa="Analisando correspondências semânticas…",
            bloco_atual=total_trechos,
            blocos_total=total_trechos,
        )
    consultas = [item["termo"] for item in termos]
    vetores_consultas = np.asarray(
        modelo.encode(
            consultas,
            batch_size=32,
            show_progress_bar=False,
            normalize_embeddings=True,
        )
    )
    similaridades = np.matmul(vetores_consultas, vetores_trechos.T)
    registros: list[dict[str, Any]] = []
    livro = v1.id_livro(caminho)
    total_consultas = len(consultas)

    for indice_consulta, consulta in enumerate(consultas):
        for indice_trecho, valor in enumerate(similaridades[indice_consulta]):
            similaridade = float(valor)
            if similaridade < limiar:
                continue
            trecho = trechos[indice_trecho]
            bloco_anterior, bloco_posterior = _blocos_adjacentes(
                blocos,
                trecho.get("indice_bloco_inicial"),
                trecho.get("indice_bloco_final"),
            )
            contexto = v1.identificar_contexto_sociologico(trecho["texto"])
            registros.append(
                {
                    "ID livro": livro,
                    "Consulta": consulta,
                    "Termo encontrado": "",
                    "Tipo de correspondência": "Semântica",
                    "Similaridade semântica": round(similaridade, 4),
                    "Página inicial": trecho["pagina_inicial"],
                    "Página final": trecho["pagina_final"],
                    "Bloco anterior": bloco_anterior,
                    "Trecho da ocorrência": trecho["texto"],
                    "Bloco posterior": bloco_posterior,
                    "Contexto sociológico": contexto,
                    "Descrição": _descricao_semantica(consulta, contexto),
                    "Validação manual": "",
                    "Observações do pesquisador": "",
                    "_quantidade_no_registro": 1,
                    "_pagina_inicial_pdf": trecho["pagina_inicial_pdf"],
                    "_pagina_final_pdf": trecho["pagina_final_pdf"],
                    "_texto_normalizado": v1.normalizar(trecho["texto"]),
                }
            )
        _emitir_progresso(
            progress_callback,
            fase="comparando_semantica",
            etapa="Analisando correspondências semânticas…",
            bloco_atual=total_trechos,
            blocos_total=total_trechos,
            consulta_atual=indice_consulta + 1,
            consultas_total=total_consultas,
        )
    return _deduplicar_semanticos(registros)


def analisar_pdf(
    caminho: Path,
    termos: list[dict[str, str]],
    configuracoes: dict[str, Any] | None = None,
    progress_callback=None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Executa a V3 sem alterar nenhuma rotina de análise das versões anteriores."""
    configuracao = _configuracao(configuracoes)
    if progress_callback is None:
        blocos, paginas_texto, paginas_ocr, idioma_ocr = v1.extrair_pdf(caminho)
        registros = (
            _registros_lexicais(caminho, blocos, termos)
            if configuracao["incluir_lexical"]
            else []
        )
    else:
        blocos, paginas_texto, paginas_ocr, idioma_ocr = v1.extrair_pdf(
            caminho, progress_callback=progress_callback
        )
        registros = (
            _registros_lexicais(
                caminho, blocos, termos, progress_callback=progress_callback
            )
            if configuracao["incluir_lexical"]
            else []
        )

    if configuracao["incluir_semantica"]:
        if progress_callback is None:
            candidatos = _registros_semanticos(
                caminho,
                criar_trechos_semanticos(blocos),
                termos,
                configuracao["limiar_semantico"],
                blocos,
            )
        else:
            candidatos = _registros_semanticos(
                caminho,
                criar_trechos_semanticos(blocos, progress_callback=progress_callback),
                termos,
                configuracao["limiar_semantico"],
                blocos,
                progress_callback=progress_callback,
                progresso_por_lote=not configuracao["incluir_lexical"],
            )
        for candidato in candidatos:
            equivalente_lexico = _registro_tem_evidencia_lexical(registros, candidato)
            if equivalente_lexico is not None:
                anterior = equivalente_lexico.get("Similaridade semântica")
                if anterior is None or candidato["Similaridade semântica"] > anterior:
                    equivalente_lexico["Similaridade semântica"] = candidato[
                        "Similaridade semântica"
                    ]
                equivalente_lexico["Tipo de correspondência"] = "Lexical + semântica"
            else:
                registros.append(candidato)

    diagnostico = {
        "arquivo": caminho.name,
        "paginas_com_texto_extraivel": paginas_texto,
        "paginas_processadas_com_OCR": paginas_ocr,
        "idioma_OCR": idioma_ocr or "não utilizado",
        "busca_lexical_ativada": configuracao["incluir_lexical"],
        "busca_semantica_ativada": configuracao["incluir_semantica"],
        "limiar_semantico": configuracao["limiar_semantico"],
    }
    return registros, diagnostico


def _resumos(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    colunas_lex = ["ID livro", "Consulta", "Tipo de correspondência", "Ocorrências"]
    colunas_sem = [
        "ID livro",
        "Consulta",
        "Correspondências semânticas",
        "Similaridade média",
        "Similaridade máxima",
    ]
    colunas_consulta = [
        "Consulta",
        "Ocorrências lexicais/morfológicas",
        "Correspondências semânticas exclusivas",
        "Resultados recuperados",
    ]
    if df.empty:
        return (
            pd.DataFrame(columns=colunas_lex),
            pd.DataFrame(columns=colunas_sem),
            pd.DataFrame(columns=colunas_consulta),
        )

    lexical = df[df["Tipo de correspondência"].isin(["Lexical", "Morfológica", "Lexical + semântica"])]
    resumo_lex = (
        lexical.groupby(["ID livro", "Consulta", "Tipo de correspondência"], as_index=False)["_quantidade_no_registro"]
        .sum()
        .rename(columns={"_quantidade_no_registro": "Ocorrências"})
    )

    semantico = df[df["Tipo de correspondência"] == "Semântica"]
    resumo_sem = (
        semantico.groupby(["ID livro", "Consulta"], as_index=False)["Similaridade semântica"]
        .agg(["count", "mean", "max"])
        .reset_index()
        .rename(
            columns={
                "count": "Correspondências semânticas",
                "mean": "Similaridade média",
                "max": "Similaridade máxima",
            }
        )
    )

    todas_consultas = sorted(df["Consulta"].dropna().unique().tolist(), key=str.casefold)
    linhas_consulta = []
    for consulta in todas_consultas:
        recorte = df[df["Consulta"] == consulta]
        recorte_lex = recorte[recorte["Tipo de correspondência"].isin(["Lexical", "Morfológica", "Lexical + semântica"])]
        recorte_sem = recorte[recorte["Tipo de correspondência"] == "Semântica"]
        linhas_consulta.append(
            {
                "Consulta": consulta,
                "Ocorrências lexicais/morfológicas": int(recorte_lex["_quantidade_no_registro"].sum()),
                "Correspondências semânticas exclusivas": int(len(recorte_sem)),
                "Resultados recuperados": int(len(recorte)),
            }
        )
    return resumo_lex, resumo_sem, pd.DataFrame(linhas_consulta, columns=colunas_consulta)


def _formatar_excel(caminho: Path) -> None:
    workbook = load_workbook(caminho)
    for aba in workbook.worksheets:
        aba.freeze_panes = "A2"
        aba.auto_filter.ref = aba.dimensions
        for celula in aba[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="1F4E78")
            celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for linha in aba.iter_rows(min_row=2):
            for celula in linha:
                celula.alignment = Alignment(vertical="top", wrap_text=True)
        for indice in range(1, aba.max_column + 1):
            titulo = aba.cell(1, indice).value
            largura = 22
            if titulo in {
                "Bloco anterior",
                "Trecho da ocorrência",
                "Bloco posterior",
                "Descrição",
                "Observações do pesquisador",
            }:
                largura = 68
            elif titulo in {"Contexto sociológico", "Tipo de correspondência"}:
                largura = 34
            elif titulo in {
                "Consulta",
                "Termo encontrado",
                "ID livro",
                "ID resultado",
            }:
                largura = 28
            aba.column_dimensions[get_column_letter(indice)].width = largura

    ocorrencias = workbook["Resultados"]
    cabecalhos = {celula.value: celula.column for celula in ocorrencias[1]}
    coluna_validacao = cabecalhos.get("Validação manual")
    if coluna_validacao:
        letra = get_column_letter(coluna_validacao)
        validacao = DataValidation(
            type="list", formula1='"incluir,excluir,revisar"', allow_blank=True
        )
        validacao.error = "Escolha: incluir, excluir ou revisar."
        validacao.errorTitle = "Valor inválido"
        ocorrencias.add_data_validation(validacao)
        # Também cobre novas linhas inseridas pelo pesquisador durante a revisão.
        validacao.add(f"{letra}2:{letra}{max(ocorrencias.max_row + 500, 1000)}")

    coluna_id = cabecalhos.get("ID resultado")
    if coluna_id:
        for linha in range(2, ocorrencias.max_row + 1):
            ocorrencias.cell(linha, coluna_id).number_format = "@"
    workbook.save(caminho)


def salvar_excel_completo(
    arquivo_saida: Path,
    df_completo: pd.DataFrame,
    df_termos: pd.DataFrame,
    df_diag: pd.DataFrame,
    configuracoes: dict[str, Any] | None = None,
    indice_livro: int = 1,
) -> None:
    """Gera a planilha V3 sem colunas ou abas herdadas das versões anteriores."""
    resultados = (
        df_completo.reindex(columns=COLUNAS_RESULTADOS).copy()
        if not df_completo.empty
        else pd.DataFrame(columns=COLUNAS_RESULTADOS)
    )
    if "Similaridade semântica" in resultados:
        resultados["Similaridade semântica"] = pd.to_numeric(
            resultados["Similaridade semântica"], errors="coerce"
        )
    # Identificadores legíveis, determinísticos e únicos dentro deste arquivo,
    # inclusive quando ele consolida mais de um PDF. Não dependem do índice do
    # DataFrame de origem, que pode ser repetido na concatenação.
    resultados["ID resultado"] = gerar_ids_resultado(
        df_completo,
        "V3",
        indice_livro,
    )
    resumo_lex, resumo_sem, resumo_consulta = _resumos(df_completo)
    termos = (
        df_termos[["Termo"]]
        .rename(columns={"Termo": "Consulta"})
        .drop_duplicates()
        .reset_index(drop=True)
        if "Termo" in df_termos
        else pd.DataFrame(columns=["Consulta"])
    )

    with pd.ExcelWriter(arquivo_saida, engine="openpyxl") as writer:
        resultados.to_excel(writer, sheet_name="Resultados", index=False)
        resumo_lex.to_excel(writer, sheet_name="Resumo lexical", index=False)
        resumo_sem.to_excel(writer, sheet_name="Resumo semântico", index=False)
        resumo_consulta.to_excel(writer, sheet_name="Resumo por consulta", index=False)
        termos.to_excel(writer, sheet_name="Termos pesquisados", index=False)
        df_diag.to_excel(writer, sheet_name="Diagnostico OCR", index=False)
    _formatar_excel(arquivo_saida)
    adicionar_aba_parametros(
        arquivo_saida,
        "v3",
        df_termos,
        df_diag,
        configuracoes,
    )


def _serie(df: pd.DataFrame, coluna: str, peso: str | None = None) -> dict[str, list[Any]]:
    if df.empty:
        return {"rotulos": [], "valores": []}
    if peso:
        agrupado = df.groupby(coluna, dropna=False)[peso].sum()
    else:
        agrupado = df.groupby(coluna, dropna=False).size()
    itens = sorted(
        ((str(chave or "Não informado"), float(valor)) for chave, valor in agrupado.items()),
        key=lambda item: (-item[1], item[0].casefold()),
    )
    return {
        "rotulos": [item[0] for item in itens],
        "valores": [int(item[1]) if item[1].is_integer() else round(item[1], 4) for item in itens],
    }


def _dados_sankey(df: pd.DataFrame) -> dict[str, Any]:
    vazio = {"labels": [], "sources": [], "targets": [], "values": [], "aviso": ""}
    if df.empty:
        return vazio
    base = df.copy()
    base["_peso"] = np.where(
        base["Tipo de correspondência"].isin(["Lexical", "Morfológica", "Lexical + semântica"]),
        pd.to_numeric(base["_quantidade_no_registro"], errors="coerce").fillna(0),
        1,
    )
    consultas = _serie(base, "Consulta", "_peso")["rotulos"][:12]
    contextos = _serie(base, "Contexto sociológico", "_peso")["rotulos"][:10]
    aviso = ""
    if base["Consulta"].nunique() > len(consultas) or base["Contexto sociológico"].nunique() > len(contextos):
        aviso = "A visualização agrupa categorias menos frequentes em “Outros”."
    base["_consulta"] = base["Consulta"].where(base["Consulta"].isin(consultas), "Outras consultas")
    base["_contexto"] = base["Contexto sociológico"].where(base["Contexto sociológico"].isin(contextos), "Outros contextos")
    consultas = _serie(base, "_consulta", "_peso")["rotulos"]
    contextos = _serie(base, "_contexto", "_peso")["rotulos"]
    livros = _serie(base, "ID livro", "_peso")["rotulos"]
    labels = [
        *[f"Consulta: {item}" for item in consultas],
        *[f"Contexto: {item}" for item in contextos],
        *[f"Livro: {item}" for item in livros],
    ]
    ind_consulta = {item: indice for indice, item in enumerate(consultas)}
    deslocamento_contexto = len(consultas)
    ind_contexto = {item: deslocamento_contexto + indice for indice, item in enumerate(contextos)}
    deslocamento_livro = deslocamento_contexto + len(contextos)
    ind_livro = {item: deslocamento_livro + indice for indice, item in enumerate(livros)}
    sources: list[int] = []
    targets: list[int] = []
    valores: list[float] = []
    for (consulta, contexto), valor in base.groupby(["_consulta", "_contexto"])["_peso"].sum().items():
        sources.append(ind_consulta[consulta])
        targets.append(ind_contexto[contexto])
        valores.append(float(valor))
    for (contexto, livro), valor in base.groupby(["_contexto", "ID livro"])["_peso"].sum().items():
        sources.append(ind_contexto[contexto])
        targets.append(ind_livro[livro])
        valores.append(float(valor))
    return {
        "labels": labels,
        "sources": sources,
        "targets": targets,
        "values": [int(valor) if valor.is_integer() else round(valor, 4) for valor in valores],
        "aviso": aviso,
    }


def criar_dashboard(resultado: dict[str, Any]) -> dict[str, Any]:
    """Prepara somente agregações V3 para o dashboard híbrido."""
    df = resultado["ocorrencias"].copy()
    if df.empty:
        df = dataframe_vazio()
    lexical = df[df["Tipo de correspondência"].isin(["Lexical", "Morfológica", "Lexical + semântica"])]
    semantico = df[df["Tipo de correspondência"] == "Semântica"]
    peso_lexical = int(pd.to_numeric(lexical["_quantidade_no_registro"], errors="coerce").fillna(0).sum())
    consultas = [item["termo"] for item in resultado["termos"]]
    livros = [Path(item["arquivo"]).stem for item in resultado["diagnosticos"].to_dict("records")]
    matriz = [[0 for _ in livros] for _ in consultas]
    if not df.empty:
        base_matriz = df.copy()
        base_matriz["_peso"] = np.where(
            base_matriz["Tipo de correspondência"].isin(["Lexical", "Morfológica", "Lexical + semântica"]),
            pd.to_numeric(base_matriz["_quantidade_no_registro"], errors="coerce").fillna(0),
            1,
        )
        pivot = base_matriz.pivot_table(index="Consulta", columns="ID livro", values="_peso", aggfunc="sum", fill_value=0)
        for indice_consulta, consulta in enumerate(consultas):
            for indice_livro, livro in enumerate(livros):
                if consulta in pivot.index and livro in pivot.columns:
                    matriz[indice_consulta][indice_livro] = int(pivot.loc[consulta, livro])

    semelhancas = pd.to_numeric(semantico["Similaridade semântica"], errors="coerce").dropna().tolist()
    total_ocr = int(resultado["diagnosticos"]["paginas_processadas_com_OCR"].sum()) if not resultado["diagnosticos"].empty else 0
    return {
        "indicadores": {
            "lexicais": peso_lexical,
            "semanticos": int(len(semantico)),
            "resultados": int(len(df)),
            "livros": resultado["quantidade_pdfs"],
            "consultas": resultado["quantidade_termos"],
            "ocr": total_ocr,
        },
        "lexical_por_consulta": _serie(lexical, "Consulta", "_quantidade_no_registro"),
        "semantico_por_consulta": _serie(semantico, "Consulta"),
        "similaridade_media_por_consulta": {
            "rotulos": sorted(semantico["Consulta"].dropna().unique().tolist(), key=str.casefold),
            "valores": [
                round(float(semantico.loc[semantico["Consulta"] == consulta, "Similaridade semântica"].mean()), 4)
                for consulta in sorted(semantico["Consulta"].dropna().unique().tolist(), key=str.casefold)
            ],
        },
        "resultados_por_livro": _serie(df, "ID livro"),
        "comparacao": {"consultas": consultas, "livros": livros, "matriz": matriz},
        "contextos": _serie(df, "Contexto sociológico"),
        "distribuicao_similaridade": [round(float(valor), 4) for valor in semelhancas],
        "sankey": _dados_sankey(df),
    }
